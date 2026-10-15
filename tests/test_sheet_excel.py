"""Excel on the user's PC as the first formula engine (through the companion), with silent fallbacks; the _xlfn. prefix fix."""
import io
import unittest
from unittest import mock

import core.agent_tools  # noqa: F401
import openpyxl

import core.doc_ops as D
from core.sheet_engine import calc, excel, office
from core.sheet_engine.formulas import to_file_form

BOOK = """## Sheet: A
| x | y |
|---|---|
| 4 | =A2*2 |
| 5 | =A3/0 |
"""


def book():
    data, _ = D.create("t.xlsx", BOOK)
    return data


class FakeBridge:
    """Stands in for core.companion_bridge: connected or not, with the ops a companion advertises."""

    def __init__(self, connected=True, ops=("excel.recalc", "fs.read"), reply=None, error=None):
        self.connected, self.ops, self.reply, self.error = connected, ops, reply, error
        self.calls = []

    def is_connected(self, uid):
        return self.connected

    async def call(self, uid, op, params, timeout=60):
        self.calls.append((op, dict(params), timeout))
        if op == "companion.capabilities":
            if self.ops is None:
                raise RuntimeError("unknown op: companion.capabilities")
            return {"ops": list(self.ops), "platform": "win32", "version": "0.2.106"}
        if self.error:
            raise self.error
        return self.reply


class Engine(unittest.TestCase):
    def setUp(self):
        excel.forget()
        from core.request_context import set_current_user
        set_current_user(31)
        self.addCleanup(set_current_user, None)
        self.addCleanup(excel.forget)
        self.lo = mock.patch.object(office, "find_soffice", return_value=None)
        self.lo.start()
        self.addCleanup(self.lo.stop)

    def patch(self, bridge):
        p = mock.patch.object(excel, "_bridge", return_value=bridge)
        p.start()
        self.addCleanup(p.stop)
        return bridge

    def test_excel_values_replace_the_builtin_ones_and_the_report_says_so(self):
        b = self.patch(FakeBridge(reply={"cells": {"A!B2": 8, "A!B3": "#DIV/0!"}, "errors": 1, "version": "16.0"}))
        res = calc.calculate(book())
        self.assertEqual(res["engine"], "Excel")
        self.assertEqual(res["values"][("A", "B2")], 8)
        self.assertTrue(any("#DIV/0!" in e for e in res["errors"]))
        self.assertIn("Checked (Excel)", calc.check_report(res))
        op, payload, timeout = [c for c in b.calls if c[0] == "excel.recalc"][0]
        self.assertTrue(payload["data_b64"])
        self.assertEqual(payload["name"], "workbook.xlsx")
        self.assertGreaterEqual(timeout, 60)

    def test_every_failure_falls_back_to_the_next_engine_silently(self):
        for bridge in (FakeBridge(connected=False), FakeBridge(ops=("fs.read",)), FakeBridge(ops=None),
                       FakeBridge(error=TimeoutError("slow")), FakeBridge(error=RuntimeError("denied by local user")),
                       FakeBridge(reply=None), FakeBridge(reply={"cells": "nope"})):
            excel.forget()
            self.patch(bridge)
            res = calc.calculate(book())
            self.assertEqual(res["engine"], "built-in calculator")
            self.assertEqual(res["values"][("A", "B2")], 8)               # the built-in answer still stands

    def test_no_signed_in_user_means_no_call(self):
        from core.request_context import set_current_user
        set_current_user(None)
        b = self.patch(FakeBridge())
        self.assertIsNone(excel.recalculate(book()))
        self.assertEqual(b.calls, [])

    def test_the_switch(self):
        b = self.patch(FakeBridge(reply={"cells": {}}))
        with mock.patch.object(excel, "enabled", return_value=False):
            self.assertIsNone(excel.recalculate(book()))
        self.assertEqual(b.calls, [])

    def test_the_op_is_never_retried(self):
        data = book()                                   # building it already tried the (absent) engine
        b = self.patch(FakeBridge(error=ConnectionError("dropped")))
        self.assertIsNone(excel.recalculate(data))
        self.assertEqual(len([c for c in b.calls if c[0] == "excel.recalc"]), 1)
        from core import companion_bridge
        self.assertNotIn("excel.recalc", companion_bridge.READ_ONLY_OPS)

    def test_capabilities_are_cached_per_user(self):
        b = self.patch(FakeBridge(reply={"cells": {"A!B2": 8}}))
        excel.recalculate(book())
        excel.recalculate(book())
        self.assertEqual(len([c for c in b.calls if c[0] == "companion.capabilities"]), 1)

    def test_workbooks_over_the_device_limit_are_not_sent(self):
        b = self.patch(FakeBridge(reply={"cells": {}}))
        with mock.patch("core.doc_tools.constants.DEVICE_MAX_BYTES", 10):
            self.assertIsNone(excel.recalculate(book()))
        self.assertEqual(b.calls, [])

    def test_returned_text_is_pan_masked(self):
        self.patch(FakeBridge(reply={"cells": {"A!B2": "card 4111 1111 1111 1111"}}))
        out = excel.recalculate(book())
        self.assertNotIn("4111 1111 1111 1111", out[("A", "B2")])

    def test_the_audit_row_holds_counts_only(self):
        self.patch(FakeBridge(reply={"cells": {"A!B2": 8}, "errors": 0, "version": "16.0"}))
        rows = []
        with mock.patch("core.audit.audit_log", lambda *a, **k: rows.append((a, k))), \
                mock.patch("core.auth._to_principal", lambda uid: object()):
            excel.recalculate(book())
        self.assertEqual(rows[0][0][1], "sheet.excel_recalc")
        self.assertEqual(set(rows[0][1]["detail"]), {"cells", "errors", "excel"})

    def test_excel_comes_before_libreoffice(self):
        self.patch(FakeBridge(reply={"cells": {"A!B2": 8}}))
        with mock.patch.object(office, "recalculate", return_value={("A", "B2"): 999}) as lo:
            res = calc.calculate(book())
        self.assertEqual(res["engine"], "Excel")
        lo.assert_not_called()


class Status(unittest.TestCase):
    def test_status_says_what_will_be_used(self):
        excel.forget()
        with mock.patch.object(excel, "_bridge", return_value=FakeBridge(connected=True, ops=("fs.read",))), \
                mock.patch.object(office, "find_soffice", return_value=None):
            st = excel.status(5)
        self.assertTrue(st["excel"]["companion_connected"])
        self.assertFalse(st["excel"]["companion_supports_excel"])
        self.assertIn("update the companion", st["excel"]["note"])
        self.assertFalse(st["libreoffice"]["found"])
        self.assertEqual(st["order"][-1], "built-in calculator")
        excel.forget()
        with mock.patch.object(excel, "_bridge", return_value=FakeBridge(connected=False)), mock.patch.object(office, "find_soffice", return_value="C:/lo/soffice.exe"):
            st = excel.status(5)
        self.assertIn("connect the companion", st["excel"]["note"])
        self.assertTrue(st["libreoffice"]["found"])

    def test_the_endpoint_exists(self):
        from pathlib import Path
        src = (Path(__file__).resolve().parents[1] / "routes" / "capabilities" / "endpoints" / "settings.py").read_text(encoding="utf-8")
        self.assertIn('"/control/sheet_engines"', src)


class XlfnPrefix(unittest.TestCase):
    def test_newer_functions_are_prefixed_and_old_ones_and_strings_are_not(self):
        self.assertEqual(to_file_form('=XLOOKUP("a",A1:A3,B1:B3)'), '=_xlfn.XLOOKUP("a",A1:A3,B1:B3)')
        self.assertEqual(to_file_form("=SUM(A1:A3)+IFS(A1>1,1,TRUE,0)"), "=SUM(A1:A3)+_xlfn.IFS(A1>1,1,TRUE,0)")
        self.assertEqual(to_file_form('=IF(A1="XLOOKUP(",1,STDEV.S(B1:B4))'), '=IF(A1="XLOOKUP(",1,_xlfn.STDEV.S(B1:B4))')
        self.assertEqual(to_file_form("=_xlfn.XLOOKUP(1,A1:A2,B1:B2)"), "=_xlfn.XLOOKUP(1,A1:A2,B1:B2)")
        self.assertEqual(to_file_form("=VLOOKUP(1,A1:B2,2,FALSE)"), "=VLOOKUP(1,A1:B2,2,FALSE)")

    def test_a_built_workbook_stores_the_prefixed_form_and_still_calculates(self):
        data, _ = D.create("x.xlsx", '## Sheet: A\n| k | v |\n|---|---|\n| USD | 120 |\n| EUR | 130 |\n| r | =XLOOKUP("EUR",A2:A3,B2:B3) |\n')
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws["B4"].value, '=_xlfn.XLOOKUP("EUR",A2:A3,B2:B3)')
        with mock.patch.object(excel, "recalculate", return_value=None), mock.patch.object(office, "find_soffice", return_value=None):
            res = calc.calculate(data)
        self.assertEqual(res["values"][("A", "B4")], 130)                 # the built-in calculator understands the prefix too

    def test_the_xml_level_writer_adds_it_too(self):
        data, spec = D.create("x.xlsx", "## Sheet: A\n| k | v |\n|---|---|\n| a | 1 |\n")
        out = D.edit(data, "x.xlsx", [{"op": "set_formula", "addr": "A!B2", "formula": '=IFNA(XLOOKUP("a",A2:A2,B2:B2),0)'}], source_spec=spec)
        self.assertIn("_xlfn.IFNA(_xlfn.XLOOKUP", openpyxl.load_workbook(io.BytesIO(out.data)).active["B2"].value)


if __name__ == "__main__":
    unittest.main()
