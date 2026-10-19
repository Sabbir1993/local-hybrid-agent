"""Tests for advanced Excel expert directives: merge, color, align, border, and enhanced conditioning."""
import io
import unittest
import openpyxl

import core.doc_ops as D
from core.sheet_engine import build


OUTLINE_EXPERT = """# Executive Financial Summary
theme: ocean
## Sheet: Sales
| Date | Region | Revenue | Status |
|---|---|---|---|
| 2026-07-01 | Dhaka | 1200 | Active |
| 2026-07-02 | Khulna | 450 | Pending |
| 2026-07-03 | Dhaka | 1800 | Active |
format: Revenue=currency:USD, Date=date
total: Revenue=sum
merge: "Q3 Dhaka Highlights"=A6:D6
color: A6:D6 fill=#0B5E80 font=#FFFFFF bold=true align=center
align: Region=center, Status=center
border: A1:D4 thin
border: Total double
highlight: Revenue between 500 and 1500 amber
highlight: Status is "Active" green
"""


class SheetExpertTest(unittest.TestCase):
    def test_expert_directives_compiled(self):
        data, _spec = D.create("expert_report.xlsx", OUTLINE_EXPERT)
        self.assertIsNotNone(data)
        self.assertGreater(len(data), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb["Sales"]

        # 1. Merge verification
        # Cell A6:D6 is merged
        merged_ranges = [str(r) for r in ws.merged_cells.ranges]
        self.assertIn("A6:D6", merged_ranges)
        self.assertEqual(ws["A6"].value, "Q3 Dhaka Highlights")
        self.assertTrue(ws["A6"].font.bold)
        self.assertEqual(ws["A6"].alignment.horizontal, "center")

        # 2. Color verification
        self.assertEqual(ws["A6"].fill.fgColor.rgb, "000B5E80")
        self.assertEqual(ws["A6"].font.color.rgb, "00FFFFFF")

        # 3. Alignment verification
        self.assertEqual(ws["B2"].alignment.horizontal, "center")
        self.assertEqual(ws["D2"].alignment.horizontal, "center")

        # 4. Border verification
        self.assertEqual(ws["A1"].border.top.style, "thin")
        self.assertEqual(ws["A5"].border.bottom.style, "double")

        # 5. Conditional formatting verification
        cf_rules = list(ws.conditional_formatting)
        self.assertGreaterEqual(len(cf_rules), 2)
        operators = [rule.operator for cf in cf_rules for rule in cf.rules if hasattr(rule, "operator")]
        self.assertIn("between", operators)
        self.assertIn("equal", operators)


if __name__ == "__main__":
    unittest.main()
