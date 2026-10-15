"""Workbook outline -> real formatted workbook; formula allow-list; built-in calculator and cached values."""
import io
import unittest
from unittest import mock

import core.agent_tools  # noqa: F401
import openpyxl

import core.doc_ops as D
from core.sheet_engine import build, calc
from core.sheet_engine.formulas import FormulaError, check_formula, hardcoded_numbers

OUTLINE = """# Q3 Sales
theme: corporate
## Sheet: Assumptions
| Item | Value |
|---|---|
| Tax rate | 15% |
name: TaxRate = Assumptions!B2
## Sheet: Sales
| Date | Region | Revenue | Tax |
|---|---|---|---|
| 2026-07-05 | Dhaka | ৳1,000 | =C2*TaxRate |
| 2026-07-20 | Khulna | ৳2,500 | =C3*TaxRate |
| 2026-08-02 | Dhaka | ৳400 | =C4*TaxRate |
format: Revenue=currency:BDT, Tax=currency:BDT
total: Revenue=sum, Tax=sum
highlight: Revenue<500 red
scale: Revenue
validate: Region in [Dhaka, Khulna]
table: SalesTbl
chart: bar | title=Revenue by region | x=Region | y=Revenue
## Sheet: Summary
| Region | Revenue |
|---|---|
| Dhaka | =SUMIFS(Sales!C2:C4,Sales!B2:B4,A2) |
| Khulna | =SUMIFS(Sales!C2:C4,Sales!B2:B4,A3) |
"""


def make(text=OUTLINE):
    data, _spec = D.create("book.xlsx", text)
    return data


class Outline(unittest.TestCase):
    def test_detection(self):
        self.assertTrue(build.is_outline(OUTLINE))
        self.assertTrue(build.is_outline("| a | b |\n|---|---|\n| 1 | 2 |\nformat: a=int"))
        self.assertFalse(build.is_outline("| a | b |\n|---|---|\n| 1 | 2 |\n"))
        self.assertFalse(build.is_outline("just some prose"))

    def test_the_plain_table_path_is_unchanged(self):
        data, _ = D.create("t.xlsx", "| a | b |\n|---|---|\n| 1 | =SUM(A2:A3) |\n")
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws.title, "Sheet1")
        self.assertEqual(ws["B2"].data_type, "s")            # the old path never writes formulas


class Workbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = make()
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.vals = openpyxl.load_workbook(io.BytesIO(cls.data), data_only=True)

    def test_sheets_names_and_tables(self):
        self.assertEqual(self.wb.sheetnames, ["Assumptions", "Sales", "Summary"])
        self.assertIn("TaxRate", self.wb.defined_names)
        self.assertIn("SalesTbl", self.wb["Sales"].tables)

    def test_real_numbers_dates_and_formats(self):
        ws = self.wb["Sales"]
        self.assertEqual(ws["C2"].value, 1000)
        self.assertEqual(ws["C2"].data_type, "n")
        self.assertIn("৳", ws["C2"].number_format)
        self.assertEqual(ws["A2"].number_format, "yyyy-mm-dd")
        self.assertEqual(self.wb["Assumptions"]["B2"].value, 0.15)
        self.assertIn("%", self.wb["Assumptions"]["B2"].number_format)

    def test_formulas_are_live_and_have_cached_results(self):
        self.assertEqual(self.wb["Sales"]["D2"].value, "=C2*TaxRate")
        v = self.vals["Sales"]
        self.assertEqual([v["D2"].value, v["D3"].value, v["D4"].value], [150, 375, 60])
        self.assertEqual(v["C5"].value, 3900)                     # SUBTOTAL total row
        self.assertEqual(self.vals["Summary"]["B2"].value, 1400)
        self.assertEqual(self.vals["Summary"]["B3"].value, 2500)

    def test_styling_features(self):
        ws = self.wb["Sales"]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(len(ws._charts), 1)
        self.assertEqual(len(ws.data_validations.dataValidation), 1)
        self.assertTrue(len(ws.conditional_formatting) >= 1)
        self.assertTrue(ws["A1"].font.bold)
        self.assertGreater(ws.column_dimensions["B"].width, 8)

    def test_recalculation_is_requested_on_open(self):
        self.assertTrue(self.wb.calculation.fullCalcOnLoad)


class Refusals(unittest.TestCase):
    def refused(self, text, needle):
        with self.assertRaises(Exception) as cm:
            D.create("x.xlsx", text)
        self.assertIn(needle, str(cm.exception))

    def test_dangerous_functions(self):
        for fn in ('HYPERLINK("http://x","y")', 'WEBSERVICE("http://x")', 'INDIRECT("A1")', "CALL(\"a\")"):
            self.refused(f"## Sheet: A\n| a |\n|---|\n| ={fn} |\n", "not allowed")

    def test_external_references_and_unknown_sheets(self):
        self.refused("## Sheet: A\n| a |\n|---|\n| =[1]Sheet1!A1 |\n", "other workbooks")
        # a sheet that does not exist yet is allowed (built step by step); the check report flags it as #REF!
        data, _ = D.create("x.xlsx", "## Sheet: A\n| a |\n|---|\n| =Nope!A1 |\n")
        res = calc.calculate(data)
        self.assertTrue(any("#REF!" in e for e in res["errors"]))
        self.assertIn("add_outline_sheet", calc.check_report(res))

    def test_unknown_function(self):
        self.refused("## Sheet: A\n| a |\n|---|\n| =FROBNICATE(1) |\n", "not supported")

    def test_bad_directives(self):
        self.refused("## Sheet: A\n| a |\n|---|\n| 1 |\nhighlight: nope\n", "highlight")
        self.refused("## Sheet: A\n| a |\n|---|\n| 1 |\ntotal: zzz=sum\n", "total")
        self.refused("## Sheet: A\n| a |\n|---|\n| 1 |\nchart: bar | x=missing | y=a\n", "x column")
        self.refused("## Sheet: A\n| a |\n|---|\n| 1 |\nname: Foo = Missing!A1\n", "existing sheet")

    def test_text_that_looks_like_a_command_stays_text(self):
        data = make("## Sheet: A\n| a |\n|---|\n| +cmd|' /C calc'!A0 |\n| @SUM(1) |\n")
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws["A3"].data_type, "s")

    def test_leading_zero_identifiers_survive(self):
        ws = openpyxl.load_workbook(io.BytesIO(make("## Sheet: A\n| code |\n|---|\n| 007 |\n"))).active
        self.assertEqual(ws["A2"].value, "007")


class Formulas(unittest.TestCase):
    def test_check(self):
        self.assertEqual(check_formula("=SUM(A1:A3)*2"), "=SUM(A1:A3)*2")
        self.assertEqual(check_formula("=IFERROR(VLOOKUP(A2,Sheet1!A:B,2,FALSE),0)", sheets=["Sheet1"]),
                         "=IFERROR(VLOOKUP(A2,Sheet1!A:B,2,FALSE),0)")
        with self.assertRaises(FormulaError):
            check_formula("SUM(A1)")
        with self.assertRaises(FormulaError):
            check_formula("=" + "1+" * 600 + "1")

    def test_magic_numbers(self):
        self.assertEqual(hardcoded_numbers("=C2*1.0825"), ["1.0825"])
        self.assertEqual(hardcoded_numbers("=ROUND(C2*TaxRate,2)"), [])
        self.assertEqual(hardcoded_numbers("=SUBTOTAL(109,C2:C9)"), [])
        self.assertEqual(hardcoded_numbers('=IF(A1>"10",1,0)'), [])


class Calculator(unittest.TestCase):
    def calc(self, rows, **kw):
        body = "\n".join("| " + " | ".join(str(c) for c in r) + " |" for r in rows)
        data = make("## Sheet: S\n| a | b | c |\n|---|---|---|\n" + body + "\n")
        return calc.calculate(data)

    def val(self, formula_cell, rows):
        res = self.calc(rows)
        return res["values"].get(("S", formula_cell))

    def test_arithmetic_and_precedence(self):
        self.assertEqual(self.val("C2", [[2, 3, "=A2+B2*4"]]), 14)
        self.assertEqual(self.val("C2", [[2, 3, "=(A2+B2)*4"]]), 20)
        self.assertEqual(self.val("C2", [[2, 3, "=A2^B2"]]), 8)
        self.assertEqual(self.val("C2", [[50, 3, "=A2%"]]), 0.5)
        self.assertEqual(self.val("C2", [[2, 3, '=A2&"-"&B2']]), "2-3")

    def test_aggregates_and_conditionals(self):
        rows = [[1, "x", "=SUM(A2:A4)"], [2, "y", "=AVERAGE(A2:A4)"], [3, "x", "=SUMIF(B2:B4,\"x\",A2:A4)"]]
        r = self.calc(rows)["values"]
        self.assertEqual((r[("S", "C2")], r[("S", "C3")], r[("S", "C4")]), (6, 2, 4))
        rows = [[1, "x", "=COUNTIFS(B2:B4,\"x\",A2:A4,\">1\")"], [2, "y", "=MEDIAN(A2:A4)"], [3, "x", "=MAX(A2:A4)-MIN(A2:A4)"]]
        r = self.calc(rows)["values"]
        self.assertEqual((r[("S", "C2")], r[("S", "C3")], r[("S", "C4")]), (1, 2, 2))

    def test_logic_rounding_lookup_dates(self):
        rows = [[10, 3, '=IF(A2>B2,"big","small")'], [10, 3, "=ROUND(A3/B3,2)"], [10, 3, "=MOD(A4,B4)"],
                [1, "k", "=INDEX(A2:A5,MATCH(1,A2:A5,0))"], [2, 3, '=TEXT(DATE(2026,3,1),"yyyy-mm-dd")'], [2, 3, "=DAY(EOMONTH(DATE(2026,2,10),0))"]]
        r = self.calc(rows)["values"]
        self.assertEqual(r[("S", "C2")], "big")
        self.assertEqual(r[("S", "C3")], 3.33)
        self.assertEqual(r[("S", "C4")], 1)
        self.assertEqual(r[("S", "C5")], 1)
        self.assertEqual(r[("S", "C6")], "2026-03-01")
        self.assertEqual(r[("S", "C7")], 28)

    def test_finance_functions(self):
        r = self.calc([[0.01, 12, "=-PMT(A2,B2,1000)"]])["values"]
        self.assertAlmostEqual(r[("S", "C2")], 88.8488, places=3)

    def test_errors_are_reported_not_hidden(self):
        res = self.calc([[1, 0, "=A2/B2"], [1, 2, "=Z99"]])
        text = " ".join(res["errors"])
        self.assertIn("#DIV/0!", text)
        self.assertEqual(res["formulas"], 2)
        self.assertIn("1 error", calc.check_report(self.calc([[1, 0, "=A2/B2"]])) .replace("error(s)", "error"))

    def test_iferror_catches(self):
        self.assertEqual(self.val("C2", [[1, 0, "=IFERROR(A2/B2,0)"]]), 0)

    def test_circular_references_are_found(self):
        res = self.calc([[1, 2, "=C2+1"]])
        self.assertGreaterEqual(res["circular"], 1)

    def test_unsupported_functions_are_never_guessed(self):
        res = self.calc([[1, 2, "=TODAY()"]])
        self.assertNotIn(("S", "C2"), res["values"])
        self.assertTrue(res["unsupported"])

    def test_the_report_is_empty_without_formulas(self):
        self.assertEqual(calc.check_report({"formulas": 0}), "")


if __name__ == "__main__":
    unittest.main()


class Intent(unittest.TestCase):
    def test_sheet_build_requests(self):
        from core import prompt_scope as ps
        for q in ("create an excel workbook for monthly budget", "make me a spreadsheet to track sales", "build a workbook with a summary sheet"):
            self.assertTrue(ps.sheet_build_wanted(q), q)
        for q in ("analyse the attached excel file", "what is an xlsx file", "fix the formula in my spreadsheet", "create a pptx about sales"):
            self.assertFalse(ps.sheet_build_wanted(q), q)

    def test_the_manual_matches_the_builder(self):
        from core import prompt_scope as ps
        # the example in the manual must itself build, or the model is taught something that fails
        import re
        example = re.search(r"```\n(.*?)```", ps.XLSX_MANUAL, re.S).group(1)
        example = re.sub(r"\s{2,}\(corporate.*?\)", "", example)
        data = make(example)
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(data)).sheetnames, ["Assumptions", "Sales"])

    def test_the_route_offers_the_tools_and_manual(self):
        from pathlib import Path
        src = (Path(__file__).resolve().parents[1] / "routes" / "chat" / "run.py").read_text(encoding="utf-8")
        self.assertIn("sheet_build = prompt_scope.sheet_build_wanted(", src)
        self.assertIn("prompt_scope.XLSX_MANUAL", src)
        self.assertIn("or sheet_build or prompt_scope.deck_brief_wanted", src)


class LibreOffice(unittest.TestCase):
    """The optional LibreOffice path: used when installed, never trusted blindly, never fatal."""

    def test_not_installed_means_the_builtin_calculator(self):
        from core.sheet_engine import office
        with mock.patch.object(office, "find_soffice", return_value=None):
            res = calc.calculate(make(OUTLINE))
        self.assertEqual(res["engine"], "built-in calculator")
        self.assertIn("built-in calculator", calc.check_report(res))

    def test_its_values_replace_ours_and_errors_are_reported(self):
        from core.sheet_engine import office
        data = make("## Sheet: A\n| x | y |\n|---|---|\n| 4 | =A2*2 |\n| 5 | =A3/0 |\n")
        fake = {("A", "B2"): 8.0, ("A", "B3"): "#DIV/0!"}
        with mock.patch.object(office, "recalculate", return_value=fake):
            res = calc.calculate(data)
        self.assertEqual(res["engine"], "LibreOffice")
        self.assertEqual(res["values"][("A", "B2")], 8.0)
        self.assertTrue(any("#DIV/0!" in e for e in res["errors"]))
        self.assertIn("LibreOffice", calc.check_report(res))

    def test_a_failing_conversion_falls_back(self):
        from core.sheet_engine import office
        with mock.patch.object(office, "find_soffice", return_value="soffice"), \
                mock.patch("subprocess.run", side_effect=OSError("boom")):
            self.assertIsNone(office.recalculate(make(OUTLINE)))
            self.assertEqual(calc.calculate(make(OUTLINE))["engine"], "built-in calculator")

    def test_the_kill_switch(self):
        from core.sheet_engine import office
        with mock.patch.object(office, "enabled", return_value=False), mock.patch.object(office, "find_soffice", return_value="x"):
            self.assertIsNone(office.recalculate(b"x"))
