"""Deck theme part, real footer / slide-number placeholders, widened master and the `font:` override."""
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from core.doc_ops import pptx_ops
from core.doc_ops.pptx_ops import outline as O
from core.doc_ops.pptx_ops import theme as TH

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
MD = "# Deck\nsubtitle: s\ntheme: forest\n---\n## One\n- alpha\n- beta\n- gamma\n---\n## Two\n- delta\n- epsilon\n- zeta\n"


def prs_of(md=MD):
    return Presentation(io.BytesIO(pptx_ops.create(md)))


class ThemePart(unittest.TestCase):
    def _theme(self, prs):
        part = prs.slide_master.part.part_related_by(RT.THEME)
        return etree.fromstring(part.blob)

    def test_palette_and_fonts_are_in_the_theme(self):
        root = self._theme(prs_of())
        t = TH.THEMES["forest"]
        accent1 = root.find(".//a:clrScheme/a:accent1/a:srgbClr", NS).get("val")
        accent2 = root.find(".//a:clrScheme/a:accent2/a:srgbClr", NS).get("val")
        self.assertEqual((accent1, accent2), (t["accent"], t["accent2"]))
        self.assertEqual(root.find(".//a:fontScheme/a:majorFont/a:latin", NS).get("typeface"), t["head"])
        self.assertEqual(root.find(".//a:fontScheme/a:minorFont/a:latin", NS).get("typeface"), t["body"])

    def test_a_different_theme_gives_a_different_palette(self):
        a = self._theme(prs_of(MD)).find(".//a:clrScheme/a:accent1/a:srgbClr", NS).get("val")
        b = self._theme(prs_of(MD.replace("forest", "royal"))).find(".//a:clrScheme/a:accent1/a:srgbClr", NS).get("val")
        self.assertNotEqual(a, b)


class Footers(unittest.TestCase):
    def test_content_slides_have_real_footer_and_slide_number_placeholders(self):
        prs = prs_of()
        sl = prs.slides[1]
        types = {str(sh.placeholder_format.type).split(".")[-1].split(" ")[0] for sh in sl.placeholders}
        self.assertIn("SLIDE_NUMBER", types)
        self.assertIn("FOOTER", types)
        xml = etree.tostring(sl._element).decode()
        self.assertIn('type="slidenum"', xml)                  # a live field, not typed text

    def test_footers_are_named_deco_so_inspection_skips_them(self):
        sl = prs_of().slides[1]
        names = [sh.name for sh in sl.placeholders if sh.placeholder_format.idx in (11, 12)]
        self.assertEqual(sorted(names), ["deco-footer", "deco-number"])

    def test_covers_have_no_footer(self):
        sl = prs_of().slides[0]
        self.assertEqual([sh for sh in sl.placeholders if sh.placeholder_format.idx in (11, 12)], [])

    def test_deck_still_opens_and_reads_back(self):
        from core import doc_ops
        text = doc_ops.inspect(pptx_ops.create(MD), "d.pptx")
        self.assertIn("alpha", text)
        self.assertNotIn("deco-footer", text)


class Master(unittest.TestCase):
    def test_master_placeholders_span_the_wide_page(self):
        prs = prs_of()
        widest = max((sh.left + sh.width) / 914400 for sh in prs.slide_master.shapes)
        self.assertGreater(widest, 12.0)                       # was 9.5 in the stock 4:3 template


class FontOverride(unittest.TestCase):
    def test_font_line_overrides_both_fonts(self):
        deck = O.parse_outline("# D\nsubtitle: s\nfont: Arial\n---\n## A\n- x\n- y\n")
        t = O.resolve_theme(deck)
        self.assertEqual((t["head"], t["body"]), ("Arial", "Arial"))

    def test_bad_font_names_are_ignored(self):
        deck = O.parse_outline('# D\nsubtitle: s\nfont: Arial"/><evil\n---\n## A\n- x\n- y\n')
        self.assertEqual(deck["font"], "")

    def test_font_reaches_the_runs_and_the_theme(self):
        prs = prs_of(MD.replace("theme: forest", "theme: forest\nfont: Arial"))
        root = etree.fromstring(prs.slide_master.part.part_related_by(RT.THEME).blob)
        self.assertEqual(root.find(".//a:fontScheme/a:majorFont/a:latin", NS).get("typeface"), "Arial")
        fonts = {r.font.name for sh in prs.slides[1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
        self.assertEqual(fonts, {"Arial"})


if __name__ == "__main__":
    unittest.main()
