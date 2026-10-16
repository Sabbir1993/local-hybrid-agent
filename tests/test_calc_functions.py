"""tests/test_calc_functions.py - built-in calculator: error tests, SWITCH/CHOOSE, text functions, date criteria."""
import datetime
import io
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from core.sheet_engine import calc


def _eval(formulas: dict, cells: dict | None = None) -> dict:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for ref, v in (cells or {}).items():
        ws[ref] = v
    for ref, f in formulas.items():
        ws[ref] = f
    buf = io.BytesIO()
    wb.save(buf)
    with mock.patch.object(calc, "_prefer_external"):      # built-in calculator only: no Excel / LibreOffice in tests
        r = calc.calculate(buf.getvalue())
    return {ref: r["values"].get(("S", ref)) for ref in formulas} | {"_unsupported": r["unsupported"], "_errors": r["errors"]}


class CalcFunctionTests(unittest.TestCase):
    def test_iserror_and_isna(self):
        r = _eval({"A1": "=ISERROR(1/0)", "A2": "=ISERROR(1)", "A3": "=ISNA(1/0)", "A4": "=ISERR(1/0)"})
        self.assertEqual((r["A1"], r["A2"], r["A3"], r["A4"]), (True, False, False, True))
        self.assertEqual(r["_unsupported"], [])

    def test_switch_and_choose(self):
        r = _eval({"A1": '=SWITCH(2,1,"one",2,"two","other")', "A2": '=SWITCH(9,1,"one","other")',
                   "A3": '=CHOOSE(2,"a","b","c")', "A4": '=SWITCH("X","x",10,20)'})
        self.assertEqual((r["A1"], r["A2"], r["A3"], r["A4"]), ("two", "other", "b", 10))

    def test_text_functions(self):
        r = _eval({"A1": '=TEXTJOIN(", ",TRUE,B1:B3)', "A2": '=SUBSTITUTE("a-b-c","-","+")',
                   "A3": '=SUBSTITUTE("a-b-c","-","+",2)', "A4": '=FIND("b","abc")', "A5": '=SEARCH("B","abc")'},
                  {"B1": "x", "B3": "z"})
        self.assertEqual((r["A1"], r["A2"], r["A3"], r["A4"], r["A5"]), ("x, z", "a+b+c", "a-b+c", 2.0, 2.0))

    def test_ceiling_floor(self):
        r = _eval({"A1": "=CEILING(4.2,0.5)", "A2": "=FLOOR(4.7,0.5)"})
        self.assertEqual((r["A1"], r["A2"]), (4.5, 4.5))

    def test_iso_date_criteria_match_date_cells(self):
        cells = {"B1": datetime.date(2026, 1, 5), "B2": datetime.date(2025, 12, 1), "B3": datetime.date(2026, 2, 1),
                 "C1": 10, "C2": 20, "C3": 30}
        r = _eval({"A1": '=COUNTIF(B1:B3,">=2026-01-01")', "A2": '=SUMIFS(C1:C3,B1:B3,">=2026-01-01")'}, cells)
        self.assertEqual((r["A1"], r["A2"]), (2.0, 40.0))


if __name__ == "__main__":
    unittest.main()
