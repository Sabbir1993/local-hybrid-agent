"""Read-only SQL over sheets: joins, stats, dates, header detection, and the sandbox."""
import datetime
import tempfile
import unittest
from pathlib import Path

import core.agent_tools  # noqa: F401
import openpyxl

from core.file_tools import profile
from core.sheet_engine import sql


def build(d: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "ACME Ltd - Q3 Sales"
    ws.merge_cells("A1:D1")
    ws["A2"] = "prepared by finance"
    for i, h in enumerate(["Date", "Region", "Revenue", "Qty"], 1):
        ws.cell(4, i, h)
    rows = [(datetime.date(2026, 7, 5), "Dhaka", 100, 1), (datetime.date(2026, 7, 20), "Dhaka", 300, 3),
            (datetime.date(2026, 8, 2), "Khulna", 50, 2), (datetime.date(2026, 8, 9), "Dhaka", 200, 2),
            (datetime.date(2026, 9, 1), "Khulna", 150, 5)]
    for r, row in enumerate(rows, 5):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    ws["C10"] = "=SUM(C5:C9)"
    t = wb.create_sheet("Targets")
    t.append(["Region", "Target"])
    t.append(["Dhaka", 500])
    t.append(["Khulna", 300])
    p = d / "sales.xlsx"
    wb.save(p)
    return p


class Loader(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = build(Path(self.tmp.name))
        sql.clear_cache()

    def test_the_header_is_found_below_title_rows(self):
        t = profile.read_tables(self.path)[0]
        self.assertEqual(t.headers, ["Date", "Region", "Revenue", "Qty"])
        self.assertIn("ACME", t.preamble[0])
        self.assertEqual(t.total_rows, 5)             # the uncached formula row below the data is blank in the data view

    def test_formula_columns_are_reported(self):
        t = profile.read_tables(self.path)[0]
        self.assertIn("Revenue", t.formulas)

    def test_a_plain_table_still_uses_its_first_row(self):
        t = profile.read_tables(self.path)[1]
        self.assertEqual(t.headers, ["Region", "Target"])
        self.assertEqual(t.preamble, [])


class Queries(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = build(Path(self.tmp.name))
        sql.clear_cache()
        self.wb = sql.load(self.path)

    def q(self, text, **kw):
        return sql.query(self.wb, text, **kw)

    def test_describe_lists_tables_and_columns(self):
        d = sql.describe(self.wb)
        self.assertIn("Sales", d)
        self.assertIn('"Revenue" REAL', d)
        self.assertIn("MEDIAN", d)

    def test_group_by_with_a_join_to_another_sheet(self):
        out = self.q('SELECT s.Region, SUM(s.Revenue) AS rev, t.Target, ROUND(100.0*SUM(s.Revenue)/t.Target) AS pct '
                     'FROM Sales s JOIN Targets t ON t.Region = s.Region WHERE s.Region IS NOT NULL GROUP BY s.Region ORDER BY s.Region')
        self.assertIn("| Dhaka | 600 | 500 | 120 |", out)
        self.assertIn("| Khulna | 200 | 300 | 67 |", out)

    def test_month_buckets_and_pivots(self):
        out = self.q("SELECT strftime('%Y-%m', Date) AS m, SUM(CASE WHEN Region='Dhaka' THEN Revenue END) AS dhaka, "
                     "SUM(CASE WHEN Region='Khulna' THEN Revenue END) AS khulna FROM Sales WHERE Date IS NOT NULL GROUP BY m ORDER BY m")
        self.assertIn("| 2026-07 | 400 |", out)
        self.assertIn("| 2026-08 | 200 | 50 |", out)

    def test_median_percentile_stdev_mode(self):
        out = self.q("SELECT MEDIAN(Revenue), PERCENTILE(Revenue, 0.5), ROUND(STDEV(Revenue), 2), MODE(Region) FROM Sales WHERE Region IS NOT NULL")
        self.assertIn("| 150 | 150 |", out)
        self.assertIn("Dhaka", out)

    def test_window_functions(self):
        out = self.q("SELECT Region, Revenue, RANK() OVER (PARTITION BY Region ORDER BY Revenue DESC) AS r FROM Sales WHERE Region IS NOT NULL ORDER BY Region, r")
        self.assertIn("| Dhaka | 300 | 1 |", out)

    def test_row_cap_and_note(self):
        out = self.q("SELECT * FROM Sales", limit=2)
        self.assertIn("more rows exist", out)


class Sandbox(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        sql.clear_cache()
        self.wb = sql.load(build(Path(self.tmp.name)))

    def refused(self, text):
        with self.assertRaises(sql.SqlError):
            sql.query(self.wb, text)

    def test_only_select_is_allowed(self):
        for bad in ("DELETE FROM Sales", "UPDATE Sales SET Revenue = 0", "DROP TABLE Sales", "INSERT INTO Sales VALUES (1,2,3,4)",
                    "PRAGMA table_info(Sales)", "ATTACH DATABASE 'x.db' AS x", "CREATE TABLE z (a)", "", "   "):
            self.refused(bad)

    def test_hidden_tricks_inside_a_select(self):
        self.refused("SELECT load_extension('x')")
        self.refused("SELECT * FROM sqlite_master")
        self.refused("WITH x AS (SELECT 1) DELETE FROM Sales")
        self.refused("SELECT 1; DROP TABLE Sales")

    def test_data_survives_the_attempts(self):
        for bad in ("DELETE FROM Sales", "DROP TABLE Sales"):
            try:
                sql.query(self.wb, bad)
            except sql.SqlError:
                pass
        self.assertIn("| 5 |", sql.query(self.wb, "SELECT COUNT(*) FROM Sales WHERE Region IS NOT NULL"))

    def test_a_runaway_query_is_stopped(self):
        old = sql.QUERY_TIMEOUT_S
        sql.QUERY_TIMEOUT_S = 0.3
        try:
            with self.assertRaises(sql.SqlError) as cm:
                sql.query(self.wb, "WITH RECURSIVE r(n) AS (SELECT 1 UNION ALL SELECT n+1 FROM r) SELECT COUNT(*) FROM r")
        finally:
            sql.QUERY_TIMEOUT_S = old
        self.assertIn("longer", str(cm.exception))

    def test_card_numbers_in_cells_are_masked(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cards"
        ws.append(["name", "card"])
        ws.append(["a", "4111 1111 1111 1111"])
        p = Path(self.tmp.name) / "cards.xlsx"
        wb.save(p)
        out = sql.query(sql.load(p), "SELECT * FROM Cards")
        self.assertNotIn("4111 1111 1111 1111", out)


class Tool(unittest.TestCase):
    def test_the_tool_is_registered_and_rejects_other_files(self):
        from core.agent_tools import TOOL_IMPLS, AGENT_TOOLS
        self.assertIn("data_sql", TOOL_IMPLS)
        self.assertTrue(any(t["function"]["name"] == "data_sql" for t in AGENT_TOOLS))
        self.assertTrue(TOOL_IMPLS["data_sql"]({}).startswith("error:"))
        self.assertTrue(TOOL_IMPLS["data_sql"]({"file": "nope.xlsx"}).startswith("error:"))


if __name__ == "__main__":
    unittest.main()
