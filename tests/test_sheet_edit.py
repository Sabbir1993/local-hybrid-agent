"""Restyling a built workbook: ops change its outline and the workbook is rebuilt from it."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import core.agent_tools  # noqa: F401
import openpyxl

import core.doc_ops as D
from core.sheet_engine import edit as se

BASE = """# Sales
## Sheet: Assumptions
| Item | Value |
|---|---|
| Rate | 10% |
name: Rate = Assumptions!B2
## Sheet: Sales
| Region | Revenue |
|---|---|
| Dhaka | 1000 |
| Khulna | 500 |
"""


def built():
    data, spec = D.create("s.xlsx", BASE)
    return data, spec


def wb_of(data, values=False):
    return openpyxl.load_workbook(io.BytesIO(data), data_only=values)


class Outline(unittest.TestCase):
    def edit(self, ops, spec=None):
        data, base_spec = built()
        return D.edit(data, "s.xlsx", ops, source_spec=spec if spec is not None else base_spec)

    def test_add_chart_and_format_and_rules(self):
        out = self.edit([{"op": "add_chart", "sheet": "Sales", "chart": "bar | title=Revenue | x=Region | y=Revenue"},
                         {"op": "set_format", "sheet": "Sales", "column": "Revenue", "format": "currency:BDT"},
                         {"op": "add_highlight", "sheet": "Sales", "rule": "Revenue<600 red"},
                         {"op": "add_databar", "sheet": "Sales", "column": "Revenue"},
                         {"op": "add_validation", "sheet": "Sales", "rule": "Region in [Dhaka, Khulna]"},
                         {"op": "set_total", "sheet": "Sales", "column": "Revenue", "fn": "sum"}])
        ws = wb_of(out.data)["Sales"]
        self.assertEqual(len(ws._charts), 1)
        self.assertIn("৳", ws["B2"].number_format)
        self.assertGreaterEqual(len(ws.conditional_formatting), 1)
        self.assertEqual(len(ws.data_validations.dataValidation), 1)
        self.assertEqual(wb_of(out.data, True)["Sales"]["B4"].value, 1500)       # total row calculated
        self.assertTrue(any("chart added" in c for c in out.changes))

    def test_theme_freeze_table_note(self):
        out = self.edit([{"op": "set_theme", "theme": "forest"}, {"op": "set_freeze", "sheet": "Sales", "cell": "B2"},
                         {"op": "set_table", "sheet": "Sales", "name": "RevTbl"}, {"op": "add_note", "sheet": "Sales", "text": "Source: CRM"}])
        wb = wb_of(out.data)
        self.assertEqual(wb["Sales"].freeze_panes, "B2")
        self.assertIn("RevTbl", wb["Sales"].tables)
        self.assertIn("theme: forest", out.source_spec)
        self.assertEqual(wb["Sales"]["A1"].fill.fgColor.rgb[-6:], "2E5E4E")

    def test_add_column_with_a_row_formula(self):
        out = self.edit([{"op": "add_column", "sheet": "Sales", "name": "Tax", "formula": "=B{r}*Rate"}])
        vals = wb_of(out.data, True)["Sales"]
        self.assertEqual([vals["C1"].value, vals["C2"].value, vals["C3"].value], ["Tax", 100, 50])
        self.assertEqual(wb_of(out.data)["Sales"]["C2"].value, "=B2*Rate")

    def test_the_new_outline_is_returned_so_edits_chain(self):
        first = self.edit([{"op": "set_format", "sheet": "Sales", "column": "Revenue", "format": "int"}])
        data, _ = built()
        second = D.edit(first.data, "s.xlsx", [{"op": "add_note", "sheet": "Sales", "text": "n"}], source_spec=first.source_spec)
        self.assertIn("format: Revenue=int", second.source_spec)
        self.assertIn("note: n", second.source_spec)

    def test_bad_ops_get_a_reason(self):
        for ops, needle in (([{"op": "add_chart", "sheet": "Nope", "chart": "bar"}], "not found"),
                            ([{"op": "set_format", "sheet": "Sales"}], "column required"),
                            ([{"op": "add_column", "sheet": "Sales", "name": "x", "formula": "B2"}], "must start with ="),
                            ([{"op": "add_column", "sheet": "Sales", "name": "x", "formula": "=HYPERLINK(\"u\")"}], "not allowed"),
                            ([{"op": "add_highlight", "sheet": "Sales", "rule": "garbage"}], "highlight")):
            with self.assertRaises(Exception) as cm:
                self.edit(ops)
            self.assertIn(needle, str(cm.exception))

    def test_a_workbook_not_built_from_an_outline_is_refused_with_advice(self):
        data, _ = D.create("t.xlsx", "| a | b |\n|---|---|\n| 1 | 2 |\n")
        with self.assertRaises(Exception) as cm:
            D.edit(data, "t.xlsx", [{"op": "set_theme", "theme": "forest"}], source_spec="| a | b |\n|---|---|\n| 1 | 2 |\n")
        self.assertIn("rebuild it with doc_create", str(cm.exception))
        with self.assertRaises(Exception):
            D.edit(data, "t.xlsx", [{"op": "set_theme", "theme": "forest"}], source_spec=None)

    def test_cell_ops_and_style_ops_cannot_be_mixed(self):
        with self.assertRaises(Exception) as cm:
            self.edit([{"op": "set_theme", "theme": "forest"}, {"op": "set_cell", "addr": "Sales!B2", "value": 5}])
        self.assertIn("cannot be mixed", str(cm.exception))

    def test_cell_ops_still_work_and_drop_the_outline(self):
        out = self.edit([{"op": "set_cell", "addr": "Sales!B2", "value": 1111}])
        self.assertEqual(wb_of(out.data)["Sales"]["B2"].value, 1111)
        self.assertIsNone(out.source_spec)            # the file no longer matches its outline


class Tools(unittest.TestCase):
    """The same through doc_create / doc_edit (version chain, outline kept)."""

    def setUp(self):
        from core import agent_tools, doc_store
        from core.request_context import set_current_user
        import tests.test_doc_ops as T
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._patches = [mock.patch.object(agent_tools, "COMMON_ROOT", self.root),
                         mock.patch.object(doc_store, "_projects_db", T.doc_files_db())]
        for p in self._patches:
            p.start()
        set_current_user(12)
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        from core.request_context import set_current_user
        for p in self._patches:
            p.stop()
        set_current_user(None)
        self.tmp.cleanup()

    def test_create_then_restyle_twice(self):
        from core import doc_tools
        out = doc_tools.tool_doc_create_common({"file": "book.xlsx", "content": BASE})
        name = out.split("[DOWNLOAD: ")[1].split("]")[0]
        e1 = doc_tools.tool_doc_edit_common({"file": name, "ops": [{"op": "add_chart", "sheet": "Sales", "chart": "pie | x=Region | y=Revenue"}]})
        self.assertIn("chart added", e1)
        name2 = e1.split("[DOWNLOAD: ")[1].split("]")[0]
        e2 = doc_tools.tool_doc_edit_common({"file": name2, "ops": [{"op": "set_theme", "theme": "sunset"}]})
        self.assertIn("theme set to sunset", e2)
        name3 = e2.split("[DOWNLOAD: ")[1].split("]")[0]
        ws = openpyxl.load_workbook(self.root / "user_12" / name3)["Sales"]
        self.assertEqual(len(ws._charts), 1)               # the chart from the first edit survived the second


if __name__ == "__main__":
    unittest.main()
