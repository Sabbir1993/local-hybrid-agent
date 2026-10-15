"""One csv-ish text with `SheetN,Title` sections becomes a workbook with real sheets (what the model meant)."""
import io
import unittest

import core.agent_tools  # noqa: F401
import openpyxl

import core.doc_ops as D
from core.sheet_engine import pseudo

TEXT = """Hercules TAP - Bangladesh
Sheet1,Executive Summary
Item,Recommendation
Go/No-Go Decision,"RECOMMENDED - well-positioned, with targeted enhancements"
Total Additional Development Effort,~17 calendar months

Sheet2,Cost Analysis
Phase,Months,Cost USD
Foundation,6,"95,000"
Mobile,4,"120,500"

Sheet3,Risks
Risk,Likelihood,Mitigation
Connectivity,High,Offline-first app

Sheet8,Board Presentation Summary Slide Content
Slide Title,Hercules TAP - Bangladesh Market Expansion
Key Points for Board:
• Current Hercules has strong core SFA features
"""


class Pseudo(unittest.TestCase):
    def test_sections_become_sheets(self):
        data, _ = D.create("tap.xlsx", TEXT)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Executive Summary", "Cost Analysis", "Risks", "Board Presentation Summary Slide Co"[:31]])
        ws = wb["Cost Analysis"]
        self.assertEqual([c.value for c in ws[1]], ["Phase", "Months", "Cost USD"])
        self.assertEqual(ws["B2"].value, 6)
        self.assertEqual(ws["C2"].value, 95000)
        self.assertEqual(wb["Executive Summary"]["B2"].value, "RECOMMENDED - well-positioned, with targeted enhancements")
        self.assertEqual(wb.properties.title, "Hercules TAP - Bangladesh")

    def test_not_triggered_by_ordinary_data(self):
        self.assertIsNone(pseudo.to_outline("Name,Role\nA,B\nC,D\n"))
        self.assertIsNone(pseudo.to_outline("Sheet1,only one\nA,B\n"))
        self.assertIsNone(pseudo.to_outline("Sheet3,x\nA,B\nSheet2,y\nC,D\n"))        # not ascending

    def test_markdown_heading_style_markers(self):
        out = pseudo.to_outline("## Sheet 1: Alpha\na,b\n1,2\n## Sheet 2: Beta\nc,d\n3,4\n")
        self.assertIn("## Sheet: Alpha", out)
        self.assertIn("## Sheet: Beta", out)

    def test_ragged_rows_and_pipes_survive(self):
        out = pseudo.to_outline("Sheet1,A\nh1,h2,h3\nx,y\nSheet2,B\nk\nv|w\n")
        data, _ = D.create("r.xlsx", out)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb["A"].max_column, 3)
        self.assertEqual(wb["B"]["A2"].value, "v|w")


if __name__ == "__main__":
    unittest.main()


class Headers(unittest.TestCase):
    SCREENSHOT = [["Bengali Language Support", "Critical", "Bangla UI and data entry support for field reps who are not comfortable with English", "Critical", "High", "6 weeks", "Core App Dev", "High"],
                  ["Offline Mode & Sync", "Mandatory", "Poor internet connectivity in rural Bangladesh requires full offline functionality", "Critical", "High", "8 weeks", "Backend + Mobile", "High"],
                  ["VAT & Tax Compliance", "Bangladesh", "Necessary to align invoice and ledger formats with Bangladesh VAT Act", "Critical", "Low", "4 weeks", "Legal Review", "Medium"]]

    def test_a_first_record_is_not_taken_as_the_header(self):
        self.assertFalse(pseudo.looks_like_header(self.SCREENSHOT))
        out = pseudo.with_header(self.SCREENSHOT)
        self.assertEqual(out[0], [f"Column{i}" for i in range(1, 9)])
        self.assertEqual(len(out), 4)

    def test_real_headers_are_kept(self):
        rows = [["Feature", "Priority", "Effort", "Owner"], ["Offline mode", "Critical", "8 weeks", "Mobile"], ["VAT", "High", "4 weeks", "Legal"]]
        self.assertTrue(pseudo.looks_like_header(rows))
        self.assertEqual(pseudo.with_header(rows), rows)
        self.assertTrue(pseudo.looks_like_header([["Name", "Qty", "Price"], ["A", "1", "2.5"], ["B", "3", "4"]]))

    def test_short_first_rows_that_repeat_down_their_column_are_data(self):
        rows = [["Critical", "High"], ["Critical", "High"], ["Critical", "Low"], ["High", "High"]]
        self.assertFalse(pseudo.looks_like_header(rows))

    def test_a_number_is_not_a_label(self):
        self.assertFalse(pseudo.looks_like_header([["2024", "10"], ["a", "b"]]))

    def test_one_row_or_empty_is_left_alone(self):
        self.assertEqual(pseudo.with_header([["only", "row"]]), [["only", "row"]])
        self.assertEqual(pseudo.with_header([]), [])

    def test_the_flat_workbook_gets_a_generic_header_and_the_writer_says_so(self):
        csv_text = "\n".join(",".join(f'"{c}"' for c in r) for r in self.SCREENSHOT)
        data, _ = D.create("f.xlsx", csv_text)
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual([c.value for c in ws[1]][:2], ["Column1", "Column2"])
        self.assertEqual(ws["A2"].value, "Bengali Language Support")
        from core.agent_tools.common_writer import _sheet_note
        self.assertIn("header row", _sheet_note(data))

    def test_a_proper_flat_table_is_untouched(self):
        data, _ = D.create("f.xlsx", "Feature,Priority\nOffline,High\nVAT,Low\n")
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual([c.value for c in ws[1]], ["Feature", "Priority"])
        from core.agent_tools.common_writer import _sheet_note
        self.assertNotIn("header row", _sheet_note(data))

    def test_pseudo_sheets_get_the_same_treatment(self):
        text = "Sheet1,A\n" + "\n".join(",".join(f'"{c}"' for c in r) for r in self.SCREENSHOT) + "\nSheet2,B\nName,Qty\nx,1\n"
        data, _ = D.create("p.xlsx", text)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb["A"]["A1"].value, "Column1")
        self.assertEqual(wb["B"]["A1"].value, "Name")
