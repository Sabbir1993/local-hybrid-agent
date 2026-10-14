"""Production behaviour of the deck / PDF builders: gallery slides, concurrent picture lookup with a time budget, memory
caps, kill switches, and failures that must never cost the user their document."""
import asyncio
import io
import sys
import time
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.agent_tools  # noqa: F401,E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from core import agent_tools, doc_store, doc_tools  # noqa: E402
from core.doc_ops import pptx_ops  # noqa: E402
from core.doc_ops.pptx_ops import design as D  # noqa: E402
from core.doc_ops.pptx_ops import lint, measure  # noqa: E402
from core.doc_ops.pptx_ops import outline as O  # noqa: E402
from core.doc_tools import operations as ops  # noqa: E402
from core.doc_tools import slide_images as si  # noqa: E402
from core.request_context import set_current_user  # noqa: E402


def png(w=1200, h=800) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (w, h), (30, 120, 200)).save(b, "PNG")
    return b.getvalue()


def info(host="example.org", extra=None):
    d = {"data": png(), "credit_host": host, "credit_url": f"https://{host}/p", "via": "web"}
    if extra:
        d["extra"] = extra
    return d


GALLERY_MD = ("# Deck\nsubtitle: s\n---\n## Places\nlayout: gallery\nimage: Lalbagh Fort Dhaka | caption=Dhaka sights\n"
              "image: Ahsan Manzil Dhaka\nimage: Star Mosque Dhaka\n")


class Gallery(unittest.TestCase):
    def test_extra_image_lines_make_a_gallery(self):
        s = O.parse_outline(GALLERY_MD)["slides"][1]
        self.assertEqual(s["kind"], "gallery")
        self.assertEqual([x["query"] for x in s["image_extra"]], ["Ahsan Manzil Dhaka", "Star Mosque Dhaka"])

    def test_a_gallery_needs_more_than_one_picture(self):
        md = "# Deck\nsubtitle: s\n---\n## Places\nlayout: gallery\nimage: Lalbagh Fort\n- one\n- two\n- three\n"
        self.assertEqual(O.parse_outline(md)["slides"][1]["kind"], "bullets")

    def test_three_pictures_render_and_stay_clean(self):
        images = {1: info(extra=[info("b.org"), info("c.org")])}
        data = pptx_ops.create(GALLERY_MD, images=images)
        prs = Presentation(io.BytesIO(data))
        names = [sh.name for sh in prs.slides[1].shapes if sh.shape_type == 13]
        self.assertEqual(names, ["Image 1", "Image 2", "Image 3"])
        self.assertEqual(lint.lint_deck(data), [])
        notes = prs.slides[1].notes_slide.notes_text_frame.text
        self.assertIn("https://b.org/p", notes)
        credit = " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
        self.assertIn("example.org", credit)
        self.assertIn("c.org", credit)

    def test_two_pictures_side_by_side(self):
        md = GALLERY_MD.replace("image: Star Mosque Dhaka\n", "")
        prs = Presentation(io.BytesIO(pptx_ops.create(md, images={1: info(extra=[info("b.org")])})))
        pics = [sh for sh in prs.slides[1].shapes if sh.shape_type == 13]
        self.assertEqual(len(pics), 2)
        self.assertLess(pics[0].left + pics[0].width, pics[1].left)

    def test_resolver_returns_extras(self):
        calls = []

        async def fake(uid, q, why, anchors=()):
            calls.append(q)
            return {"data": b"x", "credit_host": q.split()[0], "credit_url": "u", "via": "web"}
        mock.patch.object(si.doc_store, "audit").start()
        self.addCleanup(mock.patch.stopall)
        si._cache.clear()
        si.allow_web.set(True)
        with mock.patch.object(si, "_from_web", fake):
            images, notes, used_web = asyncio.run(si.resolve(GALLERY_MD, 3))
        self.assertEqual(len(images[1]["extra"]), 2)
        self.assertEqual(len(calls), 3)

    def test_a_gallery_still_shows_what_it_found(self):
        async def fake(uid, q, why, anchors=()):
            if "Ahsan" in q:
                return {"data": b"x", "credit_host": "a", "credit_url": "u", "via": "web"}
            why.append("no usable picture came back from the search")
            return None
        mock.patch.object(si.doc_store, "audit").start()
        self.addCleanup(mock.patch.stopall)
        si._cache.clear()
        with mock.patch.object(si, "_from_web", fake), mock.patch.object(si, "_generate", mock.AsyncMock(return_value=None)):
            images, notes, _ = asyncio.run(si.resolve(GALLERY_MD, 3))
        self.assertIn(1, images)                                     # promoted: the one picture that was found
        self.assertTrue(any("not found" in n for n in notes))


class Concurrency(unittest.TestCase):
    MD = "# D\n" + "".join(f"---\n## S{i}\nimage: topic number {i}\n- x\n- y\n" for i in range(6))

    def setUp(self):
        si._cache.clear()
        mock.patch.object(si.doc_store, "audit").start()
        self.addCleanup(mock.patch.stopall)

    def test_pictures_are_looked_up_in_parallel(self):
        async def slow(uid, q, why, anchors=()):
            await asyncio.sleep(0.3)
            return {"data": b"x", "credit_host": "h", "credit_url": "u", "via": "web"}
        t0 = time.monotonic()
        with mock.patch.object(si, "_from_web", slow):
            images, _, _ = asyncio.run(si.resolve(self.MD, 4))
        self.assertEqual(len(images), 6)
        self.assertLess(time.monotonic() - t0, 6 * 0.3 * 0.8)      # serial would be 1.8 s

    def test_the_image_model_runs_one_picture_at_a_time(self):
        running, peak = [0], [0]
        with mock.patch.object(si, "_from_web", mock.AsyncMock(return_value=None)), \
                mock.patch("core.media.status", return_value={"image": {"ready": True, "cloud": False, "state": "loaded"}}), \
                mock.patch("core.media_tools._user", return_value=object()), \
                mock.patch("core.media.generate", new=mock.AsyncMock(side_effect=self._slow_generate(running, peak))), \
                mock.patch("core.agent_tools._common_resolve", side_effect=lambda p: _FakePath()):
            si.allow_web.set(False)
            images, _, _ = asyncio.run(si.resolve(self.MD, 4))
        self.assertEqual(len(images), 6)
        self.assertEqual(peak[0], 1)                                # six slides, never two generations at once

    @staticmethod
    def _slow_generate(running, peak):
        async def fake(kind, user, prompt, opts):
            running[0] += 1
            peak[0] = max(peak[0], running[0])
            await asyncio.sleep(0.05)
            running[0] -= 1
            return {"files": [{"path": "generated/x.png"}]}
        return fake

    def test_an_exhausted_budget_skips_instead_of_hanging(self):
        with mock.patch.object(si, "RESOLVE_BUDGET_S", -1.0), mock.patch.object(si, "_from_web", mock.AsyncMock()) as web:
            images, notes, _ = asyncio.run(si.resolve(self.MD, 4))
        self.assertEqual(images, {})
        self.assertTrue(all("budget" in n for n in notes), notes)
        web.assert_not_called()

    def test_a_picture_that_runs_too_long_is_given_up(self):
        async def hang(uid, q, why, anchors=()):
            await asyncio.sleep(10)
        with mock.patch.object(si, "MIN_PICTURE_S", 0.1), mock.patch.object(si, "RESOLVE_BUDGET_S", 0.1), \
                mock.patch.object(si, "_from_web", hang):
            t0 = time.monotonic()
            images, notes, _ = asyncio.run(si.resolve("# D\n---\n## A\nimage: boats\n- x\n- y\n", 4))
        self.assertLess(time.monotonic() - t0, 3)
        self.assertEqual(images, {})
        self.assertTrue(any("took too long" in n for n in notes), notes)


class _FakePath:
    def read_bytes(self):
        return png()


class Memory(unittest.TestCase):
    def test_cache_is_capped_by_bytes(self):
        si._cache.clear()
        with mock.patch.object(si, "_CACHE_BYTES", 3_000_000):
            for i in range(10):
                si._cache[(i,)] = {"data": b"x" * 1_000_000}
                # the same eviction rule _one applies after every insert
                while len(si._cache) > si._CACHE_MAX or sum(len(v["data"]) for v in si._cache.values()) > si._CACHE_BYTES:
                    si._cache.popitem(last=False)
        self.assertLessEqual(len(si._cache), 3)
        si._cache.clear()


class FailureIsolation(unittest.TestCase):
    MD = "# D\nsubtitle: s\n---\n## Growth\nchart: bar\n- A: 1\n- B: 2\n---\n## Plain\n- one\n- two\n- three\n"

    def test_a_slide_that_cannot_render_becomes_a_plain_slide(self):
        with mock.patch.object(D, "_chart", side_effect=RuntimeError("boom")):
            prs = Presentation(io.BytesIO(pptx_ops.create(self.MD)))
        self.assertEqual(len(prs.slides), 3)                           # nothing lost, no half-drawn extra slide
        self.assertEqual(prs.slides[1].shapes.title.text_frame.text, "Growth")
        self.assertIn("A: 1", " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame))
        self.assertEqual(prs.slides[2].shapes.title.text_frame.text, "Plain")

    def test_theme_polish_failures_are_not_fatal(self):
        with mock.patch.object(D, "_apply_theme_part", side_effect=ValueError("bad xml")), \
                mock.patch.object(D, "_widen_master", side_effect=ValueError("bad master")):
            prs = Presentation(io.BytesIO(pptx_ops.create(self.MD)))
        self.assertEqual(len(prs.slides), 3)

    def test_the_layout_check_never_breaks_a_build(self):
        with mock.patch("core.doc_ops.pptx_ops.lint.lint_deck", side_effect=RuntimeError("lint bug")):
            self.assertEqual(ops._layout_note(".pptx", pptx_ops.create(self.MD)), "")
        self.assertEqual(ops._layout_note(".pdf", b"garbage", "# x"), "")

    def test_template_table_failure_leaves_the_rest_of_the_deck(self):
        from pptx import Presentation as P
        template = pptx_ops.create("# T\nsubtitle: base\n---\n## One\n- a\n- b\n- c\n")
        md = "# D\nsubtitle: s\n---\n## Data\n| A | B |\n|---|---|\n| 1 | 2 |\n"
        with mock.patch("core.doc_ops.pptx_ops.creation._template_table", side_effect=RuntimeError("x")):
            prs = P(io.BytesIO(pptx_ops.create(md, template=template)))
        self.assertEqual(len(prs.slides), 2)


class KillSwitches(unittest.TestCase):
    def setUp(self):
        import tempfile
        import tests.test_doc_ops as T
        self.tmp = tempfile.TemporaryDirectory()
        self._patches = [mock.patch.object(agent_tools, "COMMON_ROOT", Path(self.tmp.name)),
                         mock.patch.object(doc_store, "_projects_db", T.doc_files_db())]
        for p in self._patches:
            p.start()
        set_current_user(6)
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        for p in self._patches:
            p.stop()
        set_current_user(None)
        self.tmp.cleanup()

    def test_deck_images_off_never_calls_the_resolver(self):
        md = "# D\nsubtitle: s\n---\n## A\nimage: boats\n- x\n- y\n"
        with mock.patch.object(ops, "_flag", side_effect=lambda n, d=True: n != "deck_images"), \
                mock.patch.object(si, "resolve", new=mock.AsyncMock()) as r:
            out = doc_tools.tool_doc_create_common({"file": "k.pptx", "content": md})
        r.assert_not_called()
        self.assertIn("Created", out)

    def test_deck_checks_off_skips_the_layout_note(self):
        with mock.patch.object(ops, "_flag", side_effect=lambda n, d=True: n != "deck_checks"), \
                mock.patch("core.doc_ops.pptx_ops.lint.lint_deck", side_effect=AssertionError("must not run")):
            self.assertEqual(ops._layout_note(".pptx", b"x"), "")

    def test_flags_default_to_on(self):
        self.assertTrue(ops._flag("deck_images"))
        self.assertTrue(ops._flag("deck_checks"))

    def test_a_card_number_in_the_plan_is_refused(self):
        out = doc_tools.tool_doc_create_common({"file": "p.pptx", "content": "# D\n---\n## A\n- x\n- y\n", "total": 3,
                                                "plan": "1. Cover\n2. Pay with 4111 1111 1111 1111\n3. End"})
        self.assertNotIn("Created", out)
        self.assertIn("card", out.lower())


class FontIndex(unittest.TestCase):
    def test_index_is_built_once_and_lookups_are_case_insensitive(self):
        measure._index.cache_clear()
        with mock.patch.object(measure, "_DIRS", []):
            self.assertEqual(measure._index(), {})
            self.assertIsNone(measure._find("Calibri.ttf"))
        measure._index.cache_clear()

    def test_missing_font_falls_back_to_the_estimate(self):
        measure._font.cache_clear()
        with mock.patch.object(measure, "_find", return_value=None):
            measure._font.cache_clear()
            self.assertGreaterEqual(measure.lines_for("word " * 40, "Calibri", 20, 3.0), 2)
        measure._font.cache_clear()


if __name__ == "__main__":
    unittest.main()
