"""Shared output checks for the DOCX / PDF benchmarks (offline builder tests and the live /chat/run harness).

Each check returns a list of short findings, [] when the file looks right. Nothing here reads or stores document text beyond
what the finding sentence needs: findings carry counts and fixed phrases, never content, so a PCI-scoped log stays clean.
"""
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

_MD_LEFTOVERS = (
    (re.compile(r"\*\*[^*\s][^*]*\*\*"), "literal **bold** markers"),
    (re.compile(r"^\s*\|?\s*:?-{3,}:?\s*\|", re.M), "literal table separator row"),
    (re.compile(r"^\s*#{1,6}\s+\S", re.M), "literal # heading marker"),
    (re.compile(r"```"), "literal code fence"),
)


def check_docx(data: bytes, markdown: str = "", min_words: int = 0) -> list:
    """Findings for a built .docx: unreadable, leftover markdown, missing headings / tables / footer page field, too short."""
    import docx
    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as e:                                    # noqa: BLE001 - any parse failure is the finding
        return [f"does not open ({type(e).__name__})"]
    out = []
    paras = [p for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(p.text for p in paras)
    for t in doc.tables:
        for row in t.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    body = "\n".join(p.text for p in paras)           # a "#" column header in a table is not a heading marker
    for rx, label in _MD_LEFTOVERS:
        if rx.search(body if label.startswith("literal #") else text):
            out.append(label)
    want_heads = len(re.findall(r"^#{1,3}\s+\S", markdown or "", re.M))
    got_heads = sum(1 for p in paras if (p.style.name or "").lower().startswith(("heading", "title")))
    if want_heads and got_heads < max(1, want_heads // 2):
        out.append(f"only {got_heads} heading styles for {want_heads} markdown headings")
    want_tables = len(re.findall(r"^\s*\|\s*:?-{3,}", markdown or "", re.M))
    if want_tables and len(doc.tables) < want_tables:
        out.append(f"{len(doc.tables)} tables for {want_tables} in the outline")
    if any(len(t.rows) < 2 for t in doc.tables):
        out.append("a table with no body rows")
    words = len(text.split())
    if min_words and words < min_words:
        out.append(f"{words} words, expected at least {min_words}")
    xml = "".join(s.footer._element.xml for s in doc.sections)
    if "PAGE" not in xml:
        out.append("no page number field in the footer")
    return out


def check_pdf(data: bytes, markdown: str = "", min_pages: int = 1) -> list:
    """Findings for a built PDF: the app's own check (blank / sparse pages, stranded headings, footer) plus page count."""
    from core.agent_tools import pdf_check
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(data)
        pages = len(pdf)
        pdf.close()
    except Exception as e:                                    # noqa: BLE001
        return [f"does not open ({type(e).__name__})"]
    out = list(pdf_check.check_pdf(data, markdown))
    if pages < min_pages:
        out.append(f"{pages} pages, expected at least {min_pages}")
    return out
