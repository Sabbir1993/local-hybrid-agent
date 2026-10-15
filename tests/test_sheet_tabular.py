"""Model text -> tables: pipes without outer pipes, several tables, notes, thousands separators, first-record-as-header."""
import io
import unittest

import core.agent_tools  # noqa: F401
import openpyxl

import core.doc_ops as D
from core.sheet_engine import tabular

MIXED = """Key Highlights:
- Hercules is currently a secondary sales force automation app
- For stronger market grip in Bangladesh, we need offline capabilities

Feature Gap Analysis
Module | Feature | Status | Usage Level | Notes
Order Management | Digital Order Entry | ✅ Active | High | Core functionality working
Tour Plan | Route Optimization | ⚠️ Partial | Low | Basic route mapping only

Budget Estimate
Resource Type | Quantity | Cost (BDT) | Timeline | Notes
Software Developers | 8 FTEs | 4,800,000/year | Full project | Senior + Mid level mix
Third-party APIs | - | 600,000/year | Ongoing | SMS, MFS, Maps

KPIs
KPI Category,Metric,Current Value,Target (6 months),Measurement Frequency
Market Penetration,Active Users,TBD,+40% increase,Monthly
Reconciliation Time,Avg hours per day,8-10 hrs,3-4 hrs,Daily
"""


def book(text):
    data, _ = D.create("t.xlsx", text)
    return openpyxl.load_workbook(io.BytesIO(data))


class Tables(unittest.TestCase):
    def test_each_table_is_its_own_sheet_named_by_its_heading(self):
        wb = book(MIXED)
        self.assertEqual(wb.sheetnames, ["Overview", "Feature Gap Analysis", "Budget Estimate", "KPIs"])

    def test_columns_are_not_merged_or_split(self):
        wb = book(MIXED)
        ws = wb["Feature Gap Analysis"]
        self.assertEqual([c.value for c in ws[1]], ["Module", "Feature", "Status", "Usage Level", "Notes"])
        self.assertEqual(ws["C2"].value, "✅ Active")
        self.assertEqual(ws.max_column, 5)
        b = wb["Budget Estimate"]
        self.assertEqual(b["C2"].value, "4,800,000/year")               # the thousands separators stayed inside one cell
        self.assertEqual(b.max_column, 5)
        self.assertEqual([c.value for c in wb["KPIs"][1]][:3], ["KPI Category", "Metric", "Current Value"])

    def test_headers_are_real_and_styled(self):
        ws = book(MIXED)["Budget Estimate"]
        self.assertEqual(ws["A1"].value, "Resource Type")
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.auto_filter.ref)

    def test_the_prose_is_kept_as_an_overview_sheet_not_dropped_or_poured_into_cells(self):
        ws = book(MIXED)["Overview"]
        self.assertEqual(ws["A1"].value, "Overview")
        texts = [r[0].value for r in ws.iter_rows(min_row=2)]
        self.assertIn("Hercules is currently a secondary sales force automation app", texts)

    def test_a_single_table_stays_sheet1_and_formulas_stay_text(self):
        wb = book("Item | Qty\nA | 2\nB | =SUM(B2:B2)\n")
        self.assertEqual(wb.sheetnames, ["Sheet1"])
        self.assertEqual(wb["Sheet1"]["B3"].data_type, "s")

    def test_over_split_numbers_are_repaired_in_csv(self):
        wb = book("Name,Cost,Note\nCloud,4,800,000/year,ok\nAPIs,600,000,fine\n")
        ws = wb.active
        self.assertEqual(ws["B2"].value, "4,800,000/year")
        self.assertEqual(ws["C2"].value, "ok")
        self.assertEqual(ws.max_column, 3)

    def test_markdown_with_a_separator_keeps_its_header(self):
        ws = book("| Name | Role |\n|---|---|\n| A | Manager |\n| B | Engineer |\n").active
        self.assertEqual([c.value for c in ws[1]], ["Name", "Role"])

    def test_no_header_row_gets_generic_names_not_a_record_in_the_header(self):
        rows = "Offline Mode | Mandatory | Poor internet connectivity in rural Bangladesh requires full offline functionality | Critical\n" \
               "VAT Compliance | Bangladesh | Necessary to align invoice and ledger formats with the Bangladesh VAT Act | Critical\n"
        ws = book(rows).active
        self.assertEqual(ws["A1"].value, "Column1")
        self.assertEqual(ws["A2"].value, "Offline Mode")

    def test_long_text_wraps(self):
        ws = book("Item | Detail\nA | " + "long sentence " * 8 + "\nB | short\n").active
        self.assertTrue(ws["B2"].alignment.wrap_text)

    def test_code_and_prose_are_not_tables(self):
        self.assertIsNone(tabular.to_outline("Just a paragraph, with a comma. And another sentence."))
        self.assertIsNone(tabular.to_outline("```python\nx = a | b\ny = c | d\n```"))

    def test_ragged_rows_are_padded(self):
        ws = book("A | B | C\n1 | 2\n3 | 4 | 5 | 6\n").active
        self.assertEqual(ws.max_column, 3)
        self.assertEqual(ws["C3"].value, "5, 6")

    def test_many_tables_cap_at_twenty_sheets(self):
        text = "\n\n".join(f"Table {i}\nA | B\n{i} | x" for i in range(30))
        self.assertLessEqual(len(book(text).sheetnames), 20)


if __name__ == "__main__":
    unittest.main()
