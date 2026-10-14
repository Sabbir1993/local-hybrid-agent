"""The build plan (doc_create plan=...), section numbers, and the layout / PDF notes the build tools return."""
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.agent_tools  # noqa: F401,E402
from pptx import Presentation  # noqa: E402

from core import agent_tools, doc_store, doc_tools  # noqa: E402
from core.doc_ops import pptx_ops  # noqa: E402
from core.doc_tools import operations as ops  # noqa: E402
from core.request_context import set_current_user  # noqa: E402

PLAN = ("1. Dhaka - cover\n2. Intro - why Dhaka\n3. History - key dates\n4. Geography - river and climate\n"
        "5. Thank you - close")


class Plan(unittest.TestCase):
    def test_parse_plan_lines(self):
        got = ops._parse_plan(PLAN)
        self.assertEqual([t for t, _ in got], ["Dhaka", "Intro", "History", "Geography", "Thank you"])
        self.assertEqual(got[1][1], "why Dhaka")

    def test_parse_plan_tolerates_bullets_and_noise(self):
        got = ops._parse_plan("- **Cover** - title\n* Risks: the main ones\n\n   \n3) Wrap up")
        self.assertEqual([t for t, _ in got], ["Cover", "Risks", "Wrap up"])

    def test_plan_is_capped(self):
        self.assertEqual(len(ops._parse_plan("\n".join(f"{i}. Slide {i}" for i in range(1, 60)))), 30)

    def test_next_in_your_plan(self):
        ops._BUILD_PLAN[(1, "f")] = ops._parse_plan(PLAN)
        self.assertEqual(ops._plan_note((1, "f"), 3, [], False),
                         "Next in your plan: 4. Geography - river and climate; 5. Thank you - close.")
        self.assertEqual(ops._plan_note((9, "none"), 2, [], False), "")

    def test_done_reports_what_never_got_built(self):
        ops._BUILD_PLAN[(1, "g")] = ops._parse_plan(PLAN)
        note = ops._plan_note((1, "g"), 5, ["Dhaka", "Intro", "History", "Thank you"], True)
        self.assertIn("Geography", note)
        self.assertNotIn((1, "g"), ops._BUILD_PLAN)                # forgotten once the build is done


class ToolFlow(unittest.TestCase):
    def setUp(self):
        import tests.test_doc_ops as T
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._patches = [mock.patch.object(agent_tools, "COMMON_ROOT", self.root),
                         mock.patch.object(doc_store, "_projects_db", T.doc_files_db())]
        for p in self._patches:
            p.start()
        set_current_user(5)
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        for p in self._patches:
            p.stop()
        set_current_user(None)
        self.tmp.cleanup()

    def test_plan_is_played_back_through_the_build(self):
        first = "# Dhaka\nsubtitle: s\n---\n## Intro\n- a\n- b\n"
        out = doc_tools.tool_doc_create_common({"file": "d.pptx", "content": first, "total": 5, "plan": PLAN})
        name = out.split("[DOWNLOAD: ")[1].split("]")[0]
        more = doc_tools.tool_doc_append_common({"file": name, "content": "---\n## History\n- c\n- d\n"})
        self.assertIn("Next in your plan: 4. Geography", more)
        last = doc_tools.tool_doc_append_common({"file": name, "content": "---\n## Geography\n- e\n- f\n---\n## Thank you\n", "done": True})
        self.assertNotIn("Planned but not found", last)

    def test_deck_check_runs_on_the_last_step(self):
        first = "# Dhaka\nsubtitle: s\n---\n## Intro\n- a\n- b\n"
        out = doc_tools.tool_doc_create_common({"file": "e.pptx", "content": first, "total": 3})
        name = out.split("[DOWNLOAD: ")[1].split("]")[0]
        self.assertNotIn("Layout check", out)
        with mock.patch.object(ops, "_layout_note", return_value="Layout check found 1 issue(s): slide 2: demo.") as note:
            last = doc_tools.tool_doc_append_common({"file": name, "content": "---\n## B\n- x\n- y\n---\n## Thank you\n", "done": True})
        self.assertIn("Layout check found 1 issue(s)", last)
        self.assertTrue(note.called)


class SectionNumbers(unittest.TestCase):
    def test_sections_are_numbered_in_order(self):
        md = ("# D\nsubtitle: s\n---\n## Part one\nlayout: section\n---\n## A\n- x\n- y\n---\n## Part two\nlayout: section\n"
              "---\n## B\n- x\n- y\n")
        prs = Presentation(io.BytesIO(pptx_ops.create(md)))
        nums = []
        for sl in prs.slides:
            for sh in sl.shapes:
                if sh.name == "deco-section-number":
                    nums.append(sh.text_frame.text)
        self.assertEqual(nums, ["01", "02"])


if __name__ == "__main__":
    unittest.main()
