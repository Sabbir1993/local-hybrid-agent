"""tests/eval_kb.py - measured question set for the knowledge base tools.

Builds a synthetic company corpus (Excel with two differently shaped sheets, CSV, Word, PDF, and one very large
roster), ingests it through the real pipeline (extract -> mask -> rows + chunks -> index) into temporary
databases, then asks questions of each kind through the real tool functions and checks the answers against
facts known from the generator:

    lookup   - "everything about <name>" across documents (kb_find_mentions)
    count    - exact counts, filters, group-by over spreadsheets (kb_query)
    policy   - free-text questions answered from document chunks (search_knowledge_base, lexical-only: no embedder here)
    absence  - names that are not there must give "no mention", partial indexes must say so
    scale    - a fact in the last row of a roster far past the old 400-chunk cap

No model is involved: this measures whether the tools return the right evidence. Whether a model then calls the
right tool is a separate live check. Run:  python tests/eval_kb.py   (exit code 1 on any miss)
"""

import asyncio
import json
import random
import sys
import tempfile
import time
from contextlib import contextmanager
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

FIRST = ["Aminul", "Bithi", "Chandra", "Dilruba", "Emon", "Farhana", "Gias", "Hasna", "Imran", "Jui", "Kabir", "Lima",
         "Mahin", "Nila", "Omar", "Parvez", "Rima", "Shuvo", "Tania", "Ullash"]
LAST = ["Ahmed", "Biswas", "Chowdhury", "Das", "Hossain", "Islam", "Khan", "Mia", "Paul", "Roy", "Sarker", "Talukder"]
ROLES = ["Manager", "Engineer", "Analyst", "Support", "Accountant"]
DEPTS = ["Finance", "Engineering", "Operations", "Sales"]


def _people(rng, n, tag):
    seen, out = set(), []
    while len(out) < n:
        name = f"{rng.choice(FIRST)} {rng.choice(LAST)} {tag}{len(out)}"      # unique: the tag keeps planted names unambiguous
        if name not in seen:
            seen.add(name)
            out.append(name)
    return out


def build_corpus(root: Path, seed: int = 7) -> dict:
    """Write the files and return the facts the questions check."""
    import openpyxl
    from docx import Document
    from reportlab.pdfgen.canvas import Canvas
    rng = random.Random(seed)
    facts = {}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dhaka"
    ws.append(["Name", "Role", "Dept", "Phone", "Salary"])
    dhaka = []
    for i, name in enumerate(_people(rng, 400, "D")):
        row = [name, rng.choice(ROLES), rng.choice(DEPTS), f"017{rng.randint(10000000, 99999999)}", rng.randint(30, 120) * 1000]
        if i == 249:
            row = ["Rahim Uddin", "Manager", "Finance", "01711111111", 95000]
        if i == 120:
            row = ["Nusrat Jahan", "Analyst", "Sales", "01722222222", 52000]
        ws.append(row)
        dhaka.append(row)
    ws2 = wb.create_sheet("Khulna")
    ws2.append(["Employee", "Grade", "City", "Mobile"])
    khulna = []
    for i, name in enumerate(_people(rng, 300, "K")):
        row = [name, f"G{rng.randint(1, 9)}", "Khulna", f"018{rng.randint(10000000, 99999999)}"]
        if i == 77:
            row = ["Nusrat Jahan", "G7", "Khulna", "01833333333"]
        khulna.append(row)
        ws2.append(row)
    wb.save(root / "staff_roster.xlsx")
    managers = [r for r in dhaka if r[1] == "Manager"]
    facts["dhaka_managers"] = len(managers)
    facts["dhaka_by_dept"] = {d: sum(1 for r in dhaka if r[2] == d) for d in DEPTS}
    facts["dhaka_high_salary"] = sorted(r[0] for r in dhaka if r[4] > 100000)
    facts["khulna_g7"] = sum(1 for r in khulna if r[1] == "G7")

    (root / "contractors.csv").write_text(
        "Contractor,Agency,Rate\n" + "\n".join(
            f"{('Rahim Uddin' if i == 60 else 'Vendor Person C' + str(i))},Agency{i % 5},{rng.randint(10, 90)}"
            for i in range(120)), encoding="utf-8")

    d = Document()
    d.add_heading("Leave policy", 1)
    d.add_paragraph("Every full-time employee receives 21 days of annual leave per calendar year.")
    d.add_paragraph("Sick leave requires a medical certificate after 3 consecutive days.")
    d.add_paragraph("Leave requests are approved by the department head, currently Rahim Uddin for Finance.")
    for i in range(30):
        d.add_paragraph(f"Clause {i}: general conduct provisions apply to all staff members equally.")
    d.save(root / "hr_policy.docx")

    c = Canvas(str(root / "handbook.pdf"))
    for page in range(1, 31):
        c.drawString(72, 720, f"Employee handbook page {page}")
        if page == 17:
            c.drawString(72, 700, "Relocation allowance is BDT 50000 for grade G7 and above.")
            c.drawString(72, 680, "Approved by Rahim Uddin, Finance.")
        c.showPage()
    c.save()

    big = ["Name,Role,Dept"] + [f"Bulk Person B{i},{rng.choice(ROLES)},{rng.choice(DEPTS)}" for i in range(7000)]
    big.append("Zubayer Quartz,Manager,Operations")
    (root / "big_roster.csv").write_text("\n".join(big), encoding="utf-8")
    facts["big_rows"] = 7001
    return facts


# ------------------------------------------------------------------ environment

@contextmanager
def kb_environment(root: Path):
    """Temporary memory/auth databases, no embedder, audit and access control stubbed."""
    import core.memory  # noqa: F401
    memory_pkg = sys.modules["core.memory"]      # not `import core.memory as x`: another test may have re-imported the package
    from core.memory import indexing as indexing_mod
    from core import auth_db
    from core.memory import cache as cache_mod
    from core.memory import store as store_mod
    import core.kb_tables as kb_tables
    saved = (store_mod._conn, getattr(memory_pkg, "_conn", None), auth_db._auth_db)
    had_file = hasattr(memory_pkg, "MEMORY_DB_FILE")
    saved_file = getattr(memory_pkg, "MEMORY_DB_FILE", None)
    store_mod._conn = None
    memory_pkg._conn = None
    memory_pkg.MEMORY_DB_FILE = root / "memory.db"
    store_mod._db()
    kb_tables._ready = False
    cache_mod._set_cache({"gen": None, "entries": [], "lower": [], "dim": None, "mat": None, "pos": None})
    with mock.patch.object(auth_db, "AUTH_DB_FILE", root / "auth.db"):
        auth_db._auth_db = auth_db._init_auth_db()

    async def no_embed(texts):
        return None
    try:
        # both bindings: a real embedder here would start a llama-server on the GPU and take minutes
        import core.knowledge_router as knowledge_router    # binds _embed_texts at import and backfills vectors on search
        with mock.patch.object(memory_pkg, "_embed_texts", no_embed), mock.patch.object(indexing_mod, "_embed_texts", no_embed), \
                mock.patch.object(knowledge_router, "_embed_texts", no_embed), \
                mock.patch("routes.knowledge.helpers.audit_log"), mock.patch("core.audit.audit_log"):
            yield
    finally:
        for closer in (lambda: store_mod._db().close(), lambda: auth_db._auth_db.close()):
            try:
                closer()
            except Exception:
                pass
        store_mod._conn, memory_pkg._conn, auth_db._auth_db = saved
        if had_file:
            memory_pkg.MEMORY_DB_FILE = saved_file
        elif hasattr(memory_pkg, "MEMORY_DB_FILE"):
            delattr(memory_pkg, "MEMORY_DB_FILE")
        cache_mod._set_cache({"gen": None, "entries": [], "lower": [], "dim": None, "mat": None, "pos": None})


def ingest(root: Path) -> dict:
    """Ingest every corpus file through the real pipeline. Returns {title: source_id}."""
    from core import knowledge_ingest
    from core.auth_db import knowledge_acl
    from routes.knowledge import helpers
    ids = {}
    for path in sorted(root.iterdir()):
        if path.suffix.lower() not in knowledge_ingest.EXTRACTORS:
            continue
        sid = knowledge_acl.create_knowledge_source(path.name, "file", path.name)
        text = knowledge_ingest.extract_file(path)
        res = asyncio.run(helpers._finish_ingest(sid, text, None))
        assert res.get("ok"), res
        ids[path.name] = sid
    return ids


def call(tool, **args):
    return asyncio.run(tool(args))


# ------------------------------------------------------------------ questions

def questions(facts: dict) -> list:
    """(category, question, tool_name, args, check(output)->bool)."""
    return [
        ("lookup", "everything about Rahim Uddin: found in all four document types", "kb_find_mentions", {"name": "Rahim Uddin"},
         lambda o: all(t in o for t in ("staff_roster.xlsx", "contractors.csv", "hr_policy.docx", "handbook.pdf"))),
        ("lookup", "Rahim Uddin row is cited with sheet and row number", "kb_find_mentions", {"name": "Rahim Uddin"},
         lambda o: "Sheet Dhaka, row 251" in o and "Columns: Name | Role | Dept | Phone | Salary" in o),
        ("lookup", "Rahim Uddin in the PDF is cited by page", "kb_find_mentions", {"name": "Rahim Uddin"},
         lambda o: "p.17" in o),
        ("lookup", "a name on two sheets gets each sheet's own header", "kb_find_mentions", {"name": "Nusrat Jahan"},
         lambda o: "Columns: Name | Role" in o and "Columns: Employee | Grade" in o and "01833333333" in o),
        ("lookup", "reversed word order still finds the person", "kb_find_mentions", {"name": "Uddin Rahim"},
         lambda o: "staff_roster.xlsx" in o),
        ("absence", "a name that is not in the KB says so", "kb_find_mentions", {"name": "Zaman Nonexistent"},
         lambda o: "no document" in o),
        ("absence", "a partial first name does not match another person", "kb_find_mentions", {"name": "Rahi"},
         lambda o: "no document" in o),
        ("count", "managers in Dhaka (exact count)", "kb_query",
         {"source": "staff", "sheet": "Dhaka", "filters": [{"column": "Role", "op": "=", "value": "Manager"}], "count": True},
         lambda o: f"{facts['dhaka_managers']} matching row" in o),
        ("count", "headcount by department (exact table)", "kb_query", {"source": "staff", "sheet": "Dhaka", "group_by": "Dept"},
         lambda o: all(f"| {d} | {n} |" in o for d, n in facts["dhaka_by_dept"].items())),
        ("count", "everyone earning above 100000 (exact list)", "kb_query",
         {"source": "staff", "sheet": "Dhaka", "filters": [{"column": "Salary", "op": ">", "value": 100000}],
          "select": ["Name", "Salary"], "limit": 200},
         lambda o: all(n in o for n in facts["dhaka_high_salary"]) and f"{len(facts['dhaka_high_salary'])} matching row" in o),
        ("count", "a column that exists only on the Khulna sheet", "kb_query",
         {"source": "staff", "filters": [{"column": "Grade", "op": "=", "value": "G7"}], "count": True},
         lambda o: f"{facts['khulna_g7']} matching row" in o and "Not searched" in o),
        ("count", "partial column name resolves (Mob -> Mobile)", "kb_query",
         {"source": "staff", "filters": [{"column": "Mob", "op": "starts_with", "value": "01833"}], "select": ["Employee"]},
         lambda o: "Nusrat Jahan" in o),
        ("count", "describe lists the tables and columns", "kb_query", {"describe": True},
         lambda o: "Employee | Grade | City | Mobile" in o and "Contractor | Agency | Rate" in o),
        ("policy", "annual leave days (document text)", "search_knowledge_base", {"query": "how many days of annual leave do employees get"},
         lambda o: "21 days" in o),
        ("policy", "relocation allowance lives on a PDF page", "search_knowledge_base", {"query": "relocation allowance for grade G7"},
         lambda o: "50000" in o),
        ("scale", "last row of a 7000-row roster is searchable", "kb_find_mentions", {"name": "Zubayer Quartz"},
         lambda o: "big_roster.csv" in o and "Manager" in o),
        ("scale", "kb_query counts the whole large roster", "kb_query", {"source": "big", "filters": [{"column": "Name", "op": "not_empty"}], "count": True},
         lambda o: f"{facts['big_rows']} matching row" in o),
    ]


def run() -> dict:
    from core.agent_tools import definitions
    tools = definitions.TOOL_IMPLS
    results = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        facts = build_corpus(root)
        with kb_environment(root):
            ids = ingest(root)
            ready = set(ids.values())
            with mock.patch("core.knowledge_access.allowed_source_ids_for", return_value=ready):
                for cat, text, name, args, check in questions(facts):
                    t0 = time.time()
                    out = call(tools[name], **args)
                    results.append({"category": cat, "question": text, "tool": name, "ok": bool(check(out)),
                                    "ms": round((time.time() - t0) * 1000), "output": out[:400]})
                # a source cut at the cap must say so, and a miss must not read as absence
                from core.auth_db import knowledge_acl
                from core import auth_db
                with mock.patch.dict("core.small_model.APP_CONFIG", {"knowledge": {"max_chunks_per_doc": 40}}):
                    from routes.knowledge import helpers
                    from core import knowledge_ingest
                    sid = ids["big_roster.csv"]
                    asyncio.run(helpers._finish_ingest(sid, knowledge_ingest.extract_file(root / "big_roster.csv"), None))
                out = call(tools["kb_find_mentions"], name="Zubayer Quartz")
                results.append({"category": "absence", "question": "a partly indexed source is flagged on a miss",
                                "tool": "kb_find_mentions", "ok": "partly indexed" in out and "big_roster.csv" in out,
                                "ms": 0, "output": out[:400]})
                row = next(r for r in auth_db.list_knowledge_sources() if r["id"] == sid)
                results.append({"category": "absence", "question": "the source records indexed vs total chunks",
                                "tool": "-", "ok": row["chunks_indexed"] == 40 and row["chunks_total"] > 40, "ms": 0,
                                "output": f"{row['chunks_indexed']}/{row['chunks_total']}"})
                del knowledge_acl
    by_cat = {}
    for r in results:
        by_cat.setdefault(r["category"], []).append(r["ok"])
    return {"results": results, "by_category": {c: f"{sum(v)}/{len(v)}" for c, v in by_cat.items()},
            "passed": sum(r["ok"] for r in results), "total": len(results),
            "p95_ms": sorted(r["ms"] for r in results)[int(len(results) * 0.95) - 1]}


if __name__ == "__main__":
    report = run()
    for r in report["results"]:
        print(f"[{'PASS' if r['ok'] else 'FAIL'}] {r['category']:8} {r['question']}  ({r['ms']} ms)")
        if not r["ok"]:
            print("        output:", r["output"].replace("\n", " | ")[:300])
    print(json.dumps({k: report[k] for k in ("by_category", "passed", "total", "p95_ms")}, indent=2))
    sys.exit(0 if report["passed"] == report["total"] else 1)
