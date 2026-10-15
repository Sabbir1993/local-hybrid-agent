"""Ten realistic workbook outlines: each must build, open in openpyxl, calculate with zero errors, and keep formulas live."""
import io
import unittest

import core.agent_tools  # noqa: F401
import openpyxl

import core.doc_ops as D
from core.sheet_engine import calc

SPECS = {
    "sales_summary": """# Sales
## Sheet: Data
| Date | Region | Product | Units | Price |
|---|---|---|---|---|
| 2026-01-05 | Dhaka | A | 10 | 25 |
| 2026-01-18 | Khulna | B | 4 | 40 |
| 2026-02-02 | Dhaka | B | 7 | 40 |
| 2026-02-20 | Sylhet | A | 12 | 25 |
format: Price=currency:BDT
total: Units=sum
chart: column | title=Units | x=Region | y=Units
## Sheet: Summary
| Region | Revenue |
|---|---|
| Dhaka | =SUMPRODUCT((Data!B2:B5=A2)*Data!D2:D5*Data!E2:E5) |
| Khulna | =SUMPRODUCT((Data!B2:B5=A3)*Data!D2:D5*Data!E2:E5) |
| Sylhet | =SUMPRODUCT((Data!B2:B5=A4)*Data!D2:D5*Data!E2:E5) |
""",
    "budget": """# Budget
theme: forest
## Sheet: Assumptions
| Item | Value |
|---|---|
| Inflation | 6% |
| Growth | 10% |
name: Inflation = Assumptions!B2
name: Growth = Assumptions!B3
## Sheet: Budget
| Line | 2026 | 2027 |
|---|---|---|
| Revenue | 1000 | =B2*(1+Growth) |
| Costs | 700 | =B3*(1+Inflation) |
| Profit | =B2-B3 | =C2-C3 |
format: 2026=number, 2027=number
""",
    "monthly_pivot": """# Pivot
## Sheet: Sales
| Month | Region | Amount |
|---|---|---|
| 2026-01-01 | Dhaka | 100 |
| 2026-01-01 | Khulna | 80 |
| 2026-02-01 | Dhaka | 120 |
| 2026-02-01 | Khulna | 90 |
## Sheet: Pivot
| Month | Dhaka | Khulna | Total |
|---|---|---|---|
| 2026-01 | =SUMIFS(Sales!C2:C5,Sales!B2:B5,"Dhaka",Sales!A2:A5,">=2026-01-01",Sales!A2:A5,"<2026-02-01") | 80 | =B2+C2 |
chart: line | title=Trend | x=Month | y=Dhaka,Khulna
""",
    "kpi_scorecard": """# KPIs
## Sheet: KPIs
| KPI | Target | Actual | Attainment |
|---|---|---|---|
| Revenue | 1000 | 940 | =C2/B2 |
| NPS | 50 | 58 | =C3/B3 |
| Churn | 5 | 7 | =B4/C4 |
format: Attainment=percent
highlight: Attainment<0.95 red
highlight: Attainment>=1 green
databar: Actual
""",
    "loan_schedule": """# Loan
## Sheet: Loan
| Item | Value |
|---|---|
| Principal | 100000 |
| Annual rate | 9% |
| Months | 12 |
| Payment | =-PMT(B3/12,B4,B2) |
| Total paid | =B5*B4 |
| Interest | =B6-B2 |
""",
    "attendance": """# Attendance
## Sheet: Attendance
| Name | Mon | Tue | Wed | Days |
|---|---|---|---|---|
| Rahim | P | P | A | =COUNTIF(B2:D2,"P") |
| Karim | P | A | A | =COUNTIF(B3:D3,"P") |
validate: Mon in [P, A, L]
validate: Tue in [P, A, L]
validate: Wed in [P, A, L]
""",
    "inventory": """# Inventory
## Sheet: Stock
| SKU | Item | On hand | Reorder at | Status |
|---|---|---|---|---|
| 001 | Cable | 40 | 50 | =IF(C2<D2,"Reorder","OK") |
| 002 | Adapter | 120 | 50 | =IF(C3<D3,"Reorder","OK") |
table: StockTbl
highlight: Status=Reorder amber
""",
    "multi_currency": """# FX
## Sheet: Rates
| Currency | BDT per unit |
|---|---|
| USD | 120 |
| EUR | 130 |
## Sheet: Invoices
| Invoice | Currency | Amount | Amount BDT |
|---|---|---|---|
| INV-1 | USD | 100 | =C2*VLOOKUP(B2,Rates!A2:B3,2,FALSE) |
| INV-2 | EUR | 50 | =C3*VLOOKUP(B3,Rates!A2:B3,2,FALSE) |
total: Amount BDT=sum
""",
    "cohort": """# Cohorts
## Sheet: Cohorts
| Cohort | M0 | M1 | M2 | Retention M2 |
|---|---|---|---|---|
| Jan | 100 | 60 | 40 | =D2/B2 |
| Feb | 120 | 70 | 50 | =D3/B3 |
format: Retention M2=percent
scale: Retention M2
""",
    "import_clean_export": """# Clean
## Sheet: Clean
| Name | Email | Joined |
|---|---|---|
| Ayesha Rahman | a@example.com | 2026-03-01 |
| Tanvir Hasan | t@example.com | 2026-03-02 |
freeze: A2
filter: on
note: Source: CRM export, duplicates removed
""",
}


class Benchmark(unittest.TestCase):
    def test_every_spec_builds_calculates_and_opens(self):
        for name, text in SPECS.items():
            with self.subTest(name):
                data, spec = D.create(f"{name}.xlsx", text)
                wb = openpyxl.load_workbook(io.BytesIO(data))
                self.assertTrue(wb.sheetnames)
                res = calc.calculate(data)
                self.assertEqual(res["errors"], [], name)
                self.assertEqual(res["circular"], 0, name)
                vals = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
                for ws in wb.worksheets:
                    for row in ws.iter_rows():
                        for c in row:
                            if c.data_type == "f":
                                cached = vals[ws.title][c.coordinate].value
                                self.assertIsNotNone(cached, f"{name}: {ws.title}!{c.coordinate} has no calculated value")

    def test_known_answers(self):
        def v(name, sheet, cell):
            data, _ = D.create("x.xlsx", SPECS[name])
            return openpyxl.load_workbook(io.BytesIO(data), data_only=True)[sheet][cell].value
        self.assertEqual(v("sales_summary", "Summary", "B2"), 10 * 25 + 7 * 40)
        self.assertEqual(v("budget", "Budget", "C2"), 1100)
        self.assertAlmostEqual(v("budget", "Budget", "C3"), 742)
        self.assertAlmostEqual(v("loan_schedule", "Loan", "B5"), 8745.15, places=2)
        self.assertAlmostEqual(v("kpi_scorecard", "KPIs", "D2"), 0.94)
        self.assertEqual(v("attendance", "Attendance", "E2"), 2)
        self.assertEqual(v("inventory", "Stock", "E2"), "Reorder")
        self.assertEqual(v("multi_currency", "Invoices", "D2"), 12000)
        self.assertEqual(v("multi_currency", "Invoices", "D4"), 12000 + 6500)
        self.assertAlmostEqual(v("cohort", "Cohorts", "E3"), 50 / 120)
        self.assertEqual(v("inventory", "Stock", "A2"), "001")        # identifiers keep their zeros

    def test_a_built_workbook_is_queryable_by_the_sql_tool(self):
        import tempfile
        from pathlib import Path
        from core.sheet_engine import sql
        data, _ = D.create("x.xlsx", SPECS["sales_summary"])
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "s.xlsx"
            p.write_bytes(data)
            sql.clear_cache()
            out = sql.query(sql.load(p), "SELECT Region, SUM(Units) u FROM Data WHERE Region IS NOT NULL GROUP BY Region ORDER BY u DESC")
            self.assertIn("| Dhaka | 17 |", out)


if __name__ == "__main__":
    unittest.main()
