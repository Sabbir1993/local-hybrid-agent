"""PDF toolbox: read, tables -> Excel, page operations. Real PDFs from reportlab; files live in a temp common space."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import core.agent_tools  # noqa: F401
import openpyxl
from pypdf import PdfReader

from core.pdf_tools import common, read
from tests.pdf_fixtures import scanned_pdf, table_pdf, text_pdf


class Base(unittest.TestCase):
    def setUp(self):
        from core import agent_tools, doc_store
        from core.request_context import set_current_user
        import tests.test_doc_ops as T
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._patches = [mock.patch.object(agent_tools, "COMMON_ROOT", self.root),
                         mock.patch.object(doc_store, "_projects_db", T.doc_files_db())]
        for p in self._patches:
            p.start()
        set_current_user(21)
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        from core.request_context import set_current_user
        for p in self._patches:
            p.stop()
        set_current_user(None)
        self.tmp.cleanup()

    def put(self, name, data):
        d = self.root / "user_21"
        d.mkdir(parents=True, exist_ok=True)
        (d / name).write_bytes(data)
        return name

    def out(self, link_text):
        name = link_text.split("[DOWNLOAD: ")[-1].split("]")[0]
        return (self.root / "user_21" / name).read_bytes()

    def tool(self, name, **args):
        from core.agent_tools import TOOL_IMPLS
        return TOOL_IMPLS[name](args)


class Helpers(unittest.TestCase):
    def test_page_ranges(self):
        self.assertEqual(common.parse_pages("1-3,7,last", 9), [1, 2, 3, 7, 9])
        self.assertEqual(common.parse_pages("odd", 5), [1, 3, 5])
        self.assertEqual(common.parse_pages("even", 5), [2, 4])
        self.assertEqual(common.parse_pages(None, 3), [1, 2, 3])
        self.assertEqual(common.parse_pages("3-1", 5), [1, 2, 3])
        for bad in ("0", "9", "a-b", "1,,x"):
            with self.assertRaises(common.PdfError):
                common.parse_pages(bad, 5)


class Reading(Base):
    def test_text_per_page_with_markers(self):
        self.put("doc.pdf", text_pdf(3))
        out = self.tool("pdf_read", file="doc.pdf", pages="2-3")
        self.assertIn("PDF: 3 pages", out)
        self.assertIn("--- page 2 ---", out)
        self.assertIn("body text of page 3", out)
        self.assertNotIn("page 1 heading", out)

    def test_long_documents_are_paged(self):
        self.put("big.pdf", text_pdf(60, extra={i: "x" * 300 for i in range(1, 61)}))
        out = self.tool("pdf_read", file="big.pdf")
        self.assertIn("more pages follow: call again with offset=", out)
        nxt = int(out.split("offset=")[1].split(")")[0])
        self.assertIn(f"--- page {nxt + 1} ---", self.tool("pdf_read", file="big.pdf", offset=nxt))

    def test_scanned_pages_are_flagged(self):
        self.put("scan.pdf", scanned_pdf())
        out = self.tool("pdf_read", file="scan.pdf")
        self.assertIn("no text layer", out)
        self.assertIn("pdf_ocr", out)

    def test_card_numbers_are_masked(self):
        self.put("c.pdf", text_pdf(1, extra={1: "Card 4111 1111 1111 1111 on file"}))
        self.assertNotIn("4111 1111 1111 1111", self.tool("pdf_read", file="c.pdf"))

    def test_errors_are_plain(self):
        self.assertTrue(self.tool("pdf_read", file="missing.pdf").startswith("error: no file"))
        self.assertTrue(self.tool("pdf_read").startswith("error:"))
        self.put("doc.pdf", text_pdf(2))
        self.assertTrue(self.tool("pdf_read", file="doc.pdf", pages="9").startswith("error: page 9"))


class Tables(Base):
    ROWS = [["Item", "Qty", "Price"], ["Cable", "4", "12.50"], ["Adapter", "10", "3.20"], ["Hub", "2", "40.00"]]

    def test_markdown_and_xlsx(self):
        self.put("t.pdf", table_pdf(self.ROWS))
        md = self.tool("pdf_tables", file="t.pdf")
        self.assertIn("Found 1 table(s)", md)
        self.assertIn("| Item | Qty | Price |", md)
        self.assertIn("| Adapter | 10 | 3.20 |", md)
        res = self.tool("pdf_tables", file="t.pdf", to="xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(self.out(res)))
        ws = wb.worksheets[0]
        self.assertEqual([c.value for c in ws[1]], ["Item", "Qty", "Price"])
        self.assertEqual(ws["B2"].value, 4)                 # real numbers in the workbook
        self.assertTrue(ws["A1"].font.bold)

    def test_no_tables_says_what_to_try(self):
        self.put("p.pdf", text_pdf(1))
        self.assertIn("No tables", self.tool("pdf_tables", file="p.pdf"))

    def test_the_helper_names_sheets_by_page_and_number(self):
        t = [{"page": 2, "index": 1, "rows": self.ROWS}, {"page": 2, "index": 2, "rows": self.ROWS}]
        outline = read.tables_to_outline(t)
        self.assertIn("## Sheet: p2-t1", outline)
        self.assertIn("## Sheet: p2-t2", outline)


class Pages(Base):
    def go(self, ops, name="doc.pdf", **extra):
        return self.tool("pdf_pages", file=name, ops=ops, **extra)

    def pdf(self, link):
        return PdfReader(io.BytesIO(self.out(link)))

    def test_extract_delete_reorder_rotate(self):
        self.put("doc.pdf", text_pdf(5))
        r = self.pdf(self.go([{"op": "extract", "pages": "2-4"}]))
        self.assertEqual(len(r.pages), 3)
        self.assertIn("Page 2 heading", r.pages[0].extract_text())
        r = self.pdf(self.go([{"op": "delete", "pages": "1,5"}]))
        self.assertEqual([p.extract_text().split("\n")[0] for p in r.pages], ["Page 2 heading", "Page 3 heading", "Page 4 heading"])
        r = self.pdf(self.go([{"op": "reorder", "order": "3,1,2,4,5"}]))
        self.assertIn("Page 3 heading", r.pages[0].extract_text())
        r = self.pdf(self.go([{"op": "rotate", "pages": "2", "degrees": 90}]))
        self.assertEqual(r.pages[1].get("/Rotate"), 90)
        self.assertFalse(r.pages[0].get("/Rotate"))

    def test_merge_two_files(self):
        self.put("a.pdf", text_pdf(2, prefix="A"))
        self.put("b.pdf", text_pdf(3, prefix="B"))
        r = self.pdf(self.go([{"op": "merge", "files": ["b.pdf"]}], name="a.pdf"))
        self.assertEqual(len(r.pages), 5)
        self.assertIn("B 1 heading", r.pages[2].extract_text())

    def test_split_makes_several_files(self):
        self.put("doc.pdf", text_pdf(6))
        res = self.go([{"op": "split", "every_n": 2}])
        self.assertEqual(res.count("[DOWNLOAD:"), 3)
        res2 = self.go([{"op": "split", "ranges": "1-2;5-6"}])
        self.assertEqual(res2.count("[DOWNLOAD:"), 2)

    def test_numbers_watermark_and_metadata(self):
        self.put("doc.pdf", text_pdf(2))
        res = self.go([{"op": "add_page_numbers"}, {"op": "watermark", "text": "DRAFT"}, {"op": "set_metadata", "title": "Board pack"}])
        r = self.pdf(res)
        txt = r.pages[1].extract_text()
        self.assertIn("Page 2 of 2", txt)
        self.assertIn("DRAFT", txt.replace("\n", ""))
        self.assertEqual(r.metadata.get("/Title"), "Board pack")

    def test_the_original_is_untouched_and_the_result_is_new(self):
        data = text_pdf(3)
        self.put("doc.pdf", data)
        res = self.go([{"op": "delete", "pages": "2"}])
        self.assertEqual((self.root / "user_21" / "doc.pdf").read_bytes(), data)
        self.assertNotIn("doc.pdf]", res)

    def test_bad_requests(self):
        self.put("doc.pdf", text_pdf(2))
        for ops, needle in (([{"op": "delete", "pages": "1-2"}], "every page"), ([{"op": "reorder", "order": "1"}], "exactly once"),
                            ([{"op": "rotate", "pages": "1", "degrees": 45}], "multiple of 90"), ([{"op": "nope"}], "unknown operation"),
                            ([{"op": "watermark", "text": "বাংলা"}], "Latin"), ([], "ops required"),
                            ([{"op": "merge", "files": ["zzz.pdf"]}], "no file named")):
            self.assertIn(needle, self.go(ops), ops)


class Encrypted(Base):
    def test_a_password_is_required_and_checked(self):
        from pypdf import PdfWriter
        w = PdfWriter()
        for p in PdfReader(io.BytesIO(text_pdf(2))).pages:
            w.add_page(p)
        w.encrypt("s3cret")
        buf = io.BytesIO()
        w.write(buf)
        self.put("locked.pdf", buf.getvalue())
        self.assertIn("password protected", self.tool("pdf_read", file="locked.pdf"))
        self.assertIn("does not open", self.tool("pdf_read", file="locked.pdf", password="nope"))
        self.assertIn("Page 2 heading", self.tool("pdf_read", file="locked.pdf", password="s3cret"))


class Wiring(unittest.TestCase):
    def test_when_the_tools_are_offered(self):
        from core import prompt_scope as ps
        pdfs = ["report-1234abcd.pdf"]
        self.assertTrue(ps.pdf_tools_wanted("remove page 3 and add a watermark", pdfs))
        self.assertTrue(ps.pdf_tools_wanted("what is on page 4?", pdfs))
        self.assertTrue(ps.pdf_tools_wanted("merge these pdf files", []))
        self.assertTrue(ps.pdf_tools_wanted("redact the card numbers", pdfs))
        self.assertFalse(ps.pdf_tools_wanted("make me a budget spreadsheet", []))
        self.assertFalse(ps.pdf_tools_wanted("what is a pdf?", []))
        self.assertFalse(ps.pdf_tools_wanted("summarise the attached sheet", ["data.xlsx"]))

    def test_the_route_and_surface_know_the_tools(self):
        from pathlib import Path
        root = Path(__file__).resolve().parents[1]
        src = (root / "routes" / "chat" / "run.py").read_text(encoding="utf-8")
        self.assertIn("prompt_scope.pdf_tools_wanted(own_q, prior_files)", src)
        self.assertIn("prompt_scope.PDF_TOOLS_MANUAL", src)
        surface = (root / "core" / "tool_surface.py").read_text(encoding="utf-8")
        for name in ("pdf_read", "pdf_tables", "pdf_pages", "pdf_form", "pdf_redact", "pdf_protect", "pdf_see", "pdf_ocr"):
            self.assertIn(f'"{name}"', surface, name)
        from core.agent_tools import AGENT_TOOLS, TOOL_IMPLS
        names = {t["function"]["name"] for t in AGENT_TOOLS}
        self.assertTrue({"pdf_read", "pdf_tables", "pdf_pages", "pdf_form", "pdf_redact", "pdf_protect", "pdf_see", "pdf_ocr"} <= names)
        self.assertTrue({"pdf_read", "pdf_ocr"} <= set(TOOL_IMPLS))


if __name__ == "__main__":
    unittest.main()
