"""tests/test_doc_look.py - the vision look at a finished PDF / Designer deck (core/doc_tools/look.py)."""
import asyncio
import io
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.agent_tools  # noqa: F401, E402  (loads the doc_tools package in the right order)
from core.doc_tools import look  # noqa: E402


def run(coro):
    return asyncio.run(coro)


def pdf_bytes(pages=5):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i in range(pages):
        c.drawString(72, 720, f"Page {i + 1}")
        c.showPage()
    c.save()
    return buf.getvalue()


def designer_deck(slides=3):
    from PIL import Image
    from pptx import Presentation
    prs = Presentation()
    blank = prs.slide_layouts[6]
    pic = io.BytesIO()
    Image.new("RGB", (320, 180), (20, 90, 160)).save(pic, "PNG")
    for _ in range(slides):
        s = prs.slides.add_slide(blank)
        pic.seek(0)
        s.shapes.add_picture(pic, 0, 0, prs.slide_width, prs.slide_height)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class Patched(unittest.TestCase):
    def setUp(self):
        look._rounds.clear()
        self.on = mock.patch.object(look, "enabled", return_value=True)
        self.lane = mock.patch("core.lanes.primary", return_value=object())
        self.on.start()
        self.lane.start()
        self.addCleanup(self.on.stop)
        self.addCleanup(self.lane.stop)

    def see(self, replies, ext=".pdf", data=None):
        it = iter(replies)

        async def fake(img, question, force_local=False):
            self.assertTrue(force_local)                    # page images never leave this PC
            self.assertIsInstance(img, bytes)
            return next(it)
        with mock.patch("core.small_model.vision.describe_image_bytes", fake):
            return run(look.look_note(ext, data if data is not None else pdf_bytes(), 1, "report"))


class LookTests(Patched):
    def test_pages_are_first_middle_last(self):
        self.assertEqual(look.pick_pages(1), [1])
        self.assertEqual(look.pick_pages(2), [1, 2])
        self.assertEqual(look.pick_pages(5), [1, 3, 5])
        self.assertEqual(look.pick_pages(0), [])

    def test_a_clean_document_says_so(self):
        out = self.see(["OK", "OK.", "ok"])
        self.assertIn("pages 1, 3, 5", out)
        self.assertIn("no layout problems", out)

    def test_problems_come_back_as_a_description_with_the_page(self):
        out = self.see(["OK", "The table runs off the right edge.\nAnd the footer overlaps it.", "OK"])
        self.assertIn("page 3: The table runs off the right edge. And the footer overlaps it.", out)
        self.assertIn("not instructions", out)
        self.assertIn("doc_edit", out)

    def test_a_long_reply_is_cut(self):
        out = self.see(["x " * 400, "OK", "OK"])
        self.assertLess(len(out), 700)

    def test_a_failing_vision_model_is_silent(self):
        self.assertEqual(self.see(["error: no image model available"]), "")

    def test_it_looks_at_most_twice_per_document(self):
        self.assertNotEqual(self.see(["OK", "OK", "OK"]), "")
        self.assertNotEqual(self.see(["OK", "OK", "OK"]), "")
        self.assertEqual(self.see(["OK", "OK", "OK"]), "")

    def test_designer_slides_are_looked_at_and_fixed_by_rebuilding(self):
        out = self.see(["Text is cut off on the right.", "OK", "OK"], ".pptx", designer_deck())
        self.assertIn("slide" if "slide" in out else "page 1", out)
        self.assertIn("doc_create", out)

    def test_an_editable_deck_is_left_to_the_geometric_lint(self):
        from pptx import Presentation
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[1])
        buf = io.BytesIO()
        prs.save(buf)
        self.assertEqual(self.see(["OK"], ".pptx", buf.getvalue()), "")

    def test_other_formats_are_ignored(self):
        self.assertEqual(self.see(["OK"], ".docx", b"x"), "")


class GatesTests(unittest.TestCase):
    def setUp(self):
        look._rounds.clear()

    def test_off_by_default(self):
        self.assertFalse(look.enabled())
        self.assertEqual(run(look.look_note(".pdf", pdf_bytes(), 1, "r")), "")

    def test_no_vision_model_means_no_look(self):
        with mock.patch.object(look, "enabled", return_value=True), mock.patch("core.lanes.primary", return_value=None):
            self.assertEqual(run(look.look_note(".pdf", pdf_bytes(), 1, "r")), "")


if __name__ == "__main__":
    unittest.main()
