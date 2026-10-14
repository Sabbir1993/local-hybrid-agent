"""Pictures in generated decks: the `image:` slide line, rendering, and the web -> image-model -> none resolver."""
import asyncio
import io
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.agent_tools  # noqa: F401  (import order: agent_tools first, as in the app)
from PIL import Image
from pptx import Presentation

from core.doc_ops import pptx_ops
from core.doc_ops.pptx_ops import outline as O
from core.doc_tools import slide_images as si


def png(w=1200, h=800, color=(30, 120, 200)) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (w, h), color).save(b, "PNG")
    return b.getvalue()


DECK = ("# Deck\nsubtitle: sub\nimage: city skyline | caption=Cover\n---\n"
        "## Why\nimage: data center racks | source=search\n- a\n- b\n---\n"
        "## Big\nlayout: full\nimage: ocean waves | source=generate\n---\n"
        "## Numbers\nimage: not here\n- 99% uptime\n- 5x faster\n---\n"
        "## Thank you\n")


class Outline(unittest.TestCase):
    def test_image_line_parsed_with_options(self):
        spec = O.parse_image_spec("city skyline at night | source=web | fit=contain | caption=Skyline")
        self.assertEqual(spec, {"query": "city skyline at night", "url": "", "source": "search",
                                "fit": "contain", "caption": "Skyline"})

    def test_url_and_empty(self):
        self.assertEqual(O.parse_image_spec("https://example.com/a.jpg")["url"], "https://example.com/a.jpg")
        self.assertIsNone(O.parse_image_spec("  | source=search"))

    def test_slide_kinds_with_pictures(self):
        deck = O.parse_outline(DECK)
        got = [(s["kind"], bool(s["image"])) for s in deck["slides"]]
        self.assertEqual(got[0], ("cover", True))
        self.assertEqual(got[1], ("bullets", True))
        self.assertEqual(got[2], ("full", True))

    def test_unknown_options_ignored(self):
        spec = O.parse_image_spec("boats | source=evil | fit=stretch | color=red")
        self.assertEqual((spec["source"], spec["fit"]), ("auto", "cover"))


class Render(unittest.TestCase):
    def _build(self, images):
        return Presentation(io.BytesIO(pptx_ops.create(DECK, images=images)))

    def test_pictures_embedded_on_supported_slides(self):
        imgs = {i: {"data": png(), "credit_host": "example.com", "credit_url": "https://example.com/p"} for i in (0, 1, 2)}
        prs = self._build(imgs)
        pics = [[sh.name for sh in sl.shapes if sh.shape_type == 13] for sl in prs.slides]
        self.assertEqual(pics[0], ["Cover image"])
        self.assertEqual(pics[1], ["Image"])
        self.assertEqual(pics[2], ["Background image"])
        self.assertEqual(pics[3], [])

    def test_credit_on_slide_and_source_in_notes(self):
        prs = self._build({1: {"data": png(), "credit_host": "example.com", "credit_url": "https://example.com/p"}})
        sl = prs.slides[1]
        self.assertTrue(any("example.com" in sh.text_frame.text for sh in sl.shapes if sh.has_text_frame))
        self.assertIn("https://example.com/p", sl.notes_slide.notes_text_frame.text)

    def test_no_pictures_is_a_normal_deck(self):
        prs = self._build({})
        self.assertEqual(sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13), 0)
        self.assertEqual(len(prs.slides), 5)

    def test_slide_without_image_line_ignores_stray_bytes(self):
        prs = self._build({3: {"data": png(), "credit_host": "x"}, 4: {"data": png(), "credit_host": "x"}})
        self.assertEqual(sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13), 0)

    def test_picture_is_cropped_to_its_box(self):
        blob, dx, dy, dw, dh = __import__("core.doc_ops.pptx_ops.design", fromlist=["x"])._fit_image(png(1200, 800), 5.0, 5.0)
        im = Image.open(io.BytesIO(blob))
        self.assertEqual(im.size[0], im.size[1])
        self.assertEqual((dx, dy, dw, dh), (0.0, 0.0, 5.0, 5.0))


def run(coro):
    return asyncio.run(coro)


class Resolver(unittest.TestCase):
    def setUp(self):
        si._cache.clear()
        self.audit = mock.patch.object(si.doc_store, "audit").start()
        self.addCleanup(mock.patch.stopall)

    def _resolve(self, md=DECK):
        return run(si.resolve(md, 7))

    def test_web_first_then_generate_then_none(self):
        web = mock.AsyncMock(side_effect=lambda uid, q, why, *a: {"data": b"w", "credit_host": "h", "credit_url": "u", "via": "web"}
                             if "racks" in q else (why.append("no usable picture came back from the search") or None))
        gen = mock.AsyncMock(side_effect=lambda uid, q, asp, why, *a: {"data": b"g", "credit_host": "AI-generated",
                                                                "credit_url": "", "via": "generated"}
                             if "skyline" in q else (why.append("the image model is not loaded") or None))
        with mock.patch.object(si, "_from_web", web), mock.patch.object(si, "_generate", gen):
            images, notes, used_web = self._resolve()
        self.assertEqual(images[1]["via"], "web")                   # source=search, found on the web
        self.assertEqual(images[0]["via"], "generated")             # auto: web had nothing, image model did
        self.assertNotIn(2, images)                                 # source=generate, model not loaded
        self.assertTrue(any("slide 3" in n and "not loaded" in n for n in notes))
        self.assertTrue(any("slide 4" in n and "no room" in n for n in notes))   # numbers slide: unsupported kind
        self.assertTrue(used_web)

    def test_source_generate_never_searches_the_web(self):
        web, gen = mock.AsyncMock(return_value=None), mock.AsyncMock(return_value=None)
        with mock.patch.object(si, "_from_web", web), mock.patch.object(si, "_generate", gen):
            self._resolve("# D\n---\n## A\nimage: boats | source=generate\n- x\n")
        web.assert_not_called()
        gen.assert_called_once()

    def test_cache_means_no_second_fetch_when_rebuilding(self):
        res = {"data": b"w", "credit_host": "h", "credit_url": "u", "via": "web"}
        web = mock.AsyncMock(return_value=res)
        md = "# D\n---\n## A\nimage: boats\n- x\n"
        with mock.patch.object(si, "_from_web", web):
            self._resolve(md)
            self._resolve(md + "---\n## B\n- y\n")
        self.assertEqual(web.call_count, 1)

    def test_at_most_eight_pictures(self):
        md = "# D\n" + "".join(f"---\n## S{i}\nimage: topic {i}\n- x\n- y\n" for i in range(12))
        res = {"data": b"w", "credit_host": "h", "credit_url": "u", "via": "web"}
        with mock.patch.object(si, "_from_web", mock.AsyncMock(return_value=res)):
            images, notes, _ = self._resolve(md)
        self.assertEqual(len(images), si.MAX_IMAGES)
        self.assertTrue(any("at most" in n for n in notes))

    def test_a_crashing_slide_does_not_stop_the_deck(self):
        calls = {"n": 0}

        async def flaky(uid, q, why, *a):
            calls["n"] += 1
            if calls["n"] == 1:
                raise RuntimeError("boom")
            return {"data": b"w", "credit_host": "h", "credit_url": "u", "via": "web"}
        md = "# D\n---\n## A\nimage: one\n- x\n---\n## B\nimage: two\n- y\n"
        with mock.patch.object(si, "_from_web", flaky):
            images, notes, _ = self._resolve(md)
        self.assertEqual(len(images), 1)
        self.assertEqual(len(notes), 1)


class WebGate(unittest.TestCase):
    def test_off_when_the_web_switch_is_off(self):
        si.allow_web.set(False)
        why = []
        self.assertIsNone(run(si._from_web(7, "boats", why)))
        self.assertIn("web search is off", why[-1])

    def test_card_number_in_the_words_is_refused(self):
        si.allow_web.set(True)
        why = []
        with mock.patch("core.web_tools._pan_egress", return_value=True):
            self.assertIsNone(run(si._from_web(7, "card 4111 1111 1111 1111", why)))
        self.assertIn("card number", why[-1])

    def test_downloads_only_through_the_proxy_and_credits_the_source(self):
        si.allow_web.set(True)
        results = [{"image_url": "http://insecure/x.jpg", "source_url": "http://a", "title": "boats"},
                   {"image_url": "https://img.example.org/p.jpg", "source_url": "https://www.example.org/page",
                    "title": "Boats at the harbour"}]
        fetched = []

        def fake_fetch(uid, url):
            fetched.append(url)
            return "image/png", png(1200, 800)
        with mock.patch("core.web_tools._bing_images_search_sync", return_value=results), \
                mock.patch("core.image_proxy.fetch", side_effect=fake_fetch), \
                mock.patch("core.image_proxy.remember") as rem:
            res = run(si._from_web(7, "boats", []))
        self.assertEqual(fetched, ["https://img.example.org/p.jpg"])    # http result skipped
        rem.assert_called_once()
        self.assertEqual(res["credit_host"], "example.org")
        self.assertEqual(res["via"], "web")
        self.assertLessEqual(len(res["data"]), si.MAX_IMAGE_BYTES)

    def test_tiny_pictures_are_skipped(self):
        si.allow_web.set(True)
        with mock.patch("core.web_tools._bing_images_search_sync",
                        return_value=[{"image_url": "https://i.example.org/t.png", "source_url": "https://e.org", "title": "boats"}]), \
                mock.patch("core.image_proxy.fetch", return_value=("image/png", png(120, 90))), \
                mock.patch("core.image_proxy.remember"):
            why = []
            self.assertIsNone(run(si._from_web(7, "boats", why)))
        self.assertIn("no usable picture", why[-1])


class Relevance(unittest.TestCase):
    def test_anchors_come_from_the_deck_subject(self):
        self.assertEqual(si.anchors_for("Welcome to Dhaka"), ["dhaka"])
        self.assertEqual(si.anchors_for("Cybersecurity Risk & Resilience: Executive Board Briefing"),
                         ["cybersecurity", "risk", "resilience"])

    def test_query_gets_the_missing_subject_words(self):
        self.assertEqual(si.search_query("river banks", ["dhaka"]), "river banks dhaka")
        self.assertEqual(si.search_query("Lalbagh Fort Dhaka", ["dhaka"]), "Lalbagh Fort Dhaka")

    def test_off_topic_and_thumbnail_hosts_are_dropped(self):
        results = [
            {"title": "Aloo tulte gelam", "image_url": "https://i.ytimg.com/vi/x/hq.jpg", "source_url": "https://youtube.com/w"},
            {"title": "Ahsan calligraphy digital download dhaka", "image_url": "https://i.etsystatic.com/a.jpg",
             "source_url": "https://etsy.com/l"},
            {"title": "Man in front of a house", "image_url": "https://actu.fr/a.jpg", "source_url": "https://actu.fr/p"},
            {"title": "Lalbagh Fort, Dhaka", "image_url": "https://upload.wikimedia.org/lalbagh.jpg",
             "source_url": "https://en.wikipedia.org/wiki/Lalbagh_Fort"},
        ]
        good = si.rank_candidates(results, "Lalbagh Fort Dhaka", ["dhaka"])
        self.assertEqual([r["title"] for r in good], ["Lalbagh Fort, Dhaka"])

    def test_same_words_for_another_place_are_not_enough(self):
        results = [{"title": "Red Fort, Delhi", "image_url": "https://x.org/redfort.jpg", "source_url": "https://x.org/p"}]
        self.assertEqual(si.rank_candidates(results, "Lalbagh Fort Dhaka", ["dhaka"]), [])

    def test_best_match_first_and_nothing_good_means_none(self):
        results = [{"title": "Dhaka street", "image_url": "https://a.org/1.jpg", "source_url": "https://a.org"},
                   {"title": "Buriganga river Dhaka boats", "image_url": "https://b.org/2.jpg", "source_url": "https://b.org"}]
        good = si.rank_candidates(results, "Buriganga river Dhaka", ["dhaka"])
        self.assertEqual(good[0]["title"], "Buriganga river Dhaka boats")

    def test_web_step_reports_no_match_instead_of_using_a_random_hit(self):
        si.allow_web.set(True)
        junk = [{"title": "Man outside a house", "image_url": "https://actu.fr/a.jpg", "source_url": "https://actu.fr/p"}]
        with mock.patch("core.web_tools._bing_images_search_sync", return_value=junk), \
                mock.patch("core.image_proxy.fetch") as fetch, mock.patch("core.image_proxy.remember"):
            why = []
            self.assertIsNone(run(si._from_web(7, "Lalbagh Fort", why, ("dhaka",))))
        fetch.assert_not_called()
        self.assertIn("clearly matched", why[-1])


class GenerateGate(unittest.TestCase):
    def _gen(self, status):
        why = []
        with mock.patch("core.media.status", return_value={"image": status}):
            res = run(si._generate(7, "boats", "square", why))
        return res, why

    def test_needs_a_local_loaded_image_model(self):
        self.assertIn("no image model", self._gen({"ready": False})[1][-1])
        self.assertIn("cloud", self._gen({"ready": True, "cloud": True, "state": None})[1][-1])
        self.assertIn("not loaded", self._gen({"ready": True, "cloud": False, "state": "not_loaded"})[1][-1])
        self.assertIn("not loaded", self._gen({"ready": True, "cloud": False, "state": "loading"})[1][-1])

    def test_never_loads_the_model_itself(self):
        with mock.patch("core.media.status", return_value={"image": {"ready": True, "cloud": False, "state": "not_loaded"}}), \
                mock.patch("core.media.generate", new=mock.AsyncMock()) as gen:
            run(si._generate(7, "boats", "square", []))
        gen.assert_not_called()


class MainLoop(unittest.TestCase):
    def test_generation_runs_on_the_servers_main_loop(self):
        import threading
        seen = {}

        async def work():
            seen["thread"] = threading.current_thread().name
            return 42

        async def main():
            si.main_loop.set(asyncio.get_running_loop())
            # what the chat route does: a worker thread with its own loop asks for the work
            import contextvars
            ctx = contextvars.copy_context()           # run_in_executor_ctx does the same in the app
            return await asyncio.get_running_loop().run_in_executor(
                None, ctx.run, lambda: asyncio.run(si._on_main_loop(work())))
        self.assertEqual(asyncio.run(main()), 42)
        self.assertEqual(seen["thread"], threading.main_thread().name)

    def test_without_a_main_loop_it_just_awaits(self):
        async def work():
            return 7
        si.main_loop.set(None)
        self.assertEqual(run(si._on_main_loop(work())), 7)


class Summary(unittest.TestCase):
    def test_copyright_note_only_for_web_pictures(self):
        web = {"via": "web", "data": b""}
        gen = {"via": "generated", "data": b""}
        self.assertIn("licence", si.summary({0: web}, [], True))
        self.assertNotIn("licence", si.summary({0: gen}, [], False))
        self.assertIn("slide 2", si.summary({}, ["slide 2: no picture - web search is off for this chat"], False))


class ToolFlow(unittest.TestCase):
    """doc_create / doc_append (chat mode, common space): pictures resolved, embedded, and kept on rebuild."""

    def setUp(self):
        import tempfile
        from core import agent_tools, doc_store, doc_tools
        from core.request_context import set_current_user
        self.doc_tools, self.agent_tools = doc_tools, agent_tools
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        import tests.test_doc_ops as T
        self._patches = [mock.patch.object(agent_tools, "COMMON_ROOT", self.root),
                         mock.patch.object(doc_store, "_projects_db", T.doc_files_db()),
                         mock.patch.object(si.doc_store, "audit")]
        for p in self._patches:
            p.start()
        set_current_user(9)
        si._cache.clear()
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        from core.request_context import set_current_user
        for p in self._patches:
            p.stop()
        set_current_user(None)
        self.tmp.cleanup()

    def _picture_count(self, name):
        prs = Presentation(io.BytesIO((self.root / "user_9" / name).read_bytes()))
        return sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13)

    def test_create_embeds_pictures_and_append_reuses_them(self):
        res = {"data": png(), "credit_host": "example.org", "credit_url": "https://example.org/p", "via": "web"}
        web = mock.AsyncMock(return_value=res)
        first = "# Deck\nsubtitle: s\n---\n## Why\nimage: boats\n- a\n- b\n"
        with mock.patch.object(si, "_from_web", web):
            out = self.doc_tools.tool_doc_create_common({"file": "d.pptx", "content": first, "total": 3})
            self.assertIn("Pictures: 1 added (1 from the web)", out)
            self.assertIn("licence", out)
            name = out.split("[DOWNLOAD: ")[1].split("]")[0]
            self.assertEqual(self._picture_count(name), 1)
            more = self.doc_tools.tool_doc_append_common(
                {"file": name, "content": "## Next\nimage: ships\n- c\n- d\n", "done": False})
            self.assertEqual(self._picture_count(name), 2)
            self.assertEqual(web.call_count, 2)            # "boats" came from the cache on the rebuild
            self.assertIn("Pictures:", more)

    def test_no_image_lines_means_no_resolver_call(self):
        with mock.patch.object(si, "resolve", new=mock.AsyncMock()) as r:
            self.doc_tools.tool_doc_create_common({"file": "p.pptx", "content": "# D\n---\n## A\n- x\n- y\n"})
        r.assert_not_called()


if __name__ == "__main__":
    unittest.main()
