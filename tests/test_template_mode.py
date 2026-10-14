"""Deck built on a user template keeps the template's look but now gets real tables and pictures."""
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PIL import Image
from pptx import Presentation

from core.doc_ops import pptx_ops

TEMPLATE_MD = "# Template\nsubtitle: base\ntheme: corporate\n---\n## One\n- a\n- b\n- c\n"


def template_bytes() -> bytes:
    return pptx_ops.create(TEMPLATE_MD)


def png() -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (1200, 800), (30, 120, 200)).save(b, "PNG")
    return b.getvalue()


def build(md, images=None):
    return Presentation(io.BytesIO(pptx_ops.create(md, template=template_bytes(), images=images)))


class TemplateTables(unittest.TestCase):
    MD = ("# Deck\nsubtitle: s\n---\n## Data\n| City | People | Area |\n|---|---|---|\n| Dhaka | 22M | 306 |\n"
          "| Chittagong | 5M | 168 |\n| Khulna | 1M | 59 |\n")

    def test_a_table_is_a_real_table_not_pipe_text(self):
        prs = build(self.MD)
        tables = [sh for sl in prs.slides for sh in sl.shapes if sh.has_table]
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].table.cell(0, 0).text_frame.text, "City")
        self.assertEqual(len(tables[0].table.rows), 4)
        text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
        self.assertNotIn(" | ", text)

    def test_the_table_stays_inside_the_slide(self):
        prs = build(self.MD)
        tb = next(sh for sl in prs.slides for sh in sl.shapes if sh.has_table)
        bottom = (tb.top + sum(r.height for r in tb.table.rows)) / 914400
        self.assertLessEqual(bottom, prs.slide_height / 914400)

    def test_slides_without_tables_are_unchanged(self):
        prs = build("# Deck\nsubtitle: s\n---\n## Plain\n- one\n- two\n- three\n")
        self.assertEqual(sum(1 for sl in prs.slides for sh in sl.shapes if sh.has_table), 0)


class TemplatePictures(unittest.TestCase):
    MD = "# Deck\nsubtitle: s\n---\n## Why\nimage: boats | caption=Harbour\n- one\n- two\n- three\n"

    def test_picture_goes_on_the_right_and_text_keeps_the_left(self):
        info = {"data": png(), "credit_host": "example.org", "credit_url": "https://example.org/p", "via": "web"}
        prs = build(self.MD, images={1: info})
        sl = prs.slides[1]
        pics = [sh for sh in sl.shapes if sh.shape_type == 13]
        self.assertEqual(len(pics), 1)
        body = next(sh for sh in sl.placeholders if sh.placeholder_format.idx != 0)
        self.assertLessEqual(body.left + body.width, pics[0].left)
        self.assertIn("https://example.org/p", sl.notes_slide.notes_text_frame.text)
        self.assertTrue(any("example.org" in sh.text_frame.text for sh in sl.shapes if sh.has_text_frame))

    def test_no_picture_no_change(self):
        prs = build(self.MD)
        self.assertEqual(sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13), 0)


if __name__ == "__main__":
    unittest.main()
