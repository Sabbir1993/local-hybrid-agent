"""data_clean: operations over one sheet, written to a new workbook; the original is untouched."""
import datetime
import io
import tempfile
import unittest
from pathlib import Path

import core.agent_tools  # noqa: F401
import openpyxl

from core.sheet_engine import clean


def messy(d: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws["A1"] = "Export from CRM"
    for i, h in enumerate(["Name", "Email", "Amount", "Joined", "City"], 1):
        ws.cell(3, i, h)
    rows = [("  Ayesha   Rahman ", "A@Example.com", "1,200", "05/03/2026", "Dhaka"),
            ("Tanvir Hasan", "t@example.com", "300", "2026-03-06", None),
            ("Ayesha Rahman", "a@example.com", "1,200", "05/03/2026", "Dhaka"),
            (None, None, None, None, None),
            ("Nusrat Jahan", "n@example.com", "(50)", "07/03/2026", "Khulna")]
    for r, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    p = d / "people.xlsx"
    wb.save(p)
    return p


class Clean(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = messy(Path(self.tmp.name))

    def run_ops(self, ops):
        data, report = clean.run(self.path, ops)
        return openpyxl.load_workbook(io.BytesIO(data)), report

    def test_a_full_cleaning_pass(self):
        wb, report = self.run_ops([{"op": "trim"}, {"op": "drop_empty_rows"}, {"op": "case", "column": "Email", "to": "lower"},
                                   {"op": "dedupe", "columns": ["Email"]}, {"op": "fix_types", "columns": ["Amount", "Joined"]},
                                   {"op": "fill", "column": "City", "value": "Unknown"}])
        ws = wb["People"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        self.assertEqual(len(rows), 3)                                   # the duplicate was removed
        self.assertEqual(rows[0][:3], ["Ayesha Rahman", "a@example.com", 1200])
        self.assertEqual(rows[0][3], datetime.datetime(2026, 3, 5))       # day-first dates became real dates
        self.assertEqual(rows[1][4], "Unknown")
        self.assertEqual(rows[2][2], -50)                                  # (50) is a negative number
        self.assertIn("4 rows -> 3", report)          # the loader already drops fully blank rows

    def test_changes_sheet_documents_the_steps(self):
        wb, _ = self.run_ops([{"op": "trim"}, {"op": "dedupe"}])
        ch = [[c.value for c in r] for r in wb["Changes"].iter_rows()]
        self.assertEqual(ch[1][0], "rows before")
        self.assertEqual(ch[-1][0], "rows after")
        self.assertTrue(any(r[0] == "trim" for r in ch))

    def test_split_rename_replace_drop(self):
        wb, _ = self.run_ops([{"op": "trim"}, {"op": "split", "column": "Name", "delimiter": " ", "into": ["First", "Last"]},
                              {"op": "rename", "from": "City", "to": "Town"}, {"op": "replace", "column": "Town", "find": "Dhaka", "with": "Dhaka City"},
                              {"op": "drop_columns", "columns": ["Joined"]}])
        ws = wb["People"]
        self.assertEqual([c.value for c in ws[1]], ["Name", "Email", "Amount", "Town", "First", "Last"])
        self.assertEqual(ws["E2"].value, "Ayesha")
        self.assertEqual(ws["D2"].value, "Dhaka City")

    def test_the_original_file_is_untouched(self):
        before = self.path.read_bytes()
        self.run_ops([{"op": "dedupe"}])
        self.assertEqual(self.path.read_bytes(), before)

    def test_bad_requests_get_a_reason(self):
        for ops, needle in (([], "ops required"), ([{"op": "nope"}], "unknown operation"),
                            ([{"op": "case", "column": "Zzz", "to": "upper"}], "not found"),
                            ([{"op": "case", "column": "Name", "to": "sideways"}], "upper, lower or title"),
                            ([{"op": "trim"}] * 25, "at most")):
            with self.assertRaises(clean.CleanError) as cm:
                clean.run(self.path, ops)
            self.assertIn(needle, str(cm.exception))

    def test_the_tool_is_registered(self):
        from core.agent_tools import AGENT_TOOLS, TOOL_IMPLS
        self.assertIn("data_clean", TOOL_IMPLS)
        self.assertTrue(any(t["function"]["name"] == "data_clean" for t in AGENT_TOOLS))
        self.assertTrue(TOOL_IMPLS["data_clean"]({"file": "missing.xlsx", "ops": [{"op": "trim"}]}).startswith("error:"))


if __name__ == "__main__":
    unittest.main()
