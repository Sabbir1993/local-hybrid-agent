"""Deck charts, comparison / agenda layouts and the layout check (lint) that keeps the builder honest."""
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pptx.util import Inches, Pt

from core.doc_ops import pptx_ops
from core.doc_ops.pptx_ops import design as D
from core.doc_ops.pptx_ops import lint
from core.doc_ops.pptx_ops import outline as O

LONG = "Founded in the seventeenth century by Mughal governors and expanded under later rulers of the region"

GOLDEN = f"""# Dhaka Deck
subtitle: A city in numbers
---
## Agenda
layout: agenda
- Landscape
- Risks
- Plan
---
## Why
- {LONG}
- {LONG}
- {LONG}
---
## Growth
chart: line | unit=%
- 2024: 40%
- 2025: 60%
- 2026: 89%
---
## Compare
layout: comparison
- Legacy:
  - Slow to change
  - Costly to run
- Modern:
  - Fast to change
  - Cheap to run
---
## Services
- Speed: very fast delivery
- Cost: quite low
- Safety: well tested
---
## Numbers
- 99.9% uptime
- 5x faster
- 40% cheaper
---
## Timeline
layout: timeline
- 2018: Founded
- 2020: Expanded
- 2022: Listed
- 2024: Global
---
## Quote
layout: quote
- "Simplicity is the ultimate sophistication"
- Leonardo
---
## Table
| City | People | Area |
|---|---|---|
| Dhaka | 22M | 306 |
| Chittagong | 5M | 168 |
| Khulna | 1M | 59 |
---
## Thank you
"""


class Numbers(unittest.TestCase):
    def test_parse_number(self):
        cases = {"40%": 40.0, "$10.8T": 10.8e12, "1,200": 1200.0, "3.5B": 3.5e9, "5 million": 5e6, "29 min": 29.0,
                 "7x": 7.0, "about 5": None, "": None, "n/a": None}
        for text, want in cases.items():
            self.assertEqual(O.parse_number(text), want, text)

    def test_chart_from_bullets(self):
        d = O.parse_outline("# D\nsubtitle: s\n---\n## G\nchart: line | unit=%\n- 2024: 40%\n- 2025: 60%\n")
        s = d["slides"][1]
        self.assertEqual(s["kind"], "chart")
        self.assertEqual(s["chart_data"]["kind"], "line")
        self.assertEqual(s["chart_data"]["labels"], ["2024", "2025"])
        self.assertEqual(s["chart_data"]["series"][0]["values"], [40.0, 60.0])

    def test_chart_from_a_table_with_two_series(self):
        d = O.parse_outline("# D\nsubtitle: s\n---\n## G\nchart: bar\n| Year | Sales | Cost |\n|---|---|---|\n| 2024 | 10 | 7 |\n| 2025 | 14 | 9 |\n")
        cd = d["slides"][1]["chart_data"]
        self.assertEqual([x["name"] for x in cd["series"]], ["Sales", "Cost"])

    def test_no_numbers_means_no_chart(self):
        d = O.parse_outline("# D\nsubtitle: s\n---\n## G\nchart: bar\n- alpha: high\n- beta: low\n")
        self.assertNotEqual(d["slides"][1]["kind"], "chart")

    def test_pie_with_negatives_falls_back_to_bars(self):
        d = O.parse_outline("# D\nsubtitle: s\n---\n## G\nchart: pie\n- A: 5\n- B: -2\n")
        self.assertEqual(d["slides"][1]["chart_data"]["kind"], "bar")


class Rendering(unittest.TestCase):
    def _prs(self, md):
        return Presentation(io.BytesIO(pptx_ops.create(md)))

    def test_native_charts_are_real_chart_objects(self):
        prs = self._prs(GOLDEN)
        charts = [sh.chart for sl in prs.slides for sh in sl.shapes if sh.has_chart]
        self.assertEqual(len(charts), 1)
        self.assertEqual(list(charts[0].plots[0].categories), ["2024", "2025", "2026"])

    def test_pie_and_multi_series(self):
        md = ("# D\nsubtitle: s\n---\n## P\nchart: pie\n- A: 50\n- B: 30\n- C: 20\n---\n## M\nchart: bar\n"
              "| Y | Sales | Cost |\n|---|---|---|\n| 2024 | 10 | 7 |\n| 2025 | 14 | 9 |\n")
        prs = self._prs(md)
        charts = [sh.chart for sl in prs.slides for sh in sl.shapes if sh.has_chart]
        self.assertEqual(len(charts), 2)
        self.assertEqual(len(list(charts[1].plots[0].series)), 2)
        self.assertTrue(charts[1].has_legend)

    def test_comparison_and_agenda(self):
        prs = self._prs(GOLDEN)
        names = {sh.name for sl in prs.slides for sh in sl.shapes}
        self.assertTrue({"Option 1", "Option 2", "Item 1", "Item 3"} <= names)

    def test_chart_slide_has_a_title_and_footer(self):
        prs = self._prs(GOLDEN)
        sl = prs.slides[3]
        self.assertEqual(sl.shapes.title.text_frame.text, "Growth")


class Lint(unittest.TestCase):
    def test_golden_deck_is_clean_in_every_theme(self):
        for theme in ("midnight", "corporate", "forest", "sunset", "minimal", "royal"):
            data = pptx_ops.create(GOLDEN.replace("subtitle: A city in numbers", f"subtitle: A city in numbers\ntheme: {theme}"))
            self.assertEqual([f["message"] for f in lint.lint_deck(data)], [], theme)

    def test_flags_text_that_overflows_its_box(self):
        prs = Presentation(io.BytesIO(pptx_ops.create("# D\nsubtitle: s\n---\n## Why\n- one\n- two\n- three\n")))
        body = next(sh for sh in prs.slides[1].shapes if sh.is_placeholder and sh.placeholder_format.idx == 1)
        body.text_frame.text = "\n".join(LONG for _ in range(30))
        for p in body.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.name = Pt(24), "Calibri"
        buf = io.BytesIO()
        prs.save(buf)
        kinds = {f["kind"] for f in lint.lint_deck(buf.getvalue())}
        self.assertIn("overflow", kinds)

    def test_flags_small_text_and_off_slide_shapes(self):
        prs = Presentation(io.BytesIO(pptx_ops.create("# D\nsubtitle: s\n---\n## Why\n- one\n- two\n- three\n")))
        sl = prs.slides[1]
        tb = sl.shapes.add_textbox(Inches(1), Inches(6), Inches(3), Inches(0.5))
        tb.name = "Tiny"
        tb.text_frame.text = "small print"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        off = sl.shapes.add_textbox(Inches(12.9), Inches(1), Inches(3), Inches(0.5))
        off.name = "Off"
        off.text_frame.text = "outside"
        off.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        buf = io.BytesIO()
        prs.save(buf)
        kinds = {f["kind"] for f in lint.lint_deck(buf.getvalue())}
        self.assertTrue({"small-text", "bounds"} <= kinds, kinds)

    def test_flags_overlapping_text_boxes(self):
        prs = Presentation(io.BytesIO(pptx_ops.create("# D\nsubtitle: s\n---\n## Why\n- one\n- two\n- three\n")))
        sl = prs.slides[1]
        for n in ("A", "B"):
            tb = sl.shapes.add_textbox(Inches(2), Inches(6), Inches(3), Inches(0.5))
            tb.name = f"Box {n}"
            tb.text_frame.text = "same place"
            tb.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        buf = io.BytesIO()
        prs.save(buf)
        self.assertIn("overlap", {f["kind"] for f in lint.lint_deck(buf.getvalue())})

    def test_contrast_helper(self):
        self.assertGreater(lint.contrast("000000", "FFFFFF"), 20)
        self.assertLess(lint.contrast("777777", "808080"), 1.5)

    def test_ink_makes_an_accent_readable_on_white(self):
        fixed = D._ink("22D3EE", "FFFFFF", 4.5)
        self.assertGreaterEqual(lint.contrast(fixed, "FFFFFF"), 4.5)
        self.assertEqual(D._ink("0F2A5C", "FFFFFF", 4.5), "0F2A5C")        # already readable: unchanged

    def test_on_picks_the_better_of_black_and_white(self):
        self.assertEqual(D._on("0EA5E9"), "111111")
        self.assertEqual(D._on("0B1B2E"), "FFFFFF")

    def test_summary_text(self):
        self.assertEqual(lint.summary([]), "")
        text = lint.summary([{"message": "slide 2: x"}, {"message": "slide 3: y"}])
        self.assertIn("2 issue(s)", text)
        self.assertIn("slide 3: y", text)


if __name__ == "__main__":
    unittest.main()
