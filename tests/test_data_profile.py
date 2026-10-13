"""Attached spreadsheets: compact profile for the model's context, and the read-only data_query tool behind it."""
import asyncio
import datetime
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from core import agent_tools  # noqa: E402
from core.file_tools import extract_attachment_preview, extract_file_content  # noqa: E402
from core.file_tools.data_query import QueryError, run_data_query  # noqa: E402
from core.file_tools.profile import column_stats, parse_number, profile_file, read_tables  # noqa: E402
from core.request_context import set_current_user  # noqa: E402


def _workbook(path: Path, rows: int = 300) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pending"
    ws.append(["Billing Doc", "Customer", "Amount", "Status", "Date", "Note"])
    for i in range(1, rows + 1):
        ws.append([90000 + i, f"Cust {i % 7}", float(i), ["Pending", "Cleared", "Disputed"][i % 3],
                   datetime.date(2026, 9, 1 + i % 28), "x" if i % 4 == 0 else None])
    ws2 = wb.create_sheet("Small")
    ws2.append(["a", "b"])
    ws2.append([1, "2"])
    ws2.append(["(300)", "1,200.50"])
    wb.create_sheet("Empty")
    wb.save(path)
    return path


class ProfileTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.xlsx = _workbook(self.dir / "pending.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def test_profile_describes_the_whole_workbook_in_a_small_text(self):
        text, complete = profile_file(self.xlsx)
        self.assertLess(len(text), 4100)
        self.assertIn('Sheet "Pending" - 300 data rows x 6 columns', text)
        self.assertIn("Amount: number", text)
        self.assertIn("sum 45,150", text)                       # 1+2+...+300
        self.assertIn("Billing Doc: number (identifier)", text)  # all-different whole numbers: no sum / mean
        self.assertIn("Status: text; 3 values", text)
        self.assertIn("225 empty", text)                        # Note is filled on every 4th row
        self.assertIn("Sheet \"Small\"", text)
        self.assertNotIn("Empty", text.split("Small")[1][:5], "an empty sheet is not listed")
        self.assertFalse(complete, "rows were left out, the model is told")

    def test_numbers_in_text_cells_are_numbers(self):
        self.assertEqual(parse_number("1,200.50"), 1200.5)
        self.assertEqual(parse_number("(300)"), -300.0)
        self.assertEqual(parse_number("12%"), 12.0)
        self.assertIsNone(parse_number("abc"))
        self.assertIsNone(parse_number(True))
        st = column_stats(["1", "2,000", None, "x"])
        self.assertEqual((st["kind"], st["empty"]), ("number", 1))

    def test_csv_and_odd_encodings(self):
        f = self.dir / "d.csv"
        f.write_text("name;qty\nbolt;5\nnut;7\n", encoding="utf-8")
        t = read_tables(f)[0]
        self.assertEqual((t.headers, len(t.rows)), (["name", "qty"], 2))
        g = self.dir / "e.csv"
        g.write_bytes("city,amount\nDhaka,10\nKhulna,20\n".encode("latin-1"))
        self.assertEqual(len(read_tables(g)[0].rows), 2)

    def test_attachment_preview_is_the_profile_for_sheets_and_unchanged_for_text(self):
        text, truncated = extract_attachment_preview(self.xlsx)
        self.assertTrue(text.startswith("Data file: pending.xlsx"))
        self.assertTrue(truncated)
        raw, _ = extract_file_content(self.xlsx)
        self.assertIn("## Sheet: Pending", raw)                  # the raw extraction (read_file_chunk) is untouched
        note = self.dir / "n.txt"
        note.write_text("hello", encoding="utf-8")
        self.assertEqual(extract_attachment_preview(note), ("hello", False))


class DataQueryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.xlsx = _workbook(self.dir / "pending.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def q(self, **a):
        return run_data_query(self.xlsx, a)

    def test_counts_filters_and_aggregates_are_exact(self):
        self.assertIn("100 of 300 rows match", self.q(filters=[{"column": "Status", "op": "=", "value": "Pending"}], count=True))
        out = self.q(filters=[{"column": "Amount", "op": ">", "value": 295}], agg=[{"fn": "sum", "column": "Amount"}])
        self.assertIn("5 of 300 rows match", out)
        self.assertIn("1,490", out)                                # 296+297+298+299+300
        out = self.q(group_by="Status", agg=[{"fn": "count"}, {"fn": "avg", "column": "Amount"}])
        self.assertIn("3 groups by Status", out)
        self.assertIn("| Pending | 100 |", out)
        self.assertIn("avg(Amount)", out)

    def test_listing_sorting_distinct_and_describe(self):
        out = self.q(select=["Billing Doc", "Amount"], sort_by="Amount", desc=True, limit=3)
        self.assertIn("| 90300 | 300 |", out)
        self.assertIn("297 more rows", out)
        self.assertIn("7 distinct values of Customer", self.q(distinct="Customer"))
        self.assertIn("Sheet \"Pending\"", self.q(describe=True))
        self.assertIn("2 data rows", self.q(sheet="small", count=True) + run_data_query(self.xlsx, {"describe": True}))

    def test_bad_requests_say_what_is_wrong(self):
        with self.assertRaises(QueryError) as c:
            self.q(filters=[{"column": "Nope", "op": "="}])
        self.assertIn("Columns:", str(c.exception))
        with self.assertRaises(QueryError):
            self.q(filters=[{"column": "Status", "op": "like", "value": "x"}])
        with self.assertRaises(QueryError):
            self.q(sheet="missing")
        with self.assertRaises(QueryError):
            self.q(group_by="Status", agg=[{"fn": "median"}])

    def test_output_is_capped(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["id", "text"])
        for i in range(500):
            ws.append([i, "word " * 40])
        p = self.dir / "wide.xlsx"
        wb.save(p)
        out = run_data_query(p, {"limit": 500})
        self.assertLess(len(out), 6200)
        self.assertIn("more rows", out)


class DataToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._p = mock.patch.object(agent_tools, "COMMON_ROOT", self.root)
        self._p.start()
        set_current_user(11)
        _workbook(agent_tools.common_workspace() / "pending-1a2b3c4d.xlsx")

    def tearDown(self):
        self._p.stop()
        set_current_user(None)
        self.tmp.cleanup()

    def test_tool_reads_only_the_callers_common_space_files(self):
        from core.agent_tools.data_tools import tool_data_query
        out = tool_data_query({"file": "pending-1a2b3c4d.xlsx", "count": True})
        self.assertIn("300 of 300 rows match", out)
        self.assertTrue(tool_data_query({"file": "nope.xlsx"}).startswith("error: no attached data file"))
        set_current_user(12)
        self.assertTrue(tool_data_query({"file": "pending-1a2b3c4d.xlsx"}).startswith("error: no attached data file"),
                        "another user's upload is not reachable")
        set_current_user(11)
        (agent_tools.common_workspace() / "n.txt").write_text("x", encoding="utf-8")
        self.assertIn("reads .xlsx", tool_data_query({"file": "n.txt"}))

    def test_tool_is_registered_with_a_schema(self):
        names = {t["function"]["name"] for t in agent_tools.AGENT_TOOLS}
        self.assertIn("data_query", names)
        self.assertIn("data_query", agent_tools.TOOL_IMPLS)
        self.assertEqual(asyncio.iscoroutinefunction(agent_tools.TOOL_IMPLS["data_query"]), False)


if __name__ == "__main__":
    unittest.main()
