"""tests/test_doc_ops.py - surgical document edits keep everything else identical.

Covers the doc_ops engine (pptx/xlsx/csv/docx/pdf-from-source), its verifier,
PAN guards, per-user common storage isolation, versioning via doc_store, and
the agent-mode device path (bytes through the companion, nothing on server disk).

Run: python -m unittest tests.test_doc_ops -v
"""

import asyncio
import base64
import io
import sqlite3
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.doc_ops as D  # noqa: E402
from core import agent_tools, companion_bridge, doc_store, doc_tools  # noqa: E402
from core.doc_ops import base, csv_ops, pptx_ops, xlsx_ops  # noqa: E402
from core.doc_ops.base import DocOpError  # noqa: E402
from core.request_context import set_current_device, set_current_user  # noqa: E402

DECK_MD = """# Quarterly Review
Merchant team
---
# Revenue
- Up 12%
- New merchants: 40
---
# Risks
- FX volatility
notes: mention the central bank circular
"""


def changed_parts(a: bytes, b: bytes) -> list[str]:
    za, zb = zipfile.ZipFile(io.BytesIO(a)), zipfile.ZipFile(io.BytesIO(b))
    na, nb = set(za.namelist()), set(zb.namelist())
    return sorted(n for n in na | nb if n not in na or n not in nb or za.read(n) != zb.read(n))


def synthetic_pan() -> str:
    """A Luhn-valid 16-digit test number built here (not a real card)."""
    body = "400000000000000"
    total = 0
    for i, ch in enumerate(reversed(body)):
        d = int(ch)
        if i % 2 == 0:
            d *= 2
            d = d - 9 if d > 9 else d
        total += d
    return body + str((10 - total % 10) % 10)


def doc_files_db() -> sqlite3.Connection:
    db = sqlite3.connect(":memory:", check_same_thread=False)
    db.row_factory = sqlite3.Row
    db.execute("""CREATE TABLE doc_files (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL,
        session_id INTEGER, name TEXT NOT NULL, kind TEXT NOT NULL, location TEXT NOT NULL DEFAULT 'common',
        parent_id INTEGER, version INTEGER NOT NULL DEFAULT 1, sha256 TEXT, source_spec TEXT,
        created_at REAL NOT NULL)""")
    return db


class PptxTests(unittest.TestCase):
    def setUp(self):
        self.deck = pptx_ops.create(DECK_MD)

    def test_title_edit_touches_only_that_slide(self):
        r = pptx_ops.apply(self.deck, [{"op": "set_text", "addr": "s2/title", "text": "Revenue FY26"}])
        self.assertEqual(changed_parts(self.deck, r.data), ["ppt/slides/slide2.xml"])
        self.assertIn("Revenue FY26", D.inspect(r.data, "x.pptx"))

    def test_replace_text_keeps_other_slides(self):
        r = pptx_ops.apply(self.deck, [{"op": "replace_text", "find": "12%", "replace": "15%"}])
        self.assertEqual(changed_parts(self.deck, r.data), ["ppt/slides/slide2.xml"])

    def test_add_delete_move_slides(self):
        r = pptx_ops.apply(self.deck, [{"op": "add_slide", "after": 2, "title": "Compliance",
                                        "bullets": ["PCI DSS", "Local regulator"]}])
        titles = [e["text"] for e in pptx_ops.inspect(r.data)["elements"] if e["kind"] == "title"]
        self.assertEqual(titles, ["Quarterly Review", "Revenue", "Compliance", "Risks"])
        r2 = pptx_ops.apply(r.data, [{"op": "delete_slide", "addr": "s1"},
                                     {"op": "move_slide", "addr": "s1", "to": 3}])
        titles = [e["text"] for e in pptx_ops.inspect(r2.data)["elements"] if e["kind"] == "title"]
        self.assertEqual(titles, ["Compliance", "Risks", "Revenue"])

    def test_bullet_levels_follow_indentation(self):
        r = pptx_ops.apply(self.deck, [{"op": "set_text", "addr": "s2/sh3", "text": "A\n\tB\nC"}])
        paras = [e["text"] for e in pptx_ops.inspect(r.data)["elements"] if e["addr"].startswith("s2/sh3/p")]
        self.assertEqual(paras, ["A", "  B", "C"])

    def test_verifier_rejects_untargeted_change(self):
        def sneaky(ctx, op):
            slide = ctx.prs.slides[1]
            title, body = slide.shapes.title, [s for s in slide.shapes if s.shape_id != slide.shapes.title.shape_id][0]
            ctx.text_edit(slide, title.text_frame._txBody)
            title.text_frame.text = "asked"
            body.text_frame.text = "NOT asked"          # changed without touch()
        with mock.patch.dict(pptx_ops.OPS, {"sneaky": sneaky}):
            with self.assertRaises(DocOpError) as cm:
                pptx_ops.apply(self.deck, [{"op": "sneaky"}])
        self.assertIn("did not ask for", str(cm.exception))

    def test_bad_address(self):
        with self.assertRaises(DocOpError):
            pptx_ops.apply(self.deck, [{"op": "set_text", "addr": "s9/title", "text": "x"}])

    def test_parse_conversational_slide_outlines(self):
        outline = (
            "10 SLIDES + TITLE/THANKYOU STRUCTURE:\n"
            "Slide 1 (Title): Full-bleed deep navy. Big title \"Acme Corp\". Subtitle \"Annual Strategy\".\n"
            "Slide 2 (Executive Summary): Header \"Strategic Priorities\". Three bullets:\n"
            "- Accelerate Cloud Migration\n"
            "- Expand Global Merchant Footprint\n"
            "- Strengthen Compliance & Governance\n"
            "Slide 3 (Scale): Header \"Platform Metrics\". Bullets:\n"
            "- 50M+ active accounts\n"
            "GLOBAL THEME: Deep navy background, gold highlights.\n"
        )
        deck = pptx_ops.create(outline)
        titles = [e["text"] for e in pptx_ops.inspect(deck)["elements"] if e["kind"] == "title"]
        self.assertEqual(titles, ["Acme Corp", "Strategic Priorities", "Platform Metrics"])


class XlsxTests(unittest.TestCase):
    def setUp(self):
        import openpyxl
        from openpyxl.chart import BarChart, Reference
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Merchant", "Q1", "Q2", "Total"])
        for i, row in enumerate([("Alpha", 10, 12), ("Beta", 7, 9)], 2):
            ws.append([*row, f"=B{i}+C{i}"])
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3, max_col=3), titles_from_data=True)
        ws.add_chart(ch, "F2")
        buf = io.BytesIO()
        wb.save(buf)
        self.book = buf.getvalue()

    def test_cell_edit_preserves_chart_and_other_parts(self):
        r = xlsx_ops.apply(self.book, [{"op": "set_cell", "addr": "Sales!B3", "value": 70}])
        self.assertEqual(changed_parts(self.book, r.data), ["xl/workbook.xml", "xl/worksheets/sheet1.xml"])
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(r.data))["Sales"]
        self.assertEqual((ws["B3"].value, ws["D3"].value), (70, "=B3+C3"))

    def test_structural_edit_refused_when_lossy(self):
        with self.assertRaises(DocOpError):
            xlsx_ops.apply(self.book, [{"op": "insert_rows", "sheet": "Sales", "at": 2}])

    def test_formula_needs_explicit_op(self):
        with self.assertRaises(DocOpError):
            xlsx_ops.apply(self.book, [{"op": "set_cell", "addr": "Sales!A2", "value": "=cmd|' /C calc'!A0"}])
        r = xlsx_ops.apply(self.book, [{"op": "set_formula", "addr": "Sales!D4", "formula": "SUM(D2:D3)"}])
        self.assertIn("D4==SUM(D2:D3)", D.inspect(r.data, "x.xlsx"))

    def test_created_text_never_becomes_formula(self):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_ops.create([["Name"], ["=cmd"]])))
        self.assertEqual(wb.active["A2"].data_type, "s")


class CsvDocxPdfTests(unittest.TestCase):
    def test_csv_untouched_rows_byte_identical(self):
        src = b'id,merchant,amount\r\n1,Alpha,100\r\n2,"Beta, Ltd",200\r\n3,Gamma,"multi\r\nline"\r\n'
        r = csv_ops.apply(src, [{"op": "set_cell", "addr": "r2camount", "value": 150}])
        self.assertEqual(r.data, src.replace(b"1,Alpha,100", b"1,Alpha,150"))

    def test_csv_formula_injection_neutralised(self):
        r = csv_ops.apply(b"a,b\n1,2\n", [{"op": "set_cell", "addr": "r2c1", "value": "=HYPERLINK(1)"}])
        self.assertIn(b"'=HYPERLINK(1)", r.data)

    def test_docx_paragraph_and_table_edits(self):
        md = "# Guide\n\nIntro\n\n- KYC\n- Sign\n\n| Item | Fee |\n|---|---|\n| Card | 2% |\n"
        d, _ = D.create("g.docx", md)
        out = D.edit(d, "g.docx", [{"op": "set_text", "addr": "p3", "text": "KYC + AML"},
                                   {"op": "set_table_cell", "addr": "t1[2,2]", "text": "1.8%"},
                                   {"op": "add_table_row", "addr": "t1", "values": ["Wallet", "1.2%"]}])
        self.assertEqual(changed_parts(d, out.data), ["word/document.xml"])
        text = D.inspect(out.data, "g.docx")
        for s in ("KYC + AML", "1.8%", "Wallet", "Sign"):
            self.assertIn(s, text)

    def test_pdf_edits_patch_source_and_rerender(self):
        md = "# Report\n\nIntro\n\n## Numbers\nold\n\n## Contact\nus\n"
        pdf, spec = D.create("r.pdf", md)
        self.assertTrue(pdf.startswith(b"%PDF-"))
        out = D.edit(pdf, "r.pdf", [{"op": "replace_section", "addr": "sec2", "text": "## Numbers\nnew"}],
                     source_spec=spec)
        self.assertTrue(out.data.startswith(b"%PDF-"))
        self.assertIn("new", out.source_spec)
        self.assertIn("## Contact\nus", out.source_spec)

    def test_uploaded_pdf_is_read_only(self):
        pdf, _ = D.create("r.pdf", "# A\ntext")
        with self.assertRaises(DocOpError):
            D.edit(pdf, "r.pdf", [{"op": "replace_text", "find": "A", "replace": "B"}])

    def test_macro_and_legacy_refused(self):
        for name in ("m.pptm", "old.ppt"):
            with self.assertRaises(DocOpError):
                D.edit(b"x", name, [{"op": "set_text"}])


class PanTests(unittest.TestCase):
    def test_pan_in_op_blocked(self):
        deck = pptx_ops.create(DECK_MD)
        with self.assertRaises(DocOpError) as cm:
            D.edit(deck, "d.pptx", [{"op": "set_text", "addr": "s1/title", "text": f"card {synthetic_pan()}"}])
        self.assertIn("PCI", str(cm.exception))

    def test_pan_in_document_masked_in_outline(self):
        pan = synthetic_pan()
        deck = pptx_ops.create(f"# Payments\n- test card {pan}\n")
        text = D.inspect(deck, "d.pptx")
        self.assertNotIn(pan, text)
        self.assertIn("****" + pan[-4:], text)


class StorageAndToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._patches = [
            mock.patch.object(agent_tools, "COMMON_ROOT", self.root),
            mock.patch.object(doc_store, "_projects_db", doc_files_db()),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        set_current_user(None)
        set_current_device(None)
        self.tmp.cleanup()

    def test_users_have_separate_common_spaces(self):
        set_current_user(1)
        agent_tools.tool_write_file_common({"path": "secret.csv", "content": "a,b\n1,2"})
        set_current_user(2)
        self.assertEqual(list(agent_tools.common_workspace().iterdir()), [])
        from routes.agent import _resolve_requested_file
        name = next((self.root / "user_1").iterdir()).name
        for probe in (name, f"../user_1/{name}", str(self.root / "user_1" / name), f"..%2Fuser_1%2F{name}"):
            p = _resolve_requested_file(probe)
            self.assertTrue(p is None or not p.is_file() or "user_2" in str(p), probe)

    def test_upload_names_are_deduped(self):
        from routes.agent import _unique_dest
        (self.root / "a.pptx").write_bytes(b"1")
        self.assertEqual(_unique_dest(self.root, "a.pptx").name, "a-2.pptx")

    def test_chat_create_then_edit_keeps_original(self):
        set_current_user(3)
        res = agent_tools.tool_write_file_common({"path": "deck.pptx", "content": DECK_MD})
        first = res.split("[DOWNLOAD: ")[1].rstrip("]")
        v1 = (self.root / "user_3" / first).read_bytes()
        out = doc_tools.tool_doc_edit_common({"file": first, "ops": [
            {"op": "set_text", "addr": "s3/title", "text": "Key Risks"}]})
        self.assertIn("[DOWNLOAD: deck-v2-", out, out)
        second = out.split("[DOWNLOAD: ")[1].split("]")[0]
        v2 = (self.root / "user_3" / second).read_bytes()
        self.assertEqual((self.root / "user_3" / first).read_bytes(), v1)   # old version untouched
        self.assertEqual(changed_parts(v1, v2), ["ppt/slides/slide3.xml"])
        row = doc_store.find(3, second)
        self.assertEqual([r["version"] for r in doc_store.history(3, row["id"])], [1, 2])

    def test_chat_pdf_edit_uses_stored_source(self):
        set_current_user(4)
        res = agent_tools.tool_write_file_common({"path": "memo.pdf", "content": "# Memo\n\n## Fees\n2%\n"})
        name = res.split("[DOWNLOAD: ")[1].rstrip("]")
        out = doc_tools.tool_doc_edit_common({"file": name, "ops": [
            {"op": "replace_text", "find": "2%", "replace": "1.5%"}]})
        self.assertIn("[DOWNLOAD: memo-v2-", out, out)
        new = out.split("[DOWNLOAD: ")[1].split("]")[0]
        self.assertIn("1.5%", doc_store.find(4, new)["source_spec"])

    def test_agent_mode_edits_bytes_on_device_only(self):
        user_dir = r"C:\Users\[PLACEHOLDER]\proj"
        device = {str(Path(user_dir) / "deck.pptx"): pptx_ops.create(DECK_MD)}
        db = sqlite3.connect(":memory:")
        db.row_factory = sqlite3.Row
        db.execute("CREATE TABLE projects (name TEXT, user_id INTEGER, device_id TEXT, workspace_dir TEXT)")
        db.execute("INSERT INTO projects VALUES ('p', 5, 'dev', ?)", (user_dir,))

        async def fake_call(uid, op, params, timeout=None):
            if op == "fs.read_b64":
                b = device.get(params["path"])
                return {"data": base64.b64encode(b).decode() if b else None}
            if op == "fs.write_b64":
                device[params["path"]] = base64.b64decode(params["data"])
                return {"existed": False}
            raise AssertionError(op)

        with mock.patch.object(agent_tools, "_projects_db", db), \
                mock.patch.object(companion_bridge, "is_connected", lambda uid: uid == 5), \
                mock.patch.object(companion_bridge, "call", fake_call), \
                mock.patch.dict(agent_tools._active_project, {"5:dev": "p"}, clear=True):
            set_current_user(5)
            set_current_device("dev")
            out = asyncio.run(doc_tools.tool_doc_edit({"file": "deck.pptx", "ops": [
                {"op": "set_text", "addr": "s1/title", "text": "Q3 Review"}]}))
        self.assertIn("saved as deck-v2-", out, out)
        new = [k for k in device if "deck-v2-" in k]
        self.assertEqual(len(new), 1)
        self.assertIn("Q3 Review", D.inspect(device[new[0]], "d.pptx"))
        self.assertFalse(any(self.root.rglob("*")), "nothing may be written to the server's disk")


class PackageCheckTests(unittest.TestCase):
    def test_changed_part_detected(self):
        deck = pptx_ops.create(DECK_MD)
        parts = base.read_zip(deck)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            for n, b in parts.items():
                z.writestr(n, b.replace(b"Merchant team", b"Tampered") if n.endswith("slide1.xml") else b)
        problems = base.check_package(deck, buf.getvalue(), base.PartRules(), pptx_ops._rename_parts)
        self.assertTrue(any("slide:" in p for p in problems), problems)


if __name__ == "__main__":
    unittest.main()
