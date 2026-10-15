"""Designer style: slides drawn as HTML, printed by the browser, placed in the deck as pictures."""
import io
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.agent_tools  # noqa: F401,E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from core.agent_tools import pdf_render  # noqa: E402
from core.doc_ops import pptx_ops  # noqa: E402
from core.doc_ops.pptx_ops import designer, lint  # noqa: E402
from core.doc_ops.pptx_ops import outline as O  # noqa: E402
from core.doc_tools import operations as ops  # noqa: E402

HAVE_BROWSER = bool(pdf_render._find_chromium_binary())

MD = """# Harbour City
theme: editorial
mode: designer
subtitle: A short tour
eyebrow: Field notes
art: arches
---
## Part one
layout: section
subtitle: Where it started
---
## City in numbers
eyebrow: At a glance
layout: stats
source: Municipal archive; census 2022
notes: Say the definitions out loud.
- 36.6M | Metro population
- 1610 | Year of the charter
- 306 km2 | Area
---
## Rivers
layout: textside
- The city sits on a delta.
- Six rivers surround it.
| Month | Mean C | Rain |
|---|---|---|
| Jan | 18.4 | 7 |
| Jul | 28.9 | 391 |
---
## Dates
layout: timeline
- 1952 | A movement begins
- 1971 | A country is born
- 2022 | Metro opens
---
## Growth
chart: line | unit=%
- 2024: 40
- 2025: 60
- 2026: 89
---
## Shares
chart: stacked | unit=%
- Garments: 58.6
- Food: 13.4
- Other: 28
---
## Old vs new
layout: comparison
- Old:
  - Narrow lanes
  - Brick houses
- New:
  - Glass towers
  - Metro
---
## Thank you
"""


class Parsing(unittest.TestCase):
    def test_mode_line(self):
        self.assertEqual(O.parse_outline(MD)["mode"], "designer")
        self.assertEqual(O.parse_outline("# D\nmode: editable\n---\n## A\n- x\n- y\n")["mode"], "")
        self.assertEqual(O.parse_outline("# D\nmode: bogus\n---\n## A\n- x\n- y\n")["mode"], "")

    def test_html_has_one_section_per_slide_and_no_scripts_or_remote_assets(self):
        deck = O.parse_outline(MD)
        html = designer.deck_html(deck)
        self.assertEqual(html.count('<section class="slide"'), len(deck["slides"]))
        self.assertNotIn("<script", html.lower())
        remote = html.replace("http://www.w3.org/2000/svg", "")         # the SVG namespace is not a fetch
        self.assertNotIn("http://", remote)
        self.assertNotIn("https://", remote)

    def test_titles_are_escaped(self):
        md = '# D\nmode: designer\n---\n## <img src=x onerror=alert(1)>\n- one\n- two\n'
        html = designer.deck_html(O.parse_outline(md))
        self.assertNotIn("<img src=x", html)
        self.assertIn("&lt;img src=x", html)

    def test_theme_colours_and_fonts_reach_the_css(self):
        html = designer.deck_html(O.parse_outline(MD))
        self.assertIn("Georgia", html)
        self.assertIn("--bg:#F4EFE4", html)             # editorial content slide
        self.assertIn("--bg:#10302D", html)             # the alternate (deep) palette and the cover

    def test_cover_gets_art_and_a_date(self):
        html = designer.deck_html(O.parse_outline(MD))
        self.assertIn("<svg viewBox=\"0 0 490 540\"", html)
        self.assertIn("Prepared ", html)

    def test_source_line_becomes_the_footer(self):
        html = designer.deck_html(O.parse_outline(MD))
        self.assertIn("Source: Municipal archive; census 2022", html)

    def test_every_art_kind_draws(self):
        t = O.resolve_theme(O.parse_outline(MD))
        for kind in O.ART_KINDS:
            if kind != "none":
                self.assertIn("<svg", designer._art_svg(kind, t), kind)


class Fallback(unittest.TestCase):
    def test_no_browser_gives_the_editable_deck(self):
        with mock.patch.object(pdf_render, "_find_chromium_binary", return_value=None):
            data = pptx_ops.create(MD)
        prs = Presentation(io.BytesIO(data))
        self.assertGreater(sum(1 for sl in prs.slides for sh in sl.shapes if sh.has_text_frame), 5)   # real text, not pictures
        self.assertFalse(ops.is_designer_deck(data))

    def test_a_failed_print_gives_the_editable_deck(self):
        with mock.patch.object(pdf_render, "_render_html_to_pdf", return_value=False):
            data = pptx_ops.create(MD)
        self.assertFalse(ops.is_designer_deck(data))

    def test_templates_ignore_designer_mode(self):
        template = pptx_ops.create("# T\nsubtitle: s\n---\n## One\n- a\n- b\n- c\n")
        data = pptx_ops.create(MD, template=template)
        self.assertFalse(ops.is_designer_deck(data))


@unittest.skipUnless(HAVE_BROWSER, "no Chrome/Edge to print with")
class RealRender(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = pptx_ops.create(MD)
        cls.prs = Presentation(io.BytesIO(cls.data))
        cls.deck = O.parse_outline(MD)

    def test_every_slide_is_one_full_page_picture(self):
        self.assertTrue(ops.is_designer_deck(self.data))
        self.assertEqual(len(self.prs.slides), len(O.parse_outline(MD)["slides"]))
        for sl in self.prs.slides:
            pic = sl.shapes[0]
            self.assertEqual((pic.left, pic.top), (0, 0))
            self.assertEqual((pic.width, pic.height), (self.prs.slide_width, self.prs.slide_height))

    def test_pictures_are_sharp_and_the_file_stays_reasonable(self):
        im = Image.open(io.BytesIO(self.prs.slides[1].shapes[0].image.blob))
        self.assertGreaterEqual(im.size[0], 2400)
        self.assertEqual(round(im.size[0] / im.size[1], 2), round(13.333 / 7.5, 2))
        self.assertLess(len(self.data), 8 * 1024 * 1024)

    def test_text_is_kept_in_notes_and_alt_text(self):
        sl = self.prs.slides[2]
        self.assertIn("Say the definitions out loud.", sl.notes_slide.notes_text_frame.text)
        self.assertIn("Slide text:", sl.notes_slide.notes_text_frame.text)
        self.assertIn("36.6M", sl.notes_slide.notes_text_frame.text)
        self.assertIn("City in numbers", sl.shapes[0]._element.nvPicPr.cNvPr.get("descr"))

    def test_the_pictures_are_not_blank(self):
        for sl in self.prs.slides:
            im = Image.open(io.BytesIO(sl.shapes[0].image.blob)).convert("L").resize((160, 90))
            self.assertGreater(len(set(im.tobytes())), 6)

    def test_the_layout_check_has_nothing_to_flag_on_pictures(self):
        self.assertEqual(lint.lint_deck(self.data), [])
        self.assertIn("Designer style", ops._layout_note(".pptx", self.data, MD))

    def test_designer_deck_cannot_be_edited_in_place_and_says_how(self):
        import asyncio
        from core.doc_tools.location import _Loc
        loc = _Loc("common", Path("x.pptx"), 1)
        with self.assertRaises(ValueError) as cm:
            asyncio.run(ops._edit(loc, self.data, {"ops": [{"op": "set_text", "addr": "s1/title", "text": "x"}]}))
        self.assertIn("Designer style", str(cm.exception))

    def test_a_deck_with_pictures_gallery_and_full_slides(self):
        def pic(color):
            b = io.BytesIO()
            Image.new("RGB", (1200, 800), color).save(b, "JPEG")
            return {"data": b.getvalue(), "credit_host": "example.org", "credit_url": "https://example.org", "via": "web"}
        md = ("# D\nmode: designer\nsubtitle: s\nimage: harbour\n---\n## A\nimage: street\n- one\n- two\n- three\n---\n"
              "## G\nlayout: gallery\nimage: a\nimage: b\nimage: c\n---\n## F\nlayout: full\nimage: wide\n---\n## Thanks\n")
        images = {0: pic((40, 90, 120)), 1: pic((120, 70, 50)),
                  2: {**pic((60, 110, 60)), "extra": [pic((130, 60, 100)), pic((50, 80, 140))]}, 3: pic((20, 60, 100))}
        data = pptx_ops.create(md, images=images)
        self.assertTrue(ops.is_designer_deck(data))
        self.assertEqual(len(Presentation(io.BytesIO(data)).slides), 5)


@unittest.skipUnless(HAVE_BROWSER, "no Chrome/Edge to print with")
class BuiltInPieces(unittest.TestCase):
    """doc_create total=N + doc_append: editable while building (instant), Designer pictures on the last step."""

    def setUp(self):
        import tempfile
        from core import agent_tools, doc_store
        from core.request_context import set_current_user
        import tests.test_doc_ops as T
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._patches = [mock.patch.object(agent_tools, "COMMON_ROOT", self.root),
                         mock.patch.object(doc_store, "_projects_db", T.doc_files_db())]
        for p in self._patches:
            p.start()
        set_current_user(11)
        ops._LAST_STYLE.clear()
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        from core.request_context import set_current_user
        for p in self._patches:
            p.stop()
        set_current_user(None)
        ops._LAST_STYLE.clear()
        self.tmp.cleanup()

    def _file(self, name):
        return (self.root / "user_11" / name).read_bytes()

    def test_designer_pictures_are_drawn_only_on_the_last_step(self):
        from core import doc_tools
        first = "mode: designer\n# Harbour\nsubtitle: s\n---\n## One\n- a\n- b\n"
        out = doc_tools.tool_doc_create_common({"file": "p.pptx", "content": first, "total": 4})
        name = out.split("[DOWNLOAD: ")[1].split("]")[0]
        self.assertFalse(ops.is_designer_deck(self._file(name)))              # interim: editable and instant
        mid = doc_tools.tool_doc_append_common({"file": name, "content": "---\n## Two\n- c\n- d\n"})
        self.assertFalse(ops.is_designer_deck(self._file(name)))
        self.assertNotIn("Designer style", mid)
        last = doc_tools.tool_doc_append_common({"file": name, "content": "---\n## Thank you\n", "done": True})
        self.assertTrue(ops.is_designer_deck(self._file(name)))               # final: the pictures
        self.assertIn("Designer style", last)
        self.assertEqual(len(Presentation(io.BytesIO(self._file(name))).slides), 4)


class StyleMemory(unittest.TestCase):
    def setUp(self):
        ops._LAST_STYLE.clear()
        self.addCleanup(ops._LAST_STYLE.clear)

    def test_designer_style_is_part_of_the_session_look(self):
        ops._inherit_style(1, ".pptx", "theme: ember\nmode: designer\n# One\nsubtitle: s\n---\n## A\n- x\n- y\n")
        out = ops._inherit_style(1, ".pptx", "# Two\nsubtitle: s\n---\n## B\n- x\n- y\n")
        self.assertEqual(O.parse_outline(out)["mode"], "designer")

    def test_kill_switch_strips_the_mode_line(self):
        import asyncio
        md = "# D\nmode: designer\nsubtitle: s\n---\n## A\n- x\n- y\n"
        with mock.patch.object(ops, "_flag", side_effect=lambda n, d=True: n != "deck_designer"):
            content, _images, _note = asyncio.run(ops.prepare_build(".pptx", md, None))
        self.assertNotIn("designer", content.lower())


if __name__ == "__main__":
    unittest.main()
