"""Deck layout that is measured, not guessed: font-metric text fitting, slides split instead of overflowing, tables that fit."""
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation

from core.doc_ops import pptx_ops
from core.doc_ops.pptx_ops import design as D
from core.doc_ops.pptx_ops import measure as M

LONG = "Founded in the seventeenth century by Mughal governors and expanded under later rulers of the region"


def build(md: str):
    return Presentation(io.BytesIO(pptx_ops.create(md)))


class Measure(unittest.TestCase):
    def test_narrower_box_takes_more_lines(self):
        wide = M.lines_for(LONG, "Calibri", 24, 9.0)
        narrow = M.lines_for(LONG, "Calibri", 24, 3.0)
        self.assertLess(wide, narrow)

    def test_bigger_text_takes_more_height(self):
        a = M.height_in([LONG] * 3, "Calibri", 16, 6.0, 8)
        b = M.height_in([LONG] * 3, "Calibri", 28, 6.0, 8)
        self.assertLess(a, b)

    def test_fit_size_shrinks_until_it_fits(self):
        texts = [LONG] * 6
        size = M.fit_size(texts, "Calibri", 5.0, 4.5, 30, 12, 16)
        self.assertTrue(12 <= size < 30)
        self.assertTrue(M.fits(texts, "Calibri", size, 5.0, 4.5, 16) or size == 12)

    def test_short_text_keeps_the_base_size(self):
        self.assertEqual(M.fit_size(["Short line"], "Calibri", 8.0, 3.0, 30, 12, 6), 30)

    def test_unknown_font_falls_back_to_an_estimate(self):
        self.assertGreaterEqual(M.lines_for(LONG, "No Such Font", 20, 3.0), 2)

    def test_a_very_long_word_wraps_instead_of_counting_as_one_line(self):
        self.assertGreater(M.lines_for("x" * 200, "Calibri", 20, 3.0), 2)

    def test_bengali_lines_are_taller(self):
        self.assertGreater(M.line_height(["ঢাকা"]), M.line_height(["Dhaka"]))


class Splitting(unittest.TestCase):
    def test_a_normal_slide_is_left_alone(self):
        prs = build("# D\nsubtitle: s\n---\n## Why\n- one\n- two\n- three\n")
        self.assertEqual(len(prs.slides), 2)

    def test_too_many_long_bullets_become_continued_slides(self):
        bullets = "\n".join(f"- {LONG} number {i}" for i in range(14))
        prs = build(f"# D\nsubtitle: s\n---\n## Landmarks\n{bullets}\n")
        titles = [sl.shapes.title.text_frame.text for sl in prs.slides if sl.shapes.title is not None]
        self.assertGreater(len(prs.slides), 2)
        self.assertEqual(titles[1], "Landmarks")
        self.assertTrue(all(t.endswith("(cont.)") for t in titles[2:]), titles)

    def test_every_bullets_slide_fits_after_splitting(self):
        from core.doc_ops.pptx_ops import outline as O
        bullets = "\n".join(f"- {LONG} number {i}" for i in range(14))
        deck = O.parse_outline(f"# D\nsubtitle: s\n---\n## Landmarks\n{bullets}\n")
        t = O.resolve_theme(deck)
        for s in D._prepare(deck["slides"], t):
            if s["kind"] == "bullets":
                w, h, floor = D._bullet_box(s)
                self.assertTrue(M.fits([b.strip() for b in s["bullets"]], t["body"], floor, w - 0.4, h, 16), s["title"])

    def test_notes_stay_on_the_first_slide(self):
        bullets = "\n".join(f"- {LONG} number {i}" for i in range(14))
        prs = build(f"# D\nsubtitle: s\n---\n## Landmarks\nnotes: speaker words\n{bullets}\n")
        notes = [sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else "" for sl in prs.slides]
        self.assertEqual(notes[1], "speaker words")
        self.assertTrue(all(n == "" for n in notes[2:]))


class Variety(unittest.TestCase):
    def _kinds(self, md):
        from core.doc_ops.pptx_ops import outline as O
        return [s["kind"] for s in O.parse_outline(md)["slides"]]

    def test_third_plain_bullet_slide_in_a_row_changes_shape(self):
        slide = lambda n: f"---\n## Topic {n}\n- alpha one\n- beta two\n- gamma three\n- delta four\n"   # noqa: E731
        kinds = self._kinds("# D\nsubtitle: s\n" + "".join(slide(i) for i in range(1, 5)))
        self.assertEqual(kinds[1:3], ["bullets", "bullets"])
        self.assertEqual(kinds[3], "columns")

    def test_label_bullets_become_cards(self):
        plain = "---\n## T\n- one thing\n- two things\n"
        labelled = "---\n## L\n- Speed: very fast\n- Cost: quite low\n- Safety: well tested\n"
        kinds = self._kinds("# D\nsubtitle: s\n" + plain + plain + labelled)
        self.assertEqual(kinds[-1], "cards")

    def test_explicit_layout_and_pictures_are_respected(self):
        slide = "---\n## T\n- alpha one\n- beta two\n- gamma three\n- delta four\n"
        kinds = self._kinds("# D\nsubtitle: s\n" + slide + slide + slide.replace("## T", "## T\nimage: boats"))
        self.assertEqual(kinds[3], "bullets")


def table_md(rows: int, cols: int = 3, cell: str = "value") -> str:
    head = "| " + " | ".join(f"Col {c}" for c in range(cols)) + " |\n|" + "---|" * cols + "\n"
    body = "".join("| " + " | ".join(f"{cell} {r}-{c}" for c in range(cols)) + " |\n" for r in range(rows))
    return f"# D\nsubtitle: s\n---\n## Data\n{head}{body}"


class Tables(unittest.TestCase):
    def _tables(self, prs):
        return [sh for sl in prs.slides for sh in sl.shapes if sh.has_table]

    def test_small_table_is_one_slide_with_a_big_font(self):
        prs = build(table_md(4))
        self.assertEqual(len(self._tables(prs)), 1)
        size = self._tables(prs)[0].table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size.pt
        self.assertGreaterEqual(size, 16)

    def test_long_table_splits_and_repeats_its_header(self):
        prs = build(table_md(22, 4, "a fairly long cell value to wrap"))
        tables = self._tables(prs)
        self.assertGreater(len(tables), 1)
        for tb in tables:
            self.assertEqual(tb.table.cell(0, 0).text_frame.text, "Col 0")
        titles = [sl.shapes.title.text_frame.text for sl in prs.slides if sl.shapes.title is not None]
        self.assertTrue(any(t.endswith("(cont.)") for t in titles), titles)
        total_rows = sum(len(tb.table.rows) - 1 for tb in tables)
        self.assertEqual(total_rows, 22)                       # no row lost

    def test_every_table_stays_inside_the_slide(self):
        for rows, cols in ((6, 3), (12, 5), (22, 4)):
            prs = build(table_md(rows, cols, "some longer words in the cell"))
            for tb in self._tables(prs):
                bottom = (tb.top + sum(r.height for r in tb.table.rows)) / 914400
                self.assertLessEqual(bottom, 7.5 - 0.3, (rows, cols, bottom))

    def test_wide_column_gets_more_width(self):
        rows = [["Id", "Description of the thing"], ["1", "a long description that goes on and on and on"]]
        widths = D._table_widths(rows, 11.9)
        self.assertGreater(widths[1], widths[0] * 2)


if __name__ == "__main__":
    unittest.main()
