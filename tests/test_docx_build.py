"""tests/test_docx_build.py - the markdown -> .docx builder: cover/contents, footer fields, tables, lists, links, XML order."""
import io
import re
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx
from lxml import etree

from core.doc_ops.docx_ops import create
from core.doc_ops.docx_ops.creation import _ORDER

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
LONG = ("# Review\n\nIntro line.\n\n" + "".join(f"## Section {i}\nText {i}.\n\n" for i in range(1, 6)))


def _parts(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    return z, z.read("word/document.xml").decode()


class DocxBuildTests(unittest.TestCase):
    def test_long_document_gets_cover_and_contents(self):
        d = docx.Document(io.BytesIO(create(LONG, "Fallback")))
        styles = [p.style.name for p in d.paragraphs if p.text]
        self.assertEqual(styles[0], "Title")
        texts = [p.text for p in d.paragraphs]
        self.assertIn("Contents", texts)
        self.assertTrue(any(t.startswith("01") and "Section 1" in t for t in texts))
        _, xml = _parts(create(LONG))
        self.assertIn("TOC", xml)

    def test_short_document_stays_plain(self):
        d = docx.Document(io.BytesIO(create("# Memo\n\nHello **world**\n\n- a\n- b\n", "Memo")))
        self.assertEqual([(p.style.name, p.text) for p in d.paragraphs if p.text],
                         [("Heading 1", "Memo"), ("Normal", "Hello world"), ("List Bullet", "a"), ("List Bullet", "b")])

    def test_footer_has_page_fields(self):
        z, _ = _parts(create(LONG))
        foot = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
        self.assertIn("PAGE", foot)
        self.assertIn("NUMPAGES", foot)

    def test_second_numbered_list_restarts(self):
        _, xml = _parts(create("1. a\n2. b\n\ntext\n\n1. c\n2. d\n"))
        self.assertEqual(len(set(re.findall(r'<w:numId w:val="(\d+)"', xml))), 2)

    def test_only_safe_links_become_hyperlinks(self):
        _, xml = _parts(create("[ok](https://example.com) and [bad](javascript:alert(1)) and [file](file:///c:/x)"))
        self.assertEqual(xml.count("<w:hyperlink"), 1)
        self.assertNotIn("javascript", xml)

    def test_table_header_repeats_and_numbers_align_right(self):
        d = docx.Document(io.BytesIO(create("| Item | Fee |\n|---|---|\n| Card | 2.5% |\n| Wallet | 1,200.50 |\n")))
        t = d.tables[0]
        self.assertIn("tblHeader", t._tbl.xml)
        self.assertEqual(t.cell(1, 1).paragraphs[0].alignment, 2)   # WD_ALIGN_PARAGRAPH.RIGHT
        self.assertIsNone(t.cell(1, 0).paragraphs[0].alignment)

    def test_property_children_are_in_schema_order(self):
        z, _ = _parts(create(LONG + "| A | B |\n|---|---|\n| 1 | 2 |\n\n> quote\n\n```\ncode\n```\n"))
        for part in ("word/document.xml", "word/styles.xml", "word/footer1.xml"):
            root = etree.fromstring(z.read(part))
            for name, seq in _ORDER.items():
                for el in root.iter(W + name):
                    idx = [seq.index(c.tag.rsplit("}", 1)[1]) for c in el if c.tag.rsplit("}", 1)[1] in seq]
                    self.assertEqual(idx, sorted(idx), f"{part} <{name}> out of order")

    def test_bengali_sets_complex_script_font(self):
        z, _ = _parts(create("# শিরোনাম\n\nআমার সোনার বাংলা\n"))
        self.assertIn("Nirmala UI", z.read("word/styles.xml").decode())


def _png(w=400, h=200):
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (w, h), (30, 120, 200)).save(b, "PNG")
    return b.getvalue()


CHART_MD = "# Sales\n\nIntro.\n\nchart: bar | unit=BDT\n- Dhaka: 120\n- Khulna: 80\n\nAfter.\n"


class DocxFigureTests(unittest.TestCase):
    def test_image_lines_become_captioned_pictures(self):
        md = "# Trip\n\nimage: a river | caption=Padma river\n\nEnd.\n"
        d = docx.Document(io.BytesIO(create(md, "T", images={0: {"data": _png(), "credit_host": "example.org"}})))
        self.assertEqual(len(d.inline_shapes), 1)
        texts = [p.text for p in d.paragraphs if p.text]
        self.assertIn("Padma river - example.org", texts)
        self.assertFalse(any(t.lower().startswith("image:") for t in texts))

    def test_an_image_line_without_a_picture_is_dropped(self):
        d = docx.Document(io.BytesIO(create("# Trip\n\nimage: a river\n\nEnd.\n", "T")))
        self.assertEqual(len(d.inline_shapes), 0)
        self.assertFalse(any("image:" in p.text for p in d.paragraphs))

    def test_a_chart_becomes_a_picture(self):
        from unittest import mock
        from core.doc_ops.docx_ops import figures
        with mock.patch.object(figures, "chart_png", return_value=_png()):
            d = docx.Document(io.BytesIO(create(CHART_MD, "T")))
        self.assertEqual(len(d.inline_shapes), 1)
        texts = [p.text for p in d.paragraphs if p.text]
        self.assertNotIn("chart: bar | unit=BDT", texts)
        self.assertNotIn("Dhaka: 120", " ".join(texts))             # drawn, not repeated as text

    def test_no_browser_keeps_the_numbers_as_text(self):
        from unittest import mock
        from core.doc_ops.docx_ops import figures
        with mock.patch.object(figures, "chart_png", return_value=None):
            d = docx.Document(io.BytesIO(create(CHART_MD, "T")))
        self.assertEqual(len(d.inline_shapes), 0)
        texts = " ".join(p.text for p in d.paragraphs)
        self.assertIn("Dhaka: 120", texts)
        self.assertNotIn("chart:", texts)

    def test_picture_after_the_title_is_not_taken_for_the_cover_subtitle(self):
        md = "# Report\n\nimage: a river\n\n" + "".join(f"## S{i}\ntext\n\n" for i in range(1, 5))
        d = docx.Document(io.BytesIO(create(md, "T", images={0: {"data": _png()}})))
        self.assertEqual(len(d.inline_shapes), 1)
        self.assertFalse(any("pdfimg" in p.text for p in d.paragraphs))


if __name__ == "__main__":
    unittest.main()
