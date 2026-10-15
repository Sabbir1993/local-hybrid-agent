"""Restyle / extend a workbook that was built here: the ops change its OUTLINE (the source kept with the file) and the
workbook is rebuilt from it, so charts, formats, rules and styles that a cell-level edit cannot express are all still
possible. Cell-level ops (set_cell, set_formula, set_range, append_rows ...) stay byte-preserving and are not handled here;
after one of them the file no longer matches its outline, so these ops are refused until the workbook is rebuilt.

Ops: add_chart, set_format, add_highlight, add_scale, add_databar, add_validation, set_total, set_theme, set_freeze,
set_table, set_filter, add_note, add_column.
"""
import re
from typing import Optional

from .build import BuildError, _SHEET_RX, _split_row, build, is_outline, parse

OUTLINE_OPS = ("add_outline_sheet", "add_chart", "set_format", "add_highlight", "add_scale", "add_databar", "add_validation", "set_total",
               "set_theme", "set_freeze", "set_table", "set_filter", "add_note", "add_column")


class EditError(ValueError):
    pass


def is_outline_op(op: dict) -> bool:
    return str((op or {}).get("op") or "").lower() in OUTLINE_OPS


def _sections(lines: list) -> list:
    """[(sheet name, first line index, end index)] for each `## Sheet:` block (a single unnamed sheet counts as 'Sheet1')."""
    heads = [(i, _SHEET_RX.match(l).group(1)) for i, l in enumerate(lines) if _SHEET_RX.match(l)]
    if not heads:
        first = next((i for i, l in enumerate(lines) if l.lstrip().startswith("|")), 0)
        return [("Sheet1", first, len(lines))]
    return [(name, i, heads[k + 1][0] if k + 1 < len(heads) else len(lines)) for k, (i, name) in enumerate(heads)]


def _find(lines: list, sheet: Optional[str]) -> tuple:
    secs = _sections(lines)
    if not sheet:
        return secs[0]
    for s in secs:
        if s[0].lower() == str(sheet).strip().lower():
            return s
    raise EditError(f"sheet '{sheet}' not found (sheets: {', '.join(s[0] for s in secs)})")


def _append_directive(lines: list, sec: tuple, line: str) -> None:
    end = sec[2]
    while end > sec[1] and not lines[end - 1].strip():
        end -= 1
    lines.insert(end, line)


def _need(op: dict, key: str) -> str:
    v = op.get(key)
    if v is None or str(v).strip() == "":
        raise EditError(f"{op.get('op')}: {key} required")
    return str(v).strip()


def apply_outline_ops(outline: str, ops: list) -> tuple:
    """-> (new outline, [change descriptions]). The caller rebuilds the workbook and surfaces build errors."""
    lines = str(outline).split("\n")
    changes = []
    for op in ops:
        kind = str(op.get("op")).lower()
        sheet = op.get("sheet")
        if kind == "add_outline_sheet":
            name, body = _need(op, "name"), _need(op, "content")
            if any(n.lower() == name.lower() for n, _a, _b in _sections(lines)):
                raise EditError(f"add_outline_sheet: a sheet named '{name}' already exists")
            body_lines = body.strip().splitlines()
            if not any(ln.lstrip().startswith("|") for ln in body_lines):
                raise EditError("add_outline_sheet: content must contain a markdown table (| Col | Col | then |---|---| and rows)")
            while lines and not lines[-1].strip():
                lines.pop()
            lines += ["", f"## Sheet: {name}"] + body_lines
            changes.append(f"sheet '{name}' added")
            continue
        if kind == "set_theme":
            theme = _need(op, "theme").lower()
            idx = next((i for i, l in enumerate(lines) if re.match(r"^\s*theme\s*:", l, re.I)), None)
            if idx is not None:
                lines[idx] = f"theme: {theme}"
            else:
                at = next((i + 1 for i, l in enumerate(lines) if re.match(r"^\s*#\s+", l)), 0)
                lines.insert(at, f"theme: {theme}")
            changes.append(f"theme set to {theme}")
            continue
        sec = _find(lines, sheet)
        label = f"{sec[0]}: "
        if kind == "add_chart":
            spec = op.get("chart") or " | ".join(filter(None, [
                str(op.get("type") or "bar"), f"title={op['title']}" if op.get("title") else "",
                f"x={op['x']}" if op.get("x") else "", f"y={op['y']}" if op.get("y") else "",
                f"at={op['at']}" if op.get("at") else "", f"source={op['source']}" if op.get("source") else ""]))
            _append_directive(lines, sec, f"chart: {spec}")
            changes.append(label + f"chart added ({spec})")
        elif kind == "set_format":
            col, fmt = _need(op, "column"), _need(op, "format")
            _append_directive(lines, sec, f"format: {col}={fmt}")
            changes.append(label + f"{col} formatted as {fmt}")
        elif kind == "add_highlight":
            rule = _need(op, "rule")
            _append_directive(lines, sec, f"highlight: {rule}")
            changes.append(label + f"highlight rule '{rule}'")
        elif kind in ("add_scale", "add_databar"):
            col = _need(op, "column") + (" reverse" if op.get("reverse") else "")
            _append_directive(lines, sec, ("scale: " if kind == "add_scale" else "databar: ") + col)
            changes.append(label + f"{'colour scale' if kind == 'add_scale' else 'data bars'} on {col}")
        elif kind == "add_validation":
            rule = _need(op, "rule")
            _append_directive(lines, sec, f"validate: {rule}")
            changes.append(label + f"validation '{rule}'")
        elif kind == "set_total":
            col, fn = _need(op, "column"), _need(op, "fn")
            _append_directive(lines, sec, f"total: {col}={fn}")
            changes.append(label + f"total row: {fn} of {col}")
        elif kind == "set_freeze":
            _append_directive(lines, sec, f"freeze: {_need(op, 'cell')}")
            changes.append(label + f"freeze panes at {op['cell']}")
        elif kind == "set_table":
            _append_directive(lines, sec, f"table: {_need(op, 'name')}")
            changes.append(label + f"Excel table '{op['name']}'")
        elif kind == "set_filter":
            _append_directive(lines, sec, f"filter: {'on' if op.get('on', True) else 'off'}")
            changes.append(label + f"filter {'on' if op.get('on', True) else 'off'}")
        elif kind == "add_note":
            _append_directive(lines, sec, f"note: {_need(op, 'text')}")
            changes.append(label + "note added")
        elif kind == "add_column":
            name, formula = _need(op, "name"), _need(op, "formula")
            if not formula.startswith("="):
                raise EditError("add_column: formula must start with = and may use {r} for the row number (=C{r}*TaxRate)")
            rows = [i for i in range(sec[1], sec[2]) if lines[i].lstrip().startswith("|")]
            if len(rows) < 2:
                raise EditError(f"add_column: sheet '{sec[0]}' has no table")
            header = _split_row(lines[rows[0]])
            lines[rows[0]] = "| " + " | ".join(header + [name]) + " |"
            n = 0
            for i in rows[1:]:
                cells = _split_row(lines[i])
                if all(re.fullmatch(r":?-{2,}:?", c.strip()) or not c.strip() for c in cells):
                    continue                                              # the |---| separator
                n += 1
                lines[i] = "| " + " | ".join(cells + [formula.replace("{r}", str(n + 1))]) + " |"
            changes.append(label + f"column '{name}' added ({n} rows)")
        else:
            raise EditError(f"unknown workbook op '{kind}'")
    return "\n".join(lines), changes


def edit(outline: Optional[str], ops: list) -> tuple:
    """-> (xlsx bytes with cached values, new outline, changes). Raises EditError / BuildError with a reason for the model."""
    if not outline or not is_outline(outline):
        raise EditError("this workbook was not built from an outline here (or it was changed cell by cell since), so it cannot be "
                        "restyled - rebuild it with doc_create, or change cells with set_cell / set_formula / set_range")
    new_outline, changes = apply_outline_ops(outline, ops)
    try:
        data = build(new_outline)
    except BuildError as e:
        raise EditError(str(e))
    try:
        from . import calc
        data = calc.write_cached(data, calc.calculate(data)["values"])
    except Exception:
        pass
    parse(new_outline)
    return data, new_outline, changes



def expand_column_ops(data: bytes, ops: list) -> list:
    """Turn {op:'set_column', sheet?, column:'Cost' | 'C', values:[..] | formula:'=B{r}*Rate', rows?} into set_range ops aimed at
    exactly that column (found by its header text or its letter). Nothing else in the workbook is addressed."""
    if not any(str(o.get("op")).lower() == "set_column" for o in ops):
        return ops
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    out = []
    try:
        for op in ops:
            if str(op.get("op")).lower() != "set_column":
                out.append(op)
                continue
            col = _need(op, "column")
            want = str(op.get("sheet") or "").strip().lower()
            ws = next((w for w in wb.worksheets if w.title.lower() == want), None) if want else wb.worksheets[0]
            if ws is None:
                raise EditError(f"set_column: sheet '{op.get('sheet')}' not found (sheets: {', '.join(w.title for w in wb.worksheets)})")
            head_row, col_idx, headers = None, None, []
            for r, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True), start=1):
                cells = [str(c).strip() if c is not None else "" for c in row]
                if re.fullmatch(r"[A-Za-z]{1,3}", col) and not any(c.lower() == col.lower() for c in cells):
                    from openpyxl.utils import column_index_from_string
                    col_idx, head_row = column_index_from_string(col.upper()), 1
                    break
                hit = next((i for i, c in enumerate(cells, 1) if c and c.lower() == col.lower()), None) \
                    or next((i for i, c in enumerate(cells, 1) if c and col.lower() in c.lower()), None)
                if hit:
                    col_idx, head_row = hit, r
                    break
                if sum(1 for c in cells if c) >= 2 and not headers:
                    headers = [c for c in cells if c]
            if col_idx is None:
                raise EditError(f"set_column: no column '{col}' on sheet '{ws.title}'"
                                + (f" (its columns: {', '.join(headers[:15])})" if headers else ""))
            from openpyxl.utils import get_column_letter
            letter = get_column_letter(col_idx)
            q = f"'{ws.title}'" if re.search(r"[^A-Za-z0-9_]", ws.title) else ws.title
            if isinstance(op.get("values"), list):
                values = [[v] for v in op["values"]]
            elif op.get("formula"):
                formula = str(op["formula"])
                if not formula.startswith("="):
                    raise EditError("set_column: formula must start with = and may use {r} for the row number")
                n = int(op.get("rows") or max(1, (ws.max_row or head_row + 1) - head_row))
                values = [[formula.replace("{r}", str(head_row + 1 + i))] for i in range(n)]
            else:
                raise EditError("set_column: give values (a list, one per row) or a formula such as =B{r}*Rate")
            out.append({"op": "set_range", "addr": f"{q}!{letter}{head_row + 1}", "values": values,
                        "formulas": any(isinstance(v[0], str) and v[0].startswith("=") for v in values)})
    finally:
        wb.close()
    return out
