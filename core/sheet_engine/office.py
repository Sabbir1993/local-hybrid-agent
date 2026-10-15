"""Optional LibreOffice recalculation: the reference calculator when it is installed (no Docker, a plain local program).

Used only to COMPUTE values: the workbook is converted in a temp directory and its calculated cells are read back; our own file
(styles, charts, validation) is never replaced by LibreOffice's re-saved copy. Not installed, disabled, slow or failing -> None,
and the built-in calculator's answer stands.
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Optional

TIMEOUT_S = 60
_WINDOWS_PATHS = (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")
_ERRORS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


def enabled() -> bool:
    try:
        from ..small_model import APP_CONFIG
        return (APP_CONFIG.get("chat") or {}).get("sheet_libreoffice", True) is not False
    except Exception:
        return True


def find_soffice() -> Optional[str]:
    for name in ("soffice", "libreoffice"):
        hit = shutil.which(name)
        if hit:
            return hit
    for p in _WINDOWS_PATHS:
        if os.path.isfile(p):
            return p
    return None


def recalculate(data: bytes) -> Optional[dict]:
    """-> {(sheet, 'A1'): value} for every formula cell, as LibreOffice calculated them (errors as the error text), or None."""
    if not enabled():
        return None
    exe = find_soffice()
    if not exe:
        return None
    try:
        import openpyxl
        with tempfile.TemporaryDirectory(prefix="sheetcalc_") as d:
            src, out = Path(d) / "in.xlsx", Path(d) / "out"
            src.write_bytes(data)
            out.mkdir()
            # a private profile dir: never touches (or waits on) the user's own LibreOffice session
            profile = Path(d) / "profile"
            cmd = [exe, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", "--nolockcheck",
                   "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)]
            subprocess.run(cmd, check=True, timeout=TIMEOUT_S, capture_output=True, stdin=subprocess.DEVNULL)
            done = out / "in.xlsx"
            if not done.is_file():
                return None
            formulas = openpyxl.load_workbook(io.BytesIO(data))
            values = openpyxl.load_workbook(str(done), data_only=True)
            result = {}
            for ws in formulas.worksheets:
                if ws.title not in values.sheetnames:
                    continue
                vs = values[ws.title]
                for row in ws.iter_rows():
                    for c in row:
                        if c.data_type == "f":
                            result[(ws.title, c.coordinate)] = vs[c.coordinate].value
            return result
    except Exception as e:
        print(f"[xlsx] LibreOffice recalculation skipped ({type(e).__name__}: {e})", file=sys.stderr)
        return None


def is_error(v) -> bool:
    return isinstance(v, str) and v.startswith(_ERRORS)
