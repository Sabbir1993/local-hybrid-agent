"""Data cleaning: a list of simple operations over one sheet, written to a NEW workbook (the original is never touched).

Operations (each a dict): trim, dedupe{columns?}, drop_empty_rows, drop_empty_columns, fix_types{columns?}, rename{from,to},
fill{column, value | method:'down'}, case{column, to:'upper'|'lower'|'title'}, replace{column, find, with}, drop_columns{columns},
split{column, delimiter, into:[names]}. The result has a 'Changes' sheet that says what each step did, with counts.
"""
import datetime as _dt
import io
import re
from pathlib import Path
from typing import Optional

from ..file_tools.profile import column_stats, parse_number, read_tables

MAX_OPS = 20


class CleanError(ValueError):
    pass


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _col(headers: list, name) -> int:
    for i, h in enumerate(headers):
        if h.lower() == str(name or "").strip().lower():
            return i
    raise CleanError(f"column '{name}' not found (columns: {', '.join(headers[:15])})")


def _cols(headers: list, names) -> list:
    return [_col(headers, n) for n in names] if names else list(range(len(headers)))


def run(path, ops: list, sheet: Optional[str] = None) -> tuple:
    """-> (xlsx bytes, report text)."""
    tables = read_tables(Path(path))
    if not tables:
        raise CleanError("no data found in the file")
    t = next((x for x in tables if sheet and x.name.lower() == str(sheet).lower()), tables[0])
    if t.truncated:
        raise CleanError(f"the sheet has {t.total_rows:,} rows and only {len(t.rows):,} can be cleaned at once - "
                         "split the data or filter it first")
    if not isinstance(ops, list) or not ops:
        raise CleanError("ops required: a list like [{\"op\": \"trim\"}, {\"op\": \"dedupe\", \"columns\": [\"Email\"]}]")
    if len(ops) > MAX_OPS:
        raise CleanError(f"at most {MAX_OPS} operations per call")
    headers = list(t.headers)
    rows = [(r + [None] * len(headers))[: len(headers)] for r in t.rows]
    before = len(rows)
    log = []
    for op in ops:
        kind = str((op or {}).get("op") or "").lower()
        if kind == "trim":
            n = 0
            for r in rows:
                for i in _cols(headers, op.get("columns")):
                    if isinstance(r[i], str):
                        s = re.sub(r"\s+", " ", r[i]).strip()
                        if s != r[i]:
                            r[i], n = s, n + 1
            log.append(("trim", f"{n} cell(s) trimmed"))
        elif kind == "dedupe":
            idx = _cols(headers, op.get("columns"))
            seen, kept = set(), []
            for r in rows:
                key = tuple(str(r[i]).strip().lower() if r[i] is not None else "" for i in idx)
                if key in seen:
                    continue
                seen.add(key)
                kept.append(r)
            log.append(("dedupe", f"{len(rows) - len(kept)} duplicate row(s) removed"))
            rows = kept
        elif kind == "drop_empty_rows":
            kept = [r for r in rows if not all(_blank(v) for v in r)]
            log.append(("drop_empty_rows", f"{len(rows) - len(kept)} empty row(s) removed"))
            rows = kept
        elif kind == "drop_empty_columns":
            keep = [i for i in range(len(headers)) if any(not _blank(r[i]) for r in rows)]
            log.append(("drop_empty_columns", f"{len(headers) - len(keep)} empty column(s) removed"))
            headers, rows = [headers[i] for i in keep], [[r[i] for i in keep] for r in rows]
        elif kind == "fix_types":
            n = 0
            for i in _cols(headers, op.get("columns")):
                vals = [r[i] for r in rows]
                kind_ = column_stats(vals, headers[i])["kind"]
                for r in rows:
                    v = r[i]
                    if _blank(v):
                        continue
                    if kind_ == "number" and not isinstance(v, (int, float)):
                        num = parse_number(v)
                        if num is not None:
                            r[i], n = num, n + 1
                    elif kind_ == "date" and isinstance(v, str):
                        from .sql import _iso
                        iso = _iso(v)
                        if re.match(r"^\d{4}-\d{2}-\d{2}$", str(iso)):
                            r[i], n = _dt.date.fromisoformat(iso), n + 1
            log.append(("fix_types", f"{n} cell(s) converted to real numbers / dates"))
        elif kind == "rename":
            i = _col(headers, op.get("from"))
            headers[i] = str(op.get("to") or headers[i])
            log.append(("rename", f"{op.get('from')} -> {headers[i]}"))
        elif kind == "fill":
            i = _col(headers, op.get("column"))
            n, last = 0, None
            for r in rows:
                if _blank(r[i]):
                    fill = op.get("value") if op.get("method") != "down" else last
                    if fill is not None:
                        r[i], n = fill, n + 1
                else:
                    last = r[i]
            log.append(("fill", f"{n} empty cell(s) in {headers[i]} filled"))
        elif kind == "case":
            i = _col(headers, op.get("column"))
            fn = {"upper": str.upper, "lower": str.lower, "title": str.title}.get(str(op.get("to")).lower())
            if fn is None:
                raise CleanError("case: to must be upper, lower or title")
            for r in rows:
                if isinstance(r[i], str):
                    r[i] = fn(r[i])
            log.append(("case", f"{headers[i]} -> {op.get('to')}"))
        elif kind == "replace":
            i = _col(headers, op.get("column"))
            find, new = str(op.get("find") or ""), str(op.get("with") if op.get("with") is not None else "")
            if not find:
                raise CleanError("replace: find required")
            n = 0
            for r in rows:
                if isinstance(r[i], str) and find in r[i]:
                    r[i], n = r[i].replace(find, new), n + 1
            log.append(("replace", f"{n} cell(s) in {headers[i]} changed"))
        elif kind == "drop_columns":
            drop = {_col(headers, c) for c in (op.get("columns") or [])}
            keep = [i for i in range(len(headers)) if i not in drop]
            headers, rows = [headers[i] for i in keep], [[r[i] for i in keep] for r in rows]
            log.append(("drop_columns", f"{len(drop)} column(s) removed"))
        elif kind == "split":
            i = _col(headers, op.get("column"))
            into = [str(x) for x in (op.get("into") or [])]
            delim = str(op.get("delimiter") or " ")
            if len(into) < 2:
                raise CleanError("split: into needs at least two new column names")
            for r in rows:
                parts = str(r[i]).split(delim, len(into) - 1) if not _blank(r[i]) else []
                r.extend((parts + [None] * len(into))[: len(into)])
            headers += into
            log.append(("split", f"{headers[i]} split into {', '.join(into)}"))
        else:
            raise CleanError(f"unknown operation '{kind}' (trim, dedupe, drop_empty_rows, drop_empty_columns, fix_types, rename, "
                             "fill, case, replace, drop_columns, split)")
    return _write(t.name, headers, rows, log, before), f"{before:,} rows -> {len(rows):,}. " + "; ".join(f"{a}: {b}" for a, b in log)


def _write(name: str, headers: list, rows: list, log: list, before: int) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (name or "Clean")[:31]
    ws.append(headers)
    for r in rows:
        ws.append([None if _blank(v) else v for v in r])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for ci, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [min(len(str(r[ci - 1])), 40) for r in rows if r[ci - 1] is not None] or [8])
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = min(max(longest + 3, 9), 60)
        if any(isinstance(r[ci - 1], _dt.date) for r in rows):
            for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                for cell in row:
                    if isinstance(cell.value, _dt.date):
                        cell.number_format = "yyyy-mm-dd"
    if rows:
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(rows) + 1}"
    ch = wb.create_sheet("Changes")
    ch.append(["Step", "What changed"])
    ch.append(["rows before", before])
    for a, b in log:
        ch.append([a, b])
    ch.append(["rows after", len(rows)])
    for c in ch[1]:
        c.font = Font(bold=True)
    ch.column_dimensions["A"].width = 20
    ch.column_dimensions["B"].width = 60
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
