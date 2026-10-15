import base64
import contextvars
import re
import time
from pathlib import Path
from .. import companion_bridge, doc_ops, doc_store
from ..agent_tools import _common_resolve, _ws_resolve, require_device_workspace
from ..doc_ops.base import sha256
from . import builds, slide_images
from .location import (
    _Loc,
    _file_arg,
    _locate,
    _locate_common,
    _spec_for,
    _uid,
    _write,
    resolve_common,
)


def _inspect(loc: _Loc, data: bytes, args: dict) -> str:
    _row, spec = _spec_for(loc)
    text = doc_ops.inspect(data, loc.path.name, focus=args.get("focus"), source_spec=spec)
    where = "your device" if loc.location == "device" else "your files"
    return f"File: {loc.key} ({where})\n{text}"


async def _load_images(loc: _Loc, ops: list) -> None:
    """replace_image ops name an image file; load its bytes from the same place."""
    for op in ops:
        if str(op.get("op", "")).lower() == "replace_image" and op.get("image"):
            if loc.location == "device":
                d = await companion_bridge.call(loc.uid, "fs.read_b64", {"path": str(loc.sibling(Path(op["image"]).name))})
                op["_image_bytes"] = base64.b64decode(d["data"]) if d.get("data") else None
            else:
                p = resolve_common(op["image"])
                op["_image_bytes"] = p.read_bytes() if p else None


async def _edit(loc: _Loc, data: bytes, args: dict) -> str:
    if loc.path.suffix.lower() == ".pptx" and is_designer_deck(data):
        raise ValueError("this deck is in Designer style: its slides are pictures, so single texts cannot be edited. "
                         "Rebuild it with doc_create using the changed outline (same theme and `mode: designer`), or build "
                         "it in the Editable style.")
    ops = doc_ops.normalize_ops(args.get("ops"))
    await _load_images(loc, ops)
    row, spec = _spec_for(loc)
    out = doc_ops.edit(data, loc.path.name, ops, source_spec=spec)
    parent_id = row["id"] if row else doc_store.register(
        loc.uid, loc.key, loc.path.suffix.lower().lstrip("."), location=loc.location, sha256=sha256(data))
    version = (row["version"] if row else 1) + 1
    in_place = bool(args.get("in_place")) and loc.location == "device"
    target = loc.path if in_place else loc.sibling(doc_store.version_name(loc.path.name, version))
    await _write(loc, target, out.data)
    new_key = target.name if loc.location == "common" else str(target)
    doc_store.register(loc.uid, new_key, loc.path.suffix.lower().lstrip("."), location=loc.location,
                       parent_id=parent_id, source_spec=out.source_spec, sha256=sha256(out.data))
    doc_store.audit(loc.uid, "doc.edit", new_key, {"parent": loc.key, "ops": [o.get("op") for o in ops],
                                                   "sha256": sha256(out.data)})
    lines = [f"Edited {loc.path.name} -> saved as {target.name} (version {version}; the previous version is kept)."]
    lines += [f"- {c}" for c in out.changes]
    lines += [f"note: {n}" for n in out.notes]
    lines.append("Everything else in the document is unchanged (verified).")
    if loc.location == "common":
        lines.append(f"[DOWNLOAD: {target.name}]")
    return "\n".join(lines)


async def _create(args: dict, location: str) -> str:
    name = _file_arg(args)
    content = str(args.get("content") or args.get("spec") or "")
    template = None
    if args.get("template"):
        tloc, template = (await _locate(args["template"])) if location == "device" else _locate_common(args["template"])
    uid = _uid()
    if args.get("plan"):
        doc_ops._check_ops_for_pan([{"op": "create", "content": str(args["plan"])}])      # the plan is echoed back to the model
    ext_ = Path(name).suffix.lower()
    # a PDF built in pieces is printed on its last doc_append, so its pictures are looked up then
    content, images, pic_note = await prepare_build(ext_, content, uid, template=bool(template),
                                                    partial=bool(args.get("total")))
    data, spec = doc_ops.create(name, _interim(content, not args.get("total")), template=template, images=images)
    spec = content
    if location == "device":
        dev_uid, _ws = require_device_workspace()
        target = _ws_resolve(name)
        loc = _Loc("device", target, dev_uid)
        key = str(target)
    else:
        stem, suf = doc_store.base_stem(Path(name).name), Path(name).suffix.lower()
        target = _common_resolve(doc_store.version_name(f"{stem}{suf}", 1).replace("-v1-", "-"))
        loc = _Loc("common", target, uid)
        key = target.name
    await _write(loc, target, data)
    doc_store.register(uid, key, target.suffix.lower().lstrip("."), location=location, source_spec=spec,
                       sha256=sha256(data))
    msg = f"Created {target.name}. Later changes can be made with doc_edit without regenerating it."
    if pic_note:
        msg += " " + pic_note
    if not args.get("total"):                      # a deck built in pieces is checked once, on its last doc_append
        note = _layout_note(target.suffix.lower(), data, content)
        if note:
            msg += " " + note
    if args.get("total") and target.suffix.lower() in _APPENDABLE:
        try:
            bkey = (uid if location == "common" else dev_uid, key)
            _BUILD_TOTAL[bkey] = max(1, min(int(args["total"]), 30))
            plan = _parse_plan(args.get("plan"))
            if plan:
                _BUILD_PLAN[bkey] = plan
            builds.save(bkey[0], key, total=_BUILD_TOTAL[bkey], plan=plan or None, have=_unit_count(target.suffix.lower(), spec))
            msg += (f" It has {_unit_count(target.suffix.lower(), spec)} of {int(args['total'])}: add the rest in small "
                    f"steps with doc_append(file='{target.name}', content=<next 2-3 slides/sections>), then done=true.")
        except (TypeError, ValueError):
            pass
    return msg + (f" [DOWNLOAD: {target.name}]" if location == "common" else f" (saved on your device: {target})")


# Build plan: (uid, file key) -> [(title, purpose)], the model's own slide / section plan (doc_create plan=...). It is
# played back after every doc_append so a small model stays on it instead of drifting or repeating itself.
_BUILD_PLAN: dict = {}
_PLAN_LINE = re.compile(r"^\s*(?:\d+[.)]|[-*])\s*(.+?)\s*(?:(?:\s[-–—]\s+|:\s+)(.*))?$")


def _parse_plan(text: str) -> list:
    out = []
    for line in str(text or "").splitlines():
        m = _PLAN_LINE.match(line)
        if m and m.group(1).strip():
            out.append((re.sub(r"[*_`#]", "", m.group(1)).strip()[:80], (m.group(2) or "").strip()[:120]))
    return out[:30]


def _plan_for(key: tuple) -> list:
    if key not in _BUILD_PLAN:
        b = builds.load(key[0], str(key[1]))
        if b and b["plan"]:
            _BUILD_PLAN[key] = b["plan"]
    return _BUILD_PLAN.get(key) or []


def _total_for(key: tuple):
    if key not in _BUILD_TOTAL:
        b = builds.load(key[0], str(key[1]))
        if b and b["total"]:
            _BUILD_TOTAL[key] = b["total"]
    return _BUILD_TOTAL.get(key)


def _plan_note(key: tuple, have: int, built_titles: list, done: bool) -> str:
    plan = _plan_for(key)
    if not plan:
        return ""
    if done:
        have_l = " | ".join(t.lower() for t in built_titles)
        missing = [t for t, _ in plan if t.lower() not in have_l]
        _BUILD_PLAN.pop(key, None)
        builds.close(key[0], str(key[1]))
        return (f"Planned but not found in the deck: {', '.join(missing[:6])}." if missing else "")
    nxt = plan[have:have + 3]
    if not nxt:
        return ""
    return "Next in your plan: " + "; ".join(f"{have + i + 1}. {t}" + (f" - {p}" if p else "") for i, (t, p) in enumerate(nxt)) + "."


# Build-in-pieces bookkeeping: (uid, file key) -> slide/section count the user asked for. The dicts are a cache over
# builds.py (projects.db), so a restart no longer loses the count or the plan.
_BUILD_TOTAL: dict = {}
_APPENDABLE = (".pptx", ".docx", ".pdf")


# Style memory: the theme / accent / font of the last deck a user built, so the NEXT deck in the same working session looks
# like the last one instead of falling back to the default palette whenever the model leaves the `theme:` line out. An
# explicit `theme:` in a new deck always wins (and becomes the new memory). Forgotten after a few hours.
_LAST_STYLE: dict = {}
_STYLE_TTL_S = 8 * 3600


def _inherit_style(uid, ext: str, content: str) -> str:
    if ext != ".pptx" or uid is None or not _flag("deck_keep_style"):
        return content
    from ..doc_ops.pptx_ops import outline
    deck = outline.parse_outline(content)
    now = time.time()
    if deck.get("theme") or deck.get("hint"):
        t = outline.resolve_theme(deck)
        _LAST_STYLE[uid] = {"theme": t["name"], "accent": deck.get("accent", ""), "font": deck.get("font", ""),
                            "mode": deck.get("mode", ""), "ts": now}
        return content
    last = _LAST_STYLE.get(uid)
    if not last or now - last["ts"] > _STYLE_TTL_S:
        return content
    lines = [f"theme: {last['theme']}"]
    if last["accent"] and not deck.get("accent"):
        lines.append(f"accent: {last['accent']}")
    if last["font"] and not deck.get("font"):
        lines.append(f"font: {last['font']}")
    if last.get("mode") and not deck.get("mode"):
        lines.append(f"mode: {last['mode']}")
    return "\n".join(lines) + "\n" + content


async def prepare_build(ext: str, content: str, uid, *, template: bool = False, partial: bool = False) -> tuple:
    """What every route that writes a deck / PDF does first: keep the session's look, then look up the pictures its
    `image:` lines ask for. -> (content, images, note). Shared by doc_create and write_file, so a deck gets the same
    pictures, style and checks whichever tool the model happened to use."""
    if ext == ".pptx" and not _flag("deck_designer"):
        content = re.sub(r"(?mi)^[ \t]*mode[ \t]*:[ \t]*designer[ \t]*\r?\n?", "", content)
    if ext == ".pptx" and not template and _flag("deck_designer"):
        # the user's own answer to the questionnaire decides, not whether the model remembered the `mode:` line
        want = forced_mode.get()
        if want == "designer" and not _MODE_LINE.search(content):
            content = "mode: designer\n" + content
        elif want == "editable":
            content = _MODE_LINE.sub("", content)
    content = _inherit_style(uid, ext, content)
    images, note = {}, ""
    if uid is not None and _flag("deck_images") and slide_images.wants_images(content):
        if ext == ".pptx" and not template:
            images, notes, used_web = await slide_images.resolve(content, uid)
            note = slide_images.summary(images, notes, used_web)
        elif ext == ".pdf" and not partial:
            images, notes, used_web = await slide_images.resolve_pdf(content, uid)
            note = slide_images.summary(images, notes, used_web)
    return content, images, note


def _flag(name: str, default: bool = True) -> bool:
    """chat.<name> in config/app.json: kill switches for the optional build steps (deck_images, deck_checks)."""
    try:
        from ..small_model import APP_CONFIG
        return bool((APP_CONFIG.get("chat") or {}).get(name, default))
    except Exception:
        return default


forced_mode: contextvars.ContextVar = contextvars.ContextVar("deck_forced_mode", default="")   # "designer" | "editable" | ""
_MODE_LINE = re.compile(r"(?mi)^[ \t]*mode[ \t]*:[ \t]*designer[ \t]*\r?\n?")


def _interim(markdown: str, final: bool) -> str:
    """Designer pictures take seconds per rebuild, and doc_append rebuilds the whole deck every step. While a deck is still
    being built in pieces the file is drawn in the editable style (instant, and the live preview follows it); only the
    last step draws the Designer pictures. The stored source keeps its `mode:` line either way."""
    return markdown if final else _MODE_LINE.sub("", markdown)


def is_designer_deck(data: bytes) -> bool:
    """True for a deck built in Designer style: every slide is one full-slide picture (see pptx_ops/designer.py)."""
    try:
        import io
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        slides = list(prs.slides)
        return bool(slides) and all(len(s.shapes) == 1 and s.shapes[0].shape_type == 13 for s in slides)
    except Exception:
        return False


DESIGNER_NOTE = ("Designer style: the slides are pictures, so their text cannot be edited in place (it is in the speaker "
                 "notes and the alt text). To change something, rebuild the deck with doc_create and the changed outline, "
                 "or ask for the Editable style.")


def _layout_note(ext: str, data: bytes, markdown: str = "") -> str:
    """Layout check of a finished .pptx (text that overflows, low contrast, shapes off the slide...) as one sentence
    for the tool result; '' when clean, not a deck, or the check itself failed (it must never block a build)."""
    if ext == ".pptx" and is_designer_deck(data):
        return DESIGNER_NOTE                       # nothing to measure: the slides are pictures
    if not _flag("deck_checks"):
        return ""
    try:
        if ext == ".pptx":
            from ..doc_ops.pptx_ops import lint
            return lint.summary(lint.lint_deck(data))
        if ext == ".pdf":
            from ..agent_tools import pdf_check
            return pdf_check.render_note(data, markdown)
    except Exception:
        return ""
    return ""


def _built_titles(ext: str, spec: str) -> list:
    if ext == ".pptx":
        from ..doc_ops.pptx_ops import outline
        return [s["title"] for s in outline.parse_outline(spec)["slides"]]
    return re.findall(r"(?m)^#{1,3}\s+(.+?)\s*$", spec)


def _unit_count(ext: str, spec: str) -> int:
    """Slides in a deck, or top-level sections in a document, in the source markdown."""
    if ext == ".pptx":
        from ..doc_ops.pptx_ops import outline
        return len(outline.parse_outline(spec)["slides"])
    import re
    return len(re.findall(r"(?m)^#{1,2}\s+\S", spec))


def _leftover_markup(ext: str, spec: str) -> list:
    """Raw markdown that would show up as text in the finished file (checked on the final pass)."""
    import re
    if ext != ".pptx":
        return []
    from ..doc_ops.pptx_ops import outline
    bad = []
    for sl in outline.parse_outline(spec)["slides"]:
        for t in [sl["title"], sl.get("subtitle", "")] + [b.strip() for b in sl["bullets"]]:
            if re.search(r"\*\*|__|^#{1,6}\s|^>\s|^\[[ xX]\]", t or ""):
                bad.append(f"slide '{sl['title'][:40]}': {t[:50]}")
    return bad


async def _append(args: dict, location: str) -> str:
    """Add the next slides / sections to a document that is being built in pieces and rebuild it IN PLACE, so every
    step leaves a real, valid file (the right-hand preview refreshes from it). The markdown source is the truth:
    the final file is exactly what one write of the whole markdown would give."""
    name = _file_arg(args)
    content = str(args.get("content") or "").strip()
    done = bool(args.get("done"))
    if not content and not done:
        raise ValueError("content required (the next slides / sections as markdown)")
    loc, data = (await _locate(name)) if location == "device" else _locate_common(name)
    ext = loc.path.suffix.lower()
    if ext not in _APPENDABLE:
        raise ValueError(f"doc_append works on .pptx / .docx / .pdf, not {ext or 'this file'}")
    row, spec = _spec_for(loc)
    if not spec:
        raise ValueError("this file was not built here, so there is no source to extend - use doc_edit")
    key = (loc.uid, loc.key)
    if args.get("total"):
        try:
            _BUILD_TOTAL[key] = max(1, min(int(args["total"]), 30))
        except (TypeError, ValueError):
            pass
    new_spec = spec.rstrip()
    if content:
        new_spec += ("\n\n---\n\n" if ext == ".pptx" else "\n\n") + content
    have, total = _unit_count(ext, new_spec), _total_for(key)
    if done and total and have < total:
        raise ValueError(f"not finished: {have} of {total} {'slides' if ext == '.pptx' else 'sections'} so far - "
                         "keep adding with doc_append (done=true only after the last one)")
    images, pic_note = {}, ""
    pictures_on = _flag("deck_images")
    if pictures_on and ext == ".pptx" and slide_images.wants_images(new_spec):
        # pictures already found for earlier slides come from the resolver's cache, so a rebuild costs nothing extra
        images, pic_notes, used_web = await slide_images.resolve(new_spec, loc.uid)
        if slide_images.wants_images(content) or done:
            pic_note = slide_images.summary(images, pic_notes, used_web)
    elif pictures_on and ext == ".pdf" and done and slide_images.wants_images(new_spec):
        images, pic_notes, used_web = await slide_images.resolve_pdf(new_spec, loc.uid)
        pic_note = slide_images.summary(images, pic_notes, used_web)
    if ext == ".pdf" and not done:
        out = data              # the PDF is printed once at the end; the preview shows a live HTML draft of the spec
    else:
        out, _ = doc_ops.create(loc.path.name, _interim(new_spec, done), images=images)
    if out is not data:
        await _write(loc, loc.path, out)
    doc_store.register(loc.uid, loc.key, ext.lstrip("."), location=loc.location,
                       parent_id=row["id"] if row else None, source_spec=new_spec, sha256=sha256(out))
    if total and not done:
        builds.save(key[0], str(key[1]), total=total, have=have)
    unit = "slides" if ext == ".pptx" else "sections"
    msg = f"Added to {loc.path.name}: now {have} {unit}" + (f" of {total}." if total else ".")
    plan_note = _plan_note(key, have, _built_titles(ext, new_spec), done)
    if plan_note:
        msg += " " + plan_note
    if done:
        _BUILD_TOTAL.pop(key, None)
        builds.close(key[0], str(key[1]))
        bad = _leftover_markup(ext, new_spec)
        msg += " Finished" + (" - fix these with doc_edit set_text: " + "; ".join(bad[:5]) if bad else ".")
    elif total and have < total:
        msg += f" Continue with the next {min(3, total - have)} {unit} (doc_append), the last one is the closing."
    if pic_note:
        msg += " " + pic_note
    if done:
        note = _layout_note(ext, out, new_spec)
        if note:
            msg += " " + note
    if loc.location == "common":
        msg += f" [DOWNLOAD: {loc.path.name}]"
    return msg
