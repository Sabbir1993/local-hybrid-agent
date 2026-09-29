import re
from typing import Optional
from .constants import (
    R_NS,
    _open,
    _save,
)
from .text_ops import _set_frame_text


def parse_markdown_slides(md: str) -> list[dict]:
    slides, cur = [], {"title": "", "bullets": [], "notes": ""}

    def flush():
        if cur["title"] or cur["bullets"] or cur["notes"]:
            slides.append(dict(cur))

    slide_marker = re.compile(
        r'^(?:'
        r'---\s*$'
        r'|(?:\#{1,3}\s+(?P<h_title>.+))$'
        r'|(?:\*{0,2}\s*Slide\s+\d+\s*(?:\((?P<cat>[^)]+)\))?\s*:\*{0,2}\s*(?P<s_rest>.*))$'
        r')',
        re.IGNORECASE
    )

    for line in (md or "").strip().splitlines():
        s = line.rstrip()
        st = s.strip()
        if not st:
            continue
        if re.match(r'^\d+\s*slides\s*\+.*structure', st, re.IGNORECASE) or st.startswith("GLOBAL THEME:"):
            continue

        if st.startswith("---"):
            flush()
            cur = {"title": "", "bullets": [], "notes": ""}
            continue

        m = slide_marker.match(st)
        if m:
            h_title = m.group("h_title")
            cat = m.group("cat") or ""
            s_rest = m.group("s_rest") or ""

            if h_title:
                if cur["title"] or cur["bullets"]:
                    flush()
                    cur = {"title": "", "bullets": [], "notes": ""}
                cur["title"] = h_title.strip().lstrip("#").strip()
            else:
                flush()
                cur = {"title": "", "bullets": [], "notes": ""}
                m_title = re.search(r'(?:header(?:\s+banner)?|title)\s*(?:[^:]+:|\s*)?\s*\"([^\"]+)\"', s_rest, re.IGNORECASE)
                if m_title:
                    cur["title"] = m_title.group(1).strip()
                else:
                    q = re.findall(r'\"([^\"]+)\"', s_rest)
                    if cat:
                        cur["title"] = cat.strip()
                    elif q:
                        cur["title"] = q[0].strip()
                    elif s_rest:
                        cur["title"] = s_rest.strip()
                    else:
                        cur["title"] = "Slide"
            continue

        if st.lower().startswith(("notes:", "note:")):
            cur["notes"] = st.split(":", 1)[1].strip()
        elif st.startswith(("- ", "* ", "+ ")) or re.match(r"^\d+[.)]\s", st):
            indent = (len(s) - len(s.lstrip(" "))) // 2
            txt = re.sub(r"^([-*+]|\d+[.)])\s+", "", st)
            cur["bullets"].append("\t" * indent + txt)
        elif st:
            cur["bullets"].append(st.lstrip("#").strip())
    flush()
    return slides


def create(md: str, template: Optional[bytes] = None, title: str = "") -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    slides = parse_markdown_slides(md) or [{"title": title or "Presentation", "bullets": [], "notes": ""}]
    if template:
        prs = _open(template)
        lst = prs.slides._sldIdLst
        for el in list(lst):
            prs.part.drop_rel(el.get(f"{{{R_NS}}}id"))
            lst.remove(el)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layouts = list(prs.slide_layouts)

    def pick(title_only: bool):
        for l in layouts:
            kinds = {str(p.placeholder_format.type).split(".")[-1].split(" ")[0] for p in l.placeholders}
            if title_only and {"CENTER_TITLE", "SUBTITLE"} & kinds:
                return l
            if not title_only and "TITLE" in kinds and ({"BODY", "OBJECT"} & kinds):
                return l
        return layouts[1 if len(layouts) > 1 else 0]

    for i, sd in enumerate(slides):
        cover = i == 0 and len(sd["bullets"]) <= 1
        slide = prs.slides.add_slide(pick(cover))

        if not template:
            try:
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(11, 25, 44)
            except Exception:
                pass

        if cover:
            if slide.shapes.title is not None:
                slide.shapes.title.text_frame.text = sd["title"]
                if not template:
                    slide.shapes.title.left = Inches(1.2)
                    slide.shapes.title.top = Inches(2.2)
                    slide.shapes.title.width = Inches(10.9)
                    slide.shapes.title.height = Inches(1.5)
                    for p in slide.shapes.title.text_frame.paragraphs:
                        p.font.name = "Segoe UI"
                        p.font.size = Pt(38)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(255, 255, 255)
            body = next((p for p in slide.placeholders if p.placeholder_format.idx != 0), None)
            if body is not None:
                if sd["bullets"]:
                    _set_frame_text(body.text_frame._txBody, "\n".join(sd["bullets"]))
                    if not template:
                        body.left = Inches(1.2)
                        body.top = Inches(3.8)
                        body.width = Inches(10.9)
                        body.height = Inches(1.2)
                        for p in body.text_frame.paragraphs:
                            p.font.name = "Segoe UI"
                            p.font.size = Pt(18)
                            p.font.color.rgb = RGBColor(148, 163, 184)
                else:
                    body._element.getparent().remove(body._element)
        else:
            if slide.shapes.title is not None:
                slide.shapes.title.text_frame.text = sd["title"]
                if not template:
                    slide.shapes.title.left = Inches(0.8)
                    slide.shapes.title.top = Inches(0.6)
                    slide.shapes.title.width = Inches(11.733)
                    slide.shapes.title.height = Inches(0.9)
                    for p in slide.shapes.title.text_frame.paragraphs:
                        p.font.name = "Segoe UI"
                        p.font.size = Pt(28)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(255, 255, 255)
            body = next((p for p in slide.placeholders if p.placeholder_format.idx != 0), None)
            if body is not None:
                if sd["bullets"]:
                    _set_frame_text(body.text_frame._txBody, "\n".join(sd["bullets"]))
                    if not template:
                        body.left = Inches(0.8)
                        body.top = Inches(1.6)
                        body.width = Inches(11.733)
                        body.height = Inches(5.0)
                        for p in body.text_frame.paragraphs:
                            p.font.name = "Segoe UI"
                            p.font.size = Pt(16)
                            p.font.color.rgb = RGBColor(226, 232, 240)
                            p.space_after = Pt(12)
                else:
                    body._element.getparent().remove(body._element)

        if sd["notes"]:
            slide.notes_slide.notes_text_frame.text = sd["notes"]
    return _save(prs)
