"""Layout check for a finished deck: the closest thing to looking at it when no PowerPoint renderer is available.

It reads the saved .pptx back and measures what the builder drew with the same font metrics the builder used:
text that does not fit its box, shapes outside the slide, text boxes sitting on top of each other, text too small to
read, and text whose colour has too little contrast against the fill behind it. Decorative shapes (named "deco*") are
skipped, so only content is judged. Findings are short sentences for the tool result; an empty list means clean.
"""
import io
from typing import Optional

from pptx import Presentation
from pptx.util import Emu

from . import measure as M

EMU_IN = 914400
MIN_FONT_PT = 11.0
BOUNDS_TOL = 0.06          # inches a shape may poke out before it is reported
OVERFLOW_TOL = 1.06        # estimated text height may exceed the box by 6% (measurement is not pixel exact)
OVERLAP_FRACTION = 0.2     # of the smaller box
MAX_FINDINGS = 40


def _lum(hex_: str) -> float:
    def ch(i):
        v = int(hex_[i:i + 2], 16) / 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    return 0.2126 * ch(0) + 0.7152 * ch(2) + 0.0722 * ch(4)


def contrast(a: str, b: str) -> float:
    la, lb = _lum(a), _lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def _is_deco(shape) -> bool:
    return str(shape.name or "").lower().startswith("deco")


def _rgb_of(fill) -> Optional[str]:
    try:
        if fill.type == 1:                       # MSO_FILL.SOLID
            return str(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def _run_props(p):
    """(font name, size pt, bold, rgb or None) of a paragraph's first run (what the builder sets on every run)."""
    for r in p.runs:
        f = r.font
        rgb = None
        try:
            if f.color is not None and f.color.type is not None:
                rgb = str(f.color.rgb)
        except Exception:
            rgb = None
        return (f.name or "Calibri", float(f.size.pt) if f.size else 18.0, bool(f.bold), rgb)
    return None


def _text_need_in(shape) -> Optional[float]:
    """Estimated height (inches) the shape's text needs inside its own width, or None when it has no text."""
    tf = shape.text_frame
    width = shape.width / EMU_IN - ((tf.margin_left or 0) + (tf.margin_right or 0)) / EMU_IN
    need, any_text = 0.0, False
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs).strip()
        props = _run_props(p)
        if not text or props is None:
            continue
        any_text = True
        name, size, bold, _ = props
        indent = 0.0
        ppr = p._p.pPr
        if ppr is not None and ppr.get("marL"):
            indent = int(ppr.get("marL")) / EMU_IN
        lines = M.lines_for(text, name, size, max(width - indent, 0.3), bold)
        need += lines * size * M.line_height([text]) / 72
        if p.space_after is not None:
            need += p.space_after.pt / 72
    if not any_text:
        return None
    need += ((tf.margin_top or 0) + (tf.margin_bottom or 0)) / EMU_IN
    return need


def _box(shape):
    return (shape.left / EMU_IN, shape.top / EMU_IN, (shape.left + shape.width) / EMU_IN, (shape.top + shape.height) / EMU_IN)


def _overlap(a, b) -> float:
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return max(w, 0) * max(h, 0)


def lint_deck(data: bytes) -> list:
    """[{slide, shape, kind, message}] for the saved deck (empty when it is clean)."""
    prs = Presentation(io.BytesIO(data))
    sw, sh = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    out: list = []

    def add(i, shape, kind, msg):
        if len(out) < MAX_FINDINGS:
            out.append({"slide": i, "shape": shape, "kind": kind, "message": f"slide {i}: {msg}"})

    for i, slide in enumerate(prs.slides, 1):
        bg = None
        try:
            bg = _rgb_of(slide.background.fill)
        except Exception:
            bg = None
        text_shapes = []
        for shp in slide.shapes:
            if _is_deco(shp):
                continue
            name = shp.name or "shape"
            l, t, r, b = _box(shp)
            if l < -BOUNDS_TOL or t < -BOUNDS_TOL or r > sw + BOUNDS_TOL or b > sh + BOUNDS_TOL:
                if shp.shape_type != 13 or (shp.name or "") not in ("Background image", "Cover image"):
                    add(i, name, "bounds", f"'{name}' extends outside the slide")
            if shp.has_text_frame and shp.text_frame.text.strip():
                text_shapes.append(shp)
                need = _text_need_in(shp)
                if need is not None and need > (shp.height / EMU_IN) * OVERFLOW_TOL + 0.02:
                    add(i, name, "overflow", f"text in '{name}' needs about {need:.1f} in but the box is "
                                              f"{shp.height / EMU_IN:.1f} in")
                fill = _rgb_of(shp.fill) if hasattr(shp, "fill") else None
                back = fill or bg
                for p in shp.text_frame.paragraphs:
                    props = _run_props(p)
                    if not props or not "".join(r.text for r in p.runs).strip():
                        continue
                    _, size, _bold, rgb = props
                    if size < MIN_FONT_PT:
                        add(i, name, "small-text", f"text in '{name}' is {size:g} pt (minimum {MIN_FONT_PT:g})")
                        break
                    if rgb and back:
                        need_ratio = 3.0 if size >= 18 else 4.5
                        if contrast(rgb, back) < need_ratio:
                            add(i, name, "contrast", f"text in '{name}' has low contrast "
                                                       f"({contrast(rgb, back):.1f}:1 on #{back})")
                            break
            elif shp.has_table:
                tbl = shp.table
                total = 0.0
                for r_i, row in enumerate(tbl.rows):
                    rh = row.height / EMU_IN
                    row_need = 0.0
                    for c_i, cell in enumerate(row.cells):
                        text = cell.text_frame.text.strip()
                        props = _run_props(cell.text_frame.paragraphs[0])
                        if not text or not props:
                            continue
                        name_f, size, bold, rgb = props
                        cw = tbl.columns[c_i].width / EMU_IN - 0.3
                        n = M.lines_for(text, name_f, size, max(cw, 0.3), bold)
                        row_need = max(row_need, n * size * M.line_height([text]) / 72 + 0.2)
                        if size < MIN_FONT_PT:
                            add(i, name, "small-text", f"table text is {size:g} pt (minimum {MIN_FONT_PT:g})")
                        cell_fill = _rgb_of(cell.fill)
                        if rgb and cell_fill and contrast(rgb, cell_fill) < (3.0 if size >= 18 else 4.5):
                            add(i, name, "contrast", f"table text has low contrast ({contrast(rgb, cell_fill):.1f}:1)")
                    total += max(rh, row_need)
                if t + total > sh - 0.25:
                    add(i, name, "overflow", f"table '{name}' would run past the bottom of the slide "
                                               f"({t + total:.1f} in of {sh:.1f})")
        # text boxes sitting on top of each other
        for a in range(len(text_shapes)):
            for b in range(a + 1, len(text_shapes)):
                sa, sb = text_shapes[a], text_shapes[b]
                ba, bb = _box(sa), _box(sb)
                small = min((ba[2] - ba[0]) * (ba[3] - ba[1]), (bb[2] - bb[0]) * (bb[3] - bb[1]))
                if small > 0 and _overlap(ba, bb) / small > OVERLAP_FRACTION:
                    add(i, sa.name, "overlap", f"'{sa.name}' and '{sb.name}' overlap")
    return out


def summary(findings: list, limit: int = 6) -> str:
    """One line for the tool result ('' when clean)."""
    if not findings:
        return ""
    shown = "; ".join(f["message"] for f in findings[:limit])
    more = f" (+{len(findings) - limit} more)" if len(findings) > limit else ""
    return f"Layout check found {len(findings)} issue(s): {shown}{more}."
