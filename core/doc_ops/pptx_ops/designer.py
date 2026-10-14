"""Designer mode: slides drawn as HTML / CSS, printed by Edge / Chrome, and placed in the deck as full-slide pictures.

Native PowerPoint shapes (design.py) are editable but limited to what PowerPoint draws with the fonts the viewer has.
Designer mode trades editing for the look of a hand-designed page: real typographic control, SVG illustrations, gradients.
The pipeline reuses what the app already has for PDFs:

    outline -> one HTML page per slide (16:9) -> headless browser prints a PDF -> pypdfium2 rasterises each page ->
    JPEG -> a .pptx whose slides are those pictures, with the slide text kept in the speaker notes and in the alt text.

`build()` returns the .pptx bytes, or None when a step is not available (no browser, no rasteriser, a failed print): the
caller then builds the normal editable deck instead, so asking for Designer never costs the user their deck.
"""
import base64
import datetime
import html
import io
import math
import os
import sys
import tempfile
from pathlib import Path
from typing import Optional

from . import design as D
from . import measure as M
from . import outline as O

SCALE = 3.0                     # PDF points -> pixels: 13.333 in x 72 x 3 = 2880 x 1620
JPEG_QUALITY = 86
W, H, MX = D.W, D.H, D.MX
esc = html.escape

_CSS = """
@page { size: __W__in __H__in; margin: 0 }
* { box-sizing: border-box; margin: 0; padding: 0 }
html, body { width: __W__in; background: #fff; -webkit-print-color-adjust: exact; print-color-adjust: exact }
.slide { position: relative; width: __W__in; height: __H__in; overflow: hidden; break-after: page; page-break-after: always;
         background: var(--bg); color: var(--text); font-family: '__BODY__', 'Segoe UI', 'Nirmala UI', Arial, sans-serif }
.slide:last-child { break-after: auto; page-break-after: auto }
.abs { position: absolute }
.head { font-family: '__HEAD__', Georgia, 'Nirmala UI', serif }
.title { font-weight: 700; line-height: 1.12; letter-spacing: -0.01em; display: flex; align-items: center }
.eyebrow { position: absolute; font-weight: 700; letter-spacing: .18em; text-transform: uppercase; font-size: 11pt; color: var(--eyebrow) }
.rule { position: absolute; left: __MX__in; top: 1.55in; width: .9in; height: .07in; background: var(--accent) }
.foot { position: absolute; top: 6.98in; height: .3in; display: flex; align-items: center; font-size: 10pt; color: var(--muted) }
ul.bl { list-style: none }
ul.bl li { position: relative; padding-left: .34in; margin-bottom: 16pt; line-height: 1.28 }
ul.bl li::before { content: ''; position: absolute; left: .06in; top: .5em; width: .09in; height: .09in; border-radius: 50%; background: var(--accent) }
ul.bl li.sub { padding-left: .78in; margin-bottom: 6pt; color: var(--muted) }
ul.bl li.sub::before { left: .5in; width: .07in; height: .03in; border-radius: 0; top: .62em }
.card { background: var(--surface); border: 1px solid var(--line); border-radius: 10px; padding: .26in .28in; overflow: hidden }
.card.hi { background: var(--text); color: var(--bg); border-color: var(--text) }
.card .k { font-weight: 700; line-height: 1.1 }
.card p, .card .t { line-height: 1.32; opacity: .9 }
.credit { position: absolute; font-size: 9pt; font-style: italic; color: var(--muted) }
table.tb { border-collapse: collapse; width: 100% }
table.tb th { background: var(--accent); color: var(--onaccent); text-align: left; padding: .1in .15in; font-weight: 700 }
table.tb td { padding: .1in .15in; border-bottom: 1px solid var(--line); background: var(--surface) }
table.tb tr:nth-child(even) td { background: var(--bg) }
img.pic { width: 100%; height: 100%; object-fit: cover; display: block; border-radius: 10px }
svg { display: block }
"""


def _box(x, y, w, h, extra=""):
    return f'style="left:{x:.3f}in;top:{y:.3f}in;width:{w:.3f}in;height:{h:.3f}in;{extra}"'


def _uri(data: bytes) -> str:
    mime = "image/jpeg" if data[:3] == b"\xff\xd8\xff" else "image/png"
    return f"data:{mime};base64," + base64.b64encode(data).decode("ascii")


def _img(data: bytes, x, y, w, h, radius=True, cls="") -> str:
    r = "" if radius else "border-radius:0;"
    return f'<div class="abs {cls}" {_box(x, y, w, h)}><img class="pic" style="{r}" src="{_uri(data)}" alt=""/></div>'


def _vars(t: dict, cover: bool = False) -> str:
    bg = t["cover_bg"] if cover else t["bg"]
    text = t["cover_text"] if cover else t["text"]
    muted = t["cover_muted"] if cover else t["muted"]
    surface = t["surface"]
    return (f"--bg:#{bg};--text:#{text};--muted:#{D._ink(muted, bg, 4.5)};--surface:#{surface};--line:#{t['line']};"
            f"--accent:#{t['accent']};--accent2:#{t['accent2']};--onaccent:#{D._on(t['accent'])};"
            f"--eyebrow:#{D._ink(t['accent'], bg, 4.5)}")


def _pt(texts, font, w, h, base, floor, gap=0.0, bold=False) -> float:
    return M.fit_size([str(x) for x in texts], font, w, h, base, floor, gap, bold)


def _title(t, text, eyebrow="", color=None) -> str:
    y, h = (0.6, 0.9) if eyebrow else (0.5, 1.0)
    size = _pt([text], t["head"], W - 2 * MX, h, 38, 22, 0, True)
    out = f'<div class="eyebrow" {_box(MX, 0.3, 10.0, 0.3, "display:flex;align-items:center")}>{esc(eyebrow)}</div>' if eyebrow else ""
    c = f"color:{color};" if color else ""
    return out + (f'<h1 class="abs title head" {_box(MX, y, W - 2 * MX, h, f"font-size:{size}pt;{c}")}>{esc(text)}</h1>'
                  '<div class="rule"></div>')


def _foot(text, num) -> str:
    return (f'<div class="foot" style="left:{MX}in;width:{W - 2 * MX - 1.0}in">{esc(text)}</div>'
            f'<div class="foot" style="left:{W - MX - 0.8}in;width:.8in;justify-content:flex-end;font-weight:700">{num}</div>')


def _credit(s, x, y, w, color="") -> str:
    spec = s.get("image") or {}
    bits = [b for b in (spec.get("caption"), s.get("image_credit")) if b]
    if not bits or not s.get("image_data"):
        return ""
    c = f"color:{color};" if color else ""
    return f'<div class="credit" {_box(x, y, w, 0.3, "text-align:right;" + c)}>{esc("  ·  ".join(bits))}</div>'


# ------------------------------------------------------------------ art (SVG, drawn for the cover)
def _art_svg(kind: str, t: dict) -> str:
    a, b = f"#{t['accent']}", f"#{t['accent2'] if D._contrast(t['accent2'], t['cover_bg']) >= 2.2 else t['cover_muted']}"
    w, h = 490, 540
    out = []
    base = 350

    def arch(x, wd, ht, sw):
        r = wd / 2
        return (f'<path d="M{x},{base} V{base - ht + r} A{r},{r} 0 0 1 {x + wd},{base - ht + r} V{base}" fill="none" '
                f'stroke="{a}" stroke-width="{sw}"/>')

    def wave(y, col, amp=11, sw=1.6):
        seg = w / 8
        d = f"M0,{y} " + " ".join(f"Q{seg * (2 * i + 1):.0f},{y - amp if i % 2 == 0 else y + amp} {seg * (2 * i + 2):.0f},{y}" for i in range(4))
        return f'<path d="{d}" fill="none" stroke="{col}" stroke-width="{sw}"/>'
    if kind == "arches":
        for x, wd, ht in ((10, 140, 190), (150, 190, 300), (360, 120, 190)):
            out += [arch(x, wd, ht, 2.4), arch(x + 26, wd - 52, ht - 40, 1.4)]
        out += [wave(base + 40 + i * 34, b) for i in range(4)]
    elif kind == "waves":
        out += [wave(80 + i * 62, a if i % 2 == 0 else b, 18, 2) for i in range(8)]
    elif kind == "skyline":
        for i, ht in enumerate((130, 230, 170, 300, 200, 250, 150)):
            x = 8 + i * 68
            out.append(f'<rect x="{x}" y="{base - ht}" width="56" height="{ht}" fill="none" stroke="{a}" stroke-width="2.2"/>')
            out += [f'<line x1="{x + 14}" y1="{base - ht + 24 + r * 34}" x2="{x + 42}" y2="{base - ht + 24 + r * 34}" stroke="{b}" stroke-width="2"/>'
                    for r in range(int((ht - 30) // 34))]
        out += [f'<line x1="0" y1="{base}" x2="{w}" y2="{base}" stroke="{a}" stroke-width="2"/>'] + [wave(base + 40 + i * 34, b) for i in range(3)]
    elif kind == "geometric":
        out += [f'<rect x="0" y="{base - 190}" width="{w}" height="190" fill="{t["cover_muted"]}" fill-opacity=".14"/>',
                f'<circle cx="120" cy="{base - 100}" r="82" fill="{a}"/>',
                f'<polygon points="250,{base - 20} 330,{base - 190} 410,{base - 20}" fill="{b}"/>',
                f'<rect x="395" y="{base - 130}" width="70" height="70" fill="{a}"/>'] + [wave(base + 40 + i * 34, b) for i in range(3)]
    elif kind == "network":
        pts = [(70, 90), (240, 40), (420, 120), (140, 260), (340, 290), (60, 430), (270, 470), (440, 430)]
        for p, q in ((0, 1), (1, 2), (0, 3), (1, 3), (1, 4), (2, 4), (3, 4), (3, 5), (4, 6), (4, 7), (5, 6), (6, 7)):
            out.append(f'<line x1="{pts[p][0]}" y1="{pts[p][1]}" x2="{pts[q][0]}" y2="{pts[q][1]}" stroke="{b}" stroke-width="1.8"/>')
        out += [f'<circle cx="{x}" cy="{y}" r="{10 if i % 3 else 17}" fill="{a if i % 2 == 0 else b}"/>' for i, (x, y) in enumerate(pts)]
    return f'<svg viewBox="0 0 {w} {h}" width="4.9in" height="5.4in" xmlns="http://www.w3.org/2000/svg">{"".join(out)}</svg>'


# ------------------------------------------------------------------ slide bodies
def _bullets_html(items, t, x, y, w, h, base=26, floor=13) -> str:
    size = _pt([i.strip() for i in items], t["body"], w - 0.4, h, base, floor, 16)
    lis = "".join(f'<li class="{"sub" if i.startswith(chr(9)) else ""}">{esc(i.strip())}</li>' for i in items)
    return f'<ul class="abs bl" {_box(x, y, w, h, f"font-size:{size}pt;overflow:hidden")}>{lis}</ul>'


def _chart_html(s, t, x, y, w, h) -> str:
    from ...agent_tools.pdf_charts import chart_svg
    pal = [t["accent"], t["accent2"], "F59E0B", "10B981", "EC4899", "8B5CF6", "06B6D4", "84CC16"]
    style = {"ink": f"#{t['text']}", "ax": f"#{t['muted']}", "grid": f"#{t['line']}", "bg": f"#{t['bg']}",
             "palette": [f"#{c}" for c in pal]}
    svg = chart_svg(s["chart_data"], style).replace("<svg ", f'<svg style="max-height:{h}in" ', 1)
    return f'<div class="abs" {_box(x, y, w, h, "display:flex;align-items:center")}>{svg}</div>'


def _table_html(rows, t, x, y, w, h) -> str:
    plan = D._plan_table(rows, t["body"], total_w=w, max_h=h)[0]
    rows, size, widths, _hs = plan
    cols = "".join(f'<col style="width:{cw / w * 100:.1f}%">' for cw in widths)
    head = "".join(f"<th>{esc(c)}</th>" for c in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>" for r in rows[1:])
    return (f'<div class="abs" {_box(x, y, w, h)}><table class="tb" style="font-size:{size}pt"><colgroup>{cols}</colgroup>'
            f'<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')


def _kind_html(s: dict, ts: dict) -> str:
    kind = s["kind"]
    top = O.groups(s["bullets"])
    if kind == "bullets":
        items = s["bullets"]
        if s.get("image_data"):
            px = 7.55
            return (_bullets_html(items, ts, MX + 0.3, 1.95, px - MX - 0.7, 4.7, 26 if len(items) <= 4 else 22, 12)
                    + _img(s["image_data"], px, 1.95, W - MX - px, 4.45) + _credit(s, px, 6.42, W - MX - px))
        return _bullets_html(items, ts, MX + 0.3, 1.95, W - 2 * MX - 0.3, 4.7, 30 if len(items) <= 4 else 26, 13)
    if kind == "columns":
        half = math.ceil(len(top) / 2)
        out = ""
        cw = (W - 2 * MX - 0.6) / 2
        for i, part in enumerate((top[:half], top[half:])):
            lines = []
            for head, kids in part:
                lines.append(head)
                lines += ["\t" + k for k in kids]
            out += _bullets_html(lines, ts, MX + i * (cw + 0.6), 1.95, cw, 4.7, 21, 12)
        return out
    if kind in ("cards", "stats", "timeline"):
        if kind == "stats":
            items = [(O.stat_split(g[0]) or (g[0], "")) for g in top]
        else:
            items = [D._heading_body(g) for g in top]
        n = len(items)
        cols = n if (kind != "cards" or n <= 3) else 3
        if kind == "cards" and n == 4:
            cols = 2
        rows = math.ceil(n / cols)
        gap = 0.3
        cw = (W - 2 * MX - gap * (cols - 1)) / cols
        ch = min((4.7 - gap * (rows - 1)) / rows, 3.6 if kind == "cards" else 3.7)
        cells = []
        # one size for every number on the slide (each fitted to its own card would make them all different)
        common = min([M.fit_one_line(x[0], ts["head"], cw - 0.7, 58, 24, True) for x in items], default=40) if kind == "stats" else 0
        for i, it in enumerate(items):
            hi = kind in ("stats", "timeline") and i == n - 1 and n >= 3
            if kind == "stats":
                num, cap = it
                ns = common
                col = D._ink(ts["accent2"], ts["text"] if hi else ts["surface"], 3.0)
                cs = _pt([cap], ts["body"], cw - 0.6, ch - 1.9, 18, 12)
                cells.append(f'<div class="card{" hi" if hi else ""}" style="min-height:{min(ch, 2.4)}in"><div class="k head" style="font-size:{ns}pt;'
                             f'color:#{col}">{esc(num)}</div><div class="t" style="font-size:{cs}pt;margin-top:.14in">{esc(cap)}</div></div>')
            else:
                head, kids = it
                hs = _pt([head], ts["head"], cw - 0.56, 0.9, 24 if kind == "timeline" else 21, 14, 0, True)
                bs = _pt(kids, ts["body"], cw - 0.56, ch - 1.35, 18, 11, 6)
                hc = D._ink(ts["accent2"], ts["surface"], 3.0)
                body = "".join(f"<p style='margin-top:6pt'>{esc(k)}</p>" for k in kids)
                cells.append(f'<div class="card" style="min-height:{min(ch, 1.6)}in"><div class="k head" style="font-size:{hs}pt;color:#{hc}">{esc(head)}</div>'
                             f'<div class="t" style="font-size:{bs}pt">{body}</div></div>')
        return (f'<div class="abs" {_box(MX, 1.95, W - 2 * MX, 4.75, f"display:grid;grid-template-columns:repeat({cols},1fr);gap:{gap}in;align-content:start;align-items:stretch")}>'
                + "".join(cells) + "</div>")
    if kind == "quote":
        items = [b.strip() for b in s["bullets"]]
        quote, by = items[0].strip('"\u201c\u201d'), " \u2014 ".join(items[1:])
        size = _pt([quote], ts["head"], W - 2 * MX - 1.2, 3.0, 36, 18)
        acc = D._ink(ts["accent"], ts["bg"], 4.5)
        out = (f'<div class="abs" {_box(MX + 0.2, 2.3, 0.08, 3.4, "background:var(--accent)")}></div>'
               f'<div class="abs head" {_box(MX + 0.7, 2.2, W - 2 * MX - 1.2, 3.0, f"font-size:{size}pt;font-style:italic;display:flex;align-items:center")}>'
               f'\u201c{esc(quote)}\u201d</div>')
        if by:
            out += f'<div class="abs" {_box(MX + 0.7, 5.4, W - 2 * MX - 1.2, 0.5, f"font-size:16pt;font-weight:700;color:#{acc}")}>{esc(by)}</div>'
        return out
    if kind == "table":
        return _table_html(s["table"], ts, MX, 1.95, W - 2 * MX, D.TABLE_MAX_H)
    if kind == "chart":
        return _chart_html(s, ts, MX, 1.95, W - 2 * MX, 4.7)
    if kind == "textside":
        lw, gap = 5.8, 0.5
        rx = MX + lw + gap
        rw = W - MX - rx
        left = _bullets_html(s["bullets"], ts, MX, 1.95, lw, 4.7, 21, 12)
        right = _chart_html(s, ts, rx, 1.95, rw, 4.6) if s.get("chart_data") else _table_html(s["table"], ts, rx, 2.1, rw, 4.4)
        return left + right
    if kind == "comparison":
        gap = 0.9
        cw = (W - 2 * MX - gap) / 2
        out = ""
        for i, (head, kids) in enumerate(top[:2]):
            x = MX + i * (cw + gap)
            col = ts["accent"] if i == 0 else ts["accent2"]
            size = _pt(kids, ts["body"], cw - 0.9, 3.7, 22, 12, 12)
            hs = _pt([head], ts["head"], cw - 0.6, 0.75, 26, 14, 0, True)
            lis = "".join(f"<li>{esc(k)}</li>" for k in kids)
            out += (f'<div class="abs card" {_box(x, 1.95, cw, 4.75, "padding:0")}>'
                    f'<div style="background:#{col};color:#{D._on(col)};height:.85in;display:flex;align-items:center;padding:0 .3in;font-weight:700;'
                    f'font-size:{hs}pt" class="head">{esc(head.rstrip(":").strip())}</div>'
                    f'<ul class="bl" style="padding:.3in;font-size:{size}pt">{lis}</ul></div>')
        return out + (f'<div class="abs" {_box(W / 2 - 0.42, 1.95 + 0.425 - 0.42, 0.84, 0.84, "border-radius:50%;background:var(--bg);border:2px solid var(--muted);display:flex;align-items:center;justify-content:center;font-weight:700;font-size:16pt")}>VS</div>')
    if kind == "agenda":
        items = [h.rstrip(":").strip() for h, _ in top if h][:8]
        n = len(items)
        cols = 1 if n <= 4 else 2
        rows = math.ceil(n / cols)
        cw = (W - 2 * MX - (0.6 if cols == 2 else 0)) / cols
        rh = min(1.0, 4.6 / rows)
        out = ""
        for i, text in enumerate(items):
            c, r = divmod(i, rows)
            col = ts["accent"] if i % 2 == 0 else ts["accent2"]
            size = _pt([text], ts["body"], cw - 1.0, rh - 0.1, 26, 14)
            out += (f'<div class="abs" {_box(MX + c * (cw + 0.6), 2.0 + r * rh, cw, rh, "display:flex;align-items:center;border-bottom:1px solid var(--line)")}>'
                    f'<div style="width:.62in;height:.62in;border-radius:50%;background:#{col};color:#{D._on(col)};display:flex;align-items:center;'
                    f'justify-content:center;font-weight:700;font-size:18pt;flex:none" class="head">{i + 1}</div>'
                    f'<div style="margin-left:.23in;font-size:{size}pt">{esc(text)}</div></div>')
        return out
    if kind == "gallery":
        pics = [s["image_data"]] + [p["data"] for p in (s.get("image_extra_data") or [])]
        pics = pics[:3]
        x0, y0, w, h, gap = MX, 1.95, W - 2 * MX, 4.45, 0.3
        if len(pics) == 1:
            boxes = [(x0, y0, w, h)]
        elif len(pics) == 2:
            bw = (w - gap) / 2
            boxes = [(x0, y0, bw, h), (x0 + bw + gap, y0, bw, h)]
        else:
            lw = w * 0.58
            rw, rh = w - lw - gap, (h - gap) / 2
            boxes = [(x0, y0, lw, h), (x0 + lw + gap, y0, rw, rh), (x0 + lw + gap, y0 + rh + gap, rw, rh)]
        return "".join(_img(p, *b) for p, b in zip(pics, boxes)) + _credit(s, x0, y0 + h + 0.05, w)
    return _bullets_html(s["bullets"], ts, MX + 0.3, 1.95, W - 2 * MX - 0.3, 4.7)


def _slide_html(s: dict, t: dict, deck_title: str, i: int, sec_no: int, cn: int) -> tuple:
    """-> (html, sec_no, cn) for one slide."""
    kind = s["kind"]
    if kind in ("cover", "closing"):
        has_pic = bool(s.get("image_data")) and kind == "cover"
        art = "" if kind == "closing" or has_pic else (s.get("art") or t.get("art") or "")
        art = "" if art == "none" else art
        tw = 6.6 if has_pic else (6.8 if art else 10.4)
        size = _pt([s["title"]], t["head"], tw, 2.2, 66, 28, 0, True)
        sub = s["subtitle"]
        if kind == "closing" and s["bullets"]:
            sub = "\n".join(x for x in [sub] + [b.strip() for b in s["bullets"]] if x)
        ss = _pt(sub.split("\n"), t["body"], tw, 1.5, 22, 13, 4) if sub else 20
        eb = s.get("eyebrow", "")
        body = (f'<div class="eyebrow" {_box(1.1, 1.3, tw, 0.4, "display:flex;align-items:center")}>{esc(eb)}</div>' if eb else "")
        body += f'<h1 class="abs title head" {_box(1.1, 1.75, tw, 2.5, f"font-size:{size}pt;align-items:flex-end;color:var(--text)")}>{esc(s["title"])}</h1>'
        body += f'<div class="abs" {_box(1.1, 4.42, 1.2, 0.07, "background:var(--accent)")}></div>'
        if sub:
            body += (f'<div class="abs" {_box(1.1, 4.7, tw, 1.5, f"font-size:{ss}pt;line-height:1.35;color:var(--muted);white-space:pre-line")}>{esc(sub)}</div>')
        if kind == "cover":
            body += f'<div class="abs" {_box(1.1, 6.45, 6, 0.3, "font-size:11pt;color:var(--muted);letter-spacing:.04em")}>Prepared {datetime.date.today().strftime("%B %Y")}</div>'
        body += f'<div class="abs" {_box(0, 0, 0.28, H, "background:var(--accent)")}></div>'
        if has_pic:
            body += _img(s["image_data"], 8.0, 0, W - 8.0, H, radius=False) + _credit(s, 8.1, 7.1, W - 8.3, "#fff")
        elif art:
            body += f'<div class="abs" {_box(7.9, 0.9, 4.9, 5.4)}>{_art_svg(art, t)}</div>'
        return f'<section class="slide" style="{_vars(t, True)}">{body}</section>', sec_no, cn
    if kind == "section" and not s.get("image_data"):
        sec_no += 1
        size = _pt([s["title"]], t["head"], 9.0, 1.5, 46, 24, 0, True)
        num_col = D._ink(t["accent"], t["cover_bg"], 3.0)
        body = (f'<div class="abs head" {_box(1.1, 2.45, 3.0, 0.5, f"font-size:24pt;font-weight:700;color:#{num_col}")}>{sec_no:02d}</div>'
                f'<div class="abs" {_box(1.1, 3.0, 1.2, 0.07, "background:var(--accent)")}></div>'
                f'<h1 class="abs title head" {_box(1.1, 3.2, 9.0, 1.5, f"font-size:{size}pt;align-items:flex-start;color:var(--text)")}>{esc(s["title"])}</h1>'
                + (f'<div class="abs" {_box(1.1, 4.8, 9.0, 1.0, "font-size:20pt;color:var(--muted)")}>{esc(s["subtitle"])}</div>' if s["subtitle"] else "")
                + f'<div class="abs" {_box(0, 0, 0.28, H, "background:var(--accent)")}></div>')
        return f'<section class="slide" style="{_vars(t, True)}">{body}</section>', sec_no, cn
    if kind == "full" or kind == "section":
        size = _pt([s["title"]], t["head"], 10.5, 2.0, 50, 26, 0, True)
        body = (f'<img class="abs" style="left:0;top:0;width:{W}in;height:{H}in;object-fit:cover" src="{_uri(s["image_data"])}" alt=""/>'
                f'<div class="abs" style="left:0;top:0;width:{W}in;height:{H}in;background:linear-gradient(90deg,rgba(0,0,0,.72),rgba(0,0,0,.25))"></div>'
                f'<h1 class="abs title head" {_box(1.1, 2.4, 10.5, 2.0, f"font-size:{size}pt;align-items:flex-end;color:#fff")}>{esc(s["title"])}</h1>'
                f'<div class="abs" {_box(1.1, 4.55, 1.2, 0.07, "background:var(--accent)")}></div>'
                + (f'<div class="abs" {_box(1.1, 4.8, 10.5, 1.2, "font-size:20pt;color:#E5E7EB")}>{esc(s["subtitle"])}</div>' if s["subtitle"] else "")
                + _credit(s, W - MX - 6, 7.05, 6, "#E5E7EB"))
        return f'<section class="slide" style="{_vars(t, True)}">{body}</section>', sec_no, cn
    ts = D._variant(t, cn)
    cn += 1
    src = s.get("source") or ""
    foot = (src if src.lower().startswith(("source", "sources")) else f"Source: {src}") if src else deck_title
    body = _title(ts, s["title"], s.get("eyebrow", "")) + _kind_html(s, ts) + _foot(foot, i + 1)
    return f'<section class="slide" style="{_vars(ts)}">{body}</section>', sec_no, cn


def deck_html(deck: dict) -> str:
    t = O.resolve_theme(deck)
    css = (_CSS.replace("__W__", f"{W}").replace("__H__", f"{H}").replace("__MX__", f"{MX}")
           .replace("__BODY__", t["body"]).replace("__HEAD__", t["head"]))
    sections, sec_no, cn = [], 0, 0
    for i, s in enumerate(D._prepare(deck["slides"], t)):
        try:
            h, sec_no, cn = _slide_html(s, t, deck["title"], i, sec_no, cn)
        except Exception as e:                                 # one bad slide becomes a plain page, never a lost deck
            print(f"[designer] slide {i + 1} ({s.get('kind')}) failed: {type(e).__name__}: {e}", file=sys.stderr)
            h = (f'<section class="slide" style="{_vars(t)}">{_title(t, s.get("title", ""))}'
                 f'{_bullets_html(s.get("bullets", []), t, MX + 0.3, 1.95, W - 2 * MX - 0.3, 4.7)}</section>')
        sections.append((s, h))
    pages = "".join(h for _s, h in sections)
    return f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{esc(deck["title"])}</title><style>{css}</style></head><body>{pages}</body></html>'


def _slide_text(s: dict) -> str:
    bits = [s.get("title", "")] + [b.strip() for b in s.get("bullets", [])] + [" | ".join(r) for r in s.get("table", [])]
    return " - ".join(b for b in bits if b)


def build(deck: dict) -> Optional[bytes]:
    """The deck as picture slides (see the module docstring), or None when it cannot be made here."""
    try:
        import pypdfium2 as pdfium
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches
        from ...agent_tools import pdf_render
    except Exception as e:
        print(f"[designer] unavailable ({type(e).__name__}: {e})", file=sys.stderr)
        return None
    if not pdf_render._find_chromium_binary():
        print("[designer] no Chrome / Edge to print with", file=sys.stderr)
        return None
    t = O.resolve_theme(deck)
    slides = D._prepare(deck["slides"], t)
    with tempfile.TemporaryDirectory(prefix="designer_") as d:
        pdf_path = Path(d) / "deck.pdf"
        if not pdf_render._render_html_to_pdf(deck_html(deck), pdf_path):
            print("[designer] the browser could not print the slides", file=sys.stderr)
            return None
        pdf = pdfium.PdfDocument(str(pdf_path))
        try:
            if len(pdf) != len(slides):
                print(f"[designer] {len(pdf)} pages for {len(slides)} slides", file=sys.stderr)
                return None
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(W), Inches(H)
            blank = next(l for l in prs.slide_layouts if l.name == "Blank")
            for i, s in enumerate(slides):
                im = pdf[i].render(scale=SCALE).to_pil().convert("RGB")
                buf = io.BytesIO()
                im.save(buf, format="JPEG", quality=JPEG_QUALITY, optimize=True)
                buf.seek(0)
                slide = prs.slides.add_slide(blank)
                pic = slide.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
                pic.name = f"Slide {i + 1}"
                pic._element.nvPicPr.cNvPr.set("descr", _slide_text(s)[:600])
                notes = s.get("notes", "")
                text = f"{_slide_text(s)}"
                slide.notes_slide.notes_text_frame.text = ((notes + "\n\n") if notes else "") + "Slide text: " + text
            out = io.BytesIO()
            prs.save(out)
            return out.getvalue()
        finally:
            pdf.close()
