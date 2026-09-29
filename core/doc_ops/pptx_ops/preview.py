from typing import Optional
from .constants import (
    PREVIEW_MAX_IMAGE,
    PREVIEW_MAX_IMAGES,
    _PREVIEW_IMAGE_TYPES,
    _a,
    _clip,
    _hex,
    _open,
    _shape_kind,
)


def _solid_rgb(fill_parent) -> Optional[str]:
    """'#rrggbb' of an explicit solid fill, else None (theme/gradient/none)."""
    try:
        f = fill_parent.fill
        from pptx.enum.dml import MSO_FILL
        if f.type == MSO_FILL.SOLID and f.fore_color.type is not None:
            return _hex(f.fore_color.rgb)
    except Exception:
        pass
    return None


def _bg_rgb(cSld_owner) -> Optional[str]:
    try:
        clr = cSld_owner._element.xpath("./p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val")
        return _hex(clr[0]) if clr else None
    except Exception:
        return None


def _run_style(font) -> dict:
    st = {}
    try:
        if font.size:
            st["size"] = font.size.pt
    except Exception:
        pass
    for k in ("bold", "italic"):
        try:
            if getattr(font, k):
                st[k] = True
        except Exception:
            pass
    try:
        if font.color and font.color.type is not None:
            st["color"] = _hex(font.color.rgb)
    except Exception:
        pass
    return st


def _frame_paras(tf) -> list:
    out = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            if r.text:
                runs.append({"text": r.text.replace("\v", "\n"), **_run_style(r.font)})
        para = {"level": p.level, "runs": runs}
        try:
            if p.alignment is not None:
                para["align"] = str(p.alignment).split(".")[-1].split(" ")[0].lower()
        except Exception:
            pass
        ps = _run_style(p.font)
        if ps:
            para["style"] = ps
        out.append(para)
    return out


def _preview_shapes(shapes, tx, budget: list, skip_placeholders: bool = False) -> list:
    """Flatten shapes into positioned boxes. tx maps a shape's (x, y, w, h) EMU
    into slide EMU (identity at top level, group child-space transform inside groups)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []
    for sh in shapes:
        if skip_placeholders and sh.is_placeholder:
            continue
        try:
            x, y, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if x is None or w is None:
            continue
        box = tx(x, y or 0, w, h or 0)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                xfrm = sh._element.xpath("./p:grpSpPr/a:xfrm")[0]
                off, ext = xfrm.find(_a("off")), xfrm.find(_a("ext"))
                choff, chext = xfrm.find(_a("chOff")), xfrm.find(_a("chExt"))
                cx, cy = int(choff.get("x")), int(choff.get("y"))
                cw, ch = max(int(chext.get("cx")), 1), max(int(chext.get("cy")), 1)
                gx, gy, gw, gh = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))

                def inner(ix, iy, iw, ih, _o=tx, gx=gx, gy=gy, gw=gw, gh=gh, cx=cx, cy=cy, cw=cw, ch=ch):
                    sx, sy = gw / cw, gh / ch
                    return _o(gx + (ix - cx) * sx, gy + (iy - cy) * sy, iw * sx, ih * sy)
                out.extend(_preview_shapes(sh.shapes, inner, budget))
            except Exception:
                pass
            continue
        item = {"x": box[0], "y": box[1], "w": box[2], "h": box[3], "kind": _shape_kind(sh)}
        try:
            if sh.rotation:
                item["rot"] = sh.rotation
        except Exception:
            pass
        if item["kind"] == "picture":
            try:
                img = sh.image
                blob = img.blob
                if len(blob) <= PREVIEW_MAX_IMAGE and budget[0] + len(blob) <= PREVIEW_MAX_IMAGES \
                        and img.content_type in _PREVIEW_IMAGE_TYPES:
                    import base64
                    budget[0] += len(blob)
                    item["src"] = f"data:{img.content_type};base64," + base64.b64encode(blob).decode()
            except Exception:
                pass
            out.append(item)
            continue
        if item["kind"] == "table":
            rows = []
            for row in sh.table.rows:
                rows.append([_clip(c.text, 300) for c in row.cells])
            item["rows"] = rows
            out.append(item)
            continue
        if item["kind"] == "chart":
            try:
                ch = sh.chart
                item["title"] = ch.chart_title.text_frame.text if ch.has_title else ""
            except Exception:
                item["title"] = ""
            out.append(item)
            continue
        fill = _solid_rgb(sh) if hasattr(sh, "fill") else None
        if fill:
            item["fill"] = fill
        try:
            if sh.line.fill.type is not None and sh.line.color.type is not None:
                item["line"] = _hex(sh.line.color.rgb)
        except Exception:
            pass
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            item["paras"] = _frame_paras(sh.text_frame)
            try:
                anchor = sh.text_frame.vertical_anchor
                if anchor is not None:
                    item["anchor"] = str(anchor).split(".")[-1].split(" ")[0].lower()
            except Exception:
                pass
        if item.get("paras") or fill or item.get("line"):
            out.append(item)
    return out


def preview(data: bytes) -> dict:
    """Positioned, render-ready slide model: coordinates in EMU, fonts in pt."""
    prs = _open(data)
    w, h = int(prs.slide_width or 12192000), int(prs.slide_height or 6858000)
    ident = lambda x, y, cw, ch: (int(x), int(y), int(cw), int(ch))  # noqa: E731
    budget = [0]
    slides = []
    for slide in prs.slides:
        layout = slide.slide_layout
        master = layout.slide_master
        bg = _bg_rgb(slide) or _bg_rgb(layout) or _bg_rgb(master)
        shapes = []
        for owner in (master, layout):
            try:
                shapes.extend(_preview_shapes(owner.shapes, ident, budget, skip_placeholders=True))
            except Exception:
                pass
        shapes.extend(_preview_shapes(slide.shapes, ident, budget))
        notes = ""
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame
            notes = nt.text.strip() if nt is not None else ""
        slides.append({"bg": bg, "shapes": shapes, "notes": notes})
    return {"width": w, "height": h, "slides": slides}
