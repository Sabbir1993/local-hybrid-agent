import io
import re
from typing import Optional
from ..base import DocOpError, read_zip, require

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
RT_NOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
RT_LAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
MAX_SLIDES = 500
MAX_TABLE_CELLS_SHOWN = 400

PREVIEW_MAX_IMAGE = 3 * 1024 * 1024     # per picture
PREVIEW_MAX_IMAGES = 24 * 1024 * 1024   # whole deck
_HEX_RX = re.compile(r"^[0-9A-Fa-f]{6}$")
_PREVIEW_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/bmp", "image/webp"}

_ADDR = re.compile(r"^s(\d+)(?:/(title|notes|sh(\d+)))?(?:/(?:p(\d+)|tbl[\[(](\d+)\s*,\s*(\d+)[\])]))?$", re.I)


def _a(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


def _open(data: bytes):
    from pptx import Presentation
    read_zip(data)  # size / bomb / macro guards before python-pptx parses it
    prs = Presentation(io.BytesIO(data))
    require(len(prs.slides) <= MAX_SLIDES, f"presentation has more than {MAX_SLIDES} slides")
    return prs


def _save(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _clip(s: str, n: int = 400) -> str:
    s = s.replace("\v", " ").strip()
    return s if len(s) <= n else s[:n] + "..."


def _hex(val) -> Optional[str]:
    v = str(val or "")
    return "#" + v if _HEX_RX.match(v) else None


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)


def _shape_kind(sh) -> str:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if sh.is_placeholder:
        try:
            t = sh.placeholder_format.type
            name = str(t).split(".")[-1].split(" ")[0].lower()
            return "title" if name in ("title", "center_title") else f"placeholder:{name}"
        except Exception:
            return "placeholder"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    st = sh.shape_type
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if st == MSO_SHAPE_TYPE.TEXT_BOX:
        return "textbox"
    return "shape"


def _parse(addr: str):
    m = _ADDR.match(str(addr).strip().replace(" ", ""))
    if not m:
        raise DocOpError(f"bad address '{addr}' (use s3, s3/sh5, s3/sh5/p2, s3/sh5/tbl[1,2], s3/title, s3/notes)")
    si, part, shid, pi, tr, tc = m.groups()
    return int(si), (part or "").lower(), int(shid) if shid else None, int(pi) if pi else None, \
        (int(tr), int(tc)) if tr else None


def _slide_at(prs, n: int):
    require(1 <= n <= len(prs.slides), f"slide {n} does not exist (deck has {len(prs.slides)})")
    return prs.slides[n - 1]


def _find_shape(slide, shape_id: int):
    for sh in _iter_shapes(slide.shapes):
        if sh.shape_id == shape_id:
            return sh
    raise DocOpError(f"no shape sh{shape_id} on that slide - run doc_inspect for valid addresses")


def _title_shape(slide):
    t = slide.shapes.title
    if t is not None:
        return t
    texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    if not texts:
        return None
    return min(texts, key=lambda s: (s.top or 0, s.left or 0))


def _body_shape(slide, exclude):
    cands = [s for s in slide.shapes if s.has_text_frame and s.shape_id != exclude.shape_id]
    if not cands:
        return None
    ph = [s for s in cands if s.is_placeholder]
    pool = ph or cands
    return max(pool, key=lambda s: (s.width or 0) * (s.height or 0))
