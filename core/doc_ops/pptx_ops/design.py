"""Renders a parsed deck (outline.parse_outline) as a designed 16:9 presentation with python-pptx.

Titles and bullet bodies stay real placeholders (so doc_inspect/doc_edit address them like any deck);
everything decorative is a shape named "deco*" that inspection skips.
"""
import io
import math
import re
import sys
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from . import measure as M
from .lint import contrast as _contrast
from . import outline as O
from .text_ops import _set_frame_text

W, H = 13.333, 7.5
MX = 0.7                      # side margin
DECO = "deco"


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


def _lum(h: str) -> float:
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _on(h: str) -> str:
    """Readable text colour on top of fill h: whichever of near-black / white has the higher WCAG contrast."""
    return "111111" if _contrast("111111", h) >= _contrast("FFFFFF", h) else "FFFFFF"


def _ink(fg: str, bg: str, ratio: float = 4.5) -> str:
    """fg nudged toward black (light background) or white (dark one) until it reads on bg at `ratio`:1, so an accent
    colour used as TEXT stays on-brand without vanishing on a light card."""
    if _contrast(fg, bg) >= ratio:
        return fg
    target = (0, 0, 0) if _contrast("000000", bg) >= _contrast("FFFFFF", bg) else (255, 255, 255)
    r, g, b = (int(fg[i:i + 2], 16) for i in (0, 2, 4))
    for step in range(1, 21):
        k = step / 20
        mixed = "%02X%02X%02X" % (round(r + (target[0] - r) * k), round(g + (target[1] - g) * k),
                                  round(b + (target[2] - b) * k))
        if _contrast(mixed, bg) >= ratio:
            return mixed
    return "%02X%02X%02X" % target


def _fit(texts: list, width: float, height: float, base: float, floor: float, gap_pt: float = 6.0,
         font: str = "Calibri", bold: bool = False) -> float:
    """Largest font size (pt) in [floor, base] at which the texts fit width x height, measured on the real font."""
    return M.fit_size([str(t) for t in texts], font, width, height, base, floor, gap_pt, bold)


def _back(shape) -> None:
    el = shape._element
    tree = el.getparent()
    tree.remove(el)
    tree.insert(2, el)


def _deco(slide, kind, x, y, w, h, fill, alpha=None, back=True, name=DECO):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if alpha is not None:
        clr = sh.fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 1000)))
    if back:
        _back(sh)
    return sh


def _bg(slide, hex_):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(hex_)


def _style_runs(p, font, size, color, bold=False, italic=False):
    for r in p.runs:
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = _rgb(color)


def _frame(tf, anchor=MSO_ANCHOR.TOP, margins=(0.0, 0.0, 0.0, 0.0)):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Inches(m) for m in margins)


def _place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def _textbox(slide, x, y, w, h, paras, name=DECO + "-text", anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: [(text, font, size, color, bold, italic, space_after_pt)]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    _frame(tb.text_frame, anchor)
    for i, (text, font, size, color, bold, italic, after) in enumerate(paras):
        p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
        p.text = text
        p.alignment = align
        p.space_after = Pt(after)
        _style_runs(p, font, size, color, bold, italic)
    return tb


def _bullet(p, color, level, glyph="•"):
    ppr = p._p.get_or_add_pPr()
    marl, ind = (Inches(0.34), -Inches(0.34)) if level == 0 else (Inches(0.78), -Inches(0.28))
    ppr.set("marL", str(int(marl)))
    ppr.set("indent", str(int(ind)))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    etree.SubElement(etree.SubElement(ppr, qn("a:buClr")), qn("a:srgbClr")).set("val", color)
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(ppr, qn("a:buChar")).set("char", glyph if level == 0 else "–")


def _fill_bullets(shape, items: list, t: dict, x, y, w, h, base=22, floor=13):
    _place(shape, x, y, w, h)
    _frame(shape.text_frame, MSO_ANCHOR.TOP)
    _set_frame_text(shape.text_frame._txBody, "\n".join(items))
    size = _fit([i.strip() for i in items], w - 0.4, h, base, floor, gap_pt=16, font=t["body"])
    for p in shape.text_frame.paragraphs:
        level = p._p.pPr.get("lvl") if p._p.pPr is not None else None
        lvl = 1 if level and level != "0" else 0
        p.space_after = Pt(16 if lvl == 0 else 6)
        p.alignment = PP_ALIGN.LEFT
        _style_runs(p, t["body"], size if lvl == 0 else max(size - 3, floor), t["text"] if lvl == 0 else _ink(t["muted"], t["bg"], 4.5))
        _bullet(p, t["accent"], lvl)


def _title(slide, text, t, x=MX, y=0.5, w=W - 2 * MX, h=1.0, color=None, base=32, bold=True,
           anchor=MSO_ANCHOR.MIDDLE):
    sh = slide.shapes.title
    sh.text_frame.text = text
    _place(sh, x, y, w, h)
    _frame(sh.text_frame, anchor)
    size = _fit([text], w, h, base, 20, gap_pt=0, font=t["head"], bold=bold)
    for p in sh.text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        _style_runs(p, t["head"], size, color or t["text"], bold)
    return sh


def _drop_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx != 0 and ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def _chrome(slide, t, title, deck_title, num, eyebrow=""):
    _bg(slide, t["bg"])
    if eyebrow:                                           # small spaced capitals above the title: "AT A GLANCE"
        tb = _textbox(slide, MX, 0.3, 10.0, 0.3,
                      [(eyebrow, t["body"], 11, _ink(t["accent"], t["bg"], 4.5), True, False, 0)],
                      name=DECO + "-eyebrow", anchor=MSO_ANCHOR.MIDDLE)
        for r in tb.text_frame.paragraphs[0].runs:
            r._r.get_or_add_rPr().set("spc", "180")
        _title(slide, title, t, y=0.6, h=0.9)
    else:
        _title(slide, title, t)
    _deco(slide, MSO_SHAPE.RECTANGLE, MX, 1.55, 0.9, 0.07, t["accent"])


def _fit_image(data: bytes, w: float, h: float, fit: str = "cover"):
    """(bytes, left, top, width, height) in inches: 'cover' crops the picture to the w x h box, 'contain' keeps it whole
    and centres it. Re-encoded (JPEG for photos), so the deck stays small."""
    from PIL import Image
    im = Image.open(io.BytesIO(data))
    im.load()
    has_alpha = im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info)
    im = im.convert("RGBA" if has_alpha else "RGB")
    iw, ih = im.size
    box = w / h
    if fit == "cover":
        if iw / ih > box:
            nw = int(ih * box)
            im = im.crop(((iw - nw) // 2, 0, (iw - nw) // 2 + nw, ih))
        else:
            nh = int(iw / box)
            im = im.crop((0, (ih - nh) // 2, iw, (ih - nh) // 2 + nh))
        dx, dy, dw, dh = 0.0, 0.0, w, h
    else:
        scale = min(w / iw, h / ih)
        dw, dh = iw * scale, ih * scale
        dx, dy = (w - dw) / 2, (h - dh) / 2
    out = io.BytesIO()
    if has_alpha:
        im.save(out, format="PNG")
    else:
        im.save(out, format="JPEG", quality=85, optimize=True)
    return out.getvalue(), dx, dy, dw, dh


def _picture(slide, s, x, y, w, h, name="Image"):
    """Add the slide's resolved picture in the (x, y, w, h) box; returns the shape (None when there is no picture)."""
    data = s.get("image_data")
    if not data:
        return None
    spec = s.get("image") or {}
    blob, dx, dy, dw, dh = _fit_image(data, w, h, spec.get("fit", "cover"))
    pic = slide.shapes.add_picture(io.BytesIO(blob), Inches(x + dx), Inches(y + dy), Inches(dw), Inches(dh))
    pic.name = name
    alt = spec.get("caption") or spec.get("query") or ""
    if alt:
        pic._element.nvPicPr.cNvPr.set("descr", alt[:200])
    return pic


def _credit(slide, s, t, x, y, w, color=None, align=PP_ALIGN.RIGHT):
    """Small source line under / over a picture: the caption, then 'Image: host' for web pictures."""
    spec = s.get("image") or {}
    bits = [b for b in (spec.get("caption"), s.get("image_credit")) if b]
    if bits and s.get("image_data"):
        _textbox(slide, x, y, w, 0.3, [("  ·  ".join(bits), t["body"], 9, color or t["muted"], False, True, 0)],
                 name=DECO + "-credit", anchor=MSO_ANCHOR.MIDDLE, align=align)


def _apply_theme_part(prs, t: dict) -> None:
    """Write the deck's palette and fonts into the presentation's theme, so PowerPoint's colour / font pickers, Reset and
    any shape the user adds later use the deck's design instead of the stock Office one."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    if scheme is None:
        return
    darker, lighter = sorted([t["bg"], t["text"]], key=_lum)
    dk2, lt2 = sorted([t["cover_bg"], t["surface"]], key=_lum)
    extra = ["F59E0B", "10B981", "EC4899", "8B5CF6"]
    colors = {"dk1": darker, "lt1": lighter, "dk2": dk2, "lt2": lt2, "accent1": t["accent"], "accent2": t["accent2"],
              "accent3": extra[0], "accent4": extra[1], "accent5": extra[2], "accent6": extra[3],
              "hlink": t["accent"], "folHlink": t["muted"]}
    scheme.set("name", t.get("label", "Deck")[:30])
    for name, val in colors.items():
        el = scheme.find(f"a:{name}", ns)
        if el is None:
            continue
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, qn("a:srgbClr")).set("val", val.upper())
    for tag, face in (("majorFont", t["head"]), ("minorFont", t["body"])):
        el = root.find(f".//a:fontScheme/a:{tag}/a:latin", ns)
        if el is not None:
            el.set("typeface", face)
    fs = root.find(".//a:fontScheme", ns)
    if fs is not None:
        fs.set("name", f"{t['head']} / {t['body']}"[:40])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _widen_master(prs) -> None:
    """The stock template is 4:3 (10 in wide); the deck is 16:9. Stretch the master's and layouts' placeholders to the
    wide page, so PowerPoint's 'Reset' / 'Layout' put titles and footers where the design expects them."""
    k = W / 10.0
    for holder in [prs.slide_master, *prs.slide_layouts]:
        for shp in holder.shapes:
            try:
                shp.left, shp.width = int(shp.left * k), int(shp.width * k)
            except Exception:
                continue


def _footer_placeholders(slide, t: dict, deck_title: str, num: int) -> None:
    """Footer text and slide number as REAL placeholders (named deco-*, so inspection skips them): the number is a live
    field, so it stays right when a slide is inserted or moved in PowerPoint."""
    spTree = slide.shapes._spTree
    next_id = max([int(i) for i in spTree.xpath("//p:cNvPr/@id")] + [1]) + 1

    def sp(kind: str, idx: int, name: str, x: float, y: float, w: float, h: float, algn: str, body: str, bold: bool):
        return etree.fromstring(
            f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<p:nvSpPr><p:cNvPr id="{next_id + idx}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph type="{kind}" sz="quarter" idx="{idx}"/></p:nvPr></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{int(Inches(x))}" y="{int(Inches(y))}"/><a:ext cx="{int(Inches(w))}" cy="{int(Inches(h))}"/></a:xfrm></p:spPr>'
            f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/>'
            f'<a:p><a:pPr algn="{algn}"/>{body}</a:p></p:txBody></p:sp>')

    def run_props(size=1000, b=False):
        return (f'<a:rPr lang="en-US" sz="{size}" b="{1 if b else 0}"><a:solidFill><a:srgbClr val="{t["muted"]}"/></a:solidFill>'
                f'<a:latin typeface="{t["body"]}"/></a:rPr>')
    from xml.sax.saxutils import escape
    footer = sp("ftr", 11, "deco-footer", MX, 6.98, W - 2 * MX - 1.0, 0.3, "l",
                f'<a:r>{run_props()}<a:t>{escape(deck_title)}</a:t></a:r>', False)
    number = sp("sldNum", 12, "deco-number", W - MX - 0.8, 6.98, 0.8, 0.3, "r",
                f'<a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">{run_props(1000, True)}<a:t>{num}</a:t></a:fld>', True)
    spTree.append(footer)
    spTree.append(number)


def _gallery(slide, s, t):
    """2-3 pictures on one slide (`layout: gallery` with 2-3 `image:` lines): two side by side, or one large and two small."""
    pics = [{"data": s["image_data"], "spec": s.get("image") or {}}] + list(s.get("image_extra_data") or [])
    pics = [p for p in pics if p.get("data")][:3]
    x0, y0, w, h, gap = MX, 1.95, W - 2 * MX, 4.45, 0.3
    if len(pics) == 1:
        boxes = [(x0, y0, w, h)]
    elif len(pics) == 2:
        bw = (w - gap) / 2
        boxes = [(x0, y0, bw, h), (x0 + bw + gap, y0, bw, h)]
    else:
        lw = w * 0.58
        rw, rh = w - lw - gap, (h - gap) / 2
        boxes = [(x0, y0, lw, h), (x0 + lw + gap, y0, rw, rh), (x0 + lw + gap, y0 + rh + gap, rw, rh)]
    for k, (p, (x, y, bw, bh)) in enumerate(zip(pics, boxes)):
        _picture(slide, {"image_data": p["data"], "image": p.get("spec") or s.get("image") or {}}, x, y, bw, bh, name=f"Image {k + 1}")
    hosts = []
    for h_ in [s.get("image_credit", "")] + [p.get("credit_host", "") for p in s.get("image_extra_data") or []]:
        if h_ and h_ not in hosts:
            hosts.append(h_)
    cap = (s.get("image") or {}).get("caption")
    bits = ([cap] if cap else []) + (["Images: " + " · ".join(hosts)] if hosts else [])
    if bits:
        _textbox(slide, x0, y0 + h + 0.06, w, 0.3, [("  ·  ".join(bits), t["body"], 10, t["muted"], False, True, 0)],
                 name=DECO + "-credit", anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)


def _stroke(slide, kind, x, y, w, h, color, width=1.6, adj=None, name=DECO + "-art"):
    """An outline-only shape (no fill) pushed behind the text."""
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    sh.fill.background()
    sh.line.color.rgb = _rgb(color)
    sh.line.width = Pt(width)
    sh.shadow.inherit = False
    if adj is not None:
        sh.adjustments[0] = adj
    _back(sh)
    return sh


def _art(slide, kind: str, x: float, y: float, w: float, h: float, t: dict) -> None:
    """Drawn motif for a cover (native shapes, named deco-*, so they are skipped by inspection and the layout check):
    arches / waves / skyline / geometric / network."""
    from pptx.enum.shapes import MSO_CONNECTOR
    line, line2 = t["accent"], t["accent2"] if _contrast(t["accent2"], t["cover_bg"]) >= 2.2 else t["cover_muted"]
    base = y + h * 0.66
    if kind == "arches":
        for fx, fw, fh in ((0.02, 0.28, 0.36), (0.31, 0.40, 0.56), (0.73, 0.25, 0.36)):
            ax, aw, ah = x + w * fx, w * fw, h * fh
            _stroke(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, ax, base - ah, aw, ah, line, 1.8, adj=0.5)
            inset = min(aw * 0.2, 0.42)
            _stroke(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, ax + inset, base - ah + inset * 1.4, aw - 2 * inset,
                    ah - inset * 1.4, line, 1.1, adj=0.5)
        for i in range(4):
            _stroke(slide, MSO_SHAPE.WAVE, x, base + 0.38 + i * 0.38, w, 0.3, line2, 1.3)
    elif kind == "waves":
        for i in range(7):
            _stroke(slide, MSO_SHAPE.WAVE, x, y + h * 0.18 + i * (h * 0.11), w, 0.34, line if i % 2 == 0 else line2, 1.4)
    elif kind == "skyline":
        heights = (0.34, 0.58, 0.44, 0.78, 0.5, 0.64, 0.38)
        bw = w / (len(heights) + 0.5)
        for i, fh in enumerate(heights):
            hh = h * 0.7 * fh
            _stroke(slide, MSO_SHAPE.RECTANGLE, x + i * bw * 1.07, base - hh, bw * 0.86, hh, line, 1.5)
            for r in range(int(hh // 0.5)):
                _deco(slide, MSO_SHAPE.RECTANGLE, x + i * bw * 1.07 + bw * 0.22, base - hh + 0.2 + r * 0.5, bw * 0.4, 0.07, line2)
        _deco(slide, MSO_SHAPE.RECTANGLE, x, base, w, 0.03, line)
        for i in range(3):
            _stroke(slide, MSO_SHAPE.WAVE, x, base + 0.3 + i * 0.36, w, 0.3, line2, 1.3)
    elif kind == "geometric":
        _deco(slide, MSO_SHAPE.RECTANGLE, x, base - h * 0.42, w, h * 0.42, t["cover_muted"], alpha=22)
        _deco(slide, MSO_SHAPE.OVAL, x + w * 0.08, base - h * 0.36, w * 0.28, w * 0.28, line)
        _deco(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x + w * 0.4, base - h * 0.38, w * 0.26, h * 0.3, line2)
        _deco(slide, MSO_SHAPE.RECTANGLE, x + w * 0.74, base - h * 0.3, w * 0.16, w * 0.16, line)
        for i in range(3):
            _stroke(slide, MSO_SHAPE.WAVE, x, base + 0.3 + i * 0.36, w, 0.3, line2, 1.3)
    elif kind == "network":
        pts = [(0.15, 0.2), (0.5, 0.1), (0.85, 0.25), (0.3, 0.5), (0.7, 0.55), (0.12, 0.8), (0.55, 0.85), (0.9, 0.82)]
        edges = [(0, 1), (1, 2), (0, 3), (1, 3), (1, 4), (2, 4), (3, 4), (3, 5), (4, 6), (4, 7), (5, 6), (6, 7)]
        for a, b in edges:
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + w * pts[a][0]), Inches(y + h * pts[a][1]),
                                           Inches(x + w * pts[b][0]), Inches(y + h * pts[b][1]))
            c.name = DECO + "-art"
            c.line.color.rgb = _rgb(line2)
            c.line.width = Pt(1.2)
            _back(c)
        for i, (fx, fy) in enumerate(pts):
            d = 0.22 if i % 3 else 0.36
            _deco(slide, MSO_SHAPE.OVAL, x + w * fx - d / 2, y + h * fy - d / 2, d, d, line if i % 2 == 0 else line2)


def _variant(t: dict, n: int) -> dict:
    """Theme with its alternate palette on every second content slide (editorial / ember: cream, deep, cream, deep ...)."""
    alt = t.get("alt")
    if not alt or n % 2 == 0:
        return t
    v = dict(t)
    v.update(alt)
    return v


def _heading_body(item):
    head, kids = item
    head = head.rstrip(":").strip()
    if kids:
        return head, kids
    lb = O.label_split(head)
    return (lb[0], [lb[1]]) if lb else (head, [])


# ---------------------------------------------------------------- slide kinds

def _cover(slide, s, t, closing=False):
    _bg(slide, t["cover_bg"])
    motif = t["motif"]
    tw = 10.4 if motif == "rule" else 8.0
    has_pic = bool(s.get("image_data")) and not closing
    art = "" if closing or has_pic else (s.get("art") or t.get("art") or "")
    art = "" if art == "none" else art
    if has_pic:
        tw = 6.6
    elif art:
        tw = 6.8
    if motif == "circles":
        _deco(slide, MSO_SHAPE.OVAL, 8.9, -1.9, 7.2, 7.2, t["accent"], alpha=16)
        _deco(slide, MSO_SHAPE.OVAL, 10.6, 3.5, 5.2, 5.2, t["accent2"], alpha=28)
    elif motif == "band":
        _deco(slide, MSO_SHAPE.RECTANGLE, 10.2, 0, W - 10.2, H, t["accent2"], alpha=90)
        _deco(slide, MSO_SHAPE.RECTANGLE, 10.0, 0, 0.2, H, t["accent"])
    else:
        _deco(slide, MSO_SHAPE.RECTANGLE, 1.1, 1.4, 0.9, 0.1, t["accent"])
        _deco(slide, MSO_SHAPE.RECTANGLE, 1.1, 6.6, 11.1, 0.02, t["muted"], alpha=40)
    _deco(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.28, H, t["accent"])
    if art:
        _art(slide, art, 7.9, 0.9, 4.9, 5.4, t)
    if has_pic:                                       # full-height picture panel on the right edge
        _picture(slide, s, 8.0, 0, W - 8.0, H, name="Cover image")
        _deco(slide, MSO_SHAPE.RECTANGLE, 7.9, 0, 0.1, H, t["accent"], back=False)
        _credit(slide, s, t, 8.1, 7.1, W - 8.3, color="FFFFFF")
    _title(slide, s["title"], t, 1.1, 2.0, tw, 2.2, color=t["cover_text"], base=46,
           anchor=MSO_ANCHOR.BOTTOM)
    _deco(slide, MSO_SHAPE.RECTANGLE, 1.1, 4.42, 1.2, 0.07, t["accent"])
    sub = s["subtitle"]
    if closing and s["bullets"]:
        sub = "\n".join([x for x in [sub] + [b.strip() for b in s["bullets"]] if x])
    body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
    if body is not None and sub:
        _place(body, 1.1, 4.7, tw, 1.7)
        _frame(body.text_frame, MSO_ANCHOR.TOP)
        body.text_frame.text = sub
        size = _fit(sub.split("\n"), tw, 1.7, 20, 13, gap_pt=4, font=t["body"])
        for p in body.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(4)
            _style_runs(p, t["body"], size, t["cover_muted"])
    elif body is not None:
        body._element.getparent().remove(body._element)


def _section(slide, s, t, no: int = 0):
    _bg(slide, t["cover_bg"])
    if no:                                              # "01", "02" ... above the section title
        _textbox(slide, 1.1, 2.45, 3.0, 0.5, [(f"{no:02d}", t["head"], 24, _ink(t["accent"], t["cover_bg"], 3.0), True, False, 0)],
                 name=DECO + "-section-number", anchor=MSO_ANCHOR.MIDDLE)
    _deco(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.28, H, t["accent"])
    if t["motif"] == "circles":
        _deco(slide, MSO_SHAPE.OVAL, 9.4, 1.2, 5.2, 5.2, t["accent2"], alpha=22)
    _deco(slide, MSO_SHAPE.RECTANGLE, 1.1, 3.0, 1.2, 0.07, t["accent"])
    _title(slide, s["title"], t, 1.1, 3.2, 9.0, 1.5, color=t["cover_text"], base=40, anchor=MSO_ANCHOR.TOP)
    body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
    if body is not None and s["subtitle"]:
        _place(body, 1.1, 4.8, 9.0, 1.0)
        _frame(body.text_frame, MSO_ANCHOR.TOP)
        body.text_frame.text = s["subtitle"]
        for p in body.text_frame.paragraphs:
            _style_runs(p, t["body"], 20, t["cover_muted"])
    elif body is not None:
        body._element.getparent().remove(body._element)


def _full(slide, s, t):
    """Full-bleed picture with a dark scrim; the title (and subtitle) sit over it."""
    _bg(slide, t["cover_bg"])
    pic = _picture(slide, s, 0, 0, W, H, name="Background image")
    scrim = _deco(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, H, "000000", alpha=48, back=False)
    for shp in (scrim, pic):                           # behind the placeholders, picture lowest
        if shp is not None:
            _back(shp)
    _deco(slide, MSO_SHAPE.RECTANGLE, 1.1, 4.55, 1.2, 0.07, t["accent"], back=False)
    _title(slide, s["title"], t, 1.1, 2.4, 10.5, 2.0, color="FFFFFF", base=44, anchor=MSO_ANCHOR.BOTTOM)
    body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
    if body is not None and s["subtitle"]:
        _place(body, 1.1, 4.8, 10.5, 1.2)
        _frame(body.text_frame, MSO_ANCHOR.TOP)
        body.text_frame.text = s["subtitle"]
        for p in body.text_frame.paragraphs:
            _style_runs(p, t["body"], 20, "E5E7EB")
    elif body is not None:
        body._element.getparent().remove(body._element)
    _credit(slide, s, t, W - MX - 6, 7.05, 6, color="E5E7EB")


def _bullets(slide, s, t):
    body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
    items = s["bullets"]
    h = 4.85
    _deco(slide, MSO_SHAPE.RECTANGLE, MX, 1.95, 0.06, min(h - 0.2, 0.62 * max(len(items), 3)), t["accent"])
    if s.get("image_data"):                           # text on the left, picture on the right
        px, pw = 7.55, W - MX - 7.55
        _fill_bullets(body, items, t, MX + 0.3, 1.9, px - MX - 0.7, h, base=26 if len(items) <= 4 else 22, floor=12)
        _picture(slide, s, px, 1.95, pw, h - 0.35)
        _credit(slide, s, t, px, 1.95 + h - 0.3, pw)
        return
    _fill_bullets(body, items, t, MX + 0.3, 1.9, W - 2 * MX - 0.3, h, base=30 if len(items) <= 4 else 26)


def _columns(slide, s, t):
    tops = O.groups(s["bullets"])
    half = math.ceil(len(tops) / 2)
    parts = [tops[:half], tops[half:]]
    bodies = [p for p in slide.placeholders if p.placeholder_format.idx in (1, 2)]
    bodies.sort(key=lambda p: p.placeholder_format.idx)
    cw = (W - 2 * MX - 0.6) / 2
    for i, (ph, part) in enumerate(zip(bodies, parts)):
        lines = []
        for head, kids in part:
            lines.append(head)
            lines += ["\t" + k for k in kids]
        x = MX + i * (cw + 0.6)
        _fill_bullets(ph, lines, t, x, 1.95, cw, 4.75, base=21, floor=12)
    _deco(slide, MSO_SHAPE.RECTANGLE, W / 2 - 0.01, 2.0, 0.02, 4.4, t["line"])


def _cards(slide, s, t):
    items = [_heading_body(g) for g in O.groups(s["bullets"])]
    n = len(items)
    cols = n if n <= 4 else 3
    rows = math.ceil(n / cols)
    gap = 0.3
    area_w, area_h, y0 = W - 2 * MX, 4.75, 1.95
    cw = min((area_w - gap * (cols - 1)) / cols, 4.2 if cols == 2 else 9)
    x0 = MX + (area_w - (cw * cols + gap * (cols - 1))) / 2 if cols == 2 else MX
    ch = min((area_h - gap * (rows - 1)) / rows, 3.7)
    for i, (head, kids) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (cw + gap), y0 + r * (ch + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(ch))
        card.name = f"Card {i + 1}"
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(t["surface"])
        card.line.color.rgb = _rgb(t["line"])
        card.line.width = Pt(1)
        card.shadow.inherit = False
        accent = t["accent"] if i % 2 == 0 else t["accent2"]
        _deco(slide, MSO_SHAPE.RECTANGLE, x, y, cw, 0.08, accent, back=False)
        badge = _deco(slide, MSO_SHAPE.OVAL, x + 0.28, y + 0.32, 0.5, 0.5, accent, back=False)
        bt = badge.text_frame
        bt.text = str(i + 1)
        _frame(bt, MSO_ANCHOR.MIDDLE)
        bt.paragraphs[0].alignment = PP_ALIGN.CENTER
        _style_runs(bt.paragraphs[0], t["head"], 15, _on(accent), True)
        tf = card.text_frame
        _frame(tf, MSO_ANCHOR.TOP, (0.28, 1.0, 0.28, 0.2))
        inner_w, inner_h = cw - 0.56, ch - 1.2
        hs = _fit([head], inner_w, 0.9, 21, 14, gap_pt=0, font=t["head"], bold=True)
        bs = _fit(kids, inner_w, max(inner_h - 0.8, 0.8), 18, 11, gap_pt=6, font=t["body"])
        tf.text = head
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        tf.paragraphs[0].space_after = Pt(8)
        _style_runs(tf.paragraphs[0], t["head"], hs, t["text"], True)
        for k in kids:
            p = tf.add_paragraph()
            p.text = k
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(6)
            _style_runs(p, t["body"], bs, _ink(t["muted"], t["surface"], 4.5))


def _stats(slide, s, t):
    items = [(O.stat_split(g[0]) or (g[0], "")) for g in O.groups(s["bullets"])]
    n = len(items)
    gap = 0.35
    area_w = W - 2 * MX
    cw = (area_w - gap * (n - 1)) / n
    common_ns = min([M.fit_one_line(x[0], t["head"], cw - 0.75, 60, 24, True) for x in items], default=40)
    y, ch = 2.2, 3.7
    for i, (num, cap) in enumerate(items):
        x = MX + i * (cw + gap)
        tile = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(ch))
        tile.name = f"Stat {i + 1}"
        tile.fill.solid()
        tile.fill.fore_color.rgb = _rgb(t["surface"])
        tile.line.color.rgb = _rgb(t["line"])
        tile.line.width = Pt(1)
        tile.shadow.inherit = False
        accent = t["accent"] if i % 2 == 0 else t["accent2"]
        _deco(slide, MSO_SHAPE.RECTANGLE, x, y, 0.1, ch, accent, back=False)
        tf = tile.text_frame
        _frame(tf, MSO_ANCHOR.MIDDLE, (0.45, 0.2, 0.3, 0.2))
        ns = common_ns
        tf.text = num
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        tf.paragraphs[0].space_after = Pt(10)
        _style_runs(tf.paragraphs[0], t["head"], ns, _ink(accent, t["surface"], 3.0), True)
        p = tf.add_paragraph()
        p.text = cap
        p.alignment = PP_ALIGN.LEFT
        _style_runs(p, t["body"], _fit([cap], cw - 0.75, 1.2, 18, 12, gap_pt=0, font=t["body"]), t["text"])


def _timeline(slide, s, t):
    items = [_heading_body(g) for g in O.groups(s["bullets"])]
    n = len(items)
    area_w = W - 2 * MX
    sw = area_w / n
    ly = 3.0
    _deco(slide, MSO_SHAPE.RECTANGLE, MX + sw / 2, ly - 0.02, area_w - sw, 0.04, t["line"] if _lum(t["line"]) < 0.9 else t["muted"])
    for i, (head, kids) in enumerate(items):
        cx = MX + sw * i + sw / 2
        dot = _deco(slide, MSO_SHAPE.OVAL, cx - 0.3, ly - 0.3, 0.6, 0.6,
                    t["accent"] if i % 2 == 0 else t["accent2"], back=False)
        dt = dot.text_frame
        dt.text = str(i + 1)
        _frame(dt, MSO_ANCHOR.MIDDLE)
        dt.paragraphs[0].alignment = PP_ALIGN.CENTER
        _style_runs(dt.paragraphs[0], t["head"], 16, _on(t["accent"] if i % 2 == 0 else t["accent2"]), True)
        w = sw - 0.3
        paras = [(head, t["head"], _fit([head], w, 0.9, 22, 13, gap_pt=0, font=t["head"], bold=True), t["text"], True, False, 6)]
        bs = _fit(kids, w, 2.6, 18, 11, gap_pt=6, font=t["body"])
        paras += [(k, t["body"], bs, _ink(t["muted"], t["bg"], 4.5), False, False, 6) for k in kids]
        _textbox(slide, cx - w / 2, ly + 0.6, w, 3.2, paras, name=f"Step {i + 1}", align=PP_ALIGN.CENTER)


def _quote(slide, s, t):
    items = [b.strip() for b in s["bullets"]]
    quote = items[0].strip('"“”')
    by = " — ".join(items[1:])
    _deco(slide, MSO_SHAPE.RECTANGLE, MX + 0.2, 2.3, 0.08, 3.4, t["accent"])
    _textbox(slide, MX + 0.7, 2.2, W - 2 * MX - 1.2, 3.0,
             [("“" + quote + "”", t["head"], _fit([quote], W - 2 * MX - 1.2, 3.0, 34, 18, gap_pt=0, font=t["head"]),
               t["text"], False, True, 0)], name="Quote", anchor=MSO_ANCHOR.MIDDLE)
    if by:
        _textbox(slide, MX + 0.7, 5.4, W - 2 * MX - 1.2, 0.5,
                 [(by, t["body"], 16, _ink(t["accent"], t["bg"], 4.5), True, False, 0)], name="Attribution")


def _series_colors(t: dict) -> list:
    return [t["accent"], t["accent2"], "F59E0B", "10B981", "EC4899", "8B5CF6", "06B6D4", "84CC16"]


def _chart(slide, s, t, box=None):
    """A native, editable PowerPoint chart (bar / horizontal bar / line / pie / 100% stacked share) from the slide's numbers."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
    d = s["chart_data"]
    kinds = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "share": XL_CHART_TYPE.BAR_STACKED_100}
    cd = CategoryChartData()
    if d["kind"] == "share":                              # one stacked bar: a series per share
        cd.categories = [d.get("unit") or "Share"]
        for lab, v in zip(d["labels"], d["series"][0]["values"]):
            cd.add_series(lab, [v])
    else:
        cd.categories = d["labels"]
        for i, ser in enumerate(d["series"]):
            cd.add_series(ser["name"] or f"Series {i + 1}", ser["values"])
    bx, by, bw, bh = box or ((MX, 2.2, W - 2 * MX, 3.4) if d["kind"] == "share" else (MX, 1.95, W - 2 * MX, 4.75))
    gf = slide.shapes.add_chart(kinds[d["kind"]], Inches(bx), Inches(by), Inches(bw), Inches(bh), cd)
    gf.name = "Chart"
    ch = gf.chart
    ch.has_title = False
    ch.font.name = t["body"]
    ch.font.size = Pt(14)
    ch.font.color.rgb = _rgb(t["text"])
    colors = _series_colors(t)
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(13)
    dl.font.bold = True
    dl.font.color.rgb = _rgb(t["text"])
    dl.number_format = '#,##0.##' + (f'"{d["unit"]}"' if d["unit"] else "")
    dl.number_format_is_linked = False
    multi = len(d["series"]) > 1
    if d["kind"] == "share":
        plot.gap_width = 12
        plot.overlap = 100
        for i, ser in enumerate(plot.series):
            col = colors[i % len(colors)]
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = _rgb(col)
            ser.format.line.color.rgb = _rgb(t["bg"])
            ser.data_labels.show_value = True
            ser.data_labels.font.size = Pt(14)
            ser.data_labels.font.bold = True
            ser.data_labels.font.color.rgb = _rgb(_on(col))
            ser.data_labels.number_format = '#,##0.##' + (f'"{d["unit"]}"' if d["unit"] else "")
            ser.data_labels.number_format_is_linked = False
        ch.value_axis.visible = False
        ch.category_axis.visible = False
        ch.value_axis.has_major_gridlines = False
        ch.has_legend = True
    elif d["kind"] == "pie":
        plot.vary_by_categories = True
        for i, _ in enumerate(d["labels"]):
            pt = plot.series[0].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(colors[i % len(colors)])
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        ch.has_legend = True
    else:
        for i, ser in enumerate(plot.series):
            col = _rgb(colors[i % len(colors)])
            if d["kind"] == "line":
                ser.format.line.color.rgb = col
                ser.format.line.width = Pt(3)
                ser.smooth = False
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = col
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = col
        dl.position = XL_LABEL_POSITION.ABOVE if d["kind"] == "line" else XL_LABEL_POSITION.OUTSIDE_END
        if len(d["labels"]) > 8 or multi:
            plot.has_data_labels = len(d["labels"]) <= 8
        ch.has_legend = multi
        for ax in (ch.category_axis, ch.value_axis):
            ax.tick_labels.font.size = Pt(13)
            ax.tick_labels.font.color.rgb = _rgb(t["muted"])
            ax.format.line.color.rgb = _rgb(t["line"])
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = _rgb(t["line"])
        ch.value_axis.format.line.fill.background()
        ch.value_axis.tick_labels.number_format = '#,##0.##'
        ch.value_axis.tick_labels.number_format_is_linked = False
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(13)
        ch.legend.font.color.rgb = _rgb(t["text"])


def _comparison(slide, s, t):
    """Two options side by side (`layout: comparison`; two `Heading:` bullets, each with its own sub-bullets)."""
    top = O.groups(s["bullets"])[:2]
    gap = 0.9
    cw = (W - 2 * MX - gap) / 2
    y0, hh, ph = 1.95, 0.85, 4.75
    for i, (head, kids) in enumerate(top):
        x = MX + i * (cw + gap)
        accent = t["accent"] if i == 0 else t["accent2"]
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y0), Inches(cw), Inches(ph))
        panel.name = f"Option {i + 1}"
        panel.fill.solid()
        panel.fill.fore_color.rgb = _rgb(t["surface"])
        panel.line.color.rgb = _rgb(t["line"])
        panel.line.width = Pt(1)
        panel.shadow.inherit = False
        band = _deco(slide, MSO_SHAPE.RECTANGLE, x, y0, cw, hh, accent, back=False, name=f"Option {i + 1} heading")
        _frame(band.text_frame, MSO_ANCHOR.MIDDLE, (0.3, 0.05, 0.3, 0.05))
        band.text_frame.text = head.rstrip(":").strip()
        band.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        _style_runs(band.text_frame.paragraphs[0], t["head"],
                    _fit([head], cw - 0.6, hh - 0.1, 24, 14, gap_pt=0, font=t["head"], bold=True), _on(accent), True)
        body_h = ph - hh - 0.5
        size = _fit(kids, cw - 0.9, body_h, 22, 12, gap_pt=12, font=t["body"])
        tb = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y0 + hh + 0.3), Inches(cw - 0.6), Inches(body_h))
        tb.name = f"Option {i + 1} points"
        _frame(tb.text_frame, MSO_ANCHOR.TOP)
        for j, k in enumerate(kids):
            p = tb.text_frame.paragraphs[0] if j == 0 else tb.text_frame.add_paragraph()
            p.text = k
            p.space_after = Pt(12)
            _style_runs(p, t["body"], size, t["text"])
            _bullet(p, accent, 0)
    vs = _deco(slide, MSO_SHAPE.OVAL, W / 2 - 0.42, y0 + hh / 2 - 0.42 + 0.0, 0.84, 0.84, t["bg"], back=False, name="deco-vs")
    vs.line.color.rgb = _rgb(t["muted"])
    vs.line.width = Pt(2)
    _frame(vs.text_frame, MSO_ANCHOR.MIDDLE)
    vs.text_frame.text = "VS"
    vs.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    _style_runs(vs.text_frame.paragraphs[0], t["head"], 16, t["text"], True)


def _agenda(slide, s, t):
    """Numbered contents list (`layout: agenda`): one column up to 4 items, two columns beyond."""
    items = [h.rstrip(":").strip() for h, _ in O.groups(s["bullets"]) if h][:8]
    n = len(items)
    cols = 1 if n <= 4 else 2
    rows = math.ceil(n / cols)
    cw = (W - 2 * MX - (0.6 if cols == 2 else 0)) / cols
    rh = min(1.0, 4.6 / rows)
    for i, text in enumerate(items):
        c, r = divmod(i, rows)
        x, y = MX + c * (cw + 0.6), 2.0 + r * rh
        accent = t["accent"] if i % 2 == 0 else t["accent2"]
        dot = _deco(slide, MSO_SHAPE.OVAL, x, y + (rh - 0.62) / 2, 0.62, 0.62, accent, back=False, name=f"deco-num-{i + 1}")
        _frame(dot.text_frame, MSO_ANCHOR.MIDDLE)
        dot.text_frame.text = f"{i + 1}"
        dot.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        _style_runs(dot.text_frame.paragraphs[0], t["head"], 18, _on(accent), True)
        _textbox(slide, x + 0.85, y, cw - 0.95, rh,
                 [(text, t["body"], _fit([text], cw - 1.0, rh - 0.1, 26, 14, gap_pt=0, font=t["body"]), t["text"], False, False, 0)],
                 name=f"Item {i + 1}", anchor=MSO_ANCHOR.MIDDLE)
        if r < rows - 1:
            _deco(slide, MSO_SHAPE.RECTANGLE, x + 0.85, y + rh - 0.01, cw - 0.95, 0.01, t["line"], back=False)


TABLE_MAX_H = 4.7
TABLE_SIZES = (20, 18, 16, 14, 13, 12, 11)


def _table_widths(rows: list, total: float) -> list:
    """Column widths in inches, proportional to the longest cell of each column (clamped), not all equal."""
    cols = max(len(r) for r in rows)
    weight = [max(6, min(42, max(len(r[c]) if c < len(r) else 0 for r in rows))) for c in range(cols)]
    return [total * w / sum(weight) for w in weight]


def _table_heights(rows: list, widths: list, size: float, font: str) -> list:
    out = []
    for ri, row in enumerate(rows):
        lines = max(M.lines_for(cell, font, size, max(widths[c] - 0.34, 0.3), bold=(ri == 0)) for c, cell in enumerate(row))
        out.append(lines * size * M.line_height(row) / 72 + 0.2)
    return out


def _plan_table(rows: list, font: str, total_w: float = 0.0, max_h: float = 0.0) -> list:
    """[(rows, size, widths, heights)] : the whole table on one slide at the largest size that fits, else split after
    the header repeats on every slide. Rows are padded to the same number of cells."""
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    widths = _table_widths(rows, total_w or (W - 2 * MX))
    limit = max_h or TABLE_MAX_H

    def best(chunk):
        for size in TABLE_SIZES:
            hs = _table_heights(chunk, widths, size, font)
            if sum(hs) <= limit:
                return size, hs
        return None

    whole = best(rows)
    if whole:
        return [(rows, whole[0], widths, whole[1])]
    header, body = rows[0], rows[1:]
    chunks, cur = [], []
    for r in body:
        trial = [header] + cur + [r]
        if cur and not best(trial):
            chunks.append([header] + cur)
            cur = [r]
        else:
            cur.append(r)
    if cur:
        chunks.append([header] + cur)
    plan = []
    for ch in chunks:
        fit = best(ch)
        size, hs = fit if fit else (TABLE_SIZES[-1], _table_heights(ch, widths, TABLE_SIZES[-1], font))
        plan.append((ch, size, widths, hs))
    return plan


def _table(slide, s, t, box=None):
    bx, by, bw, bh = box or (MX, 1.95, W - 2 * MX, TABLE_MAX_H)
    rows, size, widths, heights = _plan_table(s["table"], t["body"], total_w=bw, max_h=bh)[0]
    cols = len(widths)
    gf = slide.shapes.add_table(len(rows), cols, Inches(bx), Inches(by), Inches(bw), Inches(sum(heights)))
    gf.name = "Table"
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for c, w in enumerate(widths):
        tbl.columns[c].width = Inches(w)
    for r, h in enumerate(heights):
        tbl.rows[r].height = Inches(h)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            head = r == 0
            fill = t["accent"] if head else (t["surface"] if r % 2 else t["bg"])
            cell.fill.fore_color.rgb = _rgb(fill)
            cell.margin_left = cell.margin_right = Inches(0.15)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.text = val
            _style_runs(cell.text_frame.paragraphs[0], t["body"], size, _on(fill) if head else t["text"], head)


def _textside(slide, s, t):
    """Text on the left, a table or a chart on the right (`layout: textside`): the climate-table-beside-the-story slide."""
    items = [b.strip() for b in s["bullets"]]
    body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
    lw, gap = 5.8, 0.5
    rx = MX + lw + gap
    rw = W - MX - rx
    if body is not None:
        _fill_bullets(body, items, t, MX, 1.95, lw, 4.75, base=20, floor=12)
    if s.get("chart_data"):
        _chart(slide, s, t, box=(rx, 1.95, rw, 4.6))
    else:
        _table(slide, s, t, box=(rx, 2.1, rw, 4.4))


def _bullet_box(s: dict) -> tuple:
    """(width, height, floor) of the bullet text area, as _bullets lays it out."""
    if s.get("image_data"):
        return 7.55 - MX - 0.7, 4.85, 12
    return W - 2 * MX - 0.3, 4.85, 13


def _split_bullets(s: dict, t: dict) -> list:
    """A bullets slide whose text cannot fit even at the smallest size becomes two slides ('(cont.)') instead of
    running off the box. Each half is checked again, so a very long list ends up as three or four slides."""
    w, h, floor = _bullet_box(s)
    items = [b.strip() for b in s["bullets"]]
    if M.fits(items, t["body"], floor, w - 0.4, h, gap_pt=16):
        return [s]
    top = O.groups(s["bullets"])
    if len(top) < 2:
        return [s]
    half = (len(top) + 1) // 2

    def lines(part):
        out = []
        for head, kids in part:
            if head:
                out.append(head)
            out += ["\t" + k for k in kids]
        return out
    first = {**s, "bullets": lines(top[:half])}
    rest_title = s["title"] if s["title"].endswith("(cont.)") else s["title"] + " (cont.)"
    second = {**s, "bullets": lines(top[half:]), "title": rest_title, "notes": "", "subtitle": "",
              "image": None, "image_data": None, "image_credit": ""}
    return _split_bullets(first, t) + _split_bullets(second, t)


def _prepare(slides: list, t: dict) -> list:
    """Slides after fitting: long bullet lists and long tables are split across slides."""
    out = []
    for s in slides:
        if s["kind"] == "bullets" and s["bullets"]:
            out += _split_bullets(s, t)
        elif s["kind"] == "table" and s["table"]:
            plan = _plan_table(s["table"], t["body"])
            for i, (rows, _size, _w, _h) in enumerate(plan):
                if i == 0:
                    out.append({**s, "table": rows})
                else:
                    title = s["title"] if s["title"].endswith("(cont.)") else s["title"] + " (cont.)"
                    out.append({**s, "table": rows, "title": title, "notes": "", "subtitle": ""})
        else:
            out.append(s)
    return out


_LAYOUT_FOR = {"textside": "Title and Content", "cover": "Title Slide", "closing": "Title Slide", "section": "Section Header", "full": "Section Header",
               "bullets": "Title and Content", "columns": "Two Content"}


def _plain_slide(prs, lays, s: dict, t: dict) -> None:
    """Last-resort slide for content whose designed layout failed: title + text, no decoration, so nothing is lost."""
    slide = prs.slides.add_slide(lays["Title and Content"])
    if slide.shapes.title is not None:
        slide.shapes.title.text_frame.text = s.get("title", "")
    body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
    lines = [b.strip() for b in s.get("bullets", [])] + [" | ".join(r) for r in s.get("table", [])]
    if body is not None:
        body.text_frame.text = "\n".join(lines or [s.get("subtitle", "")])
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"]


def _drop_last_slide(prs) -> None:
    lst = prs.slides._sldIdLst
    last = lst[-1]
    prs.part.drop_rel(last.rId)
    lst.remove(last)


def build(deck: dict) -> Presentation:
    t = O.resolve_theme(deck)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for step in (_widen_master, lambda p: _apply_theme_part(p, t)):         # polish only: never fatal
        try:
            step(prs)
        except Exception as e:
            print(f"[deck] theme/master step skipped ({type(e).__name__}: {e})", file=sys.stderr)
    lays = {l.name: l for l in prs.slide_layouts}
    deck_title = deck["title"]
    sec_no = cn = 0
    for i, s in enumerate(_prepare(deck["slides"], t)):
        n_before = len(prs.slides)
        try:
            sec_no, cn = _render_slide(prs, lays, s, t, deck_title, i, sec_no, cn)
        except Exception as e:                                            # one bad slide must not lose the deck
            print(f"[deck] slide {i + 1} ({s.get('kind')}) failed to render: {type(e).__name__}: {e}; "
                  "used a plain slide", file=sys.stderr)
            while len(prs.slides) > n_before:
                _drop_last_slide(prs)
            _plain_slide(prs, lays, s, t)
    return prs


def _render_slide(prs, lays, s: dict, t: dict, deck_title: str, i: int, sec_no: int, cn: int = 0) -> tuple:
    """Draw one slide; returns the running (section number, content-slide number)."""
    kind = s["kind"]
    slide = prs.slides.add_slide(lays[_LAYOUT_FOR.get(kind, "Title Only")])
    if kind in ("cover", "closing"):
        _cover(slide, s, t, closing=kind == "closing")
    elif kind == "full" or (kind == "section" and s.get("image_data")):
        _full(slide, s, t)
    elif kind == "section":
        sec_no += 1
        _section(slide, s, t, sec_no)
    else:
        ts = _variant(t, cn)                                   # alternate palette on every second content slide
        cn += 1
        _chrome(slide, ts, s["title"], deck_title, i + 1, s.get("eyebrow", ""))
        {"bullets": _bullets, "columns": _columns, "cards": _cards, "stats": _stats,
         "timeline": _timeline, "quote": _quote, "table": _table, "chart": _chart,
         "comparison": _comparison, "agenda": _agenda, "gallery": _gallery, "textside": _textside}[kind](slide, s, ts)
    _drop_empty_placeholders(slide)
    if kind not in ("cover", "closing", "section", "full"):
        src = s.get("source") or ""
        foot = (src if re.match(r"(?i)sources?\b", src) else f"Source: {src}") if src else deck_title
        _footer_placeholders(slide, ts, foot, i + 1)
    if s["notes"]:
        slide.notes_slide.notes_text_frame.text = s["notes"]
    return sec_no, cn
