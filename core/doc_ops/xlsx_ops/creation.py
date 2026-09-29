import io
from .coercion import _coerce
from .constants import num_to_col


def create(rows: list[list], sheet: str = "Sheet1") -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31] or "Sheet1"
    for r, vals in enumerate(rows, 1):
        for c, v in enumerate(vals, 1):
            kind, payload = _coerce(v, allow_formula=False) if not (isinstance(v, str) and v.startswith("=")) \
                else ("str", v)
            cell = ws.cell(row=r, column=c, value=payload if kind != "empty" else None)
            if kind == "str":
                cell.data_type = "s"   # model text like "=cmd" stays text, never a formula
    if rows:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDE7F5")
        ws.freeze_panes = "A2"
        for c in range(1, max(len(r) for r in rows) + 1):
            width = max((len(str(r[c - 1])) for r in rows if len(r) >= c and r[c - 1] is not None), default=8)
            ws.column_dimensions[num_to_col(c)].width = min(max(width + 2, 8), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
