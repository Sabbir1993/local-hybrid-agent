import io
from ..base import DocOpError, EditResult, Snapshot, check_package, digest, op_arg, require
from .book import _Book
from .coercion import _coerce
from .constants import (
    LOSSY_PARTS,
    MAX_CELLS_SCANNED,
    _split_addr,
    col_to_num,
)
from .xml_ops import XML_OPS, _Ctx, _finish_xml


def _cell_fp(cell) -> str:
    return digest(repr((cell.value, cell._style if hasattr(cell, "_style") else None)).encode())


def _structural(data: bytes, ops: list[dict]) -> EditResult:
    import openpyxl
    book = _Book(data)
    lossy = sorted({p.split("/")[1] for p in book.parts if p.startswith(LOSSY_PARTS)})
    require(not lossy, f"this workbook contains {', '.join(lossy)}, which would be lost by a structural "
                       "edit; use set_cell / set_range / append_rows instead")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    require(sum(len(ws._cells) for ws in wb.worksheets) <= MAX_CELLS_SCANNED, "workbook too large")
    snap = Snapshot(_cell_fp)
    snap.take((f"{ws.title}!{c.coordinate}", c) for ws in wb.worksheets for c in ws._cells.values())
    changes, notes = [], []

    def ws_of(name):
        if not name:
            return wb.worksheets[0]
        for ws in wb.worksheets:
            if ws.title.lower() == str(name).lower():
                return ws
        raise DocOpError(f"sheet '{name}' not found")

    def guard_shift(ws):
        require(not ws.merged_cells.ranges, f"sheet '{ws.title}' has merged cells; inserting/deleting "
                                            "rows would misalign them")
        require(not ws.tables, f"sheet '{ws.title}' has tables; use append_rows instead")
        require(not ws.conditional_formatting and not ws.data_validations.dataValidation,
                f"sheet '{ws.title}' has conditional formatting/validation that would not shift")
        require(not book.has_formulas() and not wb.defined_names,
                "workbook has formulas or named ranges that openpyxl would not re-point; use "
                "set_range / append_rows instead")

    for op in ops:
        kind = str(op.get("op", "")).lower()
        sheet = op.get("sheet") or _split_addr(op.get("addr") or "")[0]
        if kind in ("insert_rows", "delete_rows", "insert_cols", "delete_cols"):
            ws = ws_of(sheet)
            guard_shift(ws)
            at = op_arg(op, "at", "row", "col", "index")
            at = col_to_num(at) if isinstance(at, str) and at.isalpha() else int(at)
            n = int(op.get("count") or 1)
            if kind.startswith("delete"):
                rows = kind == "delete_rows"
                for c in list(ws._cells.values()):
                    if (c.row if rows else c.column) in range(at, at + n):
                        snap.touch(c)
            getattr(ws, kind)(at, n)
            if kind == "insert_rows":
                vals = op.get("values") or []
                for i, rowvals in enumerate(vals[:n]):
                    for j, v in enumerate(rowvals, 1):
                        cell = ws.cell(row=at + i, column=j)
                        kind_v, payload = _coerce(v, bool(op.get("formulas")))
                        cell.value = ("=" + payload) if kind_v == "formula" else payload
                        snap.touch(cell)
            changes.append(f"{ws.title}: {kind.replace('_', ' ')} at {at} x{n}")
        elif kind == "add_sheet":
            name = str(op_arg(op, "name", "title"))
            require(name not in wb.sheetnames, f"sheet '{name}' already exists")
            ws = wb.create_sheet(name, int(op["index"]) if op.get("index") is not None else None)
            for i, rowvals in enumerate(op.get("rows") or [], 1):
                for j, v in enumerate(rowvals, 1):
                    kind_v, payload = _coerce(v, bool(op.get("formulas")))
                    snap.touch(ws.cell(row=i, column=j, value=("=" + payload) if kind_v == "formula" else payload))
            changes.append(f"added sheet '{name}'")
        elif kind == "rename_sheet":
            ws = ws_of(op_arg(op, "sheet", "name", "from"))
            new = str(op_arg(op, "to", "new_name"))
            old = ws.title
            require(not any(isinstance(c.value, str) and c.value.startswith("=") and old in c.value
                            for w in wb.worksheets for c in w._cells.values()) and not wb.defined_names,
                    f"formulas or named ranges refer to '{old}'; renaming would break them")
            ws.title = new
            changes.append(f"renamed sheet '{old}' -> '{new}'")
        else:
            raise DocOpError(f"unknown xlsx op '{op.get('op')}'")
    problems = snap.problems((f"{ws.title}!{c.coordinate}", c) for ws in wb.worksheets for c in ws._cells.values())
    if problems:
        raise DocOpError("edit would change cells you did not ask for: " + "; ".join(problems[:8]))
    out = io.BytesIO()
    wb.save(out)
    notes.append("workbook re-saved by openpyxl (structural edit)")
    return EditResult(out.getvalue(), changes, notes)


STRUCTURAL_OPS = {"insert_rows", "delete_rows", "insert_cols", "delete_cols", "add_sheet", "rename_sheet"}


def apply(data: bytes, ops: list[dict]) -> EditResult:
    kinds = [str(o.get("op", "")).lower() for o in ops]
    unknown = [k for k in kinds if k not in XML_OPS and k not in STRUCTURAL_OPS]
    require(not unknown, f"unknown xlsx op(s) {unknown} (have: {', '.join(list(XML_OPS) + sorted(STRUCTURAL_OPS))})")
    if any(k in STRUCTURAL_OPS for k in kinds):
        # structural first (openpyxl), then cell edits on its output (XML path)
        s_ops = [o for o, k in zip(ops, kinds) if k in STRUCTURAL_OPS]
        c_ops = [o for o, k in zip(ops, kinds) if k not in STRUCTURAL_OPS]
        res = _structural(data, s_ops)
        if c_ops:
            more = apply(res.data, c_ops)
            return EditResult(more.data, res.changes + more.changes, res.notes + more.notes)
        return res
    book = _Book(data)
    ctx = _Ctx(book)
    for op, k in zip(ops, kinds):
        XML_OPS[k](ctx, op)
    out = _finish_xml(ctx)
    problems = check_package(data, out, ctx.rules)
    if problems:
        raise DocOpError("edit would change untouched workbook parts: " + "; ".join(problems[:8]))
    return EditResult(out, ctx.changes, ctx.notes)
