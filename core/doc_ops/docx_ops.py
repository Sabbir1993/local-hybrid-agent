"""Word (.docx): structured outline + in-place edit ops.

Addresses (1-based, shown by inspect):
    p12             12th body paragraph (empty paragraphs count too)
    t2              2nd table;  t2[3,1]  row 3, column 1 of it
    hdr1/p1         paragraph 1 of section 1's header (ftr1/p1 for the footer)
"""

import copy
import io
import re

from lxml import etree

from .base import (DocOpError, EditResult, PartRules, Snapshot, check_package, op_arg, read_zip,
                   require, xml_fp)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MAX_CELLS_SHOWN = 400


def _w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


def _open(data: bytes):
    import docx
    read_zip(data)
    return docx.Document(io.BytesIO(data))


def _body_children(doc):
    return [el for el in doc.element.body if etree.QName(el).localname in ("p", "tbl", "sdt")]


def _paras(doc):
    return [el for el in doc.element.body if el.tag == _w("p")]


def _tables(doc):
    return [el for el in doc.element.body if el.tag == _w("tbl")]


def _ptext(p) -> str:
    out = []
    for el in p.iter(_w("t"), _w("tab"), _w("br")):
        if el.tag == _w("t"):
            out.append(el.text or "")
        else:
            out.append("\t" if el.tag == _w("tab") else "\n")
    return "".join(out)


def _style(doc, p) -> str:
    ps = p.find(f"{_w('pPr')}/{_w('pStyle')}")
    if ps is None:
        return ""
    sid = ps.get(_w("val"))
    try:
        return doc.styles.element.get_by_id(sid).name_val or sid
    except Exception:
        return sid or ""


def _clip(s: str, n: int = 300) -> str:
    s = s.strip()
    return s if len(s) <= n else s[:n] + "..."


def _rows(tbl):
    return tbl.findall(_w("tr"))


def _cells(tr):
    return tr.findall(_w("tc"))


def inspect(data: bytes) -> dict:
    doc = _open(data)
    els = []
    paras, tables = _paras(doc), _tables(doc)
    head = f"Word document: {len(paras)} paragraphs, {len(tables)} tables, {len(doc.sections)} sections"
    pi = ti = 0
    for el in _body_children(doc):
        if el.tag == _w("p"):
            pi += 1
            t = _ptext(el)
            if t.strip():
                st = _style(doc, el)
                els.append({"addr": f"p{pi}", "kind": st or "para", "text": _clip(t)})
        elif el.tag == _w("tbl"):
            ti += 1
            rows = _rows(el)
            els.append({"addr": f"t{ti}", "kind": "table", "text": f"{len(rows)} rows"})
            n = 0
            for r, tr in enumerate(rows, 1):
                for c, tc in enumerate(_cells(tr), 1):
                    if n < MAX_CELLS_SHOWN:
                        els.append({"addr": f"t{ti}[{r},{c}]", "kind": "cell",
                                    "text": _clip(" / ".join(_ptext(p) for p in tc.iter(_w("p"))), 120)})
                    n += 1
    for si, sec in enumerate(doc.sections, 1):
        for kind, hf in (("hdr", sec.header), ("ftr", sec.footer)):
            if hf.is_linked_to_previous:
                continue
            for j, p in enumerate(hf._element.findall(_w("p")), 1):
                if _ptext(p).strip():
                    els.append({"addr": f"{kind}{si}/p{j}", "kind": "header" if kind == "hdr" else "footer",
                                "text": _clip(_ptext(p))})
    return {"kind": "docx", "header": head, "elements": els}


# ------------------------------------------------------------------
# Text writing that keeps formatting
# ------------------------------------------------------------------

_KEEP = ("pPr", "bookmarkStart", "bookmarkEnd", "commentRangeStart", "commentRangeEnd", "proofErr")


def set_para_text(p, text: str) -> None:
    """Replace the paragraph text: pPr (style, numbering) and the first run's
    rPr (font, bold, colour) are kept; other runs are removed."""
    runs = list(p.iter(_w("r")))
    first = runs[0] if runs else None
    if first is not None and first.getparent() is not p:
        # lift the run out of a hyperlink / tracked-change wrapper
        first.getparent().addprevious(first)
    for el in list(p):
        if el is first or etree.QName(el).localname in _KEEP:
            continue
        p.remove(el)
    if first is None:
        first = etree.SubElement(p, _w("r"))
    for ch in list(first):
        if ch.tag != _w("rPr"):
            first.remove(ch)
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        if i:
            etree.SubElement(first, _w("br"))
        parts = line.split("\t")
        for j, seg in enumerate(parts):
            if j:
                etree.SubElement(first, _w("tab"))
            t = etree.SubElement(first, _w("t"))
            t.text = seg
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def replace_in_para(p, find: str, repl: str) -> bool:
    full = _ptext(p)
    if find not in full:
        return False
    ts = list(p.iter(_w("t")))
    if sum((t.text or "").count(find) for t in ts) == full.count(find):
        for t in ts:
            if t.text and find in t.text:
                t.text = t.text.replace(find, repl)
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    else:
        set_para_text(p, full.replace(find, repl))
    return True


# ------------------------------------------------------------------
# Editing
# ------------------------------------------------------------------

_ADDR = re.compile(r"^(?:(hdr|ftr)(\d+)/)?(?:p(\d+)|t(\d+)(?:[\[(](\d+)\s*,\s*(\d+)[\])])?)$", re.I)


class _Ctx:
    def __init__(self, doc):
        self.doc = doc
        self.snap = Snapshot(xml_fp)
        self.snap.take(self.items())
        self.rules = PartRules()
        self.changes: list[str] = []

    def items(self):
        out = []
        for i, el in enumerate(self.doc.element.body, 1):
            out.append((f"body[{i}] {etree.QName(el).localname}", el))
        return out

    def resolve(self, addr: str):
        m = _ADDR.match(str(addr).strip().replace(" ", ""))
        require(bool(m), f"bad address '{addr}' (use p12, t2, t2[3,1], hdr1/p1)")
        hf, sec, pn, tn, r, c = m.groups()
        if hf:
            si = int(sec)
            require(1 <= si <= len(self.doc.sections), f"section {si} does not exist")
            part = self.doc.sections[si - 1].header if hf.lower() == "hdr" else self.doc.sections[si - 1].footer
            require(not part.is_linked_to_previous, f"section {si} has no own {'header' if hf.lower() == 'hdr' else 'footer'}")
            self.rules.changed.add(part.part.partname.lstrip("/"))
            container = part._element
            require(pn is not None, "headers/footers are addressed by paragraph (hdr1/p2)")
            ps = container.findall(_w("p"))
            require(1 <= int(pn) <= len(ps), f"{hf}{si} has {len(ps)} paragraphs")
            return "p", ps[int(pn) - 1]
        self.rules.changed.add("word/document.xml")
        if pn:
            ps = _paras(self.doc)
            require(1 <= int(pn) <= len(ps), f"document has {len(ps)} paragraphs, not p{pn}")
            return "p", ps[int(pn) - 1]
        ts = _tables(self.doc)
        require(1 <= int(tn) <= len(ts), f"document has {len(ts)} tables, not t{tn}")
        tbl = ts[int(tn) - 1]
        if r:
            rows = _rows(tbl)
            require(1 <= int(r) <= len(rows), f"t{tn} has {len(rows)} rows")
            cells = _cells(rows[int(r) - 1])
            require(1 <= int(c) <= len(cells), f"row {r} of t{tn} has {len(cells)} cells")
            return "cell", cells[int(c) - 1]
        return "tbl", tbl


def _set_cell_text(cell, text: str):
    ps = cell.findall(_w("p"))
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        if i < len(ps):
            p = ps[i]
        else:
            p = copy.deepcopy(ps[-1])
            cell.findall(_w("p"))[-1].addnext(p)
        set_para_text(p, line)
    for p in ps[len(lines):]:
        cell.remove(p)


def _op_set_text(ctx: _Ctx, op: dict):
    addr = op_arg(op, "addr", "target", "address")
    text = str(op_arg(op, "text", "value", default=""))
    kind, el = ctx.resolve(addr)
    require(kind in ("p", "cell"), f"'{addr}' is a table; address a cell like {addr}[1,1]")
    ctx.snap.touch(el)
    (set_para_text if kind == "p" else _set_cell_text)(el, text)
    ctx.changes.append(f"{addr}: text set")


def _op_insert_paragraph(ctx: _Ctx, op: dict):
    addr = op_arg(op, "after", "addr", "target")
    kind, el = ctx.resolve(addr)
    require(kind in ("p", "tbl"), "insert_paragraph_after needs a paragraph or table address")
    texts = op.get("texts") or [op_arg(op, "text", default="")]
    style = op.get("style")
    like = el if kind == "p" else None
    anchor = el
    for t in texts:
        if like is not None:
            new = copy.deepcopy(like)
        else:
            new = etree.Element(_w("p"))
        anchor.addnext(new)
        set_para_text(new, str(t))
        if style:
            sid = _style_id(ctx.doc, style)
            ppr = new.find(_w("pPr"))
            if ppr is None:
                ppr = etree.Element(_w("pPr"))
                new.insert(0, ppr)
            ps = ppr.find(_w("pStyle"))
            if ps is None:
                ps = etree.Element(_w("pStyle"))
                ppr.insert(0, ps)
            ps.set(_w("val"), sid)
        ctx.snap.touch(new)
        anchor = new
    ctx.changes.append(f"inserted {len(texts)} paragraph(s) after {addr}")


def _style_id(doc, name: str) -> str:
    for s in doc.styles:
        if s.name and s.name.lower() == str(name).lower():
            return s.style_id
    raise DocOpError(f"style '{name}' not found in this document")


def _op_delete(ctx: _Ctx, op: dict):
    addr = op_arg(op, "addr", "target")
    kind, el = ctx.resolve(addr)
    require(kind in ("p", "tbl"), "delete needs a paragraph or table address")
    ctx.snap.touch(el)
    el.getparent().remove(el)
    ctx.changes.append(f"deleted {addr}")


def _op_add_table_row(ctx: _Ctx, op: dict):
    addr = op_arg(op, "addr", "table", "target")
    kind, tbl = ctx.resolve(addr)
    require(kind == "tbl", "add_table_row needs a table address like t2")
    rows = _rows(tbl)
    after = int(op.get("after") or len(rows))
    require(1 <= after <= len(rows), f"'after' must be 1..{len(rows)}")
    values = op_arg(op, "values", "cells")
    src = rows[after - 1]
    new = copy.deepcopy(src)
    src.addnext(new)
    for i, tc in enumerate(_cells(new)):
        _set_cell_text(tc, str(values[i]) if i < len(values) else "")
    ctx.snap.touch(new)
    ctx.changes.append(f"{addr}: row added after row {after}")


def _op_replace_text(ctx: _Ctx, op: dict):
    find = str(op_arg(op, "find", "old"))
    repl = str(op.get("replace", op.get("new", "")))
    require(find != "", "replace_text needs a non-empty 'find'")
    hits = 0
    containers = [("word/document.xml", ctx.doc.element.body)]
    for sec in ctx.doc.sections:
        for hf in (sec.header, sec.footer):
            if not hf.is_linked_to_previous:
                containers.append((hf.part.partname.lstrip("/"), hf._element))
    for part, root in containers:
        for p in list(root.iter(_w("p"))):
            if replace_in_para(p, find, repl):
                ctx.snap.touch(p)
                ctx.rules.changed.add(part)
                hits += 1
    require(hits > 0, f"'{find}' not found in the document")
    ctx.changes.append(f"replaced '{find}' -> '{repl}' in {hits} paragraph(s)")


OPS = {
    "set_text": _op_set_text,
    "set_paragraph": _op_set_text,
    "set_table_cell": _op_set_text,
    "insert_paragraph_after": _op_insert_paragraph,
    "delete_paragraph": _op_delete,
    "delete": _op_delete,
    "add_table_row": _op_add_table_row,
    "replace_text": _op_replace_text,
}


def apply(data: bytes, ops: list[dict]) -> EditResult:
    doc = _open(data)
    ctx = _Ctx(doc)
    for op in ops:
        fn = OPS.get(str(op.get("op", "")).lower())
        require(fn is not None, f"unknown docx op '{op.get('op')}' (have: {', '.join(OPS)})")
        fn(ctx, op)
    problems = ctx.snap.problems(ctx.items())
    require(not problems, "edit would change parts you did not ask for: " + "; ".join(problems[:8]))
    buf = io.BytesIO()
    doc.save(buf)
    out = buf.getvalue()
    problems = check_package(data, out, ctx.rules)
    require(not problems, "edit would change untouched document parts: " + "; ".join(problems[:8]))
    return EditResult(out, ctx.changes, [])


# ------------------------------------------------------------------
# Creation from markdown
# ------------------------------------------------------------------

def create(md: str, title: str = "") -> bytes:
    import docx
    doc = docx.Document()
    lines = (md or "").splitlines()
    i = 0
    in_code, code = False, []
    while i < len(lines):
        line = lines[i]
        st = line.strip()
        if st.startswith("```"):
            if in_code:
                p = doc.add_paragraph("\n".join(code))
                for r in p.runs:
                    r.font.name = "Consolas"
                code, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if st.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{2,}", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^:?-{2,}:?$", cells[0] or "--"):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            for r, vals in enumerate(rows):
                for c, v in enumerate(vals):
                    t.cell(r, c).text = v.replace("**", "")
                    if r == 0:
                        for run in t.cell(r, c).paragraphs[0].runs:
                            run.bold = True
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", st)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4) if len(m.group(1)) > 1 or title else 0)
        elif re.match(r"^[-*+]\s+", st):
            _add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*+]\s+", "", st))
        elif re.match(r"^\d+[.)]\s+", st):
            _add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", st))
        elif st:
            _add_inline(doc.add_paragraph(), st)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_inline(p, text: str):
    for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            p.add_run(part).bold = bool(k % 2)
