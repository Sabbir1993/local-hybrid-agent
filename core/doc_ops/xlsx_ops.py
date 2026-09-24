"""Excel (.xlsx): structured outline + in-place edit ops.

Addresses: `Sheet1!B7`, `'My Sheet'!B2:D9`, or `B7` (first sheet).

Cell edits (set_cell / set_range / set_formula / clear / append_rows) patch the
sheet XML directly, so charts, pivots, images, conditional formats and styles
survive untouched. Structural edits (insert/delete rows or columns, add/rename
sheet) go through openpyxl and are refused on workbooks openpyxl would damage.
"""

import io
import posixpath
import re
import zipfile
from typing import Any, Optional

from lxml import etree

from .base import (DocOpError, EditResult, PartRules, Snapshot, check_package, digest, op_arg,
                   parse_xml, read_zip, require, xml_fp)

S_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_TABLE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/table"
RT_CALC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

MAX_ROWS_SHOWN = 300
MAX_CELLS_SHOWN = 4000
MAX_CELLS_SCANNED = 1_000_000
# parts openpyxl cannot round-trip; structural edits are refused when present
LOSSY_PARTS = ("xl/drawings/", "xl/charts/", "xl/media/", "xl/pivotTables/", "xl/pivotCache/",
               "xl/externalLinks/", "xl/slicers/", "xl/slicerCaches/", "xl/timelines/",
               "xl/threadedComments/", "xl/activeX/", "xl/ctrlProps/", "xl/chartsheets/")


def _s(tag: str) -> str:
    return f"{{{S_NS}}}{tag}"


# ------------------------------------------------------------------
# A1 helpers
# ------------------------------------------------------------------

def col_to_num(col: str) -> int:
    n = 0
    for ch in col.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def num_to_col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")


def split_ref(ref: str) -> tuple[int, int]:
    m = _REF.match(ref.strip())
    require(bool(m), f"bad cell reference '{ref}'")
    return int(m.group(2)), col_to_num(m.group(1))


def _split_addr(addr: str) -> tuple[Optional[str], str]:
    addr = str(addr).strip()
    if "!" in addr:
        sh, ref = addr.rsplit("!", 1)
        sh = sh.strip()
        if sh.startswith("'") and sh.endswith("'"):
            sh = sh[1:-1].replace("''", "'")
        return sh, ref.strip()
    return None, addr


# ------------------------------------------------------------------
# Package access
# ------------------------------------------------------------------

class _Book:
    def __init__(self, data: bytes):
        self.data = data
        self.parts = read_zip(data)
        require("xl/workbook.xml" in self.parts, "not an Excel workbook")
        self.wb = parse_xml(self.parts["xl/workbook.xml"])
        rels = parse_xml(self.parts.get("xl/_rels/workbook.xml.rels", b"<Relationships/>"))
        self.wb_rels = rels
        rid_target = {r.get("Id"): r.get("Target") for r in rels}
        self.sheets: list[tuple[str, str]] = []   # (name, part name)
        for sh in self.wb.iter(_s("sheet")):
            tgt = rid_target.get(sh.get(f"{{{R_NS}}}id"), "")
            part = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join("xl", tgt))
            if part in self.parts:
                self.sheets.append((sh.get("name"), part))
        require(bool(self.sheets), "workbook has no worksheets")
        self.trees: dict[str, Any] = {}
        self.dirty: set[str] = set()
        self.removed: set[str] = set()

    def sheet_part(self, name: Optional[str]) -> tuple[str, str]:
        if name is None:
            return self.sheets[0]
        for n, p in self.sheets:
            if n == name:
                return n, p
        for n, p in self.sheets:
            if n.lower() == name.lower():
                return n, p
        raise DocOpError(f"sheet '{name}' not found (have: {', '.join(n for n, _ in self.sheets)})")

    def tree(self, part: str):
        if part not in self.trees:
            self.trees[part] = parse_xml(self.parts[part])
        return self.trees[part]

    def rels_of(self, part: str) -> list:
        d, f = posixpath.split(part)
        blob = self.parts.get(f"{d}/_rels/{f}.rels")
        return list(parse_xml(blob)) if blob else []

    def has_formulas(self) -> bool:
        return any(b"<f" in self.parts[p] or b":f" in self.parts[p] for _n, p in self.sheets)

    def save(self) -> bytes:
        for part in self.dirty:
            self.parts[part] = etree.tostring(self.trees[part], xml_declaration=True,
                                              encoding="UTF-8", standalone=True)
        src = zipfile.ZipFile(io.BytesIO(self.data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for info in src.infolist():
                name = info.filename.lstrip("/")
                if name in self.removed:
                    continue
                if name in self.dirty:
                    z.writestr(info, self.parts[name], compress_type=zipfile.ZIP_DEFLATED)
                else:
                    z.writestr(info, src.read(info))   # untouched part: byte-identical
        return out.getvalue()


# ------------------------------------------------------------------
# Reading
# ------------------------------------------------------------------

def _fmt(v) -> str:
    if v is None:
        return ""
    s = str(v)
    return s if len(s) <= 80 else s[:80] + "..."


def inspect(data: bytes) -> dict:
    import openpyxl
    book = _Book(data)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    extras = sorted({p.split("/")[1] for p in book.parts if p.startswith(LOSSY_PARTS)})
    head = f"Excel workbook: sheets {[n for n, _ in book.sheets]}"
    if extras:
        head += f" (also contains: {', '.join(extras)} - preserved by cell edits)"
    els, shown = [], 0
    for ws in wb.worksheets:
        dims = ws.calculate_dimension() if hasattr(ws, "calculate_dimension") else ""
        els.append({"addr": f"{ws.title}!", "kind": "sheet", "text": f"range {dims}"})
        for r_i, row in enumerate(ws.iter_rows(), 1):
            if r_i > MAX_ROWS_SHOWN or shown >= MAX_CELLS_SHOWN:
                els.append({"addr": f"{ws.title}!", "kind": "more", "text": f"(rows after {r_i - 1} not shown - use focus)"})
                break
            real = [c for c in row if getattr(c, "value", None) is not None and hasattr(c, "row")]
            if not real:
                continue
            cells = [(c.coordinate, c.value) for c in real]
            shown += len(cells)
            els.append({"addr": f"{ws.title}!row{real[0].row}", "kind": "row",
                        "text": " | ".join(f"{k}={_fmt(v)}" for k, v in cells)})
    wb.close()
    return {"kind": "xlsx", "header": head, "elements": els}


# ------------------------------------------------------------------
# XML cell editing
# ------------------------------------------------------------------

_NUM = re.compile(r"^-?\d+(\.\d+)?([eE][-+]?\d+)?$")


def _row_el(sheet_data, r: int, create: bool = True):
    prev = None
    for row in sheet_data.iterfind(_s("row")):
        rn = int(row.get("r", "0"))
        if rn == r:
            return row
        if rn > r:
            break
        prev = row
    if not create:
        return None
    row = etree.Element(_s("row"), r=str(r))
    if prev is None:
        sheet_data.insert(0, row)
    else:
        prev.addnext(row)
    return row


def _cell_el(row, r: int, c: int, create: bool = True):
    ref = f"{num_to_col(c)}{r}"
    prev = None
    for cel in row.iterfind(_s("c")):
        cr = cel.get("r")
        cc = split_ref(cr)[1] if cr else 0
        if cc == c:
            return cel, False
        if cc > c:
            break
        prev = cel
    if not create:
        return None, False
    cel = etree.Element(_s("c"), r=ref)
    if prev is None:
        row.insert(0, cel)
    else:
        prev.addnext(cel)
    return cel, True


class _Ctx:
    def __init__(self, book: _Book):
        self.book = book
        self.snap = Snapshot(xml_fp)
        self.rules = PartRules()
        self.changes: list[str] = []
        self.notes: list[str] = []
        self.formula_changed = False
        self.cells_changed = False
        self._snapped: set[str] = set()
        self._merged: dict[str, list] = {}
        self._tables: dict[str, list] = {}

    def sheet(self, name: Optional[str]):
        n, part = self.book.sheet_part(name)
        root = self.book.tree(part)
        if part not in self._snapped:
            self._snapped.add(part)
            self.snap.take(_sheet_items(n, root))
            self._merged[part] = [_range_bounds(m.get("ref")) for m in root.iter(_s("mergeCell"))]
            self._tables[part] = self._load_tables(part)
        return n, part, root

    def _load_tables(self, part: str) -> list:
        out = []
        for r in self.book.rels_of(part):
            if r.get("Type") == RT_TABLE:
                tp = posixpath.normpath(posixpath.join(posixpath.dirname(part), r.get("Target")))
                if tp in self.book.parts:
                    t = self.book.tree(tp)
                    out.append((tp, t))
        return out

    def check_writable(self, part: str, r: int, c: int, cel):
        for (r1, c1, r2, c2) in self._merged[part]:
            if r1 <= r <= r2 and c1 <= c <= c2 and (r, c) != (r1, c1):
                raise DocOpError(f"{num_to_col(c)}{r} is inside merged range; edit the top-left cell "
                                 f"{num_to_col(c1)}{r1}")
        for _tp, t in self._tables[part]:
            r1, c1, r2, c2 = _range_bounds(t.get("ref"))
            if r == r1 and c1 <= c <= c2 and t.get("headerRowCount", "1") != "0":
                raise DocOpError(f"{num_to_col(c)}{r} is a table header (table '{t.get('displayName')}'); "
                                 "renaming table columns is not supported")
        if cel is not None:
            f = cel.find(_s("f"))
            if f is not None and f.get("ref") and f.get("t") in ("shared", "array"):
                r1, c1, r2, c2 = _range_bounds(f.get("ref"))
                if (r1, c1) != (r2, c2):
                    raise DocOpError(f"{num_to_col(c)}{r} holds a {f.get('t')} formula for {f.get('ref')}; "
                                     "editing it would break the other cells")


def _range_bounds(ref: str) -> tuple[int, int, int, int]:
    a, _, b = (ref or "A1").partition(":")
    r1, c1 = split_ref(a)
    r2, c2 = split_ref(b) if b else (r1, c1)
    return r1, c1, r2, c2


def _sheet_items(name: str, root):
    items = []
    for el in root:
        ln = etree.QName(el).localname
        if ln == "sheetData":
            for row in el.iterfind(_s("row")):
                for cel in row.iterfind(_s("c")):
                    items.append((f"{name}!{cel.get('r')}", cel))
        elif ln != "dimension":
            items.append((f"{name}!<{ln}>", el))
    return items


def _coerce(value, allow_formula: bool):
    """(kind, payload) for a cell value from the model."""
    if value is None or value == "":
        return "empty", None
    if isinstance(value, bool):
        return "bool", value
    if isinstance(value, (int, float)):
        return "num", value
    s = str(value)
    if s.startswith("="):
        if not allow_formula:
            raise DocOpError(f"value '{s[:40]}' looks like a formula - use set_formula "
                             "(or formulas: true in set_range) to write formulas")
        return "formula", s[1:]
    if _NUM.match(s.strip()):
        return "num", float(s) if any(ch in s for ch in ".eE") else int(s)
    return "str", s


def _write(ctx: _Ctx, part: str, root, r: int, c: int, value, allow_formula: bool = False,
           style_from_above: bool = True) -> None:
    kind, payload = _coerce(value, allow_formula)
    sd = root.find(_s("sheetData"))
    row = _row_el(sd, r)
    cel_existing, _ = _cell_el(row, r, c, create=False)
    ctx.check_writable(part, r, c, cel_existing)
    cel, created = _cell_el(row, r, c)
    ctx.snap.touch(cel)
    ctx.snap.touch(row)
    if created and style_from_above and r > 2:
        above = _row_el(sd, r - 1, create=False)
        if above is not None:
            a_cel, _ = _cell_el(above, r - 1, c, create=False)
            if a_cel is not None and a_cel.get("s"):
                cel.set("s", a_cel.get("s"))
    if cel.find(_s("f")) is not None:
        ctx.formula_changed = True
    for ch in list(cel):
        if etree.QName(ch).localname in ("f", "v", "is"):
            cel.remove(ch)
    cel.attrib.pop("t", None)
    if kind == "formula":
        etree.SubElement(cel, _s("f")).text = payload
        ctx.formula_changed = True
    elif kind == "bool":
        cel.set("t", "b")
        etree.SubElement(cel, _s("v")).text = "1" if payload else "0"
    elif kind == "num":
        etree.SubElement(cel, _s("v")).text = repr(payload) if isinstance(payload, float) else str(payload)
    elif kind == "str":
        cel.set("t", "inlineStr")
        is_ = etree.SubElement(cel, _s("is"))
        t = etree.SubElement(is_, _s("t"))
        t.text = payload
        if payload != payload.strip():
            t.set(XML_SPACE, "preserve")
    row.attrib.pop("spans", None)
    _grow_dimension(root, r, c)
    ctx.cells_changed = True
    ctx.book.dirty.add(part)
    ctx.rules.changed.add(part)


def _grow_dimension(root, r: int, c: int):
    dim = root.find(_s("dimension"))
    if dim is None:
        return
    r1, c1, r2, c2 = _range_bounds(dim.get("ref"))
    if r > r2 or c > c2 or r < r1 or c < c1:
        r1, c1, r2, c2 = min(r1, r), min(c1, c), max(r2, r), max(c2, c)
        dim.set("ref", f"{num_to_col(c1)}{r1}:{num_to_col(c2)}{r2}" if (r1, c1) != (r2, c2) else f"{num_to_col(c1)}{r1}")


def _op_set_cell(ctx: _Ctx, op: dict, formula: bool = False):
    sheet, ref = _split_addr(op_arg(op, "addr", "cell", "target"))
    value = op.get("formula") if formula and "formula" in op else op.get("value", op.get("text"))
    if formula and value is not None and not str(value).startswith("="):
        value = "=" + str(value)
    name, part, root = ctx.sheet(sheet)
    r, c = split_ref(ref)
    _write(ctx, part, root, r, c, value, allow_formula=formula)
    ctx.changes.append(f"{name}!{ref} = {_fmt(value)}")


def _op_set_range(ctx: _Ctx, op: dict):
    sheet, ref = _split_addr(op_arg(op, "addr", "range", "target"))
    values = op_arg(op, "values", "rows")
    require(isinstance(values, list) and all(isinstance(r, list) for r in values),
            "set_range needs values as a list of rows, e.g. [[1, 2], [3, 4]]")
    name, part, root = ctx.sheet(sheet)
    r0, c0 = split_ref(ref.split(":")[0])
    if ":" in ref:
        r1, c1, r2, c2 = _range_bounds(ref)
        require(len(values) <= r2 - r1 + 1 and all(len(v) <= c2 - c1 + 1 for v in values),
                f"values do not fit {ref}")
    n = 0
    for i, rowvals in enumerate(values):
        for j, v in enumerate(rowvals):
            _write(ctx, part, root, r0 + i, c0 + j, v, allow_formula=bool(op.get("formulas")))
            n += 1
    ctx.changes.append(f"{name}!{ref}: {n} cells set")


def _op_clear(ctx: _Ctx, op: dict):
    sheet, ref = _split_addr(op_arg(op, "addr", "range", "target"))
    name, part, root = ctx.sheet(sheet)
    r1, c1, r2, c2 = _range_bounds(ref)
    require((r2 - r1 + 1) * (c2 - c1 + 1) <= 100_000, "range too large to clear")
    sd = root.find(_s("sheetData"))
    for r in range(r1, r2 + 1):
        row = _row_el(sd, r, create=False)
        if row is None:
            continue
        for c in range(c1, c2 + 1):
            cel, _ = _cell_el(row, r, c, create=False)
            if cel is not None:
                _write(ctx, part, root, r, c, None)
    ctx.changes.append(f"{name}!{ref} cleared (formatting kept)")


def _op_append_rows(ctx: _Ctx, op: dict):
    sheet = op.get("sheet") or _split_addr(op.get("addr") or "")[0]
    rows = op_arg(op, "rows", "values")
    require(isinstance(rows, list) and all(isinstance(r, list) for r in rows),
            "append_rows needs rows as a list of lists")
    name, part, root = ctx.sheet(sheet)
    sd = root.find(_s("sheetData"))
    last = 0
    first_col = 1
    for row in sd.iterfind(_s("row")):
        if any(c.find(_s("v")) is not None or c.find(_s("is")) is not None or c.find(_s("f")) is not None
               for c in row.iterfind(_s("c"))):
            last = max(last, int(row.get("r", "0")))
    last_row = _row_el(sd, last, create=False) if last else None
    if last_row is not None:
        cols = [split_ref(c.get("r"))[1] for c in last_row.iterfind(_s("c")) if c.get("r")]
        first_col = min(cols) if cols else 1
    for i, vals in enumerate(rows):
        r = last + 1 + i
        for j, v in enumerate(vals):
            c = first_col + j
            _write(ctx, part, root, r, c, v, allow_formula=bool(op.get("formulas")), style_from_above=False)
            if last_row is not None:
                src, _ = _cell_el(last_row, last, c, create=False)
                cel, _ = _cell_el(_row_el(sd, r), r, c, create=False)
                if src is not None and src.get("s") and cel is not None:
                    cel.set("s", src.get("s"))
    # a table that ended on the old last row grows with the data
    for tp, t in ctx._tables[part]:
        r1, c1, r2, c2 = _range_bounds(t.get("ref"))
        if r2 == last:
            new_ref = f"{num_to_col(c1)}{r1}:{num_to_col(c2)}{last + len(rows)}"
            t.set("ref", new_ref)
            af = t.find(_s("autoFilter"))
            if af is not None:
                af.set("ref", new_ref)
            ctx.book.dirty.add(tp)
            ctx.rules.changed.add(tp)
            ctx.notes.append(f"table '{t.get('displayName')}' extended to {new_ref}")
    ctx.changes.append(f"{name}: appended {len(rows)} rows after row {last}")


XML_OPS = {
    "set_cell": _op_set_cell,
    "set_formula": lambda ctx, op: _op_set_cell(ctx, op, formula=True),
    "set_range": _op_set_range,
    "clear": _op_clear,
    "clear_range": _op_clear,
    "append_rows": _op_append_rows,
}


def _finish_xml(ctx: _Ctx) -> bytes:
    book = ctx.book
    if ctx.formula_changed:
        # calcChain lists formula cells; a stale one makes Excel "repair" the file
        calc = [r for r in book.wb_rels if r.get("Type") == RT_CALC]
        for r in calc:
            book.wb_rels.remove(r)
            book.removed.add("xl/calcChain.xml")
        if calc:
            book.trees["xl/_rels/workbook.xml.rels"] = book.wb_rels
            book.dirty.add("xl/_rels/workbook.xml.rels")
            ct = book.tree("[Content_Types].xml")
            for o in list(ct):
                if o.get("PartName", "").lstrip("/") == "xl/calcChain.xml":
                    ct.remove(o)
            book.dirty.add("[Content_Types].xml")
            ctx.rules.changed |= {"xl/workbook.xml:rels", "[Content_Types].xml"}
            ctx.rules.removed.add("xl/calcChain.xml")
    if ctx.cells_changed and (ctx.formula_changed or book.has_formulas()):
        calc_pr = book.wb.find(_s("calcPr"))
        if calc_pr is None:
            calc_pr = etree.SubElement(book.wb, _s("calcPr"))
            # schema order: calcPr sits before these trailing elements
            for tag in ("oleSize", "customWorkbookViews", "pivotCaches", "smartTagPr", "smartTagTypes",
                        "webPublishing", "fileRecoveryPr", "webPublishObjects", "extLst"):
                nxt = book.wb.find(_s(tag))
                if nxt is not None:
                    nxt.addprevious(calc_pr)
                    break
        calc_pr.set("fullCalcOnLoad", "1")
        book.trees["xl/workbook.xml"] = book.wb
        book.dirty.add("xl/workbook.xml")
        ctx.rules.changed.add("xl/workbook.xml")
        ctx.notes.append("formulas will recalculate when the file is opened")
    items = []
    for n, p in book.sheets:
        if p in ctx._snapped:
            items += _sheet_items(n, book.trees[p])
    problems = ctx.snap.problems(items)
    if problems:
        raise DocOpError("edit would change cells you did not ask for: " + "; ".join(problems[:8]))
    return book.save()


# ------------------------------------------------------------------
# Structural edits via openpyxl
# ------------------------------------------------------------------

def _cell_fp(cell) -> str:
    return digest(repr((cell.value, cell._style if hasattr(cell, "_style") else None)).encode())


def _structural(data: bytes, ops: list[dict]) -> EditResult:
    import openpyxl
    book = _Book(data)
    lossy = sorted({p.split("/")[1] for p in book.parts if p.startswith(LOSSY_PARTS)})
    require(not lossy, f"this workbook contains {', '.join(lossy)}, which would be lost by a structural "
                       "edit; use set_cell / set_range / append_rows instead")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    require(sum(len(ws._cells) for ws in wb.worksheets) <= MAX_CELLS_SCANNED, "workbook too large")
    snap = Snapshot(_cell_fp)
    snap.take((f"{ws.title}!{c.coordinate}", c) for ws in wb.worksheets for c in ws._cells.values())
    changes, notes = [], []

    def ws_of(name):
        if not name:
            return wb.worksheets[0]
        for ws in wb.worksheets:
            if ws.title.lower() == str(name).lower():
                return ws
        raise DocOpError(f"sheet '{name}' not found")

    def guard_shift(ws):
        require(not ws.merged_cells.ranges, f"sheet '{ws.title}' has merged cells; inserting/deleting "
                                            "rows would misalign them")
        require(not ws.tables, f"sheet '{ws.title}' has tables; use append_rows instead")
        require(not ws.conditional_formatting and not ws.data_validations.dataValidation,
                f"sheet '{ws.title}' has conditional formatting/validation that would not shift")
        require(not book.has_formulas() and not wb.defined_names,
                "workbook has formulas or named ranges that openpyxl would not re-point; use "
                "set_range / append_rows instead")

    for op in ops:
        kind = str(op.get("op", "")).lower()
        sheet = op.get("sheet") or _split_addr(op.get("addr") or "")[0]
        if kind in ("insert_rows", "delete_rows", "insert_cols", "delete_cols"):
            ws = ws_of(sheet)
            guard_shift(ws)
            at = op_arg(op, "at", "row", "col", "index")
            at = col_to_num(at) if isinstance(at, str) and at.isalpha() else int(at)
            n = int(op.get("count") or 1)
            if kind.startswith("delete"):
                rows = kind == "delete_rows"
                for c in list(ws._cells.values()):
                    if (c.row if rows else c.column) in range(at, at + n):
                        snap.touch(c)
            getattr(ws, kind)(at, n)
            if kind == "insert_rows":
                vals = op.get("values") or []
                for i, rowvals in enumerate(vals[:n]):
                    for j, v in enumerate(rowvals, 1):
                        cell = ws.cell(row=at + i, column=j)
                        kind_v, payload = _coerce(v, bool(op.get("formulas")))
                        cell.value = ("=" + payload) if kind_v == "formula" else payload
                        snap.touch(cell)
            changes.append(f"{ws.title}: {kind.replace('_', ' ')} at {at} x{n}")
        elif kind == "add_sheet":
            name = str(op_arg(op, "name", "title"))
            require(name not in wb.sheetnames, f"sheet '{name}' already exists")
            ws = wb.create_sheet(name, int(op["index"]) if op.get("index") is not None else None)
            for i, rowvals in enumerate(op.get("rows") or [], 1):
                for j, v in enumerate(rowvals, 1):
                    kind_v, payload = _coerce(v, bool(op.get("formulas")))
                    snap.touch(ws.cell(row=i, column=j, value=("=" + payload) if kind_v == "formula" else payload))
            changes.append(f"added sheet '{name}'")
        elif kind == "rename_sheet":
            ws = ws_of(op_arg(op, "sheet", "name", "from"))
            new = str(op_arg(op, "to", "new_name"))
            old = ws.title
            require(not any(isinstance(c.value, str) and c.value.startswith("=") and old in c.value
                            for w in wb.worksheets for c in w._cells.values()) and not wb.defined_names,
                    f"formulas or named ranges refer to '{old}'; renaming would break them")
            ws.title = new
            changes.append(f"renamed sheet '{old}' -> '{new}'")
        else:
            raise DocOpError(f"unknown xlsx op '{op.get('op')}'")
    problems = snap.problems((f"{ws.title}!{c.coordinate}", c) for ws in wb.worksheets for c in ws._cells.values())
    if problems:
        raise DocOpError("edit would change cells you did not ask for: " + "; ".join(problems[:8]))
    out = io.BytesIO()
    wb.save(out)
    notes.append("workbook re-saved by openpyxl (structural edit)")
    return EditResult(out.getvalue(), changes, notes)


STRUCTURAL_OPS = {"insert_rows", "delete_rows", "insert_cols", "delete_cols", "add_sheet", "rename_sheet"}


def apply(data: bytes, ops: list[dict]) -> EditResult:
    kinds = [str(o.get("op", "")).lower() for o in ops]
    unknown = [k for k in kinds if k not in XML_OPS and k not in STRUCTURAL_OPS]
    require(not unknown, f"unknown xlsx op(s) {unknown} (have: {', '.join(list(XML_OPS) + sorted(STRUCTURAL_OPS))})")
    if any(k in STRUCTURAL_OPS for k in kinds):
        # structural first (openpyxl), then cell edits on its output (XML path)
        s_ops = [o for o, k in zip(ops, kinds) if k in STRUCTURAL_OPS]
        c_ops = [o for o, k in zip(ops, kinds) if k not in STRUCTURAL_OPS]
        res = _structural(data, s_ops)
        if c_ops:
            more = apply(res.data, c_ops)
            return EditResult(more.data, res.changes + more.changes, res.notes + more.notes)
        return res
    book = _Book(data)
    ctx = _Ctx(book)
    for op, k in zip(ops, kinds):
        XML_OPS[k](ctx, op)
    out = _finish_xml(ctx)
    problems = check_package(data, out, ctx.rules)
    if problems:
        raise DocOpError("edit would change untouched workbook parts: " + "; ".join(problems[:8]))
    return EditResult(out, ctx.changes, ctx.notes)


# ------------------------------------------------------------------
# Creation
# ------------------------------------------------------------------

def create(rows: list[list], sheet: str = "Sheet1") -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31] or "Sheet1"
    for r, vals in enumerate(rows, 1):
        for c, v in enumerate(vals, 1):
            kind, payload = _coerce(v, allow_formula=False) if not (isinstance(v, str) and v.startswith("=")) \
                else ("str", v)
            cell = ws.cell(row=r, column=c, value=payload if kind != "empty" else None)
            if kind == "str":
                cell.data_type = "s"   # model text like "=cmd" stays text, never a formula
    if rows:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDE7F5")
        ws.freeze_panes = "A2"
        for c in range(1, max(len(r) for r in rows) + 1):
            width = max((len(str(r[c - 1])) for r in rows if len(r) >= c and r[c - 1] is not None), default=8)
            ws.column_dimensions[num_to_col(c)].width = min(max(width + 2, 8), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
