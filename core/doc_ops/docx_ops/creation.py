"""Markdown -> .docx with the same house style as the PDF path (core/agent_tools/doc_theme.py).

Built-in Word styles keep their names (Heading N, List Bullet, List Number, Title) so inspect/edit addressing and
Word's own navigation pane still work. Long documents (4+ sections) get a cover page, 5+ a contents field; every page
but the cover carries a footer with the title and "Page X of Y". Tables get a shaded repeating header row, banding and
right-aligned numbers. Links keep only safe schemes. Word and LibreOffice update the contents field's page numbers
on demand; until then it shows the section list.
"""
import datetime
import io
import re

_ACCENT, _INK, _MUTED, _HEAD2, _RULE, _BAND, _HEADFILL = "2563EB", "0F172A", "64748B", "1E40AF", "CBD5E1", "F8FAFC", "F1F5F9"
COVER_MIN_SECTIONS = 4
TOC_MIN_SECTIONS = 5
_BENGALI = re.compile(r"[ঀ-৿]")
_SAFE_LINK = re.compile(r"^(https?://|mailto:)", re.I)
_INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*])|`[^`]+`|\[[^\]]+\]\([^)\s]+\))")
_NUMERIC = re.compile(r"^[\(\-+−]?\s*(?:[$৳€£]|BDT|Tk\.?)?\s*[\d,]+(?:\.\d+)?\s*%?\)?$")


def _hex(c):
    from docx.shared import RGBColor
    return RGBColor.from_string(c)


def _depth(line: str) -> int:
    """Nesting level of a list line: 2 spaces (or a tab) per level, at most 2."""
    ex = line.expandtabs(2)
    return min((len(ex) - len(ex.lstrip())) // 2, 2)


def _plain(text: str) -> str:
    """Inline markdown reduced to its visible text (headings, TOC entries, table of contents cache)."""
    t = re.sub(r"\[([^\]]+)\]\([^)\s]+\)", r"\1", text or "")
    return re.sub(r"[*`]", "", t).strip()


def _add_hyperlink(p, text: str, url: str):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:color", {"w:val": _ACCENT}), ("w:u", {"w:val": "single"})):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def _add_inline(p, text: str):
    """**bold**, *italic*, ***both***, `code`, [label](https://...). Unsafe link schemes keep only the label."""
    for part in _INLINE.split(text or ""):
        if not part:
            continue
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            r = p.add_run(part[3:-3])
            r.bold = r.italic = True
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.color.rgb = _hex(_HEAD2)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = part[1:-1].split("](", 1)
            if _SAFE_LINK.match(url):
                _add_hyperlink(p, label, url)
            else:
                p.add_run(label)
        else:
            p.add_run(part)


def _field(p, instr: str, cached: str = "", size=None, color=None):
    """A complex field (PAGE, NUMPAGES, TOC ...) with cached text so it reads right before Word updates it."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def run(child):
        r = p.add_run()
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = _hex(color)
        r._r.append(child)
        return r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run(b)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run(it)
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run(s)
    t = OxmlElement("w:t")
    t.text = cached
    run(t)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run(e)


# Child order Word expects inside these property elements (it can flag a file whose children are out of order).
_ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split(),
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign "
            "hideMark".split(),
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook".split(),
}


def _put(parent, el):
    """Insert el into parent (pPr / tcPr / tblPr) at its schema position; an existing element of that name is replaced."""
    from docx.oxml.ns import qn
    seq = _ORDER[parent.tag.rsplit("}", 1)[1]]
    name = el.tag.rsplit("}", 1)[1]
    for old in parent.findall(el.tag):
        parent.remove(old)
    idx = seq.index(name)
    for child in list(parent):
        cn = child.tag.rsplit("}", 1)[1]
        if cn in seq and seq.index(cn) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def _shade(fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    return shd


def _borders(el_pr, spec: dict, tag: str):
    """spec: {'top': (size, color), 'bottom': ...}; size in eighths of a point."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    b = OxmlElement(tag)
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        if side in spec:
            sz, col = spec[side]
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), col)
        else:
            e.set(qn("w:val"), "nil")
        e.set(qn("w:space"), "0")
        b.append(e)
    _put(el_pr, b)


def _setup(doc, title: str, bengali: bool):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    st = doc.styles
    n = st["Normal"]
    n.font.name, n.font.size = "Calibri", Pt(11)
    rpr = n.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(a), "Calibri")
    fonts.set(qn("w:cs"), "Nirmala UI" if bengali else "Calibri")
    n.font.color.rgb = _hex("1E293B")
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    n.paragraph_format.widow_control = True
    for name, size, color, before, after in (("Title", 30, _INK, 0, 10), ("Heading 1", 22, _INK, 18, 8),
                                             ("Heading 2", 16, _HEAD2, 16, 6), ("Heading 3", 13, "334155", 12, 4),
                                             ("Heading 4", 11.5, "334155", 10, 3)):
        s = st[name]
        s.font.name, s.font.size, s.font.bold = "Calibri", Pt(size), True
        s.font.color.rgb = _hex(color)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
        s.paragraph_format.keep_with_next = True
        fr = s.element.get_or_add_rPr().find(qn("w:rFonts"))
        if fr is not None:                                        # drop theme fonts so the family above applies
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if fr.get(qn(a)) is not None:
                    del fr.attrib[qn(a)]
            fr.set(qn("w:cs"), "Nirmala UI" if bengali else "Calibri")
    h1 = st["Heading 2"].element.get_or_add_pPr()
    _borders(h1, {"bottom": (6, _RULE)}, "w:pBdr")
    ttl = st["Title"].element.get_or_add_pPr()
    _borders(ttl, {"bottom": (24, _ACCENT)}, "w:pBdr")
    for name in ("Quote",):
        if name not in [s.name for s in st]:
            q = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            q.base_style = st["Normal"]
    q = st["Quote"]
    q.font.italic, q.font.color.rgb = True, _hex("475569")
    q.paragraph_format.left_indent = Cm(0.6)
    qp = q.element.get_or_add_pPr()
    _borders(qp, {"left": (24, "94A3B8")}, "w:pBdr")
    _put(qp, _shade(_BAND))
    if title:                       # only non-empty values: an empty element re-serialises on the next edit
        doc.core_properties.title = title[:200]


def _footer(doc, title: str, cover: bool):
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt
    sec = doc.sections[0]
    sec.different_first_page_header_footer = cover
    p = sec.footer.paragraphs[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    size = Pt(9)
    r = p.add_run((title or "")[:80])
    r.font.size, r.font.color.rgb = size, _hex("94A3B8")
    r = p.add_run("\tPage ")
    r.font.size, r.font.color.rgb = size, _hex("94A3B8")
    _field(p, "PAGE", "1", size, "94A3B8")
    r = p.add_run(" of ")
    r.font.size, r.font.color.rgb = size, _hex("94A3B8")
    _field(p, "NUMPAGES", "1", size, "94A3B8")
    _borders(p._p.get_or_add_pPr(), {"top": (4, "E2E8F0")}, "w:pBdr")
    p.paragraph_format.space_before = Cm(0)


def _restart_numbering(doc, p):
    """Give an ordered list its own counter so a second list starts at 1 (python-docx continues one global list)."""
    try:
        from docx.oxml.ns import qn
        numbering = doc.part.numbering_part.numbering_definitions._numbering
        sid = doc.styles["List Number"].element.pPr.numPr.numId.val
        abstract = next(n.abstractNumId.val for n in numbering.num_lst if n.numId == sid)
        num = numbering.add_num(abstract)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId
    except Exception:
        return None


def _set_num(p, num_id):
    if num_id is not None:
        p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_id


def _table(doc, rows: list, aligns: list):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    tbl_pr = t._tbl.tblPr
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "5000")
    w.set(qn("w:type"), "pct")
    _put(tbl_pr, w)
    _borders(tbl_pr, {"top": (4, _RULE), "bottom": (4, _RULE), "left": (4, "E2E8F0"), "right": (4, "E2E8F0")}, "w:tblBorders")
    cm = OxmlElement("w:tblCellMar")
    for side, v in (("top", 50), ("left", 100), ("bottom", 50), ("right", 100)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        cm.append(e)
    _put(tbl_pr, cm)
    for ri, vals in enumerate(rows):
        tr = t.rows[ri]._tr
        trpr = tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            v = vals[ci] if ci < len(vals) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            _add_inline(p, v)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = _hex(_INK)
            a = aligns[ci] if ci < len(aligns) else None
            if a == "r" or (a is None and ri > 0 and _NUMERIC.match(_plain(v) or "x")):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif a == "c":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            if ri == 0:
                _borders(tcpr, {"bottom": (12, _RULE)}, "w:tcBorders")
                _put(tcpr, _shade(_HEADFILL))
            elif ri % 2 == 0:
                _put(tcpr, _shade(_BAND))
    return t


def _toc(doc, headings: list):
    """Contents heading + a TOC field whose cached result is the section list (page numbers appear on update)."""
    from docx.shared import Pt
    h = doc.add_paragraph()           # not a Heading style, so the TOC field does not list itself
    r = h.add_run("Contents")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(22), _hex(_INK)
    h.paragraph_format.space_after = Pt(12)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = doc.add_paragraph()
    r = p.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r._r.append(it)
    r = p.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r._r.append(s)
    first = True
    for i, name in enumerate(headings, 1):
        if not first:
            p = doc.add_paragraph()
        p.add_run(f"{i:02d}   {name}")
        p.paragraph_format.space_after = Pt(4)
        first = False
    r = p.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r._r.append(e)


def create(md: str, title: str = "") -> bytes:
    import docx
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    md = md or ""
    lines = md.splitlines()
    bengali = bool(_BENGALI.search(md) or _BENGALI.search(title or ""))
    sections = [_plain(m.group(1)) for m in re.finditer(r"^##\s+(.*)$", md, re.M)]
    cover = len(sections) >= COVER_MIN_SECTIONS
    toc_on = len(sections) >= TOC_MIN_SECTIONS
    doc = docx.Document()
    first_h1 = next((_plain(m.group(1)) for m in re.finditer(r"^#\s+(.*)$", md, re.M)), "")
    doc_title = first_h1 or title
    _setup(doc, doc_title, bengali)
    _footer(doc, doc_title, cover)

    i = 0
    in_code, code = False, []
    cover_done = not cover
    num_id = None
    prev_kind = ""
    while i < len(lines):
        line = lines[i]
        st = line.strip()
        if st.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code))
                r.font.name, r.font.size = "Consolas", Pt(9.5)
                ppr = p._p.get_or_add_pPr()
                _put(ppr, _shade("F1F5F9"))
                _borders(ppr, {"left": (18, _RULE)}, "w:pBdr")
                code, in_code = [], False
            else:
                in_code = True
            i += 1
            prev_kind = "code"
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if st.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{2,}", lines[i + 1].strip()):
            raw = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw.append(lines[i].strip())
                i += 1
            aligns = []
            for cell in raw[1].strip("|").split("|"):
                c = cell.strip()
                aligns.append("c" if c.startswith(":") and c.endswith(":") else "r" if c.endswith(":") else None)
            rows = [[c.strip() for c in r.strip("|").split("|")] for k, r in enumerate(raw) if k != 1]
            if rows:
                _table(doc, rows, aligns)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            prev_kind = "table"
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", st)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            if cover and not cover_done and level == 1:
                # first H1 -> cover title; the next short intro paragraph -> subtitle; then page break (+ contents)
                doc.add_paragraph(_plain(text), style="Title")
                j = i + 1
                while j < len(lines) and not lines[j].strip():
                    j += 1
                if j < len(lines) and not re.match(r"^(#|[-*+]\s|\d+[.)]\s|\||```|>)", lines[j].strip()) \
                        and len(lines[j].strip()) <= 320:
                    sp = doc.add_paragraph()
                    _add_inline(sp, lines[j].strip())
                    for r in sp.runs:
                        r.font.size, r.font.color.rgb = Pt(14), _hex("475569")
                    i = j
                dp = doc.add_paragraph(datetime.date.today().strftime("%d %B %Y"))
                for r in dp.runs:
                    r.font.size, r.font.color.rgb = Pt(10.5), _hex(_MUTED)
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                if toc_on:
                    _toc(doc, sections)
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                cover_done = True
                i += 1
                prev_kind = "cover"
                continue
            doc.add_heading(_plain(text) if level > 1 or doc_title else text,
                            level=min(level, 4) if level > 1 or doc_title else 0)
            prev_kind = "heading"
        elif re.match(r"^([-*_])(\s*\1){2,}$", st):
            p = doc.add_paragraph()
            _borders(p._p.get_or_add_pPr(), {"bottom": (6, _RULE)}, "w:pBdr")
            prev_kind = "rule"
        elif st.startswith(">"):
            _add_inline(doc.add_paragraph(style="Quote"), re.sub(r"^>\s?", "", st))
            prev_kind = "quote"
        elif re.match(r"^\s*[-*+]\s+", line):
            depth = _depth(line)
            style = "List Bullet" if depth == 0 else f"List Bullet {min(depth + 1, 3)}"
            _add_inline(doc.add_paragraph(style=style), re.sub(r"^[-*+]\s+", "", st))
            prev_kind = "bullet"
        elif re.match(r"^\s*\d+[.)]\s+", line):
            depth = _depth(line)
            if depth == 0:
                if prev_kind != "number":
                    num_id = _restart_numbering(doc, None)
                p = doc.add_paragraph(style="List Number")
                _set_num(p, num_id)
            else:
                p = doc.add_paragraph(style=f"List Number {min(depth + 1, 3)}")
            _add_inline(p, re.sub(r"^\d+[.)]\s+", "", st))
            prev_kind = "number"
        elif st:
            _add_inline(doc.add_paragraph(), st)
            prev_kind = "para"
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
