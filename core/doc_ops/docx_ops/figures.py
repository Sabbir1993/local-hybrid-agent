"""Pictures and charts for a .docx: the same `image:` and `chart:` lines a PDF takes (core/agent_tools/dom_convert.py).

Pictures are the ones the chat already fetched or generated (slide_images.resolve_pdf). A chart is drawn as SVG from numbers we
parsed ourselves (pdf_charts), printed by the local browser and read back as a PNG - Word does not embed SVG reliably. No
browser means no chart image: the chart's numbers stay in the document as the plain list or table the model wrote."""
import io
import re
import tempfile
from pathlib import Path
from typing import Optional

CHART_TOKEN = re.compile(r"^PDFCHARTTOKEN(\d+)$")
FIGURE_LINE = re.compile(r"^!\[([^\]]*)\]\(pdfimg:(\d+)\)$")
_CHART_DIRECTIVE = re.compile(r"(?mi)^[ \t]*(?:chart|graph)[ \t]*:.*\r?\n?")
MAX_CHARTS = 6


def chart_png(svg: str) -> Optional[bytes]:
    """One chart SVG -> a cropped PNG, or None when there is no browser to print it."""
    from ...agent_tools.pdf_render import _render_html_to_pdf
    page = ('<!doctype html><html><head><meta charset="utf-8"><style>@page{size:210mm 297mm;margin:0}'
            'html,body{margin:0;background:#fff}svg{display:block;width:190mm;margin:8mm}</style></head><body>'
            + svg + "</body></html>")
    with tempfile.TemporaryDirectory(prefix="docxchart_") as d:
        pdf_path = Path(d) / "chart.pdf"
        try:
            if not _render_html_to_pdf(page, pdf_path):
                return None
            import pypdfium2 as pdfium
            doc = pdfium.PdfDocument(str(pdf_path))
            try:
                img = doc[0].render(scale=3).to_pil().convert("RGB")
            finally:
                doc.close()
        except Exception:
            return None
    from PIL import Image, ImageChops
    box = ImageChops.difference(img, Image.new("RGB", img.size, (255, 255, 255))).getbbox()
    if not box:
        return None
    pad = 14
    img = img.crop((max(0, box[0] - pad), max(0, box[1] - pad), min(img.width, box[2] + pad), min(img.height, box[3] + pad)))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def prepare(md: str, images: Optional[dict] = None) -> tuple:
    """-> (markdown, {n: picture info}, {n: chart png}). `image:` lines become `![alt](pdfimg:N)` lines (dropped when no picture
    was found) and `chart:` blocks become PDFCHARTTOKEN lines. When a chart cannot be drawn the numbers stay as plain text."""
    from ...agent_tools.dom_convert import extract_charts, prepare_images
    figs: dict = {}
    if re.search(r"(?mi)^\s*(?:image|picture|photo)\s*:", md):
        md, figs = prepare_images(md, images)
    charts: dict = {}
    if re.search(r"(?mi)^\s*(?:chart|graph)\s*:", md):
        md2, svgs = extract_charts(md)
        pngs = [chart_png(s) for s in svgs[:MAX_CHARTS]] if svgs and len(svgs) <= MAX_CHARTS else []
        if svgs and len(pngs) == len(svgs) and all(pngs):
            md, charts = md2, dict(enumerate(pngs))
        else:
            md = _CHART_DIRECTIVE.sub("", md)          # keep the numbers, drop the directive
    return md, figs, charts


def add_picture(doc, data: bytes, alt: str = "", caption: str = "", max_cm: float = 15.5, max_h_cm: float = 11.0) -> bool:
    """A centred picture (and a small caption) as its own paragraph. False when the bytes are not a usable image."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image
    try:
        with Image.open(io.BytesIO(data)) as im:
            iw, ih = im.size
    except Exception:
        return False
    if not iw or not ih:
        return False
    w = max_cm
    if w * ih / iw > max_h_cm:
        w = max_h_cm * iw / ih
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = bool(caption)
    try:
        shape = p.add_run().add_picture(io.BytesIO(data), width=Cm(w))
    except Exception:
        doc._body._body.remove(p._p)
        return False
    if alt:
        shape._inline.docPr.set("descr", alt[:250])
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        r.italic, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0x64, 0x74, 0x8B)
    return True
