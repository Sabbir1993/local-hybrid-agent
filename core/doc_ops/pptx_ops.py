"""PowerPoint (.pptx): structured outline + in-place edit ops.

Addresses (1-based, shown by inspect):
    s3                  slide 3
    s3/sh5              shape with shape_id 5 on slide 3 (stable, from the file)
    s3/sh5/p2           paragraph 2 of that shape
    s3/sh7/tbl[2,3]     table cell row 2, col 3
    s3/title            the slide's title placeholder
    s3/notes            speaker notes
"""

import copy
import io
import re
from typing import Optional

from lxml import etree

from .base import (DocOpError, EditResult, PartRules, Snapshot, check_package, op_arg, parse_xml,
                   read_zip, require, xml_fp)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
RT_NOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
RT_LAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
MAX_SLIDES = 500
MAX_TABLE_CELLS_SHOWN = 400


def _a(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


def _open(data: bytes):
    from pptx import Presentation
    read_zip(data)  # size / bomb / macro guards before python-pptx parses it
    prs = Presentation(io.BytesIO(data))
    require(len(prs.slides) <= MAX_SLIDES, f"presentation has more than {MAX_SLIDES} slides")
    return prs


def _save(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------
# Reading
# ------------------------------------------------------------------

def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)


def _shape_kind(sh) -> str:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if sh.is_placeholder:
        try:
            t = sh.placeholder_format.type
            name = str(t).split(".")[-1].split(" ")[0].lower()
            return "title" if name in ("title", "center_title") else f"placeholder:{name}"
        except Exception:
            return "placeholder"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    st = sh.shape_type
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if st == MSO_SHAPE_TYPE.TEXT_BOX:
        return "textbox"
    return "shape"


def _clip(s: str, n: int = 400) -> str:
    s = s.replace("\v", " ").strip()
    return s if len(s) <= n else s[:n] + "..."


def inspect(data: bytes) -> dict:
    prs = _open(data)
    els = []
    w, h = prs.slide_width or 0, prs.slide_height or 0
    ratio = "16:9" if h and abs(w / h - 16 / 9) < 0.05 else ("4:3" if h and abs(w / h - 4 / 3) < 0.05 else "custom")
    layouts = sorted({l.name for m in prs.slide_masters for l in m.slide_layouts})
    head = f"PowerPoint: {len(prs.slides)} slides ({ratio}). Layouts: {', '.join(layouts)}"
    for i, slide in enumerate(prs.slides, 1):
        els.append({"addr": f"s{i}", "kind": "slide", "text": f"layout: {slide.slide_layout.name}"})
        for sh in _iter_shapes(slide.shapes):
            base = f"s{i}/sh{sh.shape_id}"
            kind = _shape_kind(sh)
            if kind == "table":
                n = 0
                for r, row in enumerate(sh.table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        if n < MAX_TABLE_CELLS_SHOWN:
                            els.append({"addr": f"{base}/tbl[{r},{c}]", "kind": "cell", "text": _clip(cell.text, 120)})
                        n += 1
                els.append({"addr": base, "kind": "table",
                            "text": f"{len(sh.table.rows)}x{len(sh.table.columns)} '{sh.name}'"})
                continue
            if kind == "chart":
                ch = sh.chart
                title = ch.chart_title.text_frame.text if ch.has_title else ""
                series = [s.name for p in ch.plots for s in p.series]
                els.append({"addr": base, "kind": "chart", "text": f"title='{title}' series={series} (read-only)"})
                continue
            if kind == "picture":
                alt = sh._element.xpath("./p:nvPicPr/p:cNvPr/@descr")
                els.append({"addr": base, "kind": "picture", "text": f"'{sh.name}' alt='{alt[0] if alt else ''}'"})
                continue
            if kind == "group":
                els.append({"addr": base, "kind": "group", "text": f"'{sh.name}'"})
                continue
            if sh.has_text_frame:
                paras = sh.text_frame.paragraphs
                if len(paras) <= 1:
                    els.append({"addr": base, "kind": kind, "text": _clip(sh.text_frame.text)})
                else:
                    els.append({"addr": base, "kind": kind, "text": f"{len(paras)} paragraphs"})
                    for j, p in enumerate(paras, 1):
                        lvl = "  " * p.level
                        els.append({"addr": f"{base}/p{j}", "kind": "para", "text": lvl + _clip(p.text, 200)})
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame
            if nt is not None and nt.text.strip():
                els.append({"addr": f"s{i}/notes", "kind": "notes", "text": _clip(nt.text)})
    return {"kind": "pptx", "header": head, "elements": els}


# ------------------------------------------------------------------
# Visual preview model (for the UI slide viewer)
# ------------------------------------------------------------------

PREVIEW_MAX_IMAGE = 3 * 1024 * 1024     # per picture
PREVIEW_MAX_IMAGES = 24 * 1024 * 1024   # whole deck


def _solid_rgb(fill_parent) -> Optional[str]:
    """'#rrggbb' of an explicit solid fill, else None (theme/gradient/none)."""
    try:
        f = fill_parent.fill
        from pptx.enum.dml import MSO_FILL
        if f.type == MSO_FILL.SOLID and f.fore_color.type is not None:
            return "#" + str(f.fore_color.rgb)
    except Exception:
        pass
    return None


def _bg_rgb(cSld_owner) -> Optional[str]:
    try:
        clr = cSld_owner._element.xpath("./p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val")
        return "#" + clr[0] if clr else None
    except Exception:
        return None


def _run_style(font) -> dict:
    st = {}
    try:
        if font.size:
            st["size"] = font.size.pt
    except Exception:
        pass
    for k in ("bold", "italic"):
        try:
            if getattr(font, k):
                st[k] = True
        except Exception:
            pass
    try:
        if font.color and font.color.type is not None:
            st["color"] = "#" + str(font.color.rgb)
    except Exception:
        pass
    return st


def _frame_paras(tf) -> list:
    out = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            if r.text:
                runs.append({"text": r.text.replace("\v", "\n"), **_run_style(r.font)})
        para = {"level": p.level, "runs": runs}
        try:
            if p.alignment is not None:
                para["align"] = str(p.alignment).split(".")[-1].split(" ")[0].lower()
        except Exception:
            pass
        ps = _run_style(p.font)
        if ps:
            para["style"] = ps
        out.append(para)
    return out


def _preview_shapes(shapes, tx, budget: list, skip_placeholders: bool = False) -> list:
    """Flatten shapes into positioned boxes. tx maps a shape's (x, y, w, h) EMU
    into slide EMU (identity at top level, group child-space transform inside groups)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []
    for sh in shapes:
        if skip_placeholders and sh.is_placeholder:
            continue
        try:
            x, y, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if x is None or w is None:
            continue
        box = tx(x, y or 0, w, h or 0)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                xfrm = sh._element.xpath("./p:grpSpPr/a:xfrm")[0]
                off, ext = xfrm.find(_a("off")), xfrm.find(_a("ext"))
                choff, chext = xfrm.find(_a("chOff")), xfrm.find(_a("chExt"))
                cx, cy = int(choff.get("x")), int(choff.get("y"))
                cw, ch = max(int(chext.get("cx")), 1), max(int(chext.get("cy")), 1)
                gx, gy, gw, gh = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))

                def inner(ix, iy, iw, ih, _o=tx, gx=gx, gy=gy, gw=gw, gh=gh, cx=cx, cy=cy, cw=cw, ch=ch):
                    sx, sy = gw / cw, gh / ch
                    return _o(gx + (ix - cx) * sx, gy + (iy - cy) * sy, iw * sx, ih * sy)
                out.extend(_preview_shapes(sh.shapes, inner, budget))
            except Exception:
                pass
            continue
        item = {"x": box[0], "y": box[1], "w": box[2], "h": box[3], "kind": _shape_kind(sh)}
        try:
            if sh.rotation:
                item["rot"] = sh.rotation
        except Exception:
            pass
        if item["kind"] == "picture":
            try:
                img = sh.image
                blob = img.blob
                if len(blob) <= PREVIEW_MAX_IMAGE and budget[0] + len(blob) <= PREVIEW_MAX_IMAGES \
                        and img.content_type.startswith("image/") and img.ext.lower() != "wmf":
                    import base64
                    budget[0] += len(blob)
                    item["src"] = f"data:{img.content_type};base64," + base64.b64encode(blob).decode()
            except Exception:
                pass
            out.append(item)
            continue
        if item["kind"] == "table":
            rows = []
            for row in sh.table.rows:
                rows.append([_clip(c.text, 300) for c in row.cells])
            item["rows"] = rows
            out.append(item)
            continue
        if item["kind"] == "chart":
            try:
                ch = sh.chart
                item["title"] = ch.chart_title.text_frame.text if ch.has_title else ""
            except Exception:
                item["title"] = ""
            out.append(item)
            continue
        fill = _solid_rgb(sh) if hasattr(sh, "fill") else None
        if fill:
            item["fill"] = fill
        try:
            if sh.line.fill.type is not None and sh.line.color.type is not None:
                item["line"] = "#" + str(sh.line.color.rgb)
        except Exception:
            pass
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            item["paras"] = _frame_paras(sh.text_frame)
            try:
                anchor = sh.text_frame.vertical_anchor
                if anchor is not None:
                    item["anchor"] = str(anchor).split(".")[-1].split(" ")[0].lower()
            except Exception:
                pass
        if item.get("paras") or fill or item.get("line"):
            out.append(item)
    return out


def preview(data: bytes) -> dict:
    """Positioned, render-ready slide model: coordinates in EMU, fonts in pt."""
    prs = _open(data)
    w, h = int(prs.slide_width or 12192000), int(prs.slide_height or 6858000)
    ident = lambda x, y, cw, ch: (int(x), int(y), int(cw), int(ch))  # noqa: E731
    budget = [0]
    slides = []
    for slide in prs.slides:
        layout = slide.slide_layout
        master = layout.slide_master
        bg = _bg_rgb(slide) or _bg_rgb(layout) or _bg_rgb(master)
        shapes = []
        # master/layout artwork (logos, bars) under the slide's own shapes
        for owner in (master, layout):
            try:
                shapes.extend(_preview_shapes(owner.shapes, ident, budget, skip_placeholders=True))
            except Exception:
                pass
        shapes.extend(_preview_shapes(slide.shapes, ident, budget))
        notes = ""
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame
            notes = nt.text.strip() if nt is not None else ""
        slides.append({"bg": bg, "shapes": shapes, "notes": notes})
    return {"width": w, "height": h, "slides": slides}


# ------------------------------------------------------------------
# Addressing
# ------------------------------------------------------------------

_ADDR = re.compile(r"^s(\d+)(?:/(title|notes|sh(\d+)))?(?:/(?:p(\d+)|tbl[\[(](\d+)\s*,\s*(\d+)[\])]))?$", re.I)


def _parse(addr: str):
    m = _ADDR.match(str(addr).strip().replace(" ", ""))
    if not m:
        raise DocOpError(f"bad address '{addr}' (use s3, s3/sh5, s3/sh5/p2, s3/sh5/tbl[1,2], s3/title, s3/notes)")
    si, part, shid, pi, tr, tc = m.groups()
    return int(si), (part or "").lower(), int(shid) if shid else None, int(pi) if pi else None, \
        (int(tr), int(tc)) if tr else None


def _slide_at(prs, n: int):
    require(1 <= n <= len(prs.slides), f"slide {n} does not exist (deck has {len(prs.slides)})")
    return prs.slides[n - 1]


def _find_shape(slide, shape_id: int):
    for sh in _iter_shapes(slide.shapes):
        if sh.shape_id == shape_id:
            return sh
    raise DocOpError(f"no shape sh{shape_id} on that slide - run doc_inspect for valid addresses")


def _title_shape(slide):
    t = slide.shapes.title
    if t is not None:
        return t
    texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    if not texts:
        return None
    return min(texts, key=lambda s: (s.top or 0, s.left or 0))


def _body_shape(slide, exclude):
    cands = [s for s in slide.shapes if s.has_text_frame and s.shape_id != exclude.shape_id]
    if not cands:
        return None
    ph = [s for s in cands if s.is_placeholder]
    pool = ph or cands
    return max(pool, key=lambda s: (s.width or 0) * (s.height or 0))


# ------------------------------------------------------------------
# Text writing that keeps formatting
# ------------------------------------------------------------------

def _set_para_text(p_el, text: str) -> None:
    """Replace a paragraph's text, keeping pPr (bullets/alignment) and the
    first run's rPr (font/size/colour)."""
    text = text.replace("\v", " ")
    runs = p_el.findall(_a("r"))
    for el in list(p_el):
        if el.tag in (_a("br"), _a("fld")):
            p_el.remove(el)
    if runs:
        first = runs[0]
        for r in runs[1:]:
            p_el.remove(r)
    else:
        first = etree.SubElement(p_el, _a("r"))
        end = p_el.find(_a("endParaRPr"))
        if end is not None:
            rpr = copy.deepcopy(end)
            rpr.tag = _a("rPr")
            first.insert(0, rpr)
            p_el.remove(first)
            end.addprevious(first)
    t = first.find(_a("t"))
    if t is None:
        t = etree.SubElement(first, _a("t"))
    t.text = text


def _set_frame_text(txBody, text: str) -> None:
    """Replace the text of a text body; one paragraph per line, extra lines
    clone the last paragraph's formatting. Leading tabs set the bullet level."""
    lines = str(text).split("\n")
    paras = txBody.findall(_a("p"))
    if not paras:
        paras = [etree.SubElement(txBody, _a("p"))]
    for i, line in enumerate(lines):
        # indentation mirrors inspect's display: one tab or two spaces per level
        stripped = line.lstrip("\t ")
        lead = line[:len(line) - len(stripped)]
        level = lead.count("\t") + lead.count(" ") // 2
        if i < len(paras):
            p = paras[i]
        else:
            p = copy.deepcopy(paras[-1])
            txBody.findall(_a("p"))[-1].addnext(p)
        _set_para_text(p, stripped)
        ppr = p.find(_a("pPr"))
        if level:
            if ppr is None:
                ppr = etree.Element(_a("pPr"))
                p.insert(0, ppr)
            ppr.set("lvl", str(min(level, 8)))
        elif ppr is not None and "lvl" in ppr.attrib:
            del ppr.attrib["lvl"]
    for p in paras[len(lines):]:
        txBody.remove(p)


def _txbody_of(sh):
    require(sh.has_text_frame, f"shape sh{sh.shape_id} has no text")
    return sh.text_frame._txBody


# ------------------------------------------------------------------
# Editing
# ------------------------------------------------------------------

def _snapshot_items(prs):
    items = [("slide order", prs.slides._sldIdLst)]
    for i, slide in enumerate(prs.slides, 1):
        spTree = slide.shapes._spTree
        for el in spTree:
            if etree.QName(el).localname in ("nvGrpSpPr", "grpSpPr"):
                continue
            nv = el[0].find(f"{{{P_NS}}}cNvPr") if len(el) else None
            items.append((f"s{i}/sh{nv.get('id') if nv is not None else '?'}", el))
        root = slide._element
        for el in root:
            if etree.QName(el).localname != "cSld":
                items.append((f"s{i}/{etree.QName(el).localname}", el))
        bg = root.find(f"{{{P_NS}}}cSld/{{{P_NS}}}bg")
        if bg is not None:
            items.append((f"s{i}/background", bg))
        if slide.has_notes_slide:
            items.append((f"s{i}/notes", slide.notes_slide._element))
    return items


def _rename_parts(parts: dict[str, bytes]) -> dict[str, str]:
    """python-pptx renumbers slide parts on save; key them by the slide id instead."""
    out = {}
    try:
        pres = parse_xml(parts["ppt/presentation.xml"])
        rels = parse_xml(parts["ppt/_rels/presentation.xml.rels"])
    except KeyError:
        return out
    r_target = {r.get("Id"): r.get("Target") for r in rels}
    for sld in pres.iter(f"{{{P_NS}}}sldId"):
        tgt = r_target.get(sld.get(f"{{{R_NS}}}id"))
        if not tgt:
            continue
        name = tgt.lstrip("/") if tgt.startswith("/") else "ppt/" + tgt
        name = name.replace("ppt/../", "")
        key = f"slide:{sld.get('id')}"
        out[name] = key
        d, f = name.rsplit("/", 1)
        srels = parts.get(f"{d}/_rels/{f}.rels")
        if srels:
            for r in parse_xml(srels):
                if r.get("Type") == RT_NOTES:
                    import posixpath
                    out[posixpath.normpath(posixpath.join(d, r.get("Target")))] = f"notes:{sld.get('id')}"
    return out


class _Ctx:
    def __init__(self, prs):
        self.prs = prs
        self.snap = Snapshot(xml_fp)
        self.snap.take(_snapshot_items(prs))
        self.rules = PartRules()
        self.changes: list[str] = []
        self.notes: list[str] = []
        self.images: dict[str, bytes] = {}

    def sid(self, slide) -> int:
        return slide.slide_id

    def text_edit(self, slide, el):
        self.snap.touch(el)
        self.rules.changed.add(f"slide:{slide.slide_id}")

    def structural(self):
        self.snap.touch(self.prs.slides._sldIdLst)
        self.rules.changed |= {"ppt/presentation.xml", "ppt/presentation.xml:rels", "[Content_Types].xml"}

    def new_slide(self, slide):
        k = f"slide:{slide.slide_id}"
        self.rules.added_prefixes |= {k, k + ":rels"}
        self.snap.touch(slide._element, ancestors=False)
        for el in slide._element.iter():
            self.snap.touch(el, ancestors=False)


def _target_txbody(ctx: _Ctx, addr: str):
    si, part, shid, pi, cell = _parse(addr)
    slide = _slide_at(ctx.prs, si)
    if part == "notes":
        ns = slide.notes_slide  # creates it if missing
        k = slide.slide_id
        ctx.rules.changed |= {f"notes:{k}", f"notes:{k}:rels", f"slide:{k}:rels", "[Content_Types].xml",
                              "ppt/presentation.xml", "ppt/presentation.xml:rels"}
        ctx.rules.added_prefixes |= {"ppt/notesMasters/", "ppt/theme/", f"notes:{k}", f"notes:{k}:rels"}
        ctx.snap.touch(ns._element)
        for el in ns._element.iter():
            ctx.snap.touch(el, ancestors=False)
        return slide, ns.notes_text_frame._txBody, None
    sh = _title_shape(slide) if part == "title" else (_find_shape(slide, shid) if shid else None)
    require(sh is not None, f"address '{addr}' must name a shape (s3/sh5), title or notes")
    if cell:
        require(getattr(sh, "has_table", False) and sh.has_table, f"sh{sh.shape_id} is not a table")
        r, c = cell
        tbl = sh.table
        require(1 <= r <= len(tbl.rows) and 1 <= c <= len(tbl.columns), f"cell [{r},{c}] outside the table")
        return slide, tbl.cell(r - 1, c - 1).text_frame._txBody, None
    body = _txbody_of(sh)
    if pi:
        paras = body.findall(_a("p"))
        require(1 <= pi <= len(paras), f"sh{sh.shape_id} has {len(paras)} paragraphs, not p{pi}")
        return slide, body, paras[pi - 1]
    return slide, body, None


def _op_set_text(ctx: _Ctx, op: dict):
    addr = op_arg(op, "addr", "target", "address")
    text = str(op_arg(op, "text", "value", default=""))
    slide, body, para = _target_txbody(ctx, addr)
    ctx.text_edit(slide, body)
    if para is not None:
        _set_para_text(para, text)
    else:
        _set_frame_text(body, text)
    ctx.changes.append(f"{addr}: text set")


def _all_txbodies(ctx: _Ctx, scope: Optional[str]):
    slides = list(enumerate(ctx.prs.slides, 1))
    if scope:
        si = _parse(scope)[0]
        slides = [(si, _slide_at(ctx.prs, si))]
    for i, slide in slides:
        for sh in _iter_shapes(slide.shapes):
            if sh.has_text_frame:
                yield i, slide, f"s{i}/sh{sh.shape_id}", sh.text_frame._txBody
            elif getattr(sh, "has_table", False) and sh.has_table:
                for r, row in enumerate(sh.table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        yield i, slide, f"s{i}/sh{sh.shape_id}/tbl[{r},{c}]", cell.text_frame._txBody


def _para_text(p_el) -> str:
    return "".join((t.text or "") for t in p_el.iter(_a("t")))


def replace_in_paragraph(p_el, find: str, repl: str) -> bool:
    """Replace inside single runs when possible (keeps per-run formatting);
    fall back to rewriting the paragraph when a match spans runs."""
    full = _para_text(p_el)
    if find not in full:
        return False
    ts = list(p_el.iter(_a("t")))
    if sum((t.text or "").count(find) for t in ts) == full.count(find):
        for t in ts:
            if t.text and find in t.text:
                t.text = t.text.replace(find, repl)
    else:
        _set_para_text(p_el, full.replace(find, repl))
    return True


def _op_replace_text(ctx: _Ctx, op: dict):
    find = str(op_arg(op, "find", "old", "search"))
    repl = str(op_arg(op, "replace", "new", "text", default=""))
    require(find != "", "replace_text needs a non-empty 'find'")
    hits = []
    for _i, slide, label, body in _all_txbodies(ctx, op.get("scope")):
        for p in body.findall(_a("p")):
            if replace_in_paragraph(p, find, repl):
                ctx.text_edit(slide, p)
                hits.append(label)
    require(bool(hits), f"'{find}' not found in the presentation")
    ctx.changes.append(f"replaced '{find}' -> '{repl}' in {', '.join(dict.fromkeys(hits))}")


def _move_sldId(prs, sldId_el, position: int):
    lst = prs.slides._sldIdLst
    lst.remove(sldId_el)
    lst.insert(position, sldId_el)


def _sldId_of(prs, slide):
    for el in prs.slides._sldIdLst:
        if int(el.get("id")) == slide.slide_id:
            return el
    raise DocOpError("slide id not found")


def _duplicate(ctx: _Ctx, src):
    prs = ctx.prs
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if etree.QName(el).localname not in ("nvGrpSpPr", "grpSpPr"):
            tree.remove(el)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT_LAYOUT, RT_NOTES):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree:
        if etree.QName(el).localname in ("nvGrpSpPr", "grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    src_bg = src._element.find(f"{{{P_NS}}}cSld/{{{P_NS}}}bg")
    if src_bg is not None:
        new._element.find(f"{{{P_NS}}}cSld").insert(0, copy.deepcopy(src_bg))
    for el in new._element.iter():
        for attr in list(el.attrib):
            if attr.startswith(f"{{{R_NS}}}") and el.get(attr) in rid_map:
                el.set(attr, rid_map[el.get(attr)])
    ctx.new_slide(new)
    return new


def _op_duplicate_slide(ctx: _Ctx, op: dict):
    si = _parse(op_arg(op, "addr", "slide", "target"))[0]
    src = _slide_at(ctx.prs, si)
    ctx.structural()
    new = _duplicate(ctx, src)
    after = int(op.get("after") or si)
    _move_sldId(ctx.prs, _sldId_of(ctx.prs, new), after)
    ctx.changes.append(f"duplicated s{si} as new slide {after + 1}")
    return new


def _op_add_slide(ctx: _Ctx, op: dict):
    prs = ctx.prs
    n = len(prs.slides)
    after = op.get("after")
    after = n if after in (None, "", "end") else _parse(after)[0] if str(after).lower().startswith("s") else int(after)
    require(0 <= after <= n, f"'after' must be 0..{n}")
    title = str(op.get("title") or "")
    bullets = op.get("bullets") or op.get("body") or []
    if isinstance(bullets, str):
        bullets = [b for b in bullets.split("\n")]
    ctx.structural()
    layout_name = op.get("layout")
    like = op.get("like")
    if layout_name:
        layouts = {l.name.lower(): l for m in prs.slide_masters for l in m.slide_layouts}
        lay = layouts.get(str(layout_name).lower())
        require(lay is not None, f"layout '{layout_name}' not found (have: {', '.join(layouts)})")
        new = prs.slides.add_slide(lay)
        ctx.new_slide(new)
    else:
        # Copy the look of a neighbouring content slide (theme, fonts, positions).
        if like:
            src = _slide_at(prs, _parse(like)[0])
        else:
            pool = [prs.slides[i] for i in range(max(after - 1, 0), n)] + list(prs.slides)
            src = next((s for s in pool if s.shapes.title is not None or
                        sum(1 for x in s.shapes if x.has_text_frame) >= 2), prs.slides[max(after - 1, 0)] if n else None)
        if src is None:
            new = prs.slides.add_slide(prs.slide_layouts[1 if len(prs.slide_layouts) > 1 else 0])
            ctx.new_slide(new)
        else:
            new = _duplicate(ctx, src)
    tshape = _title_shape(new)
    if tshape is not None:
        _set_frame_text(tshape.text_frame._txBody, title)
        body = _body_shape(new, tshape)
        if body is not None:
            _set_frame_text(body.text_frame._txBody, "\n".join(str(b) for b in bullets) if bullets else "")
        # other text shapes copied from the template slide would carry stale text
        for s in new.shapes:
            if s.has_text_frame and s.shape_id not in (tshape.shape_id, body.shape_id if body else -1) \
                    and not s.is_placeholder:
                _set_frame_text(s.text_frame._txBody, "")
    _move_sldId(prs, _sldId_of(prs, new), after)
    if op.get("notes"):
        new.notes_slide.notes_text_frame.text = str(op["notes"])
        k = new.slide_id
        ctx.rules.added_prefixes |= {f"notes:{k}", f"notes:{k}:rels", "ppt/notesMasters/", "ppt/theme/"}
    ctx.changes.append(f"added slide {after + 1} '{title}'")


def _op_delete_slide(ctx: _Ctx, op: dict):
    prs = ctx.prs
    si = _parse(op_arg(op, "addr", "slide", "target"))[0]
    slide = _slide_at(prs, si)
    require(len(prs.slides) > 1, "cannot delete the only slide")
    ctx.structural()
    k = slide.slide_id
    ctx.rules.removed |= {f"slide:{k}", f"slide:{k}:rels", f"notes:{k}", f"notes:{k}:rels"}
    el = _sldId_of(prs, slide)
    for x in slide._element.iter():
        ctx.snap.touch(x, ancestors=False)
    prs.part.drop_rel(el.get(f"{{{R_NS}}}id"))
    prs.slides._sldIdLst.remove(el)
    ctx.changes.append(f"deleted slide {si}")


def _op_move_slide(ctx: _Ctx, op: dict):
    prs = ctx.prs
    si = _parse(op_arg(op, "addr", "slide", "target"))[0]
    to = int(op_arg(op, "to", "position"))
    require(1 <= to <= len(prs.slides), f"'to' must be 1..{len(prs.slides)}")
    slide = _slide_at(prs, si)
    ctx.structural()
    _move_sldId(prs, _sldId_of(prs, slide), to - 1)
    ctx.changes.append(f"moved slide {si} to position {to}")


def _op_set_notes(ctx: _Ctx, op: dict):
    si = _parse(op_arg(op, "addr", "slide", "target"))[0]
    op = dict(op, addr=f"s{si}/notes")
    _op_set_text(ctx, op)


def _op_replace_image(ctx: _Ctx, op: dict):
    addr = op_arg(op, "addr", "target")
    si, _part, shid, _pi, _cell = _parse(addr)
    slide = _slide_at(ctx.prs, si)
    require(shid is not None, "replace_image needs a picture address like s2/sh4")
    sh = _find_shape(slide, shid)
    blip = sh._element.find(".//" + _a("blip"))
    require(blip is not None, f"sh{shid} is not a picture")
    img = op.get("_image_bytes")
    require(isinstance(img, (bytes, bytearray)) and img, "replace_image needs 'image' (a file you uploaded)")
    _img_part, rid = slide.part.get_or_add_image_part(io.BytesIO(img))
    ctx.text_edit(slide, sh._element)
    blip.set(f"{{{R_NS}}}embed", rid)
    k = slide.slide_id
    ctx.rules.changed |= {f"slide:{k}:rels", "[Content_Types].xml"}
    ctx.rules.added_prefixes.add("ppt/media/")
    ctx.changes.append(f"{addr}: image replaced")


OPS = {
    "set_text": _op_set_text,
    "set_table_cell": _op_set_text,
    "replace_text": _op_replace_text,
    "add_slide": _op_add_slide,
    "duplicate_slide": _op_duplicate_slide,
    "delete_slide": _op_delete_slide,
    "move_slide": _op_move_slide,
    "set_notes": _op_set_notes,
    "replace_image": _op_replace_image,
}


def apply(data: bytes, ops: list[dict]) -> EditResult:
    prs = _open(data)
    ctx = _Ctx(prs)
    for op in ops:
        fn = OPS.get(str(op.get("op", "")).lower())
        require(fn is not None, f"unknown pptx op '{op.get('op')}' (have: {', '.join(OPS)})")
        fn(ctx, op)
    problems = ctx.snap.problems(_snapshot_items(prs))
    if problems:
        raise DocOpError("edit would change parts you did not ask for: " + "; ".join(problems[:8]))
    out = _save(prs)
    problems = check_package(data, out, ctx.rules, _rename_parts)
    if problems:
        raise DocOpError("edit would change untouched package parts: " + "; ".join(problems[:8]))
    return EditResult(out, ctx.changes, ctx.notes)


# ------------------------------------------------------------------
# Creation (placeholder-based, so later edits keep the theme)
# ------------------------------------------------------------------

def parse_markdown_slides(md: str) -> list[dict]:
    slides, cur = [], {"title": "", "bullets": [], "notes": ""}

    def flush():
        if cur["title"] or cur["bullets"] or cur["notes"]:
            slides.append(dict(cur))

    slide_marker = re.compile(
        r'^(?:'
        r'---\s*$'
        r'|(?:\#{1,3}\s+(?P<h_title>.+))$'
        r'|(?:\*{0,2}\s*Slide\s+\d+\s*(?:\((?P<cat>[^)]+)\))?\s*:\*{0,2}\s*(?P<s_rest>.*))$'
        r')',
        re.IGNORECASE
    )

    for line in (md or "").strip().splitlines():
        s = line.rstrip()
        st = s.strip()
        if not st:
            continue
        # Skip deck meta header/footer outlines
        if re.match(r'^\d+\s*slides\s*\+.*structure', st, re.IGNORECASE) or st.startswith("GLOBAL THEME:"):
            continue

        if st.startswith("---"):
            flush()
            cur = {"title": "", "bullets": [], "notes": ""}
            continue

        m = slide_marker.match(st)
        if m:
            h_title = m.group("h_title")
            cat = m.group("cat") or ""
            s_rest = m.group("s_rest") or ""

            if h_title:
                if cur["title"] or cur["bullets"]:
                    flush()
                    cur = {"title": "", "bullets": [], "notes": ""}
                cur["title"] = h_title.strip().lstrip("#").strip()
            else:
                # Conversational format: Slide 1 (Title): ...
                flush()
                cur = {"title": "", "bullets": [], "notes": ""}
                # Try finding header/title in quotes first: Title "Why SSL Wireless..."
                m_title = re.search(r'(?:header(?:\s+banner)?|title)\s*(?:[^:]+:|\s*)?\s*\"([^\"]+)\"', s_rest, re.IGNORECASE)
                if m_title:
                    cur["title"] = m_title.group(1).strip()
                else:
                    q = re.findall(r'\"([^\"]+)\"', s_rest)
                    if cat:
                        cur["title"] = cat.strip()
                    elif q:
                        cur["title"] = q[0].strip()
                    elif s_rest:
                        cur["title"] = s_rest.strip()
                    else:
                        cur["title"] = "Slide"
            continue

        if st.lower().startswith(("notes:", "note:")):
            cur["notes"] = st.split(":", 1)[1].strip()
        elif st.startswith(("- ", "* ", "+ ")) or re.match(r"^\d+[.)]\s", st):
            indent = (len(s) - len(s.lstrip(" "))) // 2
            txt = re.sub(r"^([-*+]|\d+[.)])\s+", "", st)
            cur["bullets"].append("\t" * indent + txt)
        elif st:
            cur["bullets"].append(st.lstrip("#").strip())
    flush()
    return slides


def create(md: str, template: Optional[bytes] = None, title: str = "") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    slides = parse_markdown_slides(md) or [{"title": title or "Presentation", "bullets": [], "notes": ""}]
    if template:
        prs = _open(template)
        # keep the template's masters/theme, drop its example slides
        lst = prs.slides._sldIdLst
        for el in list(lst):
            prs.part.drop_rel(el.get(f"{{{R_NS}}}id"))
            lst.remove(el)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layouts = list(prs.slide_layouts)

    def pick(title_only: bool):
        for l in layouts:
            kinds = {str(p.placeholder_format.type).split(".")[-1].split(" ")[0] for p in l.placeholders}
            if title_only and {"CENTER_TITLE", "SUBTITLE"} & kinds:
                return l
            if not title_only and "TITLE" in kinds and ({"BODY", "OBJECT"} & kinds):
                return l
        return layouts[1 if len(layouts) > 1 else 0]

    for i, sd in enumerate(slides):
        cover = i == 0 and len(sd["bullets"]) <= 1
        slide = prs.slides.add_slide(pick(cover))

        if not template:
            try:
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(11, 25, 44)
            except Exception:
                pass

        if cover:
            if slide.shapes.title is not None:
                slide.shapes.title.text_frame.text = sd["title"]
                if not template:
                    slide.shapes.title.left = Inches(1.2)
                    slide.shapes.title.top = Inches(2.2)
                    slide.shapes.title.width = Inches(10.9)
                    slide.shapes.title.height = Inches(1.5)
                    for p in slide.shapes.title.text_frame.paragraphs:
                        p.font.name = "Segoe UI"
                        p.font.size = Pt(38)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(255, 255, 255)
            body = next((p for p in slide.placeholders if p.placeholder_format.idx != 0), None)
            if body is not None:
                if sd["bullets"]:
                    _set_frame_text(body.text_frame._txBody, "\n".join(sd["bullets"]))
                    if not template:
                        body.left = Inches(1.2)
                        body.top = Inches(3.8)
                        body.width = Inches(10.9)
                        body.height = Inches(1.2)
                        for p in body.text_frame.paragraphs:
                            p.font.name = "Segoe UI"
                            p.font.size = Pt(18)
                            p.font.color.rgb = RGBColor(148, 163, 184)
                else:
                    body._element.getparent().remove(body._element)
        else:
            if slide.shapes.title is not None:
                slide.shapes.title.text_frame.text = sd["title"]
                if not template:
                    slide.shapes.title.left = Inches(0.8)
                    slide.shapes.title.top = Inches(0.6)
                    slide.shapes.title.width = Inches(11.733)
                    slide.shapes.title.height = Inches(0.9)
                    for p in slide.shapes.title.text_frame.paragraphs:
                        p.font.name = "Segoe UI"
                        p.font.size = Pt(28)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(255, 255, 255)
            body = next((p for p in slide.placeholders if p.placeholder_format.idx != 0), None)
            if body is not None:
                if sd["bullets"]:
                    _set_frame_text(body.text_frame._txBody, "\n".join(sd["bullets"]))
                    if not template:
                        body.left = Inches(0.8)
                        body.top = Inches(1.6)
                        body.width = Inches(11.733)
                        body.height = Inches(5.0)
                        for p in body.text_frame.paragraphs:
                            p.font.name = "Segoe UI"
                            p.font.size = Pt(16)
                            p.font.color.rgb = RGBColor(226, 232, 240)
                            p.space_after = Pt(12)
                else:
                    body._element.getparent().remove(body._element)

        if sd["notes"]:
            slide.notes_slide.notes_text_frame.text = sd["notes"]
    return _save(prs)
